import sys
import os
import re
import json
import shutil
import sqlite3
import jdatetime
from datetime import datetime

from PyQt6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QLineEdit, QComboBox, QPushButton, QTableWidget, QTableWidgetItem,
    QHeaderView, QMessageBox, QGroupBox, QSpinBox, QFormLayout, QFileDialog,
    QTabWidget, QCheckBox, QInputDialog, QDialog, QStackedWidget, QFrame,
    QTextEdit, QScrollArea, QTreeWidget, QTreeWidgetItem, QDateEdit,
    QProgressBar, QGraphicsDropShadowEffect, QListWidget, QListWidgetItem,
    QGridLayout, QMenu, QToolButton, QSizePolicy, QAbstractItemView,
    QDialogButtonBox, QRadioButton, QButtonGroup, QSplitter
)
from PyQt6.QtCore import (Qt, QDate, QTimer, QSizeF, QRectF, QPointF, QSize, QRect,
                          pyqtSignal, QMarginsF, QByteArray, QPropertyAnimation, QEasingCurve)
from PyQt6.QtGui import (
    QFont, QIcon, QColor, QFontDatabase, QTextDocument, QPixmap, QPainter, QImage,
    QLinearGradient, QBrush, QPageSize, QPageLayout, QAction, QTextOption,
    QFontMetrics, QIntValidator, QPalette, QPen, QDesktopServices
)
from PyQt6.QtCore import QUrl
from PyQt6.QtPrintSupport import QPrinter, QPrintDialog, QPrinterInfo

# ---------------------------------------------------------------
# مسیر کش فونت matplotlib را کنار خود برنامه تنظیم می‌کنیم.
# در حالت EXE تک‌فایلی (--onefile) پوشه موقت هر بار عوض می‌شود و
# matplotlib مجبور است هر اجرا کش فونت را از نو بسازد (چند ثانیه تأخیر).
# با این تنظیم، کش یک‌بار ساخته و دفعات بعد سریع اجرا می‌شود.
# ---------------------------------------------------------------
try:
    _app_base = os.path.dirname(os.path.abspath(
        sys.executable if getattr(sys, "frozen", False) else __file__))
    os.environ.setdefault("MPLCONFIGDIR", os.path.join(_app_base, "mpl_cache"))
except Exception:
    pass

import matplotlib
matplotlib.use('QtAgg')
import matplotlib.pyplot as plt
import matplotlib.font_manager as fm
from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure

# تنظیم ID جهت آیکون ویندوز
try:
    import ctypes
    myappid = 'imartstudio.accounting.v11.0'
    ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(myappid)
except Exception:
    pass


def get_app_dir():
    """
    پوشه ثابت برنامه.
    رفع باگ اصلی «ذخیره نشدن تغییرات بعد از بستن برنامه»:
    قبلاً مسیر دیتابیس نسبی بود و با هر بار اجرا از پوشه‌ای دیگر،
    یک دیتابیس خالی جدید ساخته می‌شد و اطلاعات قبلی گم می‌شد.
    """
    try:
        if getattr(sys, "frozen", False):
            return os.path.dirname(os.path.abspath(sys.executable))
        return os.path.dirname(os.path.abspath(__file__))
    except Exception:
        return os.path.abspath(os.getcwd())


APP_DIR = get_app_dir()
DB_NAME = os.path.join(APP_DIR, "studio_accounting.db")
APP_VERSION = "11.0"
DEVELOPER_NAME = "میلاد محمدحسینی"

# ⚠️ مقدار پیش‌فرض فونت (بعد از ساخت QApplication نهایی می‌شود)
APP_FONT_FAMILY = "Tahoma"
INVOICE_FONT_FAMILY = "Tahoma"
APP_FONT_PATH = None
INVOICE_FONT_PATH = None


def setup_matplotlib_font():
    """
    پیدا کردن فونت فارسی برای matplotlib.
    ترتیب اولویت: یکان → نازنین → وزیر → فونت‌های همراه برنامه → Tahoma
    """
    candidates = []
    for fp in (APP_FONT_PATH, INVOICE_FONT_PATH):
        if fp:
            candidates.append(fp)
    candidates += [
        "C:/Windows/Fonts/BYekan.ttf", "C:/Windows/Fonts/byekan.ttf",
        "C:/Windows/Fonts/BNAZANIN.TTF", "C:/Windows/Fonts/bnazanin.ttf",
        "C:/Windows/Fonts/Vazirmatn-Regular.ttf",
        "/Library/Fonts/BYekan.ttf",
    ]
    for fname in ("BYekan.ttf", "BNAZANIN.TTF", "Vazirmatn-Regular.ttf"):
        candidates.append(resource_path(fname))
        candidates.append(resource_path(os.path.join("fonts", fname)))
    candidates += [
        "C:/Windows/Fonts/tahoma.ttf",
        "C:/Windows/Fonts/arial.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]

    chosen = None
    for fp in candidates:
        if fp and os.path.exists(fp):
            try:
                fm.fontManager.addfont(fp)
                prop = fm.FontProperties(fname=fp)
                chosen = prop.get_name()
                break
            except Exception as e:
                print(f"[MATPLOTLIB FONT] Failed {fp}: {e}")
                continue

    if not chosen:
        # آخرین تلاش: فونت‌های نصب‌شده سیستم که فارسی را پشتیبانی می‌کنند
        for fam in ("B Yekan", "B Nazanin", "Vazirmatn", "Tahoma", "DejaVu Sans"):
            try:
                if any(f.name == fam for f in fm.fontManager.ttflist):
                    chosen = fam
                    break
            except Exception:
                pass

    # فهرست جایگزین: فونت فارسی اول، بعد فونت‌هایی که حروف لاتین و اعداد را دارند
    # (رفع هشدار «Glyph missing» برای کلماتی مثل VIP و اعداد لاتین)
    fallbacks = [chosen] if chosen else []
    for fam in ("Tahoma", "DejaVu Sans", "Arial", "Vazirmatn"):
        if fam not in fallbacks and any(f.name == fam for f in fm.fontManager.ttflist):
            fallbacks.append(fam)
    if not fallbacks:
        fallbacks = ["DejaVu Sans"]
    plt.rcParams['font.family'] = 'sans-serif'
    plt.rcParams['font.sans-serif'] = fallbacks
    plt.rcParams['axes.unicode_minus'] = False
    plt.rcParams['figure.autolayout'] = True
    print(f"[MATPLOTLIB FONT] Using: {chosen} | fallbacks: {fallbacks}")
    return bool(chosen)


def resource_path(relative_path):
    """جست‌وجوی یک منبع (فونت/آیکون) در همه مسیرهای ممکن برنامه"""
    candidates = []
    try:
        candidates.append(os.path.join(sys._MEIPASS, relative_path))
    except Exception:
        pass
    candidates.append(os.path.join(APP_DIR, relative_path))
    candidates.append(os.path.join(APP_DIR, "assets", relative_path))
    candidates.append(os.path.join(APP_DIR, "fonts", relative_path))
    candidates.append(os.path.join(os.path.abspath("."), relative_path))
    for c in candidates:
        if os.path.exists(c):
            return c
    return candidates[0]


def migrate_legacy_database():
    """
    اگر نسخه قبلی برنامه دیتابیس را کنار «مسیر اجرا» یا در پوشه جاری ساخته بود،
    آن را به مسیر ثابت جدید منتقل می‌کنیم تا اطلاعات کاربر گم نشود.
    """
    try:
        if os.path.exists(DB_NAME):
            return
        legacy_candidates = [
            os.path.join(os.path.abspath("."), "studio_accounting.db"),
            os.path.join(os.getcwd(), "studio_accounting.db"),
        ]
        for legacy in legacy_candidates:
            if os.path.exists(legacy) and os.path.abspath(legacy) != os.path.abspath(DB_NAME):
                try:
                    shutil.copy2(legacy, DB_NAME)
                    print(f"[DB] Migrated legacy database: {legacy} -> {DB_NAME}")
                    return
                except Exception as e:
                    print(f"[DB] Legacy migration failed: {e}")
    except Exception as e:
        print(f"[DB] migrate_legacy_database error: {e}")


def setup_fonts():
    """تشخیص فونت مناسب فارسی - باید بعد از ساخت QApplication صدا زده شود"""
    global INVOICE_FONT_FAMILY, INVOICE_FONT_PATH, APP_FONT_PATH

    # ۱) فونت‌های همراه برنامه (پوشه fonts) را ثبت کن
    bundled_candidates = [
        "BYekan.ttf", "B Yekan.ttf", "byekan.ttf",
        "BNazanin.ttf", "B Nazanin.ttf", "bnazanin.ttf",
        "Vazirmatn-Regular.ttf", "IRANSansWeb.ttf", "Sahel.ttf",
    ]
    for fname in bundled_candidates:
        p = resource_path(fname)
        if os.path.exists(p):
            try:
                QFontDatabase.addApplicationFont(p)
            except Exception:
                pass
        p2 = resource_path(os.path.join("fonts", fname))
        if os.path.exists(p2):
            try:
                QFontDatabase.addApplicationFont(p2)
            except Exception:
                pass

    available_fonts = QFontDatabase.families()

    has_nazanin = "B Nazanin" in available_fonts
    has_yekan = "B Yekan" in available_fonts
    has_vazir = "Vazirmatn" in available_fonts or "Vazir" in available_fonts
    has_iransans = "IRANSans" in available_fonts or "IRANSansWeb" in available_fonts
    has_sahel = "Sahel" in available_fonts

    print(f"[FONT] B Nazanin installed: {has_nazanin}")
    print(f"[FONT] B Yekan installed: {has_yekan}")

    # انتخاب فونت رابط کاربری (خوانا و مدرن)
    if has_yekan:
        font_family = "B Yekan"
    elif has_vazir:
        font_family = "Vazirmatn" if "Vazirmatn" in available_fonts else "Vazir"
    elif has_nazanin:
        font_family = "B Nazanin"
    elif has_iransans:
        font_family = "IRANSans"
    elif has_sahel:
        font_family = "Sahel"
    else:
        font_family = "Tahoma"

    # انتخاب فونت فاکتور/چاپ (درخواست کاربر: یکان یا نازنین)
    if has_nazanin:
        INVOICE_FONT_FAMILY = "B Nazanin"
    elif has_yekan:
        INVOICE_FONT_FAMILY = "B Yekan"
    elif has_vazir:
        INVOICE_FONT_FAMILY = "Vazirmatn" if "Vazirmatn" in available_fonts else "Vazir"
    else:
        INVOICE_FONT_FAMILY = "Tahoma"

    INVOICE_FONT_PATH = _find_font_file(INVOICE_FONT_FAMILY)
    APP_FONT_PATH = _find_font_file(font_family)

    print(f"[FONT] Using for UI: {font_family}")
    print(f"[FONT] Using for invoice: {INVOICE_FONT_FAMILY}")
    return font_family


def _find_font_file(family_name):
    """یافتن فایل فیزیکی فونت برای استفاده در matplotlib"""
    wf = "C:/Windows/Fonts"
    named = {
        "B Nazanin": ["BNAZANIN.TTF", "bnazanin.ttf", "BNazanin.ttf"],
        "B Yekan": ["BYekan.ttf", "byekan.ttf", "B Yekan.ttf"],
        "Vazirmatn": ["Vazirmatn-Regular.ttf", "Vazirmatn.ttf"],
        "Vazir": ["Vazir.ttf", "Vazirmatn-Regular.ttf"],
        "IRANSans": ["IRANSansWeb.ttf", "IRANSans.ttf"],
        "Sahel": ["Sahel.ttf"],
        "Tahoma": ["tahoma.ttf"],
    }
    for fname in named.get(family_name, []):
        for base in (wf, os.path.join(wf, ""), APP_DIR, os.path.join(APP_DIR, "fonts")):
            p = os.path.join(base, fname)
            if os.path.exists(p):
                return p
    # جست‌وجوی عمومی بر اساس نام فونت
    try:
        for f in fm.fontManager.ttflist:
            if f.name == family_name and os.path.exists(f.fname):
                return f.fname
    except Exception:
        pass
    return None


_FA_DIGITS = str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")


def format_number(val):
    try:
        clean_val = str(val).translate(_FA_DIGITS).replace(',', '').replace(' ', '').strip()
        if not clean_val:
            return ""
        return f"{int(clean_val):,}"
    except ValueError:
        return str(val)


def parse_number(val_str):
    try:
        return int(str(val_str).translate(_FA_DIGITS).replace(',', '').replace(' ', '').strip())
    except Exception:
        return 0


def number_to_persian_words(num):
    """تبدیل عدد به حروف فارسی"""
    try:
        num = int(num)
    except Exception:
        return ""
    if num == 0:
        return "صفر"
    if num < 0:
        return "منفی " + number_to_persian_words(-num)

    yekan = ["", "یک", "دو", "سه", "چهار", "پنج", "شش", "هفت", "هشت", "نه"]
    dahgan = ["", "", "بیست", "سی", "چهل", "پنجاه", "شصت", "هفتاد", "هشتاد", "نود"]
    dah = ["ده", "یازده", "دوازده", "سیزده", "چهارده", "پانزده", "شانزده", "هفده", "هجده", "نوزده"]
    sadgan = ["", "صد", "دویست", "سیصد", "چهارصد", "پانصد", "ششصد", "هفتصد", "هشتصد", "نهصد"]
    scale = ["", " هزار", " میلیون", " میلیارد", " تریلیون"]

    def three_digit_to_words(n):
        result = []
        s = n // 100
        rem = n % 100
        if s > 0:
            result.append(sadgan[s])
        if rem >= 10 and rem < 20:
            result.append(dah[rem - 10])
        else:
            d = rem // 10
            y = rem % 10
            if d > 0:
                result.append(dahgan[d])
            if y > 0:
                result.append(yekan[y])
        return " و ".join(result)

    parts = []
    idx = 0
    while num > 0:
        chunk = num % 1000
        if chunk > 0:
            parts.insert(0, three_digit_to_words(chunk) + scale[idx])
        num //= 1000
        idx += 1

    return " و ".join(parts)


# ================================================================================
# ====================  زیرساخت مشترک نسخه ۱۰ (v10 Core)  ========================
# ================================================================================

# ---------------------------------------------------------------
# ۱) آیکون‌های گرافیکی متناسب با موضوع هر دکمه و بخش
# ---------------------------------------------------------------
ICONS = {
    "wedding": "💍", "commercial": "🎬", "staff": "👥", "expenses": "💸",
    "reports": "📈", "search": "🔍", "inventory": "📦", "banks": "🏦",
    "checks": "📑", "receipt": "🧾", "settings": "⚙️", "about": "ℹ️",
    "save": "💾", "edit": "✏️", "delete": "🗑️", "print": "🖨️", "pdf": "📄",
    "excel": "📗", "home": "🏠", "backup": "💾", "restore": "♻️", "key": "🔑",
    "add": "➕", "refresh": "🔄", "chart": "📊", "calc": "🧮", "money": "💰",
    "calendar": "📅", "next": "▶", "prev": "◀", "package": "🎁", "detail": "🔎",
    "ok": "✅", "cancel": "✖️", "people": "🧑", "camera": "🎥", "contract": "📜",
    "warn": "⚠️", "filter": "🧹", "deposit": "💵", "time": "⏳", "today": "📅",
    "word": "🔤", "list": "📋", "link": "🔗", "star": "⭐", "closed": "🔒",
}


def ico(key):
    """دریافت آیکون گرافیکی بر اساس کلید"""
    return ICONS.get(key, "")


def ico_text(key, text, icon_first=True):
    """ترکیب آیکون و متن برای دکمه‌ها"""
    i = ico(key)
    if not i:
        return text
    return f"{i} {text}" if icon_first else f"{text} {i}"


# ---------------------------------------------------------------
# ۲) تم گرافیکی سراسری — زیباسازی کل برنامه بدون تغییر در ساختار
# ---------------------------------------------------------------
GLOBAL_QSS = """
* {
    font-family: "APPFONT";
}
QWidget {
    background-color: #eef2f8;
    color: #24313f;
    selection-background-color: #2980b9;
    selection-color: #ffffff;
}
QMainWindow, QDialog {
    background-color: qlineargradient(x1:0, y1:0, x2:1, y2:1,
        stop:0 #eaf1f8, stop:0.5 #f4f7fb, stop:1 #e6eef7);
}
QLabel {
    background: transparent;
}
QLabel[heading="true"] {
    font-size: 15pt;
    font-weight: bold;
    color: #1F4E78;
}

QGroupBox {
    background-color: #ffffff;
    border: 1px solid #d6e0ec;
    border-radius: 12px;
    margin-top: 18px;
    padding: 14px 10px 10px 10px;
    font-weight: bold;
}
QGroupBox::title {
    subcontrol-origin: margin;
    subcontrol-position: top right;
    right: 16px;
    top: 2px;
    padding: 3px 14px;
    color: #ffffff;
    background-color: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #1F4E78, stop:1 #2980b9);
    border-radius: 9px;
}

QLineEdit, QTextEdit, QPlainTextEdit, QSpinBox, QDoubleSpinBox, QComboBox {
    background-color: #ffffff;
    border: 1px solid #cbd7e6;
    border-radius: 8px;
    padding: 5px 9px;
    min-height: 24px;
    color: #1d2b38;
}
QLineEdit:focus, QTextEdit:focus, QPlainTextEdit:focus,
QSpinBox:focus, QDoubleSpinBox:focus, QComboBox:focus {
    border: 2px solid #2980b9;
    background-color: #fdfefe;
}
QLineEdit:disabled, QComboBox:disabled, QSpinBox:disabled {
    background-color: #eef1f5;
    color: #99a4b0;
}
QLineEdit[readOnly="true"] {
    background-color: #f6f8fb;
}
QComboBox::drop-down {
    subcontrol-origin: padding;
    subcontrol-position: center left;
    width: 26px;
    border: none;
    background: transparent;
}
QComboBox::down-arrow {
    image: none;
    border-left: 5px solid transparent;
    border-right: 5px solid transparent;
    border-top: 6px solid #2980b9;
    margin-left: 8px;
}
QComboBox QAbstractItemView {
    background-color: #ffffff;
    border: 1px solid #cbd7e6;
    border-radius: 8px;
    selection-background-color: #2980b9;
    selection-color: #ffffff;
    outline: none;
    padding: 4px;
}

QPushButton {
    background-color: qlineargradient(x1:0, y1:0, x2:0, y2:1,
        stop:0 #3d8fd1, stop:1 #2980b9);
    color: #ffffff;
    border: none;
    border-radius: 9px;
    padding: 7px 14px;
    font-weight: bold;
    min-height: 22px;
}
QPushButton:hover {
    background-color: qlineargradient(x1:0, y1:0, x2:0, y2:1,
        stop:0 #4fa3e3, stop:1 #2f8fcc);
}
QPushButton:pressed {
    background-color: #1f6a9c;
}
QPushButton:disabled {
    background-color: #b8c3d0;
    color: #eef2f7;
}
QPushButton[flat="true"] {
    background-color: transparent;
    color: #2980b9;
}

QCheckBox, QRadioButton {
    spacing: 8px;
    background: transparent;
    padding: 3px;
}
QCheckBox::indicator, QRadioButton::indicator {
    width: 17px;
    height: 17px;
    border: 2px solid #a9b8c9;
    background-color: #ffffff;
}
QCheckBox::indicator { border-radius: 5px; }
QRadioButton::indicator { border-radius: 9px; }
QCheckBox::indicator:hover, QRadioButton::indicator:hover {
    border-color: #2980b9;
}
QCheckBox::indicator:checked {
    background-color: qlineargradient(x1:0, y1:0, x2:1, y2:1,
        stop:0 #27ae60, stop:1 #2ecc71);
    border-color: #1e8449;
    image: none;
}
QRadioButton::indicator:checked {
    background-color: qlineargradient(x1:0, y1:0, x2:1, y2:1,
        stop:0 #2980b9, stop:1 #5dade2);
    border-color: #1F4E78;
}

QTabWidget::pane {
    border: 1px solid #d6e0ec;
    border-radius: 12px;
    background-color: #fbfdff;
    top: -1px;
}
QTabBar {
    background: transparent;
}
QTabBar::tab {
    background-color: #dde6f0;
    color: #405060;
    border: 1px solid #cfdbe8;
    border-bottom: none;
    border-top-left-radius: 10px;
    border-top-right-radius: 10px;
    padding: 8px 16px;
    margin-left: 3px;
    font-weight: bold;
}
QTabBar::tab:hover {
    background-color: #e8f1fa;
    color: #1F4E78;
}
QTabBar::tab:selected {
    background-color: qlineargradient(x1:0, y1:0, x2:0, y2:1,
        stop:0 #ffffff, stop:1 #f2f8ff);
    color: #1F4E78;
    border-bottom: 3px solid #2980b9;
}

QTableWidget, QTableView, QTreeWidget, QListWidget {
    background-color: #ffffff;
    alternate-background-color: #f4f8fd;
    border: 1px solid #d6e0ec;
    border-radius: 10px;
    gridline-color: #e4ebf3;
    selection-background-color: #d6e9f8;
    selection-color: #123;
    outline: none;
}
QTableWidget::item, QTreeWidget::item {
    padding: 5px;
    border: none;
}
QTableWidget::item:selected, QTreeWidget::item:selected {
    background-color: #d6e9f8;
    color: #123;
}
QHeaderView::section {
    background-color: qlineargradient(x1:0, y1:0, x2:0, y2:1,
        stop:0 #3f7fb5, stop:1 #2c6699);
    color: #ffffff;
    border: none;
    border-left: 1px solid #35699b;
    padding: 8px 6px;
    font-weight: bold;
}
QHeaderView::section:hover {
    background-color: #4a8cc4;
}
QTableCornerButton::section {
    background-color: #2c6699;
    border: none;
}

QScrollBar:vertical {
    background: #eef2f8;
    width: 12px;
    margin: 0;
    border-radius: 6px;
}
QScrollBar::handle:vertical {
    background: #b3c4d6;
    border-radius: 6px;
    min-height: 30px;
}
QScrollBar::handle:vertical:hover { background: #8fa8c2; }
QScrollBar:horizontal {
    background: #eef2f8;
    height: 12px;
    margin: 0;
    border-radius: 6px;
}
QScrollBar::handle:horizontal {
    background: #b3c4d6;
    border-radius: 6px;
    min-width: 30px;
}
QScrollBar::handle:horizontal:hover { background: #8fa8c2; }
QScrollBar::add-line, QScrollBar::sub-line { height: 0; width: 0; }
QScrollBar::add-page, QScrollBar::sub-page { background: none; }

QProgressBar {
    background-color: #e4ebf3;
    border: 1px solid #cfdbe8;
    border-radius: 9px;
    text-align: center;
    color: #1F4E78;
    font-weight: bold;
}
QProgressBar::chunk {
    background-color: qlineargradient(x1:0, y1:0, x2:1, y2:0,
        stop:0 #27ae60, stop:0.5 #2ecc71, stop:1 #1abc9c);
    border-radius: 8px;
}

QToolTip {
    background-color: #1F4E78;
    color: #ffffff;
    border: none;
    border-radius: 6px;
    padding: 6px 9px;
}
QMenu {
    background-color: #ffffff;
    border: 1px solid #cfdbe8;
    border-radius: 8px;
    padding: 5px;
}
QMenu::item {
    padding: 7px 22px;
    border-radius: 6px;
}
QMenu::item:selected {
    background-color: #2980b9;
    color: #ffffff;
}
QSplitter::handle { background-color: #dfe8f2; }
QScrollArea { border: none; background: transparent; }
QScrollArea > QWidget > QWidget { background: transparent; }
"""


def apply_global_theme(app):
    """اعمال تم گرافیکی روی کل برنامه"""
    try:
        app.setStyle("Fusion")
    except Exception:
        pass
    qss = GLOBAL_QSS.replace('"APPFONT"', f'"{APP_FONT_FAMILY}"')
    try:
        app.setStyleSheet(qss)
    except Exception as e:
        print("[THEME] Failed to apply stylesheet:", e)


def shade_color(hex_color, amount):
    """روشن یا تیره کردن یک رنگ هگز (amount مثبت = روشن‌تر)"""
    try:
        h = str(hex_color).lstrip('#')
        if len(h) != 6:
            return hex_color
        r = max(0, min(255, int(h[0:2], 16) + amount))
        g = max(0, min(255, int(h[2:4], 16) + amount))
        b = max(0, min(255, int(h[4:6], 16) + amount))
        return f"#{r:02x}{g:02x}{b:02x}"
    except Exception:
        return hex_color


def card_shadow(widget, blur=22, dy=4, alpha=70):
    """افزودن سایه زیبا به کارت‌ها و دکمه‌های بزرگ"""
    try:
        eff = QGraphicsDropShadowEffect()
        eff.setBlurRadius(blur)
        eff.setColor(QColor(31, 78, 120, alpha))
        eff.setOffset(0, dy)
        widget.setGraphicsEffect(eff)
    except Exception:
        pass


# ---------------------------------------------------------------
# ۳) تولید کد یکتا — هیچ دو کدی هرگز یکسان نمی‌شود
# ---------------------------------------------------------------
def code_exists(table, value, column="code", exclude_id=None):
    if not value:
        return False
    try:
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        if exclude_id is None:
            cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE {column}=?", (value,))
        else:
            cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE {column}=? AND id<>?", (value, exclude_id))
        n = cursor.fetchone()[0]
        conn.close()
        return n > 0
    except Exception:
        return False


# کدهای رزروشده در همین اجرا (لایه سوم محافظت از یکتایی کدها)
_CODE_RESERVATIONS = set()


def generate_unique_code(prefix, table, column="code", start=1001, exclude_id=None, width=4):
    """
    کد یکتا می‌سازد: بیشترین عدد موجود با آن پیشوند را پیدا کرده و یکی اضافه می‌کند.
    برخلاف روش قبلی (COUNT+1) بعد از حذف رکوردها هم هرگز تکراری تولید نمی‌شود.
    سه لایه محافظت: شمارش از بیشینه دیتابیس، بررسی وجود در دیتابیس،
    و رزرو درون‌برنامه‌ای تا حتی دو فراخوانی پشت‌سرهم هم کد یکسان ندهند.
    """
    max_n = start - 1
    try:
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute(f"SELECT {column} FROM {table} WHERE {column} LIKE ?", (f"{prefix}-%",))
        for (val,) in cursor.fetchall():
            if not val:
                continue
            try:
                suffix = str(val).split("-")[-1]
                n = int(suffix)
                if n > max_n:
                    max_n = n
            except Exception:
                continue
        conn.close()
    except Exception:
        pass

    n = max_n + 1
    candidate = f"{prefix}-{n:0{width}d}"
    guard = 0
    while (code_exists(table, candidate, column, exclude_id)
           or candidate in _CODE_RESERVATIONS) and guard < 1000000:
        n += 1
        candidate = f"{prefix}-{n:0{width}d}"
        guard += 1
    _CODE_RESERVATIONS.add(candidate)
    return candidate


def ensure_unique_code_index(table, column="code"):
    """ایجاد ایندکس یگانه تا در سطح دیتابیس هم کد تکراری ممکن نباشد"""
    try:
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute(
            f"CREATE UNIQUE INDEX IF NOT EXISTS ux_{table}_{column} ON {table}({column})"
        )
        conn.commit()
        conn.close()
        return True
    except Exception as e:
        # اگر داده‌های قدیمی تکراری باشند، ایندکس ساخته نمی‌شود ولی برنامه متوقف نمی‌شود
        print(f"[DB] unique index {table}.{column} skipped: {e}")
        return False


# ---------------------------------------------------------------
# ۴) موتور چاپ و PDF — سازگار با همه نسخه‌های PyQt6
# ---------------------------------------------------------------
def build_print_document(html):
    """ساخت QTextDocument فارسی و راست‌چین با فونت یکان/نازنین"""
    doc = QTextDocument()
    font_family = INVOICE_FONT_FAMILY or APP_FONT_FAMILY or "Tahoma"
    f = QFont(font_family, 10)
    doc.setDefaultFont(f)
    try:
        opt = doc.defaultTextOption()
        opt.setTextDirection(Qt.LayoutDirection.RightToLeft)
        opt.setAlignment(Qt.AlignmentFlag.AlignRight)
        doc.setDefaultTextOption(opt)
    except Exception:
        pass
    doc.setDocumentMargin(18)
    doc.setHtml(html)
    return doc


def render_document_to_printer(doc, printer):
    """
    رندر سند روی پرینتر یا فایل PDF.
    در PyQt6 متد درست QTextDocument.print است (در PyQt5 نامش print_ بود).
    این تابع هر دو حالت و در نهایت رندر با QPainter را پشتیبانی می‌کند،
    پس خطای «object has no attribute print» به‌طور کامل برطرف می‌شود.
    """
    errors = []
    for name in ("print", "print_"):
        fn = getattr(doc, name, None)
        if callable(fn):
            try:
                fn(printer)
                return True
            except Exception as e:
                errors.append(f"{name}: {e}")
    # آخرین راه‌حل: رندر دستی چندصفحه‌ای با QPainter
    try:
        page_pt = printer.pageRect(QPrinter.Unit.Point)
        page_px = printer.pageRect(QPrinter.Unit.DevicePixel)
        doc.setPageSize(QSizeF(page_pt.width(), page_pt.height()))
        painter = QPainter(printer)
        if not painter.isActive():
            errors.append("painter not active")
            print("[PRINT] all methods failed:", errors)
            return False
        # تبدیل واحد سند (point) به پیکسل دستگاه
        sx = page_px.width() / max(1.0, float(page_pt.width()))
        sy = page_px.height() / max(1.0, float(page_pt.height()))
        painter.scale(sx, sy)
        count = max(1, doc.pageCount())
        for page in range(count):
            if page:
                printer.newPage()
            painter.save()
            painter.translate(0, -page * page_pt.height())
            doc.drawContents(painter)
            painter.restore()
        painter.end()
        return True
    except Exception as e:
        errors.append(f"painter: {e}")
        print("[PRINT] all methods failed:", errors)
        return False


def _make_printer(output_format=None, file_path=None, doc_name="سند",
                  page_size=None, margins_mm=7):
    """
    ساخت پرینتر — پیش‌فرض همه اسناد نسخه ۱۱ روی کاغذ A5 تنظیم شده است.
    """
    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setDocName(doc_name)
    try:
        printer.setPageSize(QPageSize(page_size or QPageSize.PageSizeId.A5))
    except Exception:
        pass
    try:
        printer.setPageMargins(QMarginsF(margins_mm, margins_mm, margins_mm, margins_mm),
                               QPageLayout.Unit.Millimeter)
    except Exception:
        pass
    if output_format is not None:
        printer.setOutputFormat(output_format)
    if file_path:
        printer.setOutputFileName(file_path)
    return printer


def print_html_document(html, parent=None, doc_name="چاپ سند", page_size=None):
    """چاپ واقعی سند روی یک برگ A5 با دیالوگ پرینتر ویندوز"""
    try:
        printer = _make_printer(doc_name=doc_name, page_size=page_size)
        dialog = QPrintDialog(printer, parent)
        dialog.setWindowTitle(f"چاپ {doc_name}")
        if dialog.exec() != QPrintDialog.DialogCode.Accepted:
            return False
        doc = fit_document_to_page(html, printer)
        ok = render_document_to_printer(doc, printer)
        if not ok:
            QMessageBox.critical(parent, "خطای چاپ",
                                 "چاپ انجام نشد. لطفاً از درایور چاپگر پیش‌فرض سیستم مطمئن شوید.")
        return ok
    except Exception as e:
        QMessageBox.critical(parent, "خطای چاپ", f"خطا در چاپ:\n{e}")
        return False


def save_html_pdf(html, file_path, doc_name="سند", page_size=None):
    """ذخیره PDF فارسی رنگی روی یک برگ A5 با فونت صحیح"""
    try:
        printer = _make_printer(QPrinter.OutputFormat.PdfFormat, file_path, doc_name, page_size)
        doc = fit_document_to_page(html, printer)
        ok = render_document_to_printer(doc, printer)
        if not ok:
            return False, "رندر سند روی PDF ناموفق بود."
        if not os.path.exists(file_path) or os.path.getsize(file_path) == 0:
            return False, "فایل PDF ساخته نشد."
        return True, ""
    except Exception as e:
        return False, str(e)


# ---------------------------------------------------------------
# ۵) خروجی اکسل (openpyxl) — راست‌چین و زیبا
# ---------------------------------------------------------------
def export_rows_to_excel(headers, rows, file_path, sheet_title="گزارش", report_title=None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
    except Exception as e:
        return False, f"کتابخانه openpyxl در دسترس نیست ({e})"

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = (sheet_title or "گزارش")[:31]
        ws.sheet_view.rightToLeft = True

        start_row = 1
        if report_title:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(headers)))
            c = ws.cell(row=1, column=1, value=report_title)
            c.font = Font(bold=True, size=14, color="1F4E78", name="Tahoma")
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 26
            start_row = 2

        head_fill = PatternFill("solid", fgColor="2C6699")
        head_font = Font(bold=True, color="FFFFFF", name="Tahoma", size=11)
        thin = Side(style="thin", color="B8C6D6")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col, title in enumerate(headers, start=1):
            cell = ws.cell(row=start_row, column=col, value=str(title))
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = border
        ws.row_dimensions[start_row].height = 24

        alt_fill = PatternFill("solid", fgColor="F4F8FD")
        for r, row in enumerate(rows, start=start_row + 1):
            for c_idx, val in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c_idx, value=val)
                cell.font = Font(name="Tahoma", size=10)
                cell.border = border
                if isinstance(val, (int, float)):
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                if (r - start_row) % 2 == 0:
                    cell.fill = alt_fill

        for c_idx in range(1, len(headers) + 1):
            max_len = len(str(headers[c_idx - 1]))
            for row in rows:
                if c_idx - 1 < len(row):
                    max_len = max(max_len, len(str(row[c_idx - 1])))
            ws.column_dimensions[get_column_letter(c_idx)].width = min(42, max(11, max_len + 4))

        wb.save(file_path)
        return True, ""
    except Exception as e:
        return False, str(e)


# ---------------------------------------------------------------
# ۶) شکل‌دهی حروف فارسی برای matplotlib (رفع برعکس و جدا بودن حروف)
# ---------------------------------------------------------------
_ARABIC_FORMS = {
    "\u0621": ("\uFE80", None, None, None),
    "\u0622": ("\uFE81", "\uFE82", None, None),
    "\u0623": ("\uFE83", "\uFE84", None, None),
    "\u0624": ("\uFE85", "\uFE86", None, None),
    "\u0625": ("\uFE87", "\uFE88", None, None),
    "\u0626": ("\uFE89", "\uFE8A", "\uFE8B", "\uFE8C"),
    "\u0627": ("\uFE8D", "\uFE8E", None, None),
    "\u0628": ("\uFE8F", "\uFE90", "\uFE91", "\uFE92"),
    "\u0629": ("\uFE93", "\uFE94", None, None),
    "\u062A": ("\uFE95", "\uFE96", "\uFE97", "\uFE98"),
    "\u062B": ("\uFE99", "\uFE9A", "\uFE9B", "\uFE9C"),
    "\u062C": ("\uFE9D", "\uFE9E", "\uFE9F", "\uFEA0"),
    "\u062D": ("\uFEA1", "\uFEA2", "\uFEA3", "\uFEA4"),
    "\u062E": ("\uFEA5", "\uFEA6", "\uFEA7", "\uFEA8"),
    "\u062F": ("\uFEA9", "\uFEAA", None, None),
    "\u0630": ("\uFEAB", "\uFEAC", None, None),
    "\u0631": ("\uFEAD", "\uFEAE", None, None),
    "\u0632": ("\uFEAF", "\uFEB0", None, None),
    "\u0633": ("\uFEB1", "\uFEB2", "\uFEB3", "\uFEB4"),
    "\u0634": ("\uFEB5", "\uFEB6", "\uFEB7", "\uFEB8"),
    "\u0635": ("\uFEB9", "\uFEBA", "\uFEBB", "\uFEBC"),
    "\u0636": ("\uFEBD", "\uFEBE", "\uFEBF", "\uFEC0"),
    "\u0637": ("\uFEC1", "\uFEC2", "\uFEC3", "\uFEC4"),
    "\u0638": ("\uFEC5", "\uFEC6", "\uFEC7", "\uFEC8"),
    "\u0639": ("\uFEC9", "\uFECA", "\uFECB", "\uFECC"),
    "\u063A": ("\uFECD", "\uFECE", "\uFECF", "\uFED0"),
    "\u0641": ("\uFED1", "\uFED2", "\uFED3", "\uFED4"),
    "\u0642": ("\uFED5", "\uFED6", "\uFED7", "\uFED8"),
    "\u0643": ("\uFED9", "\uFEDA", "\uFEDB", "\uFEDC"),
    "\u0644": ("\uFEDD", "\uFEDE", "\uFEDF", "\uFEE0"),
    "\u0645": ("\uFEE1", "\uFEE2", "\uFEE3", "\uFEE4"),
    "\u0646": ("\uFEE5", "\uFEE6", "\uFEE7", "\uFEE8"),
    "\u0647": ("\uFEE9", "\uFEEA", "\uFEEB", "\uFEEC"),
    "\u0648": ("\uFEED", "\uFEEE", None, None),
    "\u0649": ("\uFEEF", "\uFEF0", None, None),
    "\u064A": ("\uFEF1", "\uFEF2", "\uFEF3", "\uFEF4"),
    # حروف مخصوص فارسی
    "\u067E": ("\uFB56", "\uFB57", "\uFB58", "\uFB59"),   # پ
    "\u0686": ("\uFB7A", "\uFB7B", "\uFB7C", "\uFB7D"),   # چ
    "\u0698": ("\uFB8A", "\uFB8B", None, None),           # ژ
    "\u06A9": ("\uFB8E", "\uFB8F", "\uFB90", "\uFB91"),   # ک
    "\u06AF": ("\uFB92", "\uFB93", "\uFB94", "\uFB95"),   # گ
    "\u06CC": ("\uFBFC", "\uFBFD", "\uFBFE", "\uFBFF"),   # ی
    "\u06C0": ("\uFBA4", "\uFBA5", None, None),           # ۀ
}

# حروفی که فقط از سمت راست می‌چسبند (حرف بعدی را به خود نمی‌چسبانند)
_RIGHT_JOINING_ONLY = set("\u0622\u0623\u0624\u0625\u0627\u0629\u062F\u0630"
                          "\u0631\u0632\u0698\u0648\u0621\u0649\u06C0")

_LAM_ALEF = {
    "\u0622": ("\uFEF5", "\uFEF6"),
    "\u0623": ("\uFEF7", "\uFEF8"),
    "\u0625": ("\uFEF9", "\uFEFA"),
    "\u0627": ("\uFEFB", "\uFEFC"),
}

_DIACRITICS = set("\u064B\u064C\u064D\u064E\u064F\u0650\u0651\u0652\u0640")


def _is_arabic(ch):
    return ch in _ARABIC_FORMS


_NON_JOINING = {"\u0621"}  # ء فقط


def _can_join_prev(ch):
    """آیا این حرف از سمت راست به حرف قبلی می‌چسبد؟"""
    return ch in _ARABIC_FORMS and ch not in _NON_JOINING


def _can_join_next(ch):
    """آیا این حرف به حرف بعدی (سمت چپ) می‌چسبد؟"""
    return (ch in _ARABIC_FORMS and ch not in _NON_JOINING
            and ch not in _RIGHT_JOINING_ONLY)


def _shape_logical(text):
    """شکل‌دهی حروف بر اساس همسایه‌های منطقی (الگوریتم استاندارد اتصال عربی)"""
    chars = list(text)
    out = []
    for i, ch in enumerate(chars):
        if ch in _DIACRITICS:
            out.append(ch)
            continue
        if ch not in _ARABIC_FORMS:
            out.append(ch)
            continue

        prev_ch = None
        for j in range(i - 1, -1, -1):
            if chars[j] in _DIACRITICS:
                continue
            prev_ch = chars[j]
            break

        next_idx = None
        next_ch = None
        for j in range(i + 1, len(chars)):
            if chars[j] in _DIACRITICS:
                continue
            next_idx = j
            next_ch = chars[j]
            break

        connect_prev = bool(prev_ch) and _can_join_next(prev_ch) and _can_join_prev(ch)
        connect_next = bool(next_ch) and _can_join_next(ch) and _can_join_prev(next_ch)

        # ترکیب لام + الف به شکل لیگاتور (لا)
        if ch == "\u0644" and next_ch in _LAM_ALEF:
            isolated, final = _LAM_ALEF[next_ch]
            out.append(final if connect_prev else isolated)
            if next_idx is not None:
                chars[next_idx] = "\u0000"
            continue

        iso, fin, ini, med = _ARABIC_FORMS[ch]
        if connect_prev and connect_next and med:
            out.append(med)
        elif connect_prev and fin:
            out.append(fin)
        elif connect_next and ini:
            out.append(ini)
        else:
            out.append(iso)
    return "".join(c for c in out if c != "\u0000")


def _is_ltr_char(ch):
    return ("0" <= ch <= "9") or ("A" <= ch <= "Z") or ("a" <= ch <= "z") or ch in "./:%+-"


def _reorder_rtl(text):
    """معکوس‌سازی برای نمایش راست‌به‌چپ، با حفظ ترتیب اعداد و کلمات لاتین"""
    rev = text[::-1]
    out = []
    i = 0
    n = len(rev)
    while i < n:
        if _is_ltr_char(rev[i]):
            j = i
            while j < n and _is_ltr_char(rev[j]):
                j += 1
            out.append(rev[i:j][::-1])
            i = j
        else:
            out.append(rev[i])
            i += 1
    return "".join(out)


def fa_shape(text):
    """
    آماده‌سازی متن فارسی برای matplotlib:
    حروف را به هم می‌چسباند (رفع «جدا جدا» بودن) و ترتیب را راست‌به‌چپ می‌کند
    (رفع «برعکس» بودن مثل «عروس» → «سورع»).
    """
    if text is None:
        return ""
    s = str(text)
    if not any(_is_arabic(c) for c in s):
        return s
    # اگر کتابخانه‌های تخصصی نصب باشند، از آن‌ها استفاده می‌کنیم (کیفیت بالاتر)
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(s))
    except Exception:
        pass
    # آینه‌کردن پرانتزها تا «()» به‌جای «)(» دیده شود
    return _mirror_brackets(_reorder_rtl(_shape_logical(s)))


# ---------------------------------------------------------------
# ۶-۲) آینه‌کردن پرانتزها (رفع «)(» به‌جای «()» در نمودارها)
# ---------------------------------------------------------------
_BIDI_MIRROR = {
    "(": ")", ")": "(",
    "[": "]", "]": "[",
    "{": "}", "}": "{",
    "<": ">", ">": "<",
    "«": "»", "»": "«",
}


def _mirror_brackets(text):
    """جای پرانتزها را در متن راست‌چین‌شده اصلاح می‌کند"""
    return "".join(_BIDI_MIRROR.get(c, c) for c in text)


# ---------------------------------------------------------------
# ۶-۳) پالت رنگ اسناد: حالت رنگی (PDF) و حالت سیاه و سفید (پرینت)
# ---------------------------------------------------------------
def doc_palette(mono=False):
    """
    mono=False → رنگ‌های زیبا برای PDF و نمایش روی صفحه
    mono=True  → رنگ‌های خاکستری/سیاه‌وسفید مناسب چاپگر لیزری و چاپ سیاه‌وسفید
    """
    if mono:
        return {
            "head_bg": "#2b2b2b", "head_fg": "#ffffff",
            "table_head": "#5a5a5a", "table_head_fg": "#ffffff",
            "alt_row": "#f0f0f0", "box_bg": "#e6e6e6",
            "border": "#8a8a8a", "accent": "#1a1a1a",
            "warn_bg": "#ededed", "warn_border": "#6b6b6b",
            "recv_head": "#4a4a4a", "paid_head": "#7d7d7d",
            "staff_head": "#6a6a6a", "ok": "#1a1a1a", "bad": "#1a1a1a",
            "info_bg": "#f5f5f5", "total_bg": "#dcdcdc", "muted": "#5a5a5a",
        }
    return {
        "head_bg": "#1F4E78", "head_fg": "#ffffff",
        "table_head": "#2C6699", "table_head_fg": "#ffffff",
        "alt_row": "#f4f8fd", "box_bg": "#eef5fc",
        "border": "#9bb0c4", "accent": "#1F4E78",
        "warn_bg": "#fdedec", "warn_border": "#e74c3c",
        "recv_head": "#1e8449", "paid_head": "#a93226",
        "staff_head": "#5b2c8e", "ok": "#1e8449", "bad": "#c0392b",
        "info_bg": "#eaf6ff", "total_bg": "#ffe9a8", "muted": "#7f8c8d",
    }


# ---------------------------------------------------------------
# ۶-۴) کوچک‌سازی هوشمند سند تا در یک برگ A5 جا شود
# ---------------------------------------------------------------
def scale_html(html, factor):
    """همه اندازه‌های فونت و فاصله‌های سند را با یک ضریب کوچک/بزرگ می‌کند"""
    if abs(factor - 1.0) < 0.001:
        return html

    def pad_repl(m):
        vals = [v for v in m.group(1).split() if v.endswith("px")]
        return "padding:" + " ".join(f"{float(v[:-2]) * factor:.1f}px" for v in vals)

    out = re.sub(r"padding\s*:\s*((?:[\d.]+px\s*)+)", pad_repl, html)
    out = re.sub(r"font-size\s*:\s*([\d.]+)pt",
                 lambda m: f"font-size:{float(m.group(1)) * factor:.2f}pt", out)
    out = re.sub(r"margin-top\s*:\s*([\d.]+)px",
                 lambda m: f"margin-top:{float(m.group(1)) * factor:.1f}px", out)
    out = re.sub(r"margin-bottom\s*:\s*([\d.]+)px",
                 lambda m: f"margin-bottom:{float(m.group(1)) * factor:.1f}px", out)
    return out


def fit_document_to_page(html, printer, min_scale=0.45, max_scale=1.25, step=0.05):
    """
    سند را با «بزرگ‌ترین فونت ممکن» می‌سازد که در یک برگ جا شود،
    بدون آنکه متن‌ها بشکنند و به‌هم بریزند.
    از بزرگ‌ترین مقیاس شروع می‌کند و اولین مقیاسی که در یک صفحه جا شود
    برگردانده می‌شود؛ اگر هیچ‌کدام جا نشد، کوچک‌ترین حالت برمی‌گردد.
    """
    # اندازه صفحه باید هم‌واحد با اندازه فونت (point) باشد؛
    # در غیر این صورت شمارش صفحه‌ها درست انجام نمی‌شود.
    try:
        page = printer.pageRect(QPrinter.Unit.Point)
        size = QSizeF(page.width(), page.height())
    except Exception:
        size = QSizeF(420, 595)

    scale = max_scale
    last = None
    while scale >= min_scale - 1e-6:
        try:
            doc = build_print_document(scale_html(html, scale))
            doc.setPageSize(size)
            last = doc
            if doc.pageCount() <= 1:
                return doc
        except Exception as e:
            print("[FIT] scale", round(scale, 2), "failed:", e)
        scale -= step
    return last


# ---------------------------------------------------------------
# ۶-۵) ذخیره و بازیابی ابعاد پنجره‌ها، جدول‌ها و اسپلیترها
# ---------------------------------------------------------------
def save_geometry(key, widget):
    """ابعاد و مکان فعلی پنجره را ذخیره می‌کند"""
    try:
        set_setting(f"geom_{key}",
                    bytes(widget.saveGeometry().toBase64()).decode("ascii"))
        set_setting(f"size_{key}", f"{widget.width()}x{widget.height()}")
    except Exception as e:
        print(f"[GEOM] save {key} failed:", e)


def restore_geometry(key, widget):
    """
    ابعاد و مکان ذخیره‌شده را برمی‌گرداند.
    اندازه ذخیره‌شده صریحاً هم اعمال می‌شود تا در همه سیستم‌عامل‌ها یکسان کار کند،
    فقط اگر بزرگ‌تر از صفحه‌نمایش باشد تا اندازه صفحه محدود می‌شود.
    """
    restored = False
    try:
        data = get_setting(f"geom_{key}", "")
        if data:
            restored = widget.restoreGeometry(QByteArray.fromBase64(data.encode("ascii")))
    except Exception as e:
        print(f"[GEOM] restore {key} failed:", e)

    try:
        size = get_setting(f"size_{key}", "")
        if size and "x" in size:
            w, h = (int(v) for v in size.split("x", 1))
            try:
                screen = widget.screen() or QApplication.primaryScreen()
                if screen:
                    geo = screen.availableGeometry()
                    w = min(w, geo.width())
                    h = min(h, geo.height())
            except Exception:
                pass
            w = max(w, widget.minimumWidth())
            h = max(h, widget.minimumHeight())
            widget.resize(w, h)
            restored = True
    except Exception as e:
        print(f"[GEOM] size restore {key} failed:", e)
    return restored


def persist_dialog(dialog, key):
    """ابعاد دیالوگ را ذخیره و در اجرای بعدی همان‌طور بازیابی می‌کند"""
    restore_geometry(key, dialog)
    try:
        dialog.finished.connect(lambda _r: save_geometry(key, dialog))
    except Exception:
        pass
    return dialog


def save_table_columns(key, table):
    try:
        widths = [table.columnWidth(c) for c in range(table.columnCount())]
        set_setting(f"cols_{key}", json.dumps(widths))
    except Exception:
        pass


def restore_table_columns(key, table):
    try:
        data = get_setting(f"cols_{key}", "")
        if not data:
            return False
        widths = json.loads(data)
        for c, w in enumerate(widths[:table.columnCount()]):
            if w and int(w) > 8:
                table.setColumnWidth(c, int(w))
        return True
    except Exception as e:
        print(f"[COLS] restore {key} failed:", e)
        return False


def persist_table(table, key, resize_mode=None):
    """
    پهنای ستون‌های جدول را ذخیره و در اجرای بعدی همان‌طور بازیابی می‌کند.
    اگر قبلاً کاربر ستون‌ها را تغییر داده باشد، حالت Interactive (قابل تغییر) می‌ماند.
    """
    header = table.horizontalHeader()
    header.setSectionsMovable(True)
    header.setStretchLastSection(False)

    if restore_table_columns(key, table):
        # عرض‌های ذخیره‌شده کاربر برمی‌گردد و ستون‌ها قابل تغییر می‌مانند
        header.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
    elif resize_mode == QHeaderView.ResizeMode.Stretch:
        # جدول‌های کشسان نیازی به ذخیره عرض ستون ندارند
        header.setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        return table
    else:
        # اولین اجرا: عرض متناسب با محتوا، ولی «قابل تغییر» تا کاربر بتواند تنظیم کند
        header.setSectionResizeMode(QHeaderView.ResizeMode.Interactive)
        try:
            table.resizeColumnsToContents()
            for c in range(table.columnCount()):
                if table.columnWidth(c) < 62:
                    table.setColumnWidth(c, 62)
                elif table.columnWidth(c) > 420:
                    table.setColumnWidth(c, 420)
        except Exception as e:
            print("[COLS] initial sizing failed:", e)

    header.sectionResized.connect(lambda *_: save_table_columns(key, table))
    return table


def save_splitter(key, splitter):
    try:
        set_setting(f"split_{key}", bytes(splitter.saveState().toBase64()).decode("ascii"))
    except Exception:
        pass


def restore_splitter(key, splitter):
    try:
        data = get_setting(f"split_{key}", "")
        if data:
            return splitter.restoreState(QByteArray.fromBase64(data.encode("ascii")))
    except Exception:
        pass
    return False


def persist_splitter(splitter, key):
    restore_splitter(key, splitter)
    splitter.splitterMoved.connect(lambda *_: save_splitter(key, splitter))
    return splitter


# ---------------------------------------------------------------
# ۶-۶) ویجت‌های کمکی نسخه ۱۱
# ---------------------------------------------------------------
class ElidedCheckBox(QCheckBox):
    """
    چک‌باکسی که به‌جای ایجاد اسکرول افقی، متنش کوتاه (…) می‌شود
    تا هر تعداد آیتم در عرض موجود جا شود.
    """

    def __init__(self, text="", parent=None):
        super().__init__(text, parent)
        self._full_text = text
        self.setSizePolicy(QSizePolicy.Policy.Ignored, QSizePolicy.Policy.Fixed)
        self.setMinimumWidth(60)

    def setFullText(self, text):
        self._full_text = text
        self.setToolTip(text)
        self._apply_elide()

    def fullText(self):
        return self._full_text

    def _apply_elide(self):
        try:
            fm = QFontMetrics(self.font())
            avail = max(40, self.width() - 26)
            elided = fm.elidedText(self._full_text, Qt.TextElideMode.ElideLeft, avail)
            if super().text() != elided:
                super().setText(elided)
        except Exception:
            pass

    def resizeEvent(self, event):
        super().resizeEvent(event)
        self._apply_elide()


class HoverTileButton(QPushButton):
    """دکمه کارت داشبورد با رویداد ورود/خروج موس برای انیمیشن"""

    def __init__(self, parent=None):
        super().__init__(parent)
        self.on_enter = None
        self.on_leave = None

    def enterEvent(self, event):
        try:
            if self.on_enter:
                self.on_enter()
        finally:
            super().enterEvent(event)

    def leaveEvent(self, event):
        try:
            if self.on_leave:
                self.on_leave()
        finally:
            super().leaveEvent(event)


def days_left_short(date_str):
    d = parse_jalali(date_str)
    if d is None:
        return "-"
    try:
        left = (d - jdatetime.date.today()).days
    except Exception:
        return "-"
    if left > 0:
        return f"⏳ {left} روز مانده"
    if left == 0:
        return "🎉 امروز"
    return "✅ برگزار شده"


def settlement_state(net, paid):
    """وضعیت تسویه: کامل / ناقص / نشده"""
    net = net or 0
    paid = paid or 0
    remain = max(0, net - paid)
    if remain <= 0:
        return "تسویه کامل", remain
    if paid <= 0:
        return "تسویه نشده", remain
    return "تسویه ناقص", remain

# ---------------------------------------------------------------
# ۷) ویجت تاریخ شمسی با پرش خودکار روز → ماه → سال
# ---------------------------------------------------------------
class PersianDateEdit(QWidget):
    """
    ورودی تاریخ شمسی.
    مخاطب اول «روز» را می‌نویسد، به‌محض کامل شدن، نشانگر خودکار به «ماه»
    و سپس به «سال» می‌رود؛ نیازی به کلیک نیست.
    متد text() برای سازگاری کامل با کد قبلی رشته «YYYY/MM/DD» برمی‌گرداند.
    """

    dateChanged = pyqtSignal(str)

    def __init__(self, parent=None, date_str=None, show_today_button=True):
        super().__init__(parent)
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self._suppress = False

        lay = QHBoxLayout(self)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(3)

        self.ed_day = QLineEdit()
        self.ed_month = QLineEdit()
        self.ed_year = QLineEdit()
        for e, w, mx, ph in (
            (self.ed_day, 40, 2, "روز"),
            (self.ed_month, 40, 2, "ماه"),
            (self.ed_year, 52, 4, "سال"),
        ):
            e.setFixedWidth(w)
            e.setMaxLength(mx)
            e.setAlignment(Qt.AlignmentFlag.AlignCenter)
            e.setValidator(QIntValidator(0, 9999, self))
            e.setPlaceholderText(ph)
            e.setToolTip(ph)

        lay.addWidget(self.ed_day)
        lay.addWidget(QLabel("/"))
        lay.addWidget(self.ed_month)
        lay.addWidget(QLabel("/"))
        lay.addWidget(self.ed_year)

        if show_today_button:
            btn_today = QPushButton("امروز")
            btn_today.setFixedWidth(58)
            btn_today.setToolTip("درج تاریخ امروز")
            btn_today.clicked.connect(lambda: self.set_date(jdatetime.date.today().strftime("%Y/%m/%d")))
            lay.addWidget(btn_today)

        lay.addStretch()

        self.ed_day.textChanged.connect(lambda t: self._on_part_changed("day", t))
        self.ed_month.textChanged.connect(lambda t: self._on_part_changed("month", t))
        self.ed_year.textChanged.connect(lambda t: self._on_part_changed("year", t))
        self.ed_day.editingFinished.connect(self._normalize_all)
        self.ed_month.editingFinished.connect(self._normalize_all)
        self.ed_year.editingFinished.connect(self._normalize_all)

        self.set_date(date_str or jdatetime.date.today().strftime("%Y/%m/%d"))

    # ---------- منطق پرش خودکار ----------
    def _on_part_changed(self, part, text):
        if self._suppress:
            return
        digits = "".join(ch for ch in text if ch.isdigit())
        if digits != text:
            self._set_part(part, digits)
            return

        if part == "day":
            if len(digits) == 2 and 1 <= int(digits) <= 31:
                self.ed_month.setFocus()
                self.ed_month.selectAll()
            elif len(digits) == 1 and digits in "456789":
                self._set_part("day", f"{int(digits):02d}")
                self.ed_month.setFocus()
                self.ed_month.selectAll()
        elif part == "month":
            if len(digits) == 2 and 1 <= int(digits) <= 12:
                self.ed_year.setFocus()
                self.ed_year.selectAll()
            elif len(digits) == 1 and digits in "23456789":
                self._set_part("month", f"{int(digits):02d}")
                self.ed_year.setFocus()
                self.ed_year.selectAll()

        self._emit_change()

    def _set_part(self, part, value):
        self._suppress = True
        try:
            {"day": self.ed_day, "month": self.ed_month, "year": self.ed_year}[part].setText(value)
        finally:
            self._suppress = False

    def _normalize_all(self):
        self._suppress = True
        try:
            d = self.ed_day.text().strip()
            m = self.ed_month.text().strip()
            y = self.ed_year.text().strip()
            if d:
                d = f"{max(1, min(31, int(d))):02d}"
            if m:
                m = f"{max(1, min(12, int(m))):02d}"
            self.ed_day.setText(d)
            self.ed_month.setText(m)
            self.ed_year.setText(y)
        except Exception:
            pass
        finally:
            self._suppress = False
        self._emit_change()

    def _emit_change(self):
        try:
            self.dateChanged.emit(self.text())
        except Exception:
            pass

    # ---------- API سازگار با QLineEdit ----------
    def text(self):
        d = self.ed_day.text().strip() or "01"
        m = self.ed_month.text().strip() or "01"
        y = self.ed_year.text().strip() or str(jdatetime.date.today().year)
        try:
            d = f"{max(1, min(31, int(d))):02d}"
        except Exception:
            d = "01"
        try:
            m = f"{max(1, min(12, int(m))):02d}"
        except Exception:
            m = "01"
        return f"{y}/{m}/{d}"

    def setText(self, value):
        self.set_date(value)

    def set_date(self, value):
        value = (value or "").strip()
        y = m = d = ""
        parts = None
        for sep in ("/", "-", "."):
            if sep in value:
                parts = value.split(sep)
                break
        if parts and len(parts) == 3:
            # ورودی ممکن است Y/M/D باشد
            if len(parts[0]) == 4:
                y, m, d = parts[0], parts[1], parts[2]
            else:
                d, m, y = parts[0], parts[1], parts[2]
        if not y:
            today = jdatetime.date.today()
            y, m, d = str(today.year), f"{today.month:02d}", f"{today.day:02d}"

        self._suppress = True
        try:
            self.ed_year.setText(str(y))
            self.ed_month.setText(f"{int(m):02d}" if str(m).strip() else "")
            self.ed_day.setText(f"{int(d):02d}" if str(d).strip() else "")
        except Exception:
            today = jdatetime.date.today()
            self.ed_year.setText(str(today.year))
            self.ed_month.setText(f"{today.month:02d}")
            self.ed_day.setText(f"{today.day:02d}")
        finally:
            self._suppress = False
        self._emit_change()

    def clear(self):
        self._suppress = True
        self.ed_day.clear()
        self.ed_month.clear()
        self.ed_year.clear()
        self._suppress = False
        self._emit_change()

    def setEnabled(self, flag):
        super().setEnabled(flag)
        self.ed_day.setEnabled(flag)
        self.ed_month.setEnabled(flag)
        self.ed_year.setEnabled(flag)


def jalali_today_str():
    return jdatetime.date.today().strftime("%Y/%m/%d")


def parse_jalali(date_str):
    """تبدیل رشته تاریخ شمسی به jdatetime.date با تحمل خطا"""
    if not date_str:
        return None
    for sep in ("/", "-", "."):
        if sep in str(date_str):
            parts = str(date_str).split(sep)
            if len(parts) == 3:
                try:
                    y, m, d = int(parts[0]), int(parts[1]), int(parts[2])
                    if y < 100:
                        y, d = d, y
                    return jdatetime.date(y, m, d)
                except Exception:
                    return None
    return None


def jalali_month_name(month):
    names = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
             "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]
    try:
        return names[int(month) - 1]
    except Exception:
        return ""


def ask_security_password(parent):
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("SELECT value FROM settings WHERE key='app_password'")
    row = cursor.fetchone()
    saved_pass = row[0] if row else "123"
    conn.close()

    entered_pass, ok = QInputDialog.getText(
        parent, "تایید امنیتی", "لطفاً رمز عبور را وارد کنید:", QLineEdit.EchoMode.Password
    )
    if ok and entered_pass == saved_pass:
        return True
    elif ok:
        QMessageBox.critical(parent, "خطا", "رمز عبور نادرست است!")
        return False
    return False


def get_setting(key, default=""):
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("SELECT value FROM settings WHERE key=?", (key,))
    row = cursor.fetchone()
    conn.close()
    return row[0] if row else default


def set_setting(key, value):
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()
    cursor.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, ?)", (key, str(value)))
    conn.commit()
    conn.close()


def get_current_year():
    """سال کاری فعلی"""
    return int(get_setting("current_work_year", jdatetime.date.today().year))


# ---------------------------------------------------------------
# ۸) ابزارهای مهاجرت پایگاه داده و هماهنگ‌سازی انبار
# ---------------------------------------------------------------
def _column_exists(cursor, table, column):
    try:
        cursor.execute(f"PRAGMA table_info({table})")
        return any(r[1] == column for r in cursor.fetchall())
    except Exception:
        return False


def _add_column_if_missing(cursor, table, column, decl):
    if not _column_exists(cursor, table, column):
        try:
            cursor.execute(f"ALTER TABLE {table} ADD COLUMN {column} {decl}")
            print(f"[DB] Added column {table}.{column}")
        except Exception as e:
            print(f"[DB] add column failed {table}.{column}: {e}")


def _flag_is_set(cursor, key):
    try:
        cursor.execute("SELECT value FROM settings WHERE key=?", (key,))
        row = cursor.fetchone()
        return bool(row) and str(row[0]) == "1"
    except Exception:
        return False


def _set_flag(cursor, key):
    try:
        cursor.execute("INSERT OR REPLACE INTO settings (key, value) VALUES (?, '1')", (key,))
    except Exception:
        pass


def _dedupe_codes(cursor, table, prefix, width=4):
    """
    اگر در نسخه قبل به‌خاطر باگ COUNT+1 کد تکراری ساخته شده باشد،
    همه کدهای تکراری اصلاح می‌شوند تا «هیچ دو کدی یکسان» نباشد.
    """
    try:
        cursor.execute(f"SELECT id, {('code')} FROM {table} WHERE code IS NOT NULL AND code<>'' ORDER BY id")
        rows = cursor.fetchall()
    except Exception:
        return
    seen = {}
    max_n = 1000
    for _id, code in rows:
        try:
            n = int(str(code).split("-")[-1])
            if n > max_n:
                max_n = n
        except Exception:
            pass
    for _id, code in rows:
        if code in seen:
            max_n += 1
            new_code = f"{prefix}-{max_n:0{width}d}"
            while any(new_code == c for c in seen):
                max_n += 1
                new_code = f"{prefix}-{max_n:0{width}d}"
            try:
                cursor.execute(f"UPDATE {table} SET code=? WHERE id=?", (new_code, _id))
                seen[new_code] = True
            except Exception:
                pass
        else:
            seen[code] = True


def recalc_inventory_usage():
    """
    «استفاده شده» هر قلم انبار را از روی قراردادهای واقعی بازمحاسبه می‌کند.
    بنابراین ویرایش، حذف یا افزودن قرارداد بلافاصله در انبار هم منعکس می‌شود
    و انبار و فاکتور همیشه با هم هماهنگ می‌مانند.
    """
    try:
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        counter = {}
        cursor.execute("SELECT selected_items FROM wedding_contracts WHERE selected_items IS NOT NULL")
        for (items,) in cursor.fetchall():
            for it in str(items).split(','):
                it = it.strip()
                if it:
                    counter[it] = counter.get(it, 0) + 1
        cursor.execute("SELECT item_name FROM inventory")
        names = [r[0] for r in cursor.fetchall()]
        for name in names:
            used = counter.get(name, 0)
            cursor.execute("UPDATE inventory SET used_count=? WHERE item_name=?", (used, name))
        conn.commit()
        conn.close()
        return True
    except Exception as e:
        print("[INV] recalc_inventory_usage error:", e)
        return False


def sync_item_everywhere(old_name, new_name, code=None, price=None, is_package=None, parent=None):
    """
    هماهنگ‌سازی یک کالا/خدمت در همه بخش‌ها:
    item_prices، inventory، قراردادها و فاکتورها.
    """
    try:
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        if new_name and new_name != old_name:
            cursor.execute("UPDATE inventory SET item_name=? WHERE item_name=?", (new_name, old_name))
            cursor.execute("SELECT id, selected_items FROM wedding_contracts WHERE selected_items LIKE ?",
                           (f"%{old_name}%",))
            for cid, items in cursor.fetchall():
                lst = [x.strip() for x in str(items or "").split(',') if x.strip()]
                lst = [new_name if x == old_name else x for x in lst]
                cursor.execute("UPDATE wedding_contracts SET selected_items=? WHERE id=?",
                               (",".join(lst), cid))
            cursor.execute("UPDATE item_prices SET item_name=? WHERE item_name=?", (new_name, old_name))
            cursor.execute("SELECT item_name FROM item_prices WHERE parent_package=?", (old_name,))
            cursor.execute("UPDATE item_prices SET parent_package=? WHERE parent_package=?", (new_name, old_name))

        if code is not None:
            cursor.execute("UPDATE item_prices SET code=? WHERE item_name=?", (code, new_name or old_name))
        if price is not None:
            cursor.execute("UPDATE item_prices SET price=? WHERE item_name=?", (price, new_name or old_name))
        if is_package is not None:
            cursor.execute("UPDATE item_prices SET is_package=? WHERE item_name=?", (is_package, new_name or old_name))
        if parent is not None:
            cursor.execute("UPDATE item_prices SET parent_package=? WHERE item_name=?", (parent, new_name or old_name))

        conn.commit()
        conn.close()
        return True
    except Exception as e:
        print("[SYNC] sync_item_everywhere error:", e)
        return False


def sync_contract_staff(contract, work_year=None):
    """
    حقوق نیروهای انتخاب‌شده در یک قرارداد را در بخش کارکنان ثبت/هماهنگ می‌کند
    تا در «کارکنان» هم مراسم‌ها و هم حقوق دریافتی هر فرد دیده شود.
    """
    if not contract:
        return
    try:
        cid = contract.get("id")
        staff = contract.get("staff") or []
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM transactions WHERE contract_id=?", (cid,))
        for member in staff:
            pid = member.get("id")
            wage = int(member.get("wage") or 0)
            if not pid:
                continue
            d = parse_jalali(contract.get("ceremony_date")) or jdatetime.date.today()
            cursor.execute('''
                INSERT INTO transactions
                (trans_type, category, person_id, amount, year, month, day, description, work_year, contract_id)
                VALUES ('expense', 'حقوق قرارداد', ?, ?, ?, ?, ?, ?, ?, ?)
            ''', (pid, wage, d.year, d.month, d.day,
                  f"قرارداد {contract.get('code','')} - {contract.get('groom_name','')} و {contract.get('bride_name','')}",
                  work_year if work_year is not None else get_current_year(), cid))
        conn.commit()
        conn.close()
    except Exception as e:
        print("[STAFF] sync_contract_staff error:", e)


def init_db():
    conn = sqlite3.connect(DB_NAME)
    cursor = conn.cursor()

    cursor.execute('''CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT)''')
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('app_password', '123')")
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('studio_name', 'IMART STUDIO')")
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('studio_phone', '09173736618')")
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('studio_address', '')")
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('current_work_year', ?)", (str(jdatetime.date.today().year),))
    cursor.execute("INSERT OR IGNORE INTO settings (key, value) VALUES ('backup_dir', ?)", (os.path.join(APP_DIR, "backups"),))

    cursor.execute('''CREATE TABLE IF NOT EXISTS persons (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT,
        name TEXT NOT NULL,
        role TEXT NOT NULL,
        phone TEXT,
        work_year INTEGER
    )''')

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS transactions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            trans_type TEXT NOT NULL,
            category TEXT NOT NULL,
            person_id INTEGER,
            amount INTEGER NOT NULL,
            year INTEGER NOT NULL,
            month INTEGER NOT NULL,
            day INTEGER NOT NULL,
            description TEXT,
            work_year INTEGER,
            FOREIGN KEY (person_id) REFERENCES persons(id)
        )
    ''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS expenses (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        code TEXT, title TEXT, category TEXT, amount INTEGER,
        date_str TEXT, description TEXT, work_year INTEGER
    )''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS bank_cards (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        bank_name TEXT NOT NULL, card_number TEXT, sheba_number TEXT
    )''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS item_prices (
        item_name TEXT PRIMARY KEY,
        code TEXT, price INTEGER DEFAULT 0,
        is_package INTEGER DEFAULT 0, parent_package TEXT
    )''')

    default_items = [
        ("کلیپ فرمالیته", 15000000, 0, ""), ("کلیپ باغ", 8000000, 0, ""),
        ("یک دوربین", 5000000, 0, ""), ("دو دوربین", 9000000, 0, ""),
        ("کرین", 6000000, 0, ""), ("هلی شات", 7000000, 0, ""),
        ("عکاس مجلس", 4000000, 0, ""), ("پکیج طلایی VIP", 35000000, 1, "")
    ]
    # ⚠️ نسخه ۱۱: اقلام پیش‌فرض فقط در «اولین اجرا» ساخته می‌شوند.
    # در نسخه ۱۰ هر بار اجرای برنامه این اقلام دوباره درج می‌شدند و
    # اقلامی که کاربر حذف کرده بود، بعد از بستن و اجرای مجدد برمی‌گشتند.
    already_seeded = _flag_is_set(cursor, "seeded_defaults_v11")
    try:
        cursor.execute("SELECT COUNT(*) FROM item_prices")
        has_items = cursor.fetchone()[0] > 0
    except Exception:
        has_items = False

    if not already_seeded:
        if not has_items:
            srv_n = 1000
            pkg_n = 2000
            for idx, (item, price, is_pkg, parent) in enumerate(default_items):
                if is_pkg:
                    pkg_n += 1
                    code = f"PKG-{pkg_n}"
                else:
                    srv_n += 1
                    code = f"SRV-{srv_n}"
                cursor.execute(
                    "INSERT OR IGNORE INTO item_prices (item_name, code, price, is_package, parent_package) "
                    "VALUES (?, ?, ?, ?, ?)",
                    (item, code, price, is_pkg, parent))
        _set_flag(cursor, "seeded_defaults_v11")

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS wedding_contracts (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            code TEXT,
            groom_name TEXT, bride_name TEXT, groom_phone TEXT, bride_phone TEXT,
            contract_date TEXT, ceremony_date TEXT, selected_items TEXT,
            total_amount INTEGER, discount INTEGER, paid_amount INTEGER,
            bank_id INTEGER, is_settled INTEGER DEFAULT 0, description TEXT,
            work_year INTEGER
        )
    ''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS wedding_deposits (
        id INTEGER PRIMARY KEY AUTOINCREMENT, contract_id INTEGER, amount INTEGER,
        deposit_date TEXT, bank_name TEXT, work_year INTEGER,
        FOREIGN KEY (contract_id) REFERENCES wedding_contracts(id)
    )''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS commercial_projects (
        id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT, title TEXT, project_type TEXT,
        camera_count INTEGER, total_amount INTEGER, paid_amount INTEGER,
        project_date TEXT, description TEXT, client_name TEXT, client_phone TEXT,
        work_year INTEGER
    )''')

    cursor.execute('''CREATE TABLE IF NOT EXISTS inventory (
        item_name TEXT PRIMARY KEY, total_count INTEGER DEFAULT 0, used_count INTEGER DEFAULT 0
    )''')

    # موجودی اولیه هم فقط یک‌بار برای اقلام پیش‌فرض ساخته می‌شود
    if not already_seeded and not has_items:
        for item, _, _, _ in default_items:
            cursor.execute("INSERT OR IGNORE INTO inventory (item_name, total_count, used_count) "
                           "VALUES (?, 10, 0)", (item,))

    cursor.execute('''
        CREATE TABLE IF NOT EXISTS checks (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            check_number TEXT, bank_name TEXT, amount INTEGER,
            due_date TEXT, is_passed INTEGER DEFAULT 0, description TEXT,
            issuer_name TEXT, check_type TEXT DEFAULT 'دریافتی',
            work_year INTEGER
        )
    ''')

    # ============================================================
    # ==========  مهاجرت داده‌های نسخه ۹ به نسخه ۱۰  ==============
    # ============================================================
    _add_column_if_missing(cursor, "wedding_contracts", "staff_ids", "TEXT")
    _add_column_if_missing(cursor, "wedding_contracts", "items_detail", "TEXT")
    _add_column_if_missing(cursor, "transactions", "contract_id", "INTEGER")
    _add_column_if_missing(cursor, "checks", "contract_id", "INTEGER")
    _add_column_if_missing(cursor, "inventory", "category", "TEXT DEFAULT ''")
    _add_column_if_missing(cursor, "wedding_contracts", "venue", "TEXT")
    _add_column_if_missing(cursor, "inventory", "note", "TEXT DEFAULT ''")

    # ۰) پر کردن جزئیات قیمت هر قرارداد تا فاکتور دقیقاً همان مبالغ را نشان دهد
    if not _flag_is_set(cursor, "migration_v10_items_detail"):
        try:
            cursor.execute("SELECT id, selected_items FROM wedding_contracts")
            for cid, items in cursor.fetchall():
                if not items:
                    continue
                detail = {}
                for it in str(items).split(','):
                    it = it.strip()
                    if not it:
                        continue
                    cursor.execute("SELECT price FROM item_prices WHERE item_name=?", (it,))
                    r = cursor.fetchone()
                    detail[it] = (r[0] if r else 0)
                cursor.execute("UPDATE wedding_contracts SET items_detail=? WHERE id=?",
                               (json.dumps(detail, ensure_ascii=False), cid))
            _set_flag(cursor, "migration_v10_items_detail")
            print("[DB] Migration: items_detail filled from item_prices")
        except Exception as e:
            print("[DB] items_detail migration failed:", e)

    # ۱) واژه «صادراتی» چک به «پرداختی» تغییر می‌کند
    try:
        cursor.execute("UPDATE checks SET check_type='پرداختی' WHERE check_type='صادراتی'")
    except Exception:
        pass

    # ۲) جمع کل فاکتور باید «قیمت خام» باشد و تخفیف جداگانه نمایش داده شود.
    #    در نسخه ۹ مقدار total_amount بعد از کسر تخفیف ذخیره می‌شد و باعث
    #    دوباره کم شدن تخفیف در فاکتور می‌گردید. این اصلاح فقط یک‌بار انجام می‌شود.
    if not _flag_is_set(cursor, "migration_v10_raw_total"):
        try:
            cursor.execute("UPDATE wedding_contracts SET total_amount = COALESCE(total_amount,0) + COALESCE(discount,0)")
            _set_flag(cursor, "migration_v10_raw_total")
            print("[DB] Migration: wedding_contracts.total_amount converted to raw (before discount)")
        except Exception as e:
            print("[DB] raw total migration failed:", e)

    # ۳) ردیف بیعانه‌ها و پرداخت‌ها همیشه درست شماره‌گذاری شود (فقط نمایشی است)

    # ۴) کدهای تکراری نسخه قبل اصلاح شود
    _dedupe_codes(cursor, "wedding_contracts", "W")
    _dedupe_codes(cursor, "commercial_projects", "C")
    _dedupe_codes(cursor, "expenses", "E")
    _dedupe_codes(cursor, "persons", "P")
    _dedupe_codes(cursor, "item_prices", "SRV")

    conn.commit()
    conn.close()

    # ۵) ایندکس یگانه کدها (لایه دوم محافظت در سطح دیتابیس)
    for tbl, col in (("wedding_contracts", "code"), ("commercial_projects", "code"),
                     ("expenses", "code"), ("persons", "code")):
        ensure_unique_code_index(tbl, col)

    # ۶) هماهنگ‌سازی انبار با قراردادهای واقعی
    recalc_inventory_usage()


# ============================================================
# ======================  LoadingScreen  =====================
# ============================================================
# ============================================================
# ============  PrintChoiceDialog (چاپ / PDF)  ===============
# ============================================================
class PrintChoiceDialog(QDialog):
    """
    انتخاب نوع خروجی برای اسناد:
      • چاپ روی چاپگر (سیاه و سفید، مناسب چاپگر لیزری)  یا  ذخیره PDF رنگی
      • با ریز پکیج (زیرمجموعه‌ها چاپ شوند)  یا  بدون ریز پکیج
      • با نیروی کار  یا  بدون نیروی کار
    همه اسناد روی یک برگ A5 چاپ می‌شوند.
    """

    def __init__(self, parent=None, title="چاپ / خروجی PDF",
                 allow_details=True, allow_staff=False):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))
        self.setMinimumWidth(470)
        self.mode = "print"
        self.with_details = True
        self.with_staff = True
        persist_dialog(self, "dlg_print_choice")

        lay = QVBoxLayout(self)
        lay.setSpacing(9)

        head = QLabel(f"{ico('print')}  {title}")
        head.setProperty("heading", True)
        head.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lay.addWidget(head)

        size_hint = QLabel(f"{ico('pdf')} همه اسناد روی یک برگ <b>A5</b> چاپ می‌شوند.")
        size_hint.setAlignment(Qt.AlignmentFlag.AlignCenter)
        size_hint.setStyleSheet("background-color:#eaf6ff; border:1px solid #a9d3ee; "
                                "border-radius:8px; padding:6px;")
        lay.addWidget(size_hint)

        box1 = QGroupBox("نوع خروجی")
        v1 = QVBoxLayout()
        self.rb_print = QRadioButton(f"{ico('print')}  چاپ روی چاپگر (سیاه و سفید)")
        self.rb_pdf = QRadioButton(f"{ico('pdf')}  ذخیره فایل PDF (رنگی)")
        self.rb_print.setChecked(True)
        v1.addWidget(self.rb_print)
        v1.addWidget(self.rb_pdf)
        box1.setLayout(v1)
        lay.addWidget(box1)

        if allow_details:
            box2 = QGroupBox("میزان جزئیات")
            v2 = QVBoxLayout()
            self.rb_with = QRadioButton(f"{ico('package')}  با ریز پکیج (زیرمجموعه‌ها نمایش داده شوند)")
            self.rb_without = QRadioButton(f"{ico('list')}  بدون ریز پکیج (فقط سرتیترها)")
            self.rb_with.setChecked(True)
            v2.addWidget(self.rb_with)
            v2.addWidget(self.rb_without)
            box2.setLayout(v2)
            lay.addWidget(box2)

        if allow_staff:
            box3 = QGroupBox("نیروی کار")
            v3 = QVBoxLayout()
            self.rb_staff_yes = QRadioButton(f"{ico('staff')}  با نیروی کار (چاپ شود)")
            self.rb_staff_no = QRadioButton(f"{ico('staff')}  بدون نیروی کار (چاپ نشود)")
            self.rb_staff_yes.setChecked(True)
            v3.addWidget(self.rb_staff_yes)
            v3.addWidget(self.rb_staff_no)
            box3.setLayout(v3)
            lay.addWidget(box3)

        btn_row = QHBoxLayout()
        btn_ok = QPushButton(ico_text("ok", "تایید و ادامه"))
        btn_ok.setStyleSheet("background-color:#27ae60; color:white; font-weight:bold; padding:9px;")
        btn_ok.clicked.connect(self._accept)
        btn_cancel = QPushButton(ico_text("cancel", "انصراف"))
        btn_cancel.clicked.connect(self.reject)
        btn_row.addWidget(btn_ok, 2)
        btn_row.addWidget(btn_cancel, 1)
        lay.addLayout(btn_row)

    def _accept(self):
        self.mode = "pdf" if getattr(self, "rb_pdf", None) and self.rb_pdf.isChecked() else "print"
        if hasattr(self, "rb_with"):
            self.with_details = self.rb_with.isChecked()
        if hasattr(self, "rb_staff_yes"):
            self.with_staff = self.rb_staff_yes.isChecked()
        self.accept()


# ============================================================
# ======  PackageDetailsDialog (ویرایشگر کامل زیرمجموعه)  ====
# ============================================================
class PackageDetailsDialog(QDialog):
    """
    مدیریت کامل یک پکیج و زیرمجموعه‌هایش:
      • افزودن آیتم جدید با قیمت
      • افزودن از کالاها/خدمات موجود
      • حذف آیتم از پکیج یا حذف کامل کالا
      • تغییر نام و قیمت خود پکیج و همه آیتم‌ها
      • جمع قیمت زیرمجموعه‌ها به‌صورت خودکار قیمت کلی پکیج می‌شود
    """

    def __init__(self, item_name, parent=None):
        super().__init__(parent)
        self.old_name = item_name
        self.pkg_name = item_name
        self.setWindowTitle(f"مدیریت پکیج: {item_name}")
        self.resize(820, 700)
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))
        persist_dialog(self, "dlg_package")
        self._delete_list = []

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT code, price, is_package FROM item_prices WHERE item_name=?", (item_name,))
        row = cursor.fetchone() or ("", 0, 1)
        conn.close()

        self._code = row[0] or ""
        self._price = row[1] or 0

        lay = QVBoxLayout(self)
        lay.setSpacing(8)

        # ---------- مشخصات پکیج ----------
        pkg_box = QGroupBox("مشخصات پکیج")
        pkg_form = QFormLayout()
        self.txt_name = QLineEdit(item_name)
        self.txt_code = QLineEdit(self._code)
        self.chk_auto = QCheckBox("قیمت پکیج = جمع قیمت همه زیرمجموعه‌ها (خودکار)")
        self.chk_auto.setChecked(True)
        self.txt_price = QLineEdit(f"{self._price:,}")
        self.txt_price.textChanged.connect(lambda t: self.txt_price.setText(format_number(t)))
        self.txt_price.setEnabled(False)

        self.chk_auto.toggled.connect(lambda on: self.txt_price.setEnabled(not on))
        self.chk_auto.toggled.connect(lambda _: self.refresh_total())

        pkg_form.addRow(f"{ico('package')} نام پکیج:", self.txt_name)
        pkg_form.addRow(f"{ico('list')} کد پکیج (یکتا):", self.txt_code)
        pkg_form.addRow("", self.chk_auto)
        pkg_form.addRow(f"{ico('money')} قیمت پکیج (تومان):", self.txt_price)
        pkg_box.setLayout(pkg_form)
        lay.addWidget(pkg_box)

        # ---------- جدول زیرمجموعه‌ها ----------
        self.table = QTableWidget()
        self.table.setColumnCount(4)
        self.table.setHorizontalHeaderLabels(["ردیف", "کد", "نام آیتم / خدمت", "قیمت (تومان)"])
        self.table.setAlternatingRowColors(True)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.DoubleClicked |
                                   QAbstractItemView.EditTrigger.EditKeyPressed |
                                   QAbstractItemView.EditTrigger.AnyKeyPressed)
        self.table.itemChanged.connect(lambda _i: self.refresh_total())
        persist_table(self.table, "pkg_subs")
        lay.addWidget(self.table, 1)

        # ---------- دکمه‌های عملیات ----------
        row1 = QHBoxLayout()
        btn_add_new = QPushButton(ico_text("add", "افزودن آیتم جدید"))
        btn_add_new.setStyleSheet("background-color:#27ae60; color:white; font-weight:bold; padding:8px;")
        btn_add_new.clicked.connect(self.add_new_row)

        btn_add_existing = QPushButton(ico_text("link", "افزودن از کالاهای موجود"))
        btn_add_existing.setStyleSheet("background-color:#2980b9; color:white; font-weight:bold; padding:8px;")
        btn_add_existing.clicked.connect(self.add_existing_item)

        btn_detach = QPushButton(ico_text("cancel", "خارج کردن از پکیج"))
        btn_detach.setStyleSheet("background-color:#e67e22; color:white; font-weight:bold; padding:8px;")
        btn_detach.setToolTip("آیتم حذف نمی‌شود، فقط از این پکیج خارج می‌شود")
        btn_detach.clicked.connect(self.detach_row)

        btn_del = QPushButton(ico_text("delete", "حذف کامل کالا"))
        btn_del.setStyleSheet("background-color:#c0392b; color:white; font-weight:bold; padding:8px;")
        btn_del.setToolTip("کالا از کل برنامه حذف می‌شود")
        btn_del.clicked.connect(self.delete_row)

        for b in (btn_add_new, btn_add_existing, btn_detach, btn_del):
            row1.addWidget(b)
        lay.addLayout(row1)

        self.lbl_total = QLabel("")
        self.lbl_total.setWordWrap(True)
        self.lbl_total.setStyleSheet("background-color:#fff9e6; border:1px solid #f39c12; "
                                     "border-radius:8px; padding:10px; font-size:11pt;")
        lay.addWidget(self.lbl_total)

        row2 = QHBoxLayout()
        btn_save = QPushButton(ico_text("save", "ذخیره همه تغییرات"))
        btn_save.setStyleSheet("background-color:#16a085; color:white; font-weight:bold; padding:10px;")
        btn_save.clicked.connect(self.save_all)
        btn_print = QPushButton(ico_text("print", "چاپ / PDF جزئیات"))
        btn_print.clicked.connect(self.print_details)
        btn_close = QPushButton(ico_text("cancel", "بستن"))
        btn_close.clicked.connect(self.reject)
        row2.addWidget(btn_save, 2)
        row2.addWidget(btn_print, 2)
        row2.addWidget(btn_close, 1)
        lay.addLayout(row2)

        self.load_subs()

    # ---------------- بارگذاری و محاسبه ----------------
    def load_subs(self):
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT item_name, code, price FROM item_prices "
                       "WHERE parent_package=? ORDER BY item_name", (self.pkg_name,))
        subs = cursor.fetchall()
        conn.close()
        self.table.blockSignals(True)
        self.table.setRowCount(0)
        for i, (n, c, p) in enumerate(subs):
            self._insert_row(i + 1, c or "", n, p or 0)
        self.table.blockSignals(False)
        self.refresh_total()

    def _insert_row(self, idx, code, name, price):
        r = self.table.rowCount()
        self.table.insertRow(r)
        it_idx = QTableWidgetItem(str(idx))
        it_idx.setFlags(Qt.ItemFlag.ItemIsEnabled)
        it_idx.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        self.table.setItem(r, 0, it_idx)
        self.table.setItem(r, 1, QTableWidgetItem(code or ""))
        self.table.setItem(r, 2, QTableWidgetItem(name))
        it_price = QTableWidgetItem(f"{int(price or 0):,}")
        it_price.setTextAlignment(Qt.AlignmentFlag.AlignCenter)
        self.table.setItem(r, 3, it_price)

    def _renumber(self):
        for r in range(self.table.rowCount()):
            it = self.table.item(r, 0)
            if it:
                it.setText(str(r + 1))

    def current_subs(self):
        subs = []
        for r in range(self.table.rowCount()):
            name_it = self.table.item(r, 2)
            name = name_it.text().strip() if name_it else ""
            if not name:
                continue
            code_it = self.table.item(r, 1)
            price_it = self.table.item(r, 3)
            subs.append((name,
                         (code_it.text().strip() if code_it else ""),
                         parse_number(price_it.text() if price_it else "0")))
        return subs

    def refresh_total(self):
        subs = self.current_subs()
        total = sum(p for _, _, p in subs)
        if self.chk_auto.isChecked():
            self.txt_price.setText(f"{total:,}")
            price = total
        else:
            price = parse_number(self.txt_price.text())
        self.lbl_total.setText(
            f"<b>تعداد زیرمجموعه‌ها:</b> {len(subs)}  |  "
            f"<b>جمع قیمت زیرمجموعه‌ها:</b> {total:,} تومان<br>"
            f"<b>قیمت نهایی این پکیج:</b> {price:,} تومان  |  "
            f"<b>به حروف:</b> {number_to_persian_words(price)} تومان"
        )

    # ---------------- عملیات جدول ----------------
    def add_new_row(self):
        self.table.blockSignals(True)
        self._insert_row(self.table.rowCount() + 1, "", "", 0)
        self.table.blockSignals(False)
        r = self.table.rowCount() - 1
        self.table.setCurrentCell(r, 2)
        self.table.editItem(self.table.item(r, 2))
        self.refresh_total()

    def add_existing_item(self):
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT item_name FROM item_prices "
                       "WHERE (parent_package IS NULL OR parent_package='' OR parent_package<>?) "
                       "AND item_name<>? ORDER BY item_name", (self.pkg_name, self.pkg_name))
        names = [r[0] for r in cursor.fetchall()]
        conn.close()
        if not names:
            QMessageBox.information(self, "افزودن آیتم",
                                    "کالای آزاد دیگری برای افزودن وجود ندارد.\n"
                                    "ابتدا از بخش انبار یا مدیریت کالاها آیتم جدید بسازید.")
            return
        name, ok = QInputDialog.getItem(self, "افزودن از کالاهای موجود",
                                        "کالا / خدمت را انتخاب کنید:", names, 0, False)
        if not ok or not name:
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT code, price FROM item_prices WHERE item_name=?", (name,))
        r = cursor.fetchone()
        conn.close()
        code = r[0] if r else ""
        price = r[1] if r else 0
        self.table.blockSignals(True)
        self._insert_row(self.table.rowCount() + 1, code, name, price)
        self.table.blockSignals(False)
        self.refresh_total()

    def detach_row(self):
        r = self.table.currentRow()
        if r < 0:
            QMessageBox.warning(self, "خطا", "لطفاً یک ردیف را انتخاب کنید.")
            return
        self.table.removeRow(r)
        self._renumber()
        self.refresh_total()

    def delete_row(self):
        r = self.table.currentRow()
        if r < 0:
            QMessageBox.warning(self, "خطا", "لطفاً یک ردیف را انتخاب کنید.")
            return
        name_it = self.table.item(r, 2)
        if name_it and name_it.text().strip():
            self._delete_list.append(name_it.text().strip())
        self.table.removeRow(r)
        self._renumber()
        self.refresh_total()

    # ---------------- ذخیره ----------------
    def save_all(self):
        if not ask_security_password(self):
            return
        new_name = self.txt_name.text().strip()
        if not new_name:
            QMessageBox.warning(self, "خطا", "نام پکیج نمی‌تواند خالی باشد.")
            return

        subs = self.current_subs()
        subs_total = sum(p for _, _, p in subs)
        price = subs_total if self.chk_auto.isChecked() else parse_number(self.txt_price.text())
        code = self.txt_code.text().strip() or generate_unique_code(
            "PKG", "item_prices", "code", start=2001)

        # یکتا بودن کد
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT item_name FROM item_prices WHERE code=? AND item_name<>?",
                       (code, self.old_name))
        dup = cursor.fetchone()
        cursor.execute("SELECT item_name FROM item_prices WHERE parent_package=?",
                       (self.old_name,))
        old_subs = {r[0] for r in cursor.fetchall()}
        conn.close()
        if dup:
            QMessageBox.warning(self, "کد تکراری",
                                f"کد «{code}» قبلاً برای «{dup[0]}» ثبت شده است.")
            return

        # نام و قیمت خود پکیج
        if new_name != self.old_name:
            sync_item_everywhere(self.old_name, new_name, code, price, 1, "")
        else:
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("UPDATE item_prices SET code=?, price=?, is_package=1, parent_package='' "
                           "WHERE item_name=?", (code, price, self.old_name))
            conn.commit()
            conn.close()

        self.pkg_name = new_name
        self.old_name = new_name
        keep = {n for n, _, _ in subs}

        # زیرمجموعه‌هایی که از جدول حذف شده‌اند، از پکیج خارج می‌شوند
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        for name in old_subs - keep:
            cursor.execute("UPDATE item_prices SET parent_package='' WHERE item_name=?", (name,))
        # کالاهایی که کاربر «حذف کامل» را زده
        for name in self._delete_list:
            cursor.execute("DELETE FROM item_prices WHERE item_name=?", (name,))
            cursor.execute("DELETE FROM inventory WHERE item_name=?", (name,))
        self._delete_list = []

        # درج یا به‌روزرسانی زیرمجموعه‌ها
        for n, c, p in subs:
            if n == new_name:
                continue
            cursor.execute('''
                INSERT INTO item_prices (item_name, code, price, is_package, parent_package)
                VALUES (?, ?, ?, 0, ?)
                ON CONFLICT(item_name) DO UPDATE SET
                    code=excluded.code, price=excluded.price,
                    is_package=0, parent_package=excluded.parent_package
            ''', (n, c or generate_unique_code("SRV", "item_prices", "code", start=1001), p, new_name))
            cursor.execute("INSERT OR IGNORE INTO inventory (item_name, total_count, used_count) "
                           "VALUES (?, 0, 0)", (n,))

        # قیمت پکیج هم برابر جمع زیرمجموعه‌ها
        cursor.execute("UPDATE item_prices SET price=? WHERE item_name=?", (price, new_name))
        cursor.execute("INSERT OR IGNORE INTO inventory (item_name, total_count, used_count) "
                       "VALUES (?, 0, 0)", (new_name,))
        conn.commit()
        conn.close()

        recalc_inventory_usage()
        self.setWindowTitle(f"مدیریت پکیج: {new_name}")
        self.chk_auto.setChecked(True)
        self.load_subs()
        QMessageBox.information(
            self, "موفقیت",
            f"تغییرات پکیج «{new_name}» ذخیره شد.\n"
            f"قیمت پکیج برابر جمع {len(subs)} زیرمجموعه = {price:,} تومان شد.")

    # ---------------- چاپ و PDF ----------------
    def _html(self, mono=False):
        subs = self.current_subs()
        total = sum(p for _, _, p in subs)
        price = total if self.chk_auto.isChecked() else parse_number(self.txt_price.text())
        p = doc_palette(mono)
        studio_name, studio_phone, _a = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()

        rows = ""
        for i, (n, c, pr) in enumerate(subs, start=1):
            rows += InvoiceBuilder._row([
                InvoiceBuilder._td(str(i), size="8pt"),
                InvoiceBuilder._td(c or "-", size="8pt"),
                InvoiceBuilder._td(n, align="right", size="8pt"),
                InvoiceBuilder._td(f"{pr:,}", size="8pt"),
            ])
        if not rows:
            rows = f"<tr>{InvoiceBuilder._td('زیرمجموعه‌ای ثبت نشده', colspan=4, size='9pt')}</tr>"

        return f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {InvoiceBuilder._title_band(
              p, studio_name,
              "<div style='font-size:10pt;'>جزئیات پکیج</div>",
              [f"تاریخ: {now.strftime('%Y/%m/%d')}"])}
          <table width="100%" style="border-collapse:collapse; margin-top:5px;">
            {InvoiceBuilder._row([
                InvoiceBuilder._td("<b>نام پکیج:</b>", align="right", bg=p['box_bg'], size="9pt", width="22%"),
                InvoiceBuilder._td(self.txt_name.text(), align="right", size="9pt"),
                InvoiceBuilder._td("<b>کد:</b>", align="right", bg=p['box_bg'], size="9pt", width="14%"),
                InvoiceBuilder._td(self.txt_code.text() or "-", align="right", size="9pt"),
            ])}
          </table>
          <table width="100%" style="border-collapse:collapse; margin-top:5px;">
            <thead>{InvoiceBuilder._head(['ردیف', 'کد', 'عنوان زیرمجموعه', 'قیمت (تومان)'],
                                         bg=p['table_head'], fg=p['table_head_fg'])}</thead>
            <tbody>{rows}</tbody>
            <tfoot>
              {InvoiceBuilder._row([
                  InvoiceBuilder._td("<b>جمع قیمت زیرمجموعه‌ها</b>", align="right",
                                     colspan=3, bg=p['alt_row'], size="9pt", bold=True),
                  InvoiceBuilder._td(f"<b>{total:,}</b>", bg=p['alt_row'], size="9pt", bold=True),
              ])}
              {InvoiceBuilder._row([
                  InvoiceBuilder._td("<b>قیمت نهایی پکیج</b>", align="right",
                                     colspan=3, bg=p['total_bg'], size="10pt", bold=True),
                  InvoiceBuilder._td(f"<b>{price:,}</b>", bg=p['total_bg'], size="10pt", bold=True),
              ])}
            </tfoot>
          </table>
          <div style="margin-top:6px; font-size:8.5pt;">
            <b>قیمت پکیج به حروف:</b> {number_to_persian_words(price)} تومان
          </div>
          <div style="text-align:center; margin-top:9px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}
          </div>
        </div>
        """

    def print_details(self):
        dlg = PrintChoiceDialog(self, "چاپ / PDF جزئیات پکیج",
                                allow_details=False, allow_staff=False)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        if dlg.mode == "pdf":
            path, _ = QFileDialog.getSaveFileName(
                self, "ذخیره PDF",
                f"جزئیات_پکیج_{self.txt_name.text()}.pdf", "PDF Files (*.pdf)")
            if not path:
                return
            ok, err = save_html_pdf(self._html(mono=False), path, "جزئیات پکیج")
            if ok:
                QMessageBox.information(self, "موفقیت", f"فایل PDF ذخیره شد:\n{path}")
            else:
                QMessageBox.critical(self, "خطا", f"ساخت PDF ناموفق بود:\n{err}")
        else:
            print_html_document(self._html(mono=True), self, "جزئیات پکیج")


# ============================================================
# ================  ChartDialog (نمایش نمودار)  ==============
# ============================================================
class ChartDialog(QDialog):
    """نمایش نمودار با عنوان فارسی درست و امکان چاپ سیاه‌وسفید / ذخیره رنگی"""

    def __init__(self, figure, title, parent=None):
        super().__init__(parent)
        self.figure = figure
        self.chart_title = title
        self.setWindowTitle(title)
        self.resize(980, 660)
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))
        persist_dialog(self, "dlg_chart")

        lay = QVBoxLayout(self)
        head = QLabel(f"{ico('chart')}  {title}")
        head.setProperty("heading", True)
        head.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lay.addWidget(head)

        self.canvas = FigureCanvas(figure)
        lay.addWidget(self.canvas, 1)

        row = QHBoxLayout()
        btn_png = QPushButton(ico_text("save", "ذخیره تصویر"))
        btn_png.setStyleSheet("background-color:#16a085; color:white; font-weight:bold; padding:8px;")
        btn_png.clicked.connect(self.save_png)

        btn_pdf = QPushButton(ico_text("pdf", "ذخیره PDF"))
        btn_pdf.setStyleSheet("background-color:#c0392b; color:white; font-weight:bold; padding:8px;")
        btn_pdf.clicked.connect(self.save_pdf)

        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color:#1e8449; color:white; font-weight:bold; padding:8px;")
        btn_excel.clicked.connect(self.save_excel)

        btn_print = QPushButton(ico_text("print", "چاپ (سیاه و سفید)"))
        btn_print.setStyleSheet("background-color:#2980b9; color:white; font-weight:bold; padding:8px;")
        btn_print.clicked.connect(self.print_chart)

        btn_close = QPushButton(ico_text("cancel", "بستن"))
        btn_close.clicked.connect(self.accept)

        for b in (btn_png, btn_pdf, btn_excel, btn_print):
            row.addWidget(b)
        row.addStretch()
        row.addWidget(btn_close)
        lay.addLayout(row)

    def _safe_name(self):
        return "".join(ch for ch in self.chart_title if ch not in '\\/:*?"<>|').strip() or "نمودار"

    def save_png(self):
        path, _ = QFileDialog.getSaveFileName(self, "ذخیره تصویر نمودار",
                                              f"{self._safe_name()}.png", "PNG (*.png)")
        if not path:
            return
        try:
            self.figure.savefig(path, dpi=170, bbox_inches="tight",
                                facecolor=self.figure.get_facecolor())
            QMessageBox.information(self, "موفقیت", "تصویر نمودار ذخیره شد.")
        except Exception as e:
            QMessageBox.critical(self, "خطا", str(e))

    def save_pdf(self):
        path, _ = QFileDialog.getSaveFileName(self, "ذخیره PDF نمودار",
                                              f"{self._safe_name()}.pdf", "PDF (*.pdf)")
        if not path:
            return
        try:
            self.figure.savefig(path, format="pdf", bbox_inches="tight",
                                facecolor=self.figure.get_facecolor())
            QMessageBox.information(self, "موفقیت", "PDF نمودار ذخیره شد.")
        except Exception as e:
            QMessageBox.critical(self, "خطا", str(e))

    def save_excel(self):
        path, _ = QFileDialog.getSaveFileName(self, "خروجی اکسل نمودار",
                                              f"{self._safe_name()}.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        headers = ["عنوان", "مقدار"]
        rows = []
        for ax in self.figure.axes:
            for patch, label in zip(ax.patches, ax.get_xticklabels()):
                try:
                    rows.append([label.get_text(), int(patch.get_height())])
                except Exception:
                    pass
        ok, err = export_rows_to_excel(headers, rows, path, "نمودار", self.chart_title)
        if ok:
            QMessageBox.information(self, "موفقیت", "خروجی اکسل ذخیره شد.")
        else:
            QMessageBox.critical(self, "خطا", err)

    def _png_bytes(self, gray=False, dpi=150):
        import io as _io
        buf = _io.BytesIO()
        self.figure.savefig(buf, format="png", dpi=dpi, bbox_inches="tight",
                            facecolor="white" if gray else self.figure.get_facecolor())
        buf.seek(0)
        data = buf.read()
        if gray:
            try:
                img = QImage.fromData(data, "PNG")
                if not img.isNull():
                    img = img.convertToFormat(QImage.Format.Format_Grayscale8)
                    b2 = QByteArray()
                    from PyQt6.QtCore import QBuffer
                    qb = QBuffer(b2)
                    qb.open(QBuffer.OpenModeFlag.WriteOnly)
                    img.save(qb, "PNG")
                    qb.close()
                    data = bytes(b2)
            except Exception as e:
                print("[CHART] gray convert failed:", e)
        return data

    def _html_with_image(self, gray=False):
        import base64
        b64 = base64.b64encode(self._png_bytes(gray=gray)).decode("ascii")
        return (f"<div dir='rtl' style=\"font-family:'{INVOICE_FONT_FAMILY}', Tahoma;\">"
                f"<h2 style='text-align:center; color:#1F4E78;'>{self.chart_title}</h2>"
                f"<img src='data:image/png;base64,{b64}' style='width:100%;'/>"
                f"</div>")

    def print_chart(self):
        # پرینت سیاه و سفید می‌شود تا روی چاپگر لیزری تمیز باشد
        print_html_document(self._html_with_image(gray=True), self, self.chart_title)
# ============================================================
# ======  SettlementStatusDialog (وضعیت تسویه پروژه‌ها)  ======
# ============================================================
class SettlementStatusDialog(QDialog):
    """نمایش لیست پروژه‌ها و قراردادها بر اساس وضعیت تسویه"""

    KIND_TITLES = {
        "all": "کل پروژه‌ها و قراردادها",
        "full": "تسویه کامل",
        "partial": "تسویه ناقص",
        "none": "تسویه نشده",
    }

    def __init__(self, kind, parent=None):
        super().__init__(parent)
        title = self.KIND_TITLES.get(kind, "وضعیت تسویه")
        self.setWindowTitle(f"{title} - سال کاری {get_current_year()}")
        self.resize(980, 620)
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))
        persist_dialog(self, "dlg_settlement")
        self.kind = kind

        work_year = get_current_year()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        items = []
        cursor.execute("""SELECT id, code, groom_name, bride_name, ceremony_date,
                          total_amount, discount, paid_amount FROM wedding_contracts
                          WHERE work_year=?""", (work_year,))
        for r in cursor.fetchall():
            net = max(0, (r[5] or 0) - (r[6] or 0))
            state, remain = settlement_state(net, r[7])
            items.append(("قرارداد عروسی", r[1], f"{r[2]} و {r[3]}", r[4], net, r[7] or 0, remain, state))
        cursor.execute("""SELECT id, code, title, client_name, project_date,
                          total_amount, paid_amount FROM commercial_projects
                          WHERE work_year=?""", (work_year,))
        for r in cursor.fetchall():
            net = r[5] or 0
            state, remain = settlement_state(net, r[6])
            items.append(("پروژه تبلیغاتی", r[1], f"{r[2]} - {r[3] or '-'}", r[4], net, r[6] or 0, remain, state))
        conn.close()

        if kind == "full":
            items = [i for i in items if i[7] == "تسویه کامل"]
        elif kind == "partial":
            items = [i for i in items if i[7] == "تسویه ناقص"]
        elif kind == "none":
            items = [i for i in items if i[7] == "تسویه نشده"]

        lay = QVBoxLayout(self)
        self.lbl_sum = QLabel("")
        self.lbl_sum.setWordWrap(True)
        self.lbl_sum.setStyleSheet("background-color:#eaf6ff; border:1px solid #a9d3ee; "
                                   "border-radius:8px; padding:10px; font-size:11pt;")
        lay.addWidget(self.lbl_sum)

        self.table = QTableWidget()
        self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels([
            "نوع", "کد", "عنوان / نام", "تاریخ", "مبلغ نهایی", "پرداختی", "مانده", "وضعیت"
        ])
        self.table.setAlternatingRowColors(True)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.verticalHeader().setDefaultSectionSize(30)
        persist_table(self.table, f"settle_{kind}", QHeaderView.ResizeMode.Stretch)
        self.table.setRowCount(len(items))
        total_net = total_paid = total_remain = 0
        for i, it in enumerate(items):
            total_net += it[4]
            total_paid += it[5]
            total_remain += it[6]
            self.table.setItem(i, 0, QTableWidgetItem(it[0]))
            self.table.setItem(i, 1, QTableWidgetItem(str(it[1] or "-")))
            self.table.setItem(i, 2, QTableWidgetItem(str(it[2])))
            self.table.setItem(i, 3, QTableWidgetItem(str(it[3] or "-")))
            self.table.setItem(i, 4, QTableWidgetItem(f"{it[4]:,}"))
            self.table.setItem(i, 5, QTableWidgetItem(f"{it[5]:,}"))
            remain_item = QTableWidgetItem(f"{it[6]:,}")
            remain_item.setForeground(QColor("#c0392b" if it[6] > 0 else "#27ae60"))
            self.table.setItem(i, 6, remain_item)
            st_item = QTableWidgetItem(it[7])
            st_item.setForeground(QColor(
                "#27ae60" if it[7] == "تسویه کامل"
                else ("#e67e22" if it[7] == "تسویه ناقص" else "#c0392b")))
            f = st_item.font()
            f.setBold(True)
            st_item.setFont(f)
            self.table.setItem(i, 7, st_item)
        lay.addWidget(self.table)

        self.lbl_sum.setText(
            f"<b>{title}</b> - سال کاری {work_year} | "
            f"<b>تعداد:</b> {len(items)} | "
            f"<b>جمع مبلغ نهایی:</b> {total_net:,} تومان | "
            f"<b>جمع پرداختی:</b> {total_paid:,} تومان | "
            f"<b style='color:#c0392b;'>جمع مانده:</b> {total_remain:,} تومان"
        )

        row = QHBoxLayout()
        btn_pdf = QPushButton(ico_text("pdf", "PDF / چاپ"))
        btn_pdf.setStyleSheet("background-color:#2980b9; color:white; font-weight:bold; padding:8px;")
        btn_pdf.clicked.connect(lambda: self.output(title, items, total_net, total_paid, total_remain))
        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color:#1e8449; color:white; font-weight:bold; padding:8px;")
        btn_excel.clicked.connect(lambda: self.export_excel(title))
        btn_close = QPushButton(ico_text("ok", "بستن"))
        btn_close.clicked.connect(self.accept)
        row.addWidget(btn_pdf)
        row.addWidget(btn_excel)
        row.addStretch()
        row.addWidget(btn_close)
        lay.addLayout(row)

    def export_excel(self, title):
        path, _ = QFileDialog.getSaveFileName(self, "خروجی اکسل", f"{title}.xlsx", "Excel (*.xlsx)")
        if not path:
            return
        headers = [self.table.horizontalHeaderItem(c).text() for c in range(self.table.columnCount())]
        rows = [[self.table.item(r, c).text() if self.table.item(r, c) else ""
                 for c in range(self.table.columnCount())] for r in range(self.table.rowCount())]
        ok, err = export_rows_to_excel(headers, rows, path, title, title)
        if ok:
            QMessageBox.information(self, "موفقیت", f"خروجی اکسل ذخیره شد:\n{path}")
        else:
            QMessageBox.critical(self, "خطا", err)

    def output(self, title, items, total_net, total_paid, total_remain):
        dlg = PrintChoiceDialog(self, title, allow_details=False, allow_staff=False)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        mono = (dlg.mode == "print")
        html = self._html(title, items, total_net, total_paid, total_remain, mono)
        if mono:
            print_html_document(html, self, title)
        else:
            path, _ = QFileDialog.getSaveFileName(self, "ذخیره PDF", f"{title}.pdf", "PDF Files (*.pdf)")
            if not path:
                return
            ok, err = save_html_pdf(html, path, title)
            if ok:
                QMessageBox.information(self, "موفقیت", f"فایل PDF ذخیره شد:\n{path}")
            else:
                QMessageBox.critical(self, "خطا", err)

    def _html(self, title, items, total_net, total_paid, total_remain, mono=False):
        p = doc_palette(mono)
        studio_name, studio_phone, _a = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()
        rows = ""
        for i, it in enumerate(items, start=1):
            rows += InvoiceBuilder._row([
                InvoiceBuilder._td(str(i), size="8pt"),
                InvoiceBuilder._td(str(it[1] or "-"), size="8pt"),
                InvoiceBuilder._td(str(it[2]), align="right", size="8pt"),
                InvoiceBuilder._td(str(it[3] or "-"), size="8pt"),
                InvoiceBuilder._td(f"{it[4]:,}", size="8pt"),
                InvoiceBuilder._td(f"{it[5]:,}", size="8pt"),
                InvoiceBuilder._td(f"{it[6]:,}", size="8pt"),
                InvoiceBuilder._td(it[7], size="8pt"),
            ])
        if not rows:
            rows = f"<tr>{InvoiceBuilder._td('موردی در این دسته وجود ندارد', colspan=8, size='9pt')}</tr>"

        return f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {InvoiceBuilder._title_band(
              p, studio_name,
              f"<div style='font-size:10pt;'>{title} - سال کاری {get_current_year()}</div>",
              [f"تاریخ: {now.strftime('%Y/%m/%d')}", f"تعداد موارد: {len(items)}"])}
          <table width="100%" style="border-collapse:collapse; margin-top:6px;">
            <thead>{InvoiceBuilder._head(['ردیف', 'کد', 'عنوان / نام', 'تاریخ', 'مبلغ نهایی',
                                          'پرداختی', 'مانده', 'وضعیت'],
                                         bg=p['table_head'], fg=p['table_head_fg'])}</thead>
            <tbody>{rows}</tbody>
            <tfoot>
              {InvoiceBuilder._row([
                  InvoiceBuilder._td('<b>جمع کل</b>', align='right', colspan=4,
                                     bg=p['total_bg'], size='9pt', bold=True),
                  InvoiceBuilder._td(f"<b>{total_net:,}</b>", bg=p['total_bg'], size='9pt', bold=True),
                  InvoiceBuilder._td(f"<b>{total_paid:,}</b>", bg=p['total_bg'], size='9pt', bold=True),
                  InvoiceBuilder._td(f"<b>{total_remain:,}</b>", bg=p['total_bg'], size='9pt', bold=True),
                  InvoiceBuilder._td("", bg=p['total_bg'], size='9pt'),
              ])}
            </tfoot>
          </table>
          <div style="text-align:center; margin-top:8px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}
          </div>
        </div>
        """


class LoadingScreen(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("IMART STUDIO - در حال بارگذاری...")
        self.setFixedSize(550, 340)
        self.setWindowFlags(Qt.WindowType.FramelessWindowHint | Qt.WindowType.WindowStaysOnTopHint)
        self.setAttribute(Qt.WidgetAttribute.WA_TranslucentBackground)

        self.container = QFrame(self)
        self.container.setGeometry(0, 0, 550, 340)
        self.container.setStyleSheet("""
            QFrame {
                background-color: qlineargradient(x1:0, y1:0, x2:1, y2:1,
                    stop:0 #1F4E78, stop:0.5 #2c3e50, stop:1 #34495e);
                border-radius: 18px;
                border: 2px solid #2980b9;
            }
        """)

        shadow = QGraphicsDropShadowEffect()
        shadow.setBlurRadius(30)
        shadow.setColor(QColor(0, 0, 0, 180))
        shadow.setOffset(0, 6)
        self.container.setGraphicsEffect(shadow)

        layout = QVBoxLayout(self.container)
        layout.setContentsMargins(35, 30, 35, 30)
        layout.setSpacing(15)

        self.lbl_logo = QLabel("🎬 IMART STUDIO")
        self.lbl_logo.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_logo.setStyleSheet("color: white; font-size: 26pt; font-weight: bold; background: transparent;")
        layout.addWidget(self.lbl_logo)

        self.lbl_sub = QLabel(f"سیستم جامع مدیریت مالی و حسابداری - نسخه {APP_VERSION}")
        self.lbl_sub.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_sub.setStyleSheet("color: #aed6f1; font-size: 11pt; background: transparent;")
        layout.addWidget(self.lbl_sub)

        layout.addSpacing(10)

        self.lbl_percent = QLabel("0%")
        self.lbl_percent.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_percent.setStyleSheet("color: #f1c40f; font-size: 32pt; font-weight: bold; background: transparent;")
        layout.addWidget(self.lbl_percent)

        self.progress = QProgressBar()
        self.progress.setFixedHeight(20)
        self.progress.setTextVisible(False)
        self.progress.setRange(0, 100)
        self.progress.setValue(0)
        self.progress.setStyleSheet("""
            QProgressBar {
                background-color: rgba(255, 255, 255, 30);
                border-radius: 10px;
                border: 1px solid rgba(255, 255, 255, 60);
            }
            QProgressBar::chunk {
                background-color: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                    stop:0 #27ae60, stop:0.5 #2ecc71, stop:1 #1abc9c);
                border-radius: 10px;
            }
        """)
        layout.addWidget(self.progress)

        self.lbl_status = QLabel("در حال آماده‌سازی...")
        self.lbl_status.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_status.setStyleSheet("color: white; font-size: 11pt; background: transparent;")
        layout.addWidget(self.lbl_status)

        layout.addStretch()

        self.lbl_footer = QLabel("📞 09173736618")
        self.lbl_footer.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_footer.setStyleSheet("color: #bdc3c7; font-size: 9pt; background: transparent;")
        layout.addWidget(self.lbl_footer)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.addWidget(self.container)

        self.center_on_screen()

    def center_on_screen(self):
        screen = QApplication.primaryScreen().geometry()
        x = (screen.width() - self.width()) // 2
        y = (screen.height() - self.height()) // 2
        self.move(x, y)

    def set_progress(self, value, status_text=""):
        self.progress.setValue(value)
        self.lbl_percent.setText(f"{value}%")
        if status_text:
            self.lbl_status.setText(status_text)
        QApplication.processEvents()


# ============================================================
# =============  ManageItemsDialog (پکیج‌ها)  ================
# ============================================================
class ManageItemsDialog(QDialog):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("مدیریت پکیج‌ها، زیرمجموعه‌ها، کالاها و انبار")
        self.resize(880, 720)
        persist_dialog(self, "dlg_manage_items")
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))

        layout = QVBoxLayout()
        form_layout = QFormLayout()

        self.txt_item_code = QLineEdit()
        self.txt_item_code.setPlaceholderText("خالی بگذارید تا کد یکتا خودکار ساخته شود")
        self.txt_item_name = QLineEdit()
        self.txt_item_price = QLineEdit()
        self.txt_item_price.textChanged.connect(lambda t: self.txt_item_price.setText(format_number(t)))

        self.txt_item_category = QLineEdit()
        self.txt_item_category.setPlaceholderText("مثلاً: دوربین، لنز، نور، خدمات")
        self.spin_total_count = QSpinBox()
        self.spin_total_count.setRange(0, 100000)
        self.spin_total_count.setValue(10)
        self.spin_total_count.setToolTip("تعداد کل موجودی این قلم در انبار")
        self.txt_item_note = QLineEdit()
        self.txt_item_note.setPlaceholderText("مشخصات فنی، سریال، توضیح انبار...")

        self.chk_is_package = QCheckBox("این مورد یک پکیج اصلی است")
        self.chk_is_package.stateChanged.connect(self.toggle_package_mode)

        self.combo_parent_package = QComboBox()
        self.load_packages_combo()

        form_layout.addRow(f"{ico('add')} کد کالا / خدمات:", self.txt_item_code)
        form_layout.addRow(f"{ico('list')} نام مورد / پکیج:", self.txt_item_name)
        form_layout.addRow(f"{ico('money')} قیمت (تومان):", self.txt_item_price)
        form_layout.addRow(f"{ico('inventory')} دسته‌بندی انبار:", self.txt_item_category)
        form_layout.addRow(f"{ico('inventory')} تعداد کل در انبار:", self.spin_total_count)
        form_layout.addRow(f"{ico('detail')} مشخصات / توضیحات:", self.txt_item_note)
        form_layout.addRow("", self.chk_is_package)
        form_layout.addRow(f"{ico('package')} پکیج مادر (برای زیرمجموعه):", self.combo_parent_package)

        btn_row = QHBoxLayout()
        btn_add = QPushButton(ico_text("save", "افزودن / بروزرسانی"))
        btn_add.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 8px;")
        btn_add.clicked.connect(self.add_or_update_item)

        btn_auto_code = QPushButton(ico_text("refresh", "تولید کد یکتای خودکار"))
        btn_auto_code.setStyleSheet("background-color: #2980b9; color: white; font-weight: bold; padding: 8px;")
        btn_auto_code.clicked.connect(self.generate_auto_code)

        btn_row.addWidget(btn_add, 2)
        btn_row.addWidget(btn_auto_code, 1)
        form_layout.addRow(btn_row)
        layout.addLayout(form_layout)

        lbl_hint = QLabel(
            f"{ico('detail')} برای دیدن <b>جزئیات کامل یک پکیج</b> روی آن دوبار کلیک کنید."
        )
        lbl_hint.setStyleSheet("background-color:#eaf6ff; border:1px solid #a9d3ee; "
                               "border-radius:8px; padding:8px;")
        layout.addWidget(lbl_hint)

        self.tree = QTreeWidget()
        self.tree.setHeaderLabels(["کد", "نام آیتم / زیرمجموعه", "قیمت (تومان)", "موجودی انبار", "نوع"])
        self.tree.setColumnWidth(0, 110)
        self.tree.setColumnWidth(2, 130)
        self.tree.setColumnWidth(3, 100)
        self.tree.setAlternatingRowColors(True)
        self.tree.itemDoubleClicked.connect(self.on_tree_double_click)
        layout.addWidget(self.tree)

        bottom_row = QHBoxLayout()
        btn_details = QPushButton(ico_text("detail", "جزئیات پکیج انتخاب‌شده"))
        btn_details.setStyleSheet("background-color: #16a085; color: white; font-weight: bold; padding: 8px;")
        btn_details.clicked.connect(self.show_package_details)

        btn_edit = QPushButton(ico_text("edit", "ویرایش (کد / نام / قیمت / تعداد)"))
        btn_edit.setStyleSheet("background-color: #f39c12; color: white; font-weight: bold; padding: 8px;")
        btn_edit.clicked.connect(self.edit_selected_item)

        btn_delete = QPushButton(ico_text("delete", "حذف مورد انتخاب‌شده"))
        btn_delete.setStyleSheet("background-color: #c0392b; color: white; font-weight: bold; padding: 8px;")
        btn_delete.clicked.connect(self.delete_item)

        bottom_row.addWidget(btn_details, 2)
        bottom_row.addWidget(btn_edit, 2)
        bottom_row.addWidget(btn_delete, 1)
        layout.addLayout(bottom_row)

        self.load_tree_items()
        self.setLayout(layout)

    def generate_auto_code(self):
        """تولید کد یکتا — هرگز با کد هیچ کالا/خدمت دیگری یکسان نمی‌شود"""
        prefix = "PKG" if self.chk_is_package.isChecked() else "SRV"
        start = 2001 if prefix == "PKG" else 1001
        self.txt_item_code.setText(generate_unique_code(prefix, "item_prices", "code", start=start))

    def toggle_package_mode(self, state):
        self.combo_parent_package.setEnabled(not self.chk_is_package.isChecked())
        if self.chk_is_package.isChecked():
            self.combo_parent_package.setCurrentIndex(0)
        if not self.txt_item_code.text().strip():
            self.generate_auto_code()

    def load_packages_combo(self):
        self.combo_parent_package.clear()
        self.combo_parent_package.addItem(f"{ico('list')} --- بدون پکیج مادر ---", "")
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT item_name FROM item_prices WHERE is_package=1 ORDER BY item_name")
        for (pkg_name,) in cursor.fetchall():
            self.combo_parent_package.addItem(f"{ico('package')} {pkg_name}", pkg_name)
        conn.close()

    def load_tree_items(self):
        self.tree.clear()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT item_name, code, price, is_package, parent_package FROM item_prices")
        rows = cursor.fetchall()
        cursor.execute("SELECT item_name, total_count, used_count FROM inventory")
        inv = {r[0]: (r[1], r[2]) for r in cursor.fetchall()}
        conn.close()

        packages = [r for r in rows if r[3] == 1]
        singles = [r for r in rows if r[3] == 0]

        def stock_text(name):
            total, used = inv.get(name, (0, 0))
            remain = max(0, total - used)
            return f"{remain} از {total}"

        for pkg in packages:
            pkg_item = QTreeWidgetItem(self.tree, [
                pkg[1] or "-", pkg[0], f"{pkg[2]:,}", stock_text(pkg[0]), "پکیج اصلی"
            ])
            pkg_item.setFont(1, QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
            pkg_item.setForeground(1, QColor("#c0392b"))
            pkg_item.setData(0, Qt.ItemDataRole.UserRole, pkg[0])

            sub_items = [s for s in singles if s[4] == pkg[0]]
            for sub in sub_items:
                child = QTreeWidgetItem(pkg_item, [
                    sub[1] or "-", f"  └ {sub[0]}", f"{sub[2]:,}", stock_text(sub[0]), "زیرمجموعه"
                ])
                child.setData(0, Qt.ItemDataRole.UserRole, sub[0])

        independent = [s for s in singles if not s[4]]
        for ind in independent:
            node = QTreeWidgetItem(self.tree, [
                ind[1] or "-", ind[0], f"{ind[2]:,}", stock_text(ind[0]), "مستقل"
            ])
            node.setData(0, Qt.ItemDataRole.UserRole, ind[0])

        self.tree.expandAll()

    def _selected_item_name(self):
        selected = self.tree.currentItem()
        if not selected:
            return None
        return selected.data(0, Qt.ItemDataRole.UserRole) or selected.text(1).replace("  └ ", "").strip()

    def on_tree_double_click(self, item, column):
        name = item.data(0, Qt.ItemDataRole.UserRole) or item.text(1).replace("  └ ", "").strip()
        if not name:
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT is_package FROM item_prices WHERE item_name=?", (name,))
        r = cursor.fetchone()
        conn.close()
        if r and r[0] == 1:
            dlg = PackageDetailsDialog(name, self)
            dlg.exec()
        else:
            self.show_package_details()

    def show_package_details(self):
        """نمایش جزئیات کامل پکیج انتخاب‌شده با یک کلیک"""
        name = self._selected_item_name()
        if not name:
            QMessageBox.warning(self, "خطا", "لطفاً یک پکیج یا مورد را انتخاب کنید.")
            return
        dlg = PackageDetailsDialog(name, self)
        dlg.exec()

    def add_or_update_item(self):
        if not ask_security_password(self):
            return

        code = self.txt_item_code.text().strip()
        name = self.txt_item_name.text().strip()
        price = parse_number(self.txt_item_price.text())
        is_pkg = 1 if self.chk_is_package.isChecked() else 0
        parent = self.combo_parent_package.currentData() if not is_pkg else ""
        category = self.txt_item_category.text().strip()
        total_count = self.spin_total_count.value()
        note = self.txt_item_note.text().strip()

        if not name:
            QMessageBox.warning(self, "خطا", "لطفاً نام را وارد کنید.")
            return

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        # ---- تضمین یکتا بودن کد (هیچ دو کالا/خدمتی کد یکسان ندارند) ----
        if not code:
            conn.close()
            prefix = "PKG" if is_pkg else "SRV"
            code = generate_unique_code(prefix, "item_prices", "code",
                                        start=2001 if is_pkg else 1001)
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
        else:
            cursor.execute("SELECT item_name FROM item_prices WHERE code=? AND item_name<>?", (code, name))
            other = cursor.fetchone()
            if other:
                conn.close()
                QMessageBox.warning(
                    self, "کد تکراری",
                    f"کد «{code}» قبلاً برای «{other[0]}» ثبت شده است.\n"
                    "هر کالا و خدمت باید کد یکتا داشته باشد. لطفاً کد دیگری وارد کنید."
                )
                return

        cursor.execute('''
            INSERT INTO item_prices (item_name, code, price, is_package, parent_package)
            VALUES (?, ?, ?, ?, ?)
            ON CONFLICT(item_name) DO UPDATE SET
                code=excluded.code, price=excluded.price,
                is_package=excluded.is_package, parent_package=excluded.parent_package
        ''', (name, code, price, is_pkg, parent))

        # ---- هماهنگ‌سازی با انبار (تعداد کل، دسته‌بندی و مشخصات) ----
        cursor.execute('''
            INSERT INTO inventory (item_name, total_count, used_count, category, note)
            VALUES (?, ?, 0, ?, ?)
            ON CONFLICT(item_name) DO UPDATE SET
                total_count=excluded.total_count,
                category=excluded.category,
                note=excluded.note
        ''', (name, total_count, category, note))

        conn.commit()
        conn.close()

        recalc_inventory_usage()

        self.txt_item_code.clear()
        self.txt_item_name.clear()
        self.txt_item_price.clear()
        self.txt_item_category.clear()
        self.txt_item_note.clear()
        self.chk_is_package.setChecked(False)
        self.load_packages_combo()
        self.load_tree_items()

    def edit_selected_item(self):
        name = self._selected_item_name()
        if not name:
            QMessageBox.warning(self, "خطا", "لطفاً یک آیتم را انتخاب کنید.")
            return
        if not ask_security_password(self):
            return
        dlg = ItemEditDialog(name, self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            self.load_packages_combo()
            self.load_tree_items()

    def delete_item(self):
        name = self._selected_item_name()
        if not name:
            QMessageBox.warning(self, "خطا", "لطفاً یک آیتم را انتخاب کنید.")
            return

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT selected_items FROM wedding_contracts WHERE selected_items LIKE ?", (f"%{name}%",))
        used_in = len(cursor.fetchall())
        cursor.execute("SELECT COUNT(*) FROM item_prices WHERE parent_package=?", (name,))
        child_count = cursor.fetchone()[0]
        conn.close()

        warn = f"آیا از حذف «{name}» مطمئن هستید؟"
        if used_in:
            warn += f"\n\n⚠️ این مورد در {used_in} قرارداد استفاده شده است."
        if child_count:
            warn += f"\n\n⚠️ این پکیج {child_count} زیرمجموعه دارد؛ زیرمجموعه‌ها مستقل باقی می‌مانند."

        if QMessageBox.question(self, "تایید حذف", warn,
                                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
                                ) != QMessageBox.StandardButton.Yes:
            return

        if not ask_security_password(self):
            return

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("UPDATE item_prices SET parent_package='' WHERE parent_package=?", (name,))
        cursor.execute("DELETE FROM item_prices WHERE item_name=?", (name,))
        cursor.execute("DELETE FROM inventory WHERE item_name=?", (name,))
        conn.commit()
        conn.close()

        # نام حذف‌شده از مفاد قراردادهای قبلی هم پاک شود تا فاکتور تازه بماند
        try:
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT id, selected_items FROM wedding_contracts WHERE selected_items LIKE ?", (f"%{name}%",))
            for cid, items in cursor.fetchall():
                lst = [x.strip() for x in str(items or "").split(',') if x.strip() and x.strip() != name]
                cursor.execute("UPDATE wedding_contracts SET selected_items=? WHERE id=?", (",".join(lst), cid))
            conn.commit()
            conn.close()
        except Exception as e:
            print("[ITEMS] cleanup error:", e)

        recalc_inventory_usage()
        self.load_packages_combo()
        self.load_tree_items()


# ============================================================
# ===============  ItemEditDialog (ویرایش کالا)  =============
# ============================================================
class ItemEditDialog(QDialog):
    """ویرایش کامل یک کالا/خدمت/پکیج همراه با مشخصات انبار"""

    def __init__(self, item_name, parent=None):
        super().__init__(parent)
        self.old_name = item_name
        self.setWindowTitle(f"ویرایش: {item_name}")
        self.resize(560, 480)
        persist_dialog(self, "dlg_item_edit")
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT code, price, is_package, parent_package FROM item_prices WHERE item_name=?",
                       (item_name,))
        row = cursor.fetchone() or ("", 0, 0, "")
        cursor.execute("SELECT total_count, category, note FROM inventory WHERE item_name=?", (item_name,))
        inv = cursor.fetchone() or (0, "", "")
        cursor.execute("SELECT item_name FROM item_prices WHERE is_package=1 AND item_name<>?", (item_name,))
        packages = [r[0] for r in cursor.fetchall()]
        conn.close()

        lay = QVBoxLayout(self)
        form = QFormLayout()

        self.txt_code = QLineEdit(row[0] or "")
        self.txt_name = QLineEdit(item_name)
        self.txt_price = QLineEdit(f"{row[1]:,}")
        self.txt_price.textChanged.connect(lambda t: self.txt_price.setText(format_number(t)))
        self.txt_category = QLineEdit(inv[1] or "")
        self.spin_total = QSpinBox()
        self.spin_total.setRange(0, 100000)
        self.spin_total.setValue(int(inv[0] or 0))
        self.txt_note = QLineEdit(inv[2] or "")
        self.chk_pkg = QCheckBox("این مورد یک پکیج اصلی است")
        self.chk_pkg.setChecked(bool(row[2]))
        self.combo_parent = QComboBox()
        self.combo_parent.addItem("--- بدون پکیج مادر ---", "")
        for p in packages:
            self.combo_parent.addItem(p, p)
        idx = self.combo_parent.findData(row[3] or "")
        if idx >= 0:
            self.combo_parent.setCurrentIndex(idx)
        self.combo_parent.setEnabled(not bool(row[2]))

        form.addRow("کد:", self.txt_code)
        form.addRow("نام:", self.txt_name)
        form.addRow("قیمت (تومان):", self.txt_price)
        form.addRow("دسته‌بندی انبار:", self.txt_category)
        form.addRow("تعداد کل انبار:", self.spin_total)
        form.addRow("مشخصات / توضیحات:", self.txt_note)
        form.addRow("", self.chk_pkg)
        form.addRow("پکیج مادر:", self.combo_parent)
        lay.addLayout(form)

        note = QLabel(
            f"{ico('warn')} با تغییر نام یا کد، همه بخش‌ها (قراردادها، فاکتورها و انبار) "
            "به‌صورت خودکار هماهنگ می‌شوند."
        )
        note.setWordWrap(True)
        note.setStyleSheet("background-color:#fff9e6; border:1px solid #f39c12; "
                           "border-radius:8px; padding:9px;")
        lay.addWidget(note)

        row_btn = QHBoxLayout()
        btn_save = QPushButton(ico_text("save", "ذخیره تغییرات"))
        btn_save.setStyleSheet("background-color:#27ae60; color:white; font-weight:bold; padding:9px;")
        btn_save.clicked.connect(self.save)
        btn_cancel = QPushButton(ico_text("cancel", "لغو"))
        btn_cancel.clicked.connect(self.reject)
        row_btn.addWidget(btn_save, 2)
        row_btn.addWidget(btn_cancel, 1)
        lay.addLayout(row_btn)

    def save(self):
        new_name = self.txt_name.text().strip()
        code = self.txt_code.text().strip()
        if not new_name:
            QMessageBox.warning(self, "خطا", "نام نمی‌تواند خالی باشد.")
            return
        if not code:
            code = generate_unique_code("PKG" if self.chk_pkg.isChecked() else "SRV",
                                        "item_prices", "code",
                                        start=2001 if self.chk_pkg.isChecked() else 1001)

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT item_name FROM item_prices WHERE code=? AND item_name<>?", (code, self.old_name))
        other = cursor.fetchone()
        conn.close()
        if other:
            QMessageBox.warning(self, "کد تکراری",
                                f"کد «{code}» قبلاً برای «{other[0]}» ثبت شده است.")
            return

        is_pkg = 1 if self.chk_pkg.isChecked() else 0
        parent = "" if is_pkg else self.combo_parent.currentData()

        ok = sync_item_everywhere(self.old_name, new_name, code,
                                  parse_number(self.txt_price.text()), is_pkg, parent)
        if not ok:
            QMessageBox.critical(self, "خطا", "ذخیره تغییرات ناموفق بود.")
            return

        try:
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute('''
                INSERT INTO inventory (item_name, total_count, used_count, category, note)
                VALUES (?, ?, 0, ?, ?)
                ON CONFLICT(item_name) DO UPDATE SET
                    total_count=excluded.total_count,
                    category=excluded.category,
                    note=excluded.note
            ''', (new_name, self.spin_total.value(),
                  self.txt_category.text().strip(), self.txt_note.text().strip()))
            cursor.execute("DELETE FROM inventory WHERE item_name=? AND item_name<>?", (self.old_name, new_name))
            conn.commit()
            conn.close()
        except Exception as e:
            print("[ITEMS] inventory update error:", e)

        recalc_inventory_usage()
        QMessageBox.information(self, "موفقیت", "تغییرات ذخیره و در همه بخش‌ها هماهنگ شد.")
        self.accept()


# ============================================================
# ==================  DepositsDialog  ========================
# ============================================================
class DepositsDialog(QDialog):
    def __init__(self, contract_id, parent=None):
        super().__init__(parent)
        self.contract_id = contract_id
        self.setWindowTitle(f"مدیریت بیعانه‌ها و پرداخت‌های قرارداد #{contract_id}")
        self.resize(680, 520)
        persist_dialog(self, "dlg_deposits")
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))

        layout = QVBoxLayout()
        form_layout = QFormLayout()

        self.txt_amount = QLineEdit()
        self.txt_amount.textChanged.connect(lambda t: self.txt_amount.setText(format_number(t)))
        self.txt_date = PersianDateEdit(date_str=jalali_today_str())
        self.combo_bank = QComboBox()
        self.load_banks()

        form_layout.addRow(f"{ico('money')} مبلغ پرداخت (تومان):", self.txt_amount)
        form_layout.addRow(f"{ico('calendar')} تاریخ دریافت (روز/ماه/سال):", self.txt_date)
        form_layout.addRow(f"{ico('banks')} بانک واریزی:", self.combo_bank)

        btn_add = QPushButton(ico_text("add", "ثبت بیعانه / پرداختی جدید"))
        btn_add.setStyleSheet("background-color: #2980b9; color: white; font-weight: bold; padding: 8px;")
        btn_add.clicked.connect(self.add_deposit)
        form_layout.addRow(btn_add)
        layout.addLayout(form_layout)

        self.lbl_summary = QLabel("")
        self.lbl_summary.setWordWrap(True)
        self.lbl_summary.setStyleSheet("background-color:#fff9e6; border:1px solid #f39c12; "
                                       "border-radius:8px; padding:10px; font-size:11pt;")
        layout.addWidget(self.lbl_summary)

        self.table = QTableWidget()
        self.table.setColumnCount(5)
        self.table.setHorizontalHeaderLabels(["ردیف", "مبلغ (تومان)", "تاریخ", "بانک", "حذف"])
        self.table.setAlternatingRowColors(True)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
        self.table.horizontalHeader().setSectionResizeMode(4, QHeaderView.ResizeMode.ResizeToContents)
        layout.addWidget(self.table)

        btn_close = QPushButton(ico_text("ok", "بستن"))
        btn_close.clicked.connect(self.accept)
        layout.addWidget(btn_close)

        self.load_deposits()
        self.setLayout(layout)

    def load_banks(self):
        self.combo_bank.clear()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT bank_name FROM bank_cards")
        rows = cursor.fetchall()
        conn.close()
        if not rows:
            self.combo_bank.addItem("--- بدون بانک ثبت‌شده ---")
        for (b_name,) in rows:
            self.combo_bank.addItem(b_name)

    def load_deposits(self):
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT id, amount, deposit_date, bank_name FROM wedding_deposits "
                       "WHERE contract_id=? ORDER BY id", (self.contract_id,))
        rows = cursor.fetchall()
        cursor.execute("SELECT total_amount, discount, paid_amount, groom_name, bride_name "
                       "FROM wedding_contracts WHERE id=?", (self.contract_id,))
        c = cursor.fetchone()
        conn.close()

        self.table.setRowCount(0)
        running = 0
        for r_idx, (d_id, amount, d_date, bank) in enumerate(rows):
            running += amount or 0
            self.table.insertRow(r_idx)
            # ردیف همیشه از ۱ شروع می‌شود (قبلاً شناسه دیتابیس نمایش داده می‌شد و عدد ۲ می‌داد)
            self.table.setItem(r_idx, 0, QTableWidgetItem(str(r_idx + 1)))
            self.table.setItem(r_idx, 1, QTableWidgetItem(f"{amount:,}"))
            self.table.setItem(r_idx, 2, QTableWidgetItem(d_date or "-"))
            self.table.setItem(r_idx, 3, QTableWidgetItem(bank if bank else "-"))

            btn_del = QPushButton(ico("delete"))
            btn_del.setToolTip("حذف این پرداختی")
            btn_del.setStyleSheet("background-color: #e74c3c; color: white;")
            btn_del.clicked.connect(lambda _, did=d_id: self.delete_deposit(did))
            self.table.setCellWidget(r_idx, 4, btn_del)

        if c:
            raw, discount, paid, groom, bride = c
            net = max(0, (raw or 0) - (discount or 0))
            remain = max(0, net - running)
            settle = "✅ تسویه شده" if remain <= 0 else "⏳ در جریان"
            self.lbl_summary.setText(
                f"<b>زوجین:</b> {groom} و {bride}<br>"
                f"<b>جمع خام:</b> {raw:,} تومان  |  "
                f"<b>تخفیف:</b> {discount:,}  |  "
                f"<b>مبلغ نهایی:</b> {net:,} تومان<br>"
                f"<b>جمع پرداختی‌ها:</b> {running:,} تومان  |  "
                f"<b style='color:{'#27ae60' if remain <= 0 else '#c0392b'};'>مانده: {remain:,} تومان</b> "
                f" |  {settle}"
            )
        elif not rows:
            self.lbl_summary.setText("هنوز پرداختی برای این قرارداد ثبت نشده است.")

    def delete_deposit(self, d_id):
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM wedding_deposits WHERE id=?", (d_id,))
        cursor.execute("SELECT SUM(amount) FROM wedding_deposits WHERE contract_id=?", (self.contract_id,))
        total_paid = cursor.fetchone()[0] or 0
        cursor.execute("SELECT total_amount, discount FROM wedding_contracts WHERE id=?", (self.contract_id,))
        result = cursor.fetchone()
        if result:
            tot, disc = result
            is_settled = 1 if ((tot or 0) - (disc or 0) - total_paid) <= 0 else 0
            cursor.execute("UPDATE wedding_contracts SET paid_amount=?, is_settled=? WHERE id=?",
                           (total_paid, is_settled, self.contract_id))
        conn.commit()
        conn.close()
        self.load_deposits()

    def add_deposit(self):
        amount = parse_number(self.txt_amount.text())
        if amount <= 0:
            QMessageBox.warning(self, "خطا", "لطفاً مبلغ معتبر وارد کنید.")
            return

        d_date = self.txt_date.text().strip()
        bank = self.combo_bank.currentText()

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("INSERT INTO wedding_deposits (contract_id, amount, deposit_date, bank_name, work_year) VALUES (?, ?, ?, ?, ?)",
                       (self.contract_id, amount, d_date, bank, get_current_year()))
        cursor.execute("SELECT SUM(amount) FROM wedding_deposits WHERE contract_id=?", (self.contract_id,))
        total_paid = cursor.fetchone()[0] or 0
        cursor.execute("SELECT total_amount, discount FROM wedding_contracts WHERE id=?", (self.contract_id,))
        result = cursor.fetchone()
        if result:
            tot, disc = result
            is_settled = 1 if ((tot or 0) - (disc or 0) - total_paid) <= 0 else 0
            cursor.execute("UPDATE wedding_contracts SET paid_amount=?, is_settled=? WHERE id=?",
                           (total_paid, is_settled, self.contract_id))

        conn.commit()
        conn.close()

        self.txt_amount.clear()
        self.load_deposits()
        QMessageBox.information(self, "موفقیت", "بیعانه با موفقیت ثبت شد.")


# ============================================================
# ================  ContractDetailsDialog  ==================
# ============================================================
class ContractDetailsDialog(QDialog):
    def __init__(self, contract_id, parent=None):
        super().__init__(parent)
        self.setWindowTitle(f"جزئیات کامل قرارداد #{contract_id}")
        self.resize(780, 680)
        persist_dialog(self, "dlg_contract_details")
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))

        layout = QVBoxLayout()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT groom_name, bride_name, groom_phone, bride_phone,
                          contract_date, ceremony_date, selected_items, total_amount,
                          discount, paid_amount, description, code, staff_ids
                          FROM wedding_contracts WHERE id=?""", (contract_id,))
        row = cursor.fetchone()
        items_prices = {}
        if row and row[6]:
            for it in str(row[6]).split(','):
                it = it.strip()
                if not it:
                    continue
                cursor.execute("SELECT code, price, is_package FROM item_prices WHERE item_name=?", (it,))
                f = cursor.fetchone()
                items_prices[it] = f if f else ("-", 0, 0)
        conn.close()

        if row:
            info_text = f"<b>کد قرارداد:</b> {row[11] or '-'}<br>"
            info_text += f"<b>زوجین:</b> {row[0]} و {row[1]}<br>"
            info_text += f"<b>تلفن داماد:</b> {row[2] or '-'} | <b>تلفن عروس:</b> {row[3] or '-'}<br>"
            info_text += f"<b>تاریخ قرارداد:</b> {row[4]} | <b>تاریخ مراسم:</b> {row[5]}<br>"
            info_text += f"<b>توضیحات:</b> {row[10] if row[10] else 'ندارد'}"
            lbl_info = QLabel(info_text)
            lbl_info.setStyleSheet("background-color: #f8f9fa; padding: 12px; "
                                   "border-radius: 8px; font-size: 11pt; border:1px solid #d6e0ec;")
            lbl_info.setWordWrap(True)
            layout.addWidget(lbl_info)

            lbl_items = QLabel(f"{ico('list')} <b>مفاد و خدمات انتخاب شده:</b>")
            layout.addWidget(lbl_items)

            table = QTableWidget()
            table.setColumnCount(4)
            table.setHorizontalHeaderLabels(["ردیف", "کد", "عنوان خدمت / پکیج", "قیمت (تومان)"])
            table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
            table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
            table.setAlternatingRowColors(True)
            table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)

            items = [x.strip() for x in (row[6].split(',') if row[6] else []) if x.strip()]
            table.setRowCount(len(items))
            for idx, item in enumerate(items):
                info = items_prices.get(item, ("-", 0, 0))
                table.setItem(idx, 0, QTableWidgetItem(str(idx + 1)))
                table.setItem(idx, 1, QTableWidgetItem(info[0] or "-"))
                table.setItem(idx, 2, QTableWidgetItem(
                    f"{ico('package')} {item}" if info[2] else item))
                table.setItem(idx, 3, QTableWidgetItem(f"{info[1]:,}"))
            layout.addWidget(table)

            # ---- نیروی کار و حقوق این قرارداد ----
            staff = []
            try:
                staff = json.loads(row[12]) if row[12] else []
            except Exception:
                staff = []

            lbl_staff = QLabel(f"{ico('staff')} <b>نیروی کار این قرارداد:</b>")
            layout.addWidget(lbl_staff)
            st_table = QTableWidget()
            st_table.setColumnCount(3)
            st_table.setHorizontalHeaderLabels(["ردیف", "نام و نقش", "حقوق این قرارداد (تومان)"])
            st_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
            st_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.ResizeToContents)
            st_table.setAlternatingRowColors(True)
            st_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
            st_table.setRowCount(len(staff))
            staff_total = 0
            for idx, m in enumerate(staff):
                wage = int(m.get("wage") or 0)
                staff_total += wage
                st_table.setItem(idx, 0, QTableWidgetItem(str(idx + 1)))
                st_table.setItem(idx, 1, QTableWidgetItem(f"{m.get('name','-')} ({m.get('role','-')})"))
                st_table.setItem(idx, 2, QTableWidgetItem(f"{wage:,}"))
            layout.addWidget(st_table)
            if not staff:
                lbl_no_staff = QLabel("نیرویی برای این قرارداد ثبت نشده است.")
                lbl_no_staff.setStyleSheet("color:#7f8c8d; padding:6px;")
                layout.addWidget(lbl_no_staff)

            raw = row[7] or 0
            discount = row[8] or 0
            paid = row[9] or 0
            net = max(0, raw - discount)
            remain = max(0, net - paid)
            words = number_to_persian_words(remain)

            summary = (
                f"<div style='font-size:11pt; line-height:1.9;'>"
                f"<b>جمع قیمت خام:</b> {raw:,} تومان  |  "
                f"<b>تخفیف:</b> {discount:,} تومان<br>"
                f"<b>جمع کل بعد از تخفیف:</b> <span style='color:#2980b9;'>{net:,} تومان</span>  |  "
                f"<b>پرداختی:</b> {paid:,} تومان<br>"
                f"<b style='color:#c0392b;'>مانده:</b> {remain:,} تومان<br>"
                f"<b>مبلغ مانده به حروف:</b> {words} تومان  |  "
                f"<b>حقوق نیروی کار:</b> {staff_total:,} تومان"
                f"</div>"
            )
            lbl_sum = QLabel(summary)
            lbl_sum.setStyleSheet("font-size: 11pt; background-color: #fff9e6; padding: 10px; "
                                  "border-radius: 8px; border: 1px solid #f39c12;")
            lbl_sum.setWordWrap(True)
            layout.addWidget(lbl_sum)

        btn_row = QHBoxLayout()
        btn_print = QPushButton(ico_text("print", "چاپ / PDF این قرارداد"))
        btn_print.setStyleSheet("background-color:#2980b9; color:white; font-weight:bold; padding:8px;")
        btn_print.clicked.connect(lambda: self._print(contract_id))
        btn_close = QPushButton(ico_text("cancel", "بستن"))
        btn_close.clicked.connect(self.accept)
        btn_row.addWidget(btn_print, 2)
        btn_row.addWidget(btn_close, 1)
        layout.addLayout(btn_row)

        self.setLayout(layout)

    def _print(self, contract_id):
        parent = self.parent()
        if parent is not None and hasattr(parent, "_contract_print_flow"):
            parent._contract_print_flow(contract_id)
        else:
            html, _ = InvoiceBuilder.build_wedding_invoice(contract_id, with_details=True)
            if html:
                print_html_document(html, self, "فاکتور قرارداد")


# ============================================================
# ================  EditContractDialog  =====================
# ============================================================
class EditContractDialog(QDialog):
    def __init__(self, contract_id, parent=None):
        super().__init__(parent)
        self.contract_id = contract_id
        self.setWindowTitle(f"ویرایش قرارداد #{contract_id}")
        self.resize(780, 860)
        persist_dialog(self, "dlg_edit_contract")
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))

        layout = QVBoxLayout()
        form_layout = QFormLayout()

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT groom_name, bride_name, groom_phone, bride_phone,
                          contract_date, ceremony_date, selected_items, total_amount,
                          discount, description, staff_ids, paid_amount, venue
                          FROM wedding_contracts WHERE id=?""", (contract_id,))
        row = cursor.fetchone()

        cursor.execute("SELECT item_name, code, price, is_package, parent_package "
                       "FROM item_prices ORDER BY is_package DESC, parent_package, item_name")
        all_items = cursor.fetchall()

        cursor.execute("SELECT id, code, name, role FROM persons ORDER BY name")
        all_persons = cursor.fetchall()
        conn.close()

        self.groom = QLineEdit(row[0] if row else "")
        self.bride = QLineEdit(row[1] if row else "")
        self.groom_phone = QLineEdit(row[2] if row else "")
        self.bride_phone = QLineEdit(row[3] if row else "")
        self.contract_date = PersianDateEdit(date_str=row[4] if row and row[4] else None)
        self.ceremony_date = PersianDateEdit(date_str=row[5] if row and row[5] else None)
        self.discount = QLineEdit(f"{row[8]:,}" if row else "0")
        self.discount.textChanged.connect(lambda t: self.discount.setText(format_number(t)))
        self.discount.textChanged.connect(lambda _: self.recalc())
        self.venue = QLineEdit(row[12] if row and len(row) > 12 else "")
        self.venue.setPlaceholderText("نام تالار / باغ / مکان برگزاری مراسم")
        self.desc = QTextEdit(row[9] if row else "")
        self.desc.setFixedHeight(60)

        form_layout.addRow(f"{ico('wedding')} نام داماد:", self.groom)
        form_layout.addRow(f"{ico('wedding')} نام عروس:", self.bride)
        form_layout.addRow(f"{ico('people')} تلفن داماد:", self.groom_phone)
        form_layout.addRow(f"{ico('people')} تلفن عروس:", self.bride_phone)
        form_layout.addRow(f"{ico('calendar')} تاریخ قرارداد (روز/ماه/سال):", self.contract_date)
        form_layout.addRow(f"{ico('calendar')} تاریخ مراسم (روز/ماه/سال):", self.ceremony_date)
        form_layout.addRow(f"{ico('home')} مکان مراسم:", self.venue)
        form_layout.addRow(f"{ico('money')} تخفیف (تومان):", self.discount)

        selected = [x.strip() for x in (row[6].split(',') if row and row[6] else []) if x.strip()]
        items_box = QGroupBox("مفاد فاکتور (تیک بزنید — روی پکیج دوبار کلیک کنید تا جزئیاتش را ببینید)")
        items_box_layout = QVBoxLayout()
        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        scroll.setFixedHeight(240)

        inner = QWidget()
        inner_layout = QVBoxLayout()
        self.item_checks = {}
        self.item_prices = {}
        self.item_is_pkg = {}
        for item_name, code, price, is_pkg, parent in all_items:
            row_w = QWidget()
            row_l = QHBoxLayout()
            row_w.setFixedHeight(26)
            row_l.setContentsMargins(16 if parent else 0, 0, 0, 0)
            row_l.setSpacing(5)
            cb = ElidedCheckBox(f"{'• ' if parent else ''}[{code or '-'}] {item_name}")
            cb.setFont(QFont(APP_FONT_FAMILY, 9))
            cb.setChecked(item_name in selected)
            cb.stateChanged.connect(self.recalc)
            txt = QLineEdit(f"{price:,}")
            txt.setFixedWidth(120)
            txt.textChanged.connect(lambda t, p=txt: p.setText(format_number(t)))
            txt.textChanged.connect(self.recalc)
            self.item_checks[item_name] = cb
            self.item_prices[item_name] = txt
            self.item_is_pkg[item_name] = is_pkg
            row_l.addWidget(cb)
            row_l.addStretch()
            if is_pkg:
                btn_det = QPushButton(ico("detail"))
                btn_det.setFixedWidth(34)
                btn_det.setToolTip("نمایش جزئیات کامل این پکیج")
                btn_det.clicked.connect(lambda _, n=item_name: self.show_pkg(n))
                row_l.addWidget(btn_det)
            row_l.addWidget(QLabel("قیمت:"))
            row_l.addWidget(txt)
            row_w.setLayout(row_l)
            inner_layout.addWidget(row_w)
        inner.setLayout(inner_layout)
        scroll.setWidget(inner)
        items_box_layout.addWidget(scroll)
        items_box.setLayout(items_box_layout)
        form_layout.addRow(items_box)

        # ---------- نیروی کار این قرارداد ----------
        saved_staff = {}
        try:
            for m in (json.loads(row[10]) if row and row[10] else []):
                saved_staff[int(m.get("id"))] = int(m.get("wage") or 0)
        except Exception:
            saved_staff = {}

        staff_box = QGroupBox("نیروی کار این قرارداد (تیک بزنید و حقوق همان قرارداد را وارد کنید)")
        staff_box_layout = QVBoxLayout()
        staff_scroll = QScrollArea()
        staff_scroll.setWidgetResizable(True)
        staff_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        staff_scroll.setFixedHeight(150)
        staff_inner = QWidget()
        staff_inner_layout = QVBoxLayout()
        self.staff_checks = {}
        self.staff_wages = {}
        self.staff_meta = {}
        for pid, pcode, pname, prole in all_persons:
            row_w = QWidget()
            row_l = QHBoxLayout()
            row_l.setContentsMargins(0, 0, 0, 0)
            cb = QCheckBox(f"[{pcode or '-'}] {pname} ({prole})")
            cb.setChecked(pid in saved_staff)
            wage = QLineEdit(f"{saved_staff.get(pid, 0):,}")
            wage.setFixedWidth(130)
            wage.textChanged.connect(lambda t, p=wage: p.setText(format_number(t)))
            wage.textChanged.connect(self.recalc)
            self.staff_checks[pid] = cb
            self.staff_wages[pid] = wage
            self.staff_meta[pid] = (pname, prole)
            row_l.addWidget(cb)
            row_l.addStretch()
            row_l.addWidget(QLabel("حقوق این قرارداد:"))
            row_l.addWidget(wage)
            row_w.setLayout(row_l)
            staff_inner_layout.addWidget(row_w)
        if not all_persons:
            staff_inner_layout.addWidget(QLabel("هنوز هیچ نیرویی در بخش «کارکنان» تعریف نشده است."))
        staff_inner.setLayout(staff_inner_layout)
        staff_scroll.setWidget(staff_inner)
        staff_box_layout.addWidget(staff_scroll)
        staff_box.setLayout(staff_box_layout)
        form_layout.addRow(staff_box)

        self.lbl_calc = QLabel("")
        self.lbl_calc.setWordWrap(True)
        self.lbl_calc.setStyleSheet("background-color:#e8f8f5; border:1px solid #16a085; "
                                    "border-radius:8px; padding:10px; font-size:11pt;")
        form_layout.addRow(self.lbl_calc)

        form_layout.addRow(f"{ico('edit')} توضیحات:", self.desc)
        layout.addLayout(form_layout)

        btn_row = QHBoxLayout()
        btn_save = QPushButton(ico_text("save", "ذخیره تغییرات"))
        btn_save.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 10px;")
        btn_save.clicked.connect(self.save_changes)

        btn_cancel = QPushButton(ico_text("cancel", "لغو"))
        btn_cancel.clicked.connect(self.reject)

        btn_row.addWidget(btn_save, 2)
        btn_row.addWidget(btn_cancel, 1)
        layout.addLayout(btn_row)

        self.setLayout(layout)
        self.recalc()

    def show_pkg(self, name):
        PackageDetailsDialog(name, self).exec()

    def recalc(self):
        if not hasattr(self, "lbl_calc"):
            return
        raw = 0
        for item_name, cb in self.item_checks.items():
            if cb.isChecked():
                raw += parse_number(self.item_prices[item_name].text())
        discount = parse_number(self.discount.text())
        net = max(0, raw - discount)
        staff_total = 0
        count_staff = 0
        for pid, cb in getattr(self, "staff_checks", {}).items():
            if cb.isChecked():
                count_staff += 1
                staff_total += parse_number(self.staff_wages[pid].text())

        self.lbl_calc.setText(
            f"<b>جمع قیمت خام:</b> {raw:,} تومان  |  "
            f"<b>تخفیف:</b> {discount:,} تومان  |  "
            f"<b>جمع کل بعد از تخفیف:</b> <span style='color:#2980b9;'>{net:,} تومان</span><br>"
            f"<b>مبلغ به حروف:</b> {number_to_persian_words(net)} تومان<br>"
            f"<b>نیروی کار انتخاب‌شده:</b> {count_staff} نفر  |  "
            f"<b>جمع حقوق این قرارداد:</b> {staff_total:,} تومان"
        )

    def save_changes(self):
        if not ask_security_password(self):
            return

        selected_items = []
        raw_total = 0
        for item_name, cb in self.item_checks.items():
            if cb.isChecked():
                selected_items.append(item_name)
                raw_total += parse_number(self.item_prices[item_name].text())

        discount = parse_number(self.discount.text())
        net_total = max(0, raw_total - discount)

        staff_list = []
        for pid, cb in self.staff_checks.items():
            if cb.isChecked():
                pname, prole = self.staff_meta.get(pid, ("-", "-"))
                staff_list.append({
                    "id": pid, "name": pname, "role": prole,
                    "wage": parse_number(self.staff_wages[pid].text()),
                })

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT paid_amount FROM wedding_contracts WHERE id=?", (self.contract_id,))
        r = cursor.fetchone()
        paid = (r[0] or 0) if r else 0
        is_settled = 1 if (net_total - paid) <= 0 else 0

        # جمع کل = قیمت خام (بدون تخفیف) ذخیره می‌شود تا فاکتور بتواند
        # هم «جمع خام» و هم «جمع بعد از تخفیف» را جداگانه نشان دهد.
        cursor.execute("""UPDATE wedding_contracts SET
            groom_name=?, bride_name=?, groom_phone=?, bride_phone=?,
            contract_date=?, ceremony_date=?, selected_items=?,
            total_amount=?, discount=?, description=?, staff_ids=?, is_settled=?, venue=?
            WHERE id=?""",
            (self.groom.text(), self.bride.text(), self.groom_phone.text(),
             self.bride_phone.text(), self.contract_date.text(), self.ceremony_date.text(),
             ",".join(selected_items), raw_total, discount,
             self.desc.toPlainText(), json.dumps(staff_list, ensure_ascii=False),
             is_settled, self.venue.text().strip(), self.contract_id))
        conn.commit()
        conn.close()

        # انبار و کارکنان بلافاصله با قرارداد هماهنگ می‌شوند
        recalc_inventory_usage()
        sync_contract_staff({
            "id": self.contract_id,
            "code": self._contract_code(),
            "groom_name": self.groom.text(),
            "bride_name": self.bride.text(),
            "ceremony_date": self.ceremony_date.text(),
            "staff": staff_list,
        })

        QMessageBox.information(self, "موفقیت", "تغییرات با موفقیت ذخیره و همه بخش‌ها هماهنگ شد.")
        self.accept()

    def _contract_code(self):
        try:
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT code FROM wedding_contracts WHERE id=?", (self.contract_id,))
            r = cursor.fetchone()
            conn.close()
            return r[0] if r else ""
        except Exception:
            return ""


# ============================================================
# =============  EditCommercialDialog  ======================
# ============================================================
class EditCommercialDialog(QDialog):
    def __init__(self, project_id, parent=None):
        super().__init__(parent)
        self.project_id = project_id
        self.setWindowTitle(f"ویرایش پروژه #{project_id}")
        self.resize(500, 500)
        persist_dialog(self, "dlg_edit_commercial")
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT title, project_type, camera_count, total_amount,
                          paid_amount, project_date, description, client_name, client_phone
                          FROM commercial_projects WHERE id=?""", (project_id,))
        row = cursor.fetchone()
        conn.close()

        layout = QVBoxLayout()
        form = QFormLayout()

        self.title = QLineEdit(row[0] if row else "")
        self.ptype = QComboBox()
        self.ptype.addItems(["تبلیغاتی", "بیوتی", "تولدی", "عروسی", "عقد"])
        if row:
            self.ptype.setCurrentText(row[1])
        self.cameras = QSpinBox()
        self.cameras.setRange(1, 20)
        self.cameras.setValue(row[2] if row else 1)
        self.amount = QLineEdit(f"{row[3]:,}" if row else "0")
        self.amount.textChanged.connect(lambda t: self.amount.setText(format_number(t)))
        self.paid = QLineEdit(f"{row[4]:,}" if row else "0")
        self.paid.textChanged.connect(lambda t: self.paid.setText(format_number(t)))
        self.date = PersianDateEdit(date_str=row[5] if row and row[5] else None)
        self.client_name = QLineEdit(row[7] if row else "")
        self.client_phone = QLineEdit(row[8] if row else "")
        self.desc = QTextEdit(row[6] if row else "")
        self.desc.setFixedHeight(80)

        form.addRow(f"{ico('commercial')} عنوان پروژه:", self.title)
        form.addRow("نوع پروژه:", self.ptype)
        form.addRow(f"{ico('camera')} تعداد دوربین:", self.cameras)
        form.addRow(f"{ico('money')} مبلغ کل:", self.amount)
        form.addRow(f"{ico('deposit')} پرداختی:", self.paid)
        form.addRow(f"{ico('calendar')} تاریخ (روز/ماه/سال):", self.date)
        form.addRow(f"{ico('people')} نام مشتری:", self.client_name)
        form.addRow(f"{ico('people')} تلفن مشتری:", self.client_phone)
        form.addRow("توضیحات:", self.desc)
        layout.addLayout(form)

        btn_row = QHBoxLayout()
        btn_save = QPushButton(ico_text("save", "ذخیره تغییرات"))
        btn_save.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 8px;")
        btn_save.clicked.connect(self.save_changes)
        btn_cancel = QPushButton(ico_text("cancel", "لغو"))
        btn_cancel.clicked.connect(self.reject)
        btn_row.addWidget(btn_save, 2)
        btn_row.addWidget(btn_cancel, 1)
        layout.addLayout(btn_row)

        self.setLayout(layout)

    def save_changes(self):
        if not ask_security_password(self):
            return

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""UPDATE commercial_projects SET
            title=?, project_type=?, camera_count=?, total_amount=?,
            paid_amount=?, project_date=?, description=?,
            client_name=?, client_phone=?
            WHERE id=?""",
            (self.title.text(), self.ptype.currentText(), self.cameras.value(),
             parse_number(self.amount.text()), parse_number(self.paid.text()),
             self.date.text(), self.desc.toPlainText(),
             self.client_name.text(), self.client_phone.text(), self.project_id))
        conn.commit()
        conn.close()

        QMessageBox.information(self, "موفقیت", "تغییرات ذخیره شد.")
        self.accept()


# ============================================================
# ====================  InvoiceBuilder  =====================
# ============================================================
class InvoiceBuilder:
    """
    سازنده اسناد چاپی و PDF فارسی — نسخه ۱۱
    • همه اسناد روی یک برگ A5 و با بزرگ‌ترین فونت ممکن
    • خروجی PDF رنگی و خروجی پرینت سیاه‌وسفید (mono)
    • ترتیب ستون‌ها اصلاح‌شده برای خواندن راست‌به‌چپ (ردیف اولین ستون از راست)
    • بخش چک‌ها فقط در صورت وجود چک چاپ می‌شود
    • بخش نیروی کار قابل نمایش/حذف در چاپ است
    """

    # ============================================================
    # ====================  ابزارهای داخلی  ======================
    # ============================================================
    @staticmethod
    def _studio_info():
        return (
            get_setting("studio_name", "IMART STUDIO"),
            get_setting("studio_phone", "09173736618"),
            get_setting("studio_address", ""),
        )

    @staticmethod
    def _rev(cells):
        """
        موتور متن Qt ترتیب ستون‌های جدول را برای RTL برنمی‌گرداند،
        پس سلول‌ها را به ترتیب معکوس می‌نویسیم تا
        ستون اول منطقی (مثل «ردیف») در سمت راست دیده شود.
        """
        return "".join(reversed(cells))

    @staticmethod
    def _row(cells, style="", height=None):
        h = f" height='{height}'" if height else ""
        st = f" style='{style}'" if style else ""
        return f"<tr{st}{h}>{InvoiceBuilder._rev(cells)}</tr>"

    @staticmethod
    def _head(cells, bg="#2C6699", fg="#ffffff", size="7.5pt"):
        parts = []
        for c in cells:
            parts.append(
                f"<th style='border:1px solid #8a8a8a; padding:3px 4px; "
                f"background-color:{bg}; color:{fg}; font-size:{size};'>{c}</th>")
        return f"<tr>{InvoiceBuilder._rev(parts)}</tr>"

    @staticmethod
    def _td(text, align="center", size="7.5pt", bg=None, color=None, bold=False,
            colspan=None, width=None):
        st = ["border:1px solid #8a8a8a", "padding:3px 4px",
              f"text-align:{align}", f"font-size:{size}"]
        if bg:
            st.append(f"background-color:{bg}")
        if color:
            st.append(f"color:{color}")
        if bold:
            st.append("font-weight:bold")
        cs = f" colspan='{colspan}'" if colspan else ""
        # عرض ستون باید اتریبیوت HTML باشد؛ موتور متن Qt مقدار CSS آن را نادیده می‌گیرد
        ws = f" width='{width}'" if width else ""
        return f"<td{cs}{ws} style='{';'.join(st)};'>{text}</td>"

    @staticmethod
    def _items_breakdown(selected_items, items_detail, with_details=True):
        """سطرهای مفاد فاکتور همراه با زیرمجموعه پکیج‌ها (در صورت انتخاب)"""
        rows = []
        idx = 0
        try:
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            for item in selected_items:
                cursor.execute("SELECT code, price, is_package FROM item_prices WHERE item_name=?", (item,))
                f = cursor.fetchone()
                code = f[0] if f else "-"
                is_pkg = f[2] if f else 0
                price = items_detail.get(item)
                if price is None:
                    price = f[1] if f else 0
                idx += 1
                rows.append({"idx": idx, "code": code or "-", "name": item,
                             "price": int(price or 0), "sub": False})
                if with_details and is_pkg:
                    cursor.execute("SELECT item_name, code, price FROM item_prices "
                                   "WHERE parent_package=? ORDER BY item_name", (item,))
                    for sname, scode, sprice in cursor.fetchall():
                        idx += 1
                        rows.append({"idx": idx, "code": scode or "-",
                                     "name": f"• {sname}", "price": int(sprice or 0),
                                     "sub": True})
            conn.close()
        except Exception as e:
            print("[INVOICE] breakdown error:", e)
        return rows

    @staticmethod
    def _safe_json(text, default):
        try:
            return json.loads(text) if text else default
        except Exception:
            return default

    @staticmethod
    def _title_band(p, right_title, right_sub, left_lines):
        """سربرگ رنگی/خاکستری سند"""
        left = "".join(f"<div>{ln}</div>" for ln in left_lines if ln)
        return (
            f"<table width='100%' style='border-collapse:collapse;'>"
            f"<tr style='background-color:{p['head_bg']};'>"
            f"{InvoiceBuilder._rev([
                InvoiceBuilder._td(right_title + right_sub, align='center',
                                   size='10pt', color=p['head_fg'], bold=True),
                InvoiceBuilder._td(left, align='left', size='7pt', color=p['head_fg']),
            ])}"
            f"</tr></table>")

    # ============================================================
    # ================  فاکتور عروس و داماد  =====================
    # ============================================================
    @staticmethod
    def build_wedding_invoice(contract_id, with_details=True, with_staff=True, mono=False):
        p = doc_palette(mono)
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT id, code, groom_name, bride_name, groom_phone, bride_phone,
                          contract_date, ceremony_date, selected_items, total_amount,
                          discount, paid_amount, description, items_detail, staff_ids, venue
                          FROM wedding_contracts WHERE id=?""", (contract_id,))
        c = cursor.fetchone()

        cursor.execute("SELECT amount, deposit_date, bank_name FROM wedding_deposits "
                       "WHERE contract_id=? ORDER BY id", (contract_id,))
        deposits = cursor.fetchall()

        cursor.execute("""SELECT check_number, bank_name, amount, due_date, is_passed,
                          issuer_name, check_type FROM checks
                          WHERE contract_id=? ORDER BY check_type, due_date""", (contract_id,))
        checks = cursor.fetchall()
        conn.close()

        if not c:
            return "", ""

        studio_name, studio_phone, studio_address = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()
        date_str = now.strftime("%Y/%m/%d")
        time_str = now.strftime("%H:%M")
        invoice_code = f"INV-{c[0]:05d}"
        venue = c[15] if len(c) > 15 else ""

        items_list = [x.strip() for x in (c[8].split(',') if c[8] else []) if x.strip()]
        items_detail = InvoiceBuilder._safe_json(c[13], {})
        staff_list = InvoiceBuilder._safe_json(c[14], [])

        rows = InvoiceBuilder._items_breakdown(items_list, items_detail, with_details)

        # ---------- سربرگ ----------
        header = InvoiceBuilder._title_band(
            p,
            studio_name,
            "<div style='font-size:9pt;'>فاکتور فروش خدمات فیلم و عکس</div>",
            [f"تاریخ: {date_str}   ساعت: {time_str}",
             f"فاکتور: {invoice_code}",
             f"قرارداد: {c[1] or '-'}"],
        )

        # ---------- اطلاعات مشتری ----------
        info_rows = (
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>صورتحساب آقای / خانم:</b>", align="right",
                                   bg=p["box_bg"], size="8pt", width="26%"),
                InvoiceBuilder._td(f"{c[2]} و {c[3]}", align="right", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>تلفن همراه:</b>", align="right",
                                   bg=p["box_bg"], size="8pt", width="26%"),
                InvoiceBuilder._td(f"{c[4] or '-'}   |   {c[5] or '-'}", align="right", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>تاریخ قرارداد:</b>", align="right",
                                   bg=p["box_bg"], size="8pt", width="26%"),
                InvoiceBuilder._td(f"{c[6] or '-'}", align="right", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>تاریخ مراسم:</b>", align="right",
                                   bg=p["box_bg"], size="8pt", width="26%"),
                InvoiceBuilder._td(f"{c[7] or '-'}", align="right", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>مکان مراسم:</b>", align="right",
                                   bg=p["box_bg"], size="8pt", width="26%"),
                InvoiceBuilder._td(f"{venue or '-'}", align="right", size="8pt"),
            ])
        )
        info_table = (f"<table width='100%' style='border-collapse:collapse; "
                      f"margin-top:3px;'>{info_rows}</table>")

        # ---------- مفاد ----------
        body_rows = ""
        for r in rows:
            bg = p["alt_row"] if r["sub"] else None
            body_rows += InvoiceBuilder._row([
                InvoiceBuilder._td(str(r['idx']), bg=bg, size="7.5pt"),
                InvoiceBuilder._td(r['code'], bg=bg, size="7pt"),
                InvoiceBuilder._td(r['name'], align="right", bg=bg, size="7.5pt"),
                InvoiceBuilder._td("1", bg=bg, size="7.5pt"),
                InvoiceBuilder._td(f"{r['price']:,}", bg=bg, size="7.5pt"),
                InvoiceBuilder._td(f"{r['price']:,}", bg=bg, size="7.5pt"),
            ])
        if not body_rows:
            body_rows = (f"<tr>{InvoiceBuilder._td('موردی ثبت نشده است', colspan=6, size='8pt')}</tr>")

        items_table = (
            f"<table width='100%' style='border-collapse:collapse; margin-top:4px;'>"
            f"<thead>{InvoiceBuilder._head(['ردیف', 'کد', 'شرح خدمات / پکیج', 'تعداد', 'بهای واحد (تومان)', 'مبلغ کل (تومان)'], bg=p['table_head'], fg=p['table_head_fg'])}</thead>"
            f"<tbody>{body_rows}</tbody></table>")

        # ---------- محاسبات ----------
        raw_total = c[9] or 0
        discount = c[10] or 0
        net_total = max(0, raw_total - discount)
        paid = c[11] or 0
        remain = max(0, net_total - paid)
        words_net = number_to_persian_words(net_total)
        words_remain = number_to_persian_words(remain)

        total_rows = (
            InvoiceBuilder._row([
                InvoiceBuilder._td("جمع کل (قیمت خام)", align="right", bg=p["box_bg"], size="8pt"),
                InvoiceBuilder._td(f"{raw_total:,}", align="center", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("تخفیف", align="right", bg=p["box_bg"], size="8pt"),
                InvoiceBuilder._td(f"{discount:,}", align="center", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>جمع کل بعد از تخفیف</b>", align="right",
                                   bg=p["alt_row"], size="8.5pt", bold=True),
                InvoiceBuilder._td(f"<b>{net_total:,}</b>", align="center", size="8.5pt", bold=True),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("جمع بیعانه‌های دریافتی", align="right", bg=p["box_bg"], size="8pt"),
                InvoiceBuilder._td(f"{paid:,}", align="center", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>مانده قابل پرداخت</b>", align="right",
                                   bg=p["total_bg"], size="9pt", bold=True),
                InvoiceBuilder._td(f"<b>{remain:,}</b>", align="center", size="9pt", bold=True),
            ])
        )
        totals_table = f"<table width='100%' style='border-collapse:collapse;'>{total_rows}</table>"

        words_cell = (
            f"<div style='font-size:8pt;'><b>مبلغ کل به حروف:</b> {words_net} تومان</div>"
            f"<div style='font-size:8pt; margin-top:2px;'><b>مانده به حروف:</b> {words_remain} تومان</div>"
            f"<div style='font-size:7.5pt; margin-top:4px; color:{p['muted']};'>"
            f"<b>توضیحات:</b> {c[12] if c[12] else 'ندارد'}</div>"
        )
        calc_table = (
            f"<table width='100%' style='border-collapse:collapse; margin-top:4px;'>"
            f"{InvoiceBuilder._row([InvoiceBuilder._td(words_cell, align='right', size='8pt'), totals_table])}"
            f"</table>")

        # ---------- بیعانه‌ها (فقط اگر وجود داشته باشد) ----------
        deposits_block = ""
        if deposits:
            dep_rows = ""
            for idx, (amt, dt, bank) in enumerate(deposits):
                dep_rows += InvoiceBuilder._row([
                    InvoiceBuilder._td(str(idx + 1), size="7.5pt"),
                    InvoiceBuilder._td(f"{amt:,}", size="7.5pt"),
                    InvoiceBuilder._td(dt or "-", size="7.5pt"),
                    InvoiceBuilder._td(bank or "-", size="7.5pt"),
                ])
            deposits_block = (
                f"<div style='margin-top:5px; font-size:8.5pt; font-weight:bold; color:{p['accent']};'>"
                f"بیعانه‌ها و پرداخت‌های دریافتی</div>"
                f"<table width='100%' style='border-collapse:collapse;'>"
                f"<thead>{InvoiceBuilder._head(['ردیف', 'مبلغ (تومان)', 'تاریخ', 'بانک'], bg=p['table_head'], fg=p['table_head_fg'])}</thead>"
                f"<tbody>{dep_rows}</tbody></table>")

        # ---------- چک‌ها: فقط بخشی که چک دارد چاپ می‌شود ----------
        # اگر هیچ چک دریافتی یا پرداختی وجود نداشته باشد، آن بخش چاپ نمی‌شود.
        recv_rows = ""
        paid_rows = ""
        ri = pi = 0
        for chk in checks:
            status = "پاس شده" if chk[4] else "پاس نشده"
            if (chk[6] or "دریافتی") == "دریافتی":
                ri += 1
                recv_rows += InvoiceBuilder._row([
                    InvoiceBuilder._td(str(ri), size="7pt"),
                    InvoiceBuilder._td(chk[0], size="7pt"),
                    InvoiceBuilder._td(chk[1] or "-", size="7pt"),
                    InvoiceBuilder._td(f"{chk[2]:,}", size="7pt"),
                    InvoiceBuilder._td(chk[3] or "-", size="7pt"),
                    InvoiceBuilder._td(chk[5] or "-", size="7pt"),
                    InvoiceBuilder._td(status, size="7pt"),
                ])
            else:
                pi += 1
                paid_rows += InvoiceBuilder._row([
                    InvoiceBuilder._td(str(pi), size="7pt"),
                    InvoiceBuilder._td(chk[0], size="7pt"),
                    InvoiceBuilder._td(chk[1] or "-", size="7pt"),
                    InvoiceBuilder._td(f"{chk[2]:,}", size="7pt"),
                    InvoiceBuilder._td(chk[3] or "-", size="7pt"),
                    InvoiceBuilder._td(chk[5] or "-", size="7pt"),
                    InvoiceBuilder._td(status, size="7pt"),
                ])

        checks_block = ""
        if recv_rows or paid_rows:
            checks_block = (
                f"<div style='margin-top:5px; font-size:8.5pt; font-weight:bold; "
                f"color:{p['accent']};'>چک‌های مرتبط با این قرارداد</div>")
            if recv_rows:
                checks_block += (
                    f"<table width='100%' style='border-collapse:collapse;'>"
                    f"<thead>"
                    f"<tr style='background-color:{p['recv_head']};'>"
                    f"{InvoiceBuilder._rev([InvoiceBuilder._td('<b>چک‌های دریافتی (دریافتی از)</b>', align='center', colspan=7, color='#ffffff', size='8pt', bold=True)])}"
                    f"</tr>"
                    f"{InvoiceBuilder._head(['ردیف', 'شماره چک', 'بانک', 'مبلغ (تومان)', 'سررسید', 'دریافتی از', 'وضعیت'], bg=p['recv_head'], fg='#ffffff')}"
                    f"</thead><tbody>{recv_rows}</tbody></table>")
            if paid_rows:
                checks_block += (
                    f"<table width='100%' style='border-collapse:collapse; margin-top:2px;'>"
                    f"<thead>"
                    f"<tr style='background-color:{p['paid_head']};'>"
                    f"{InvoiceBuilder._rev([InvoiceBuilder._td('<b>چک‌های پرداختی (در وجه)</b>', align='center', colspan=7, color='#ffffff', size='8pt', bold=True)])}"
                    f"</tr>"
                    f"{InvoiceBuilder._head(['ردیف', 'شماره چک', 'بانک', 'مبلغ (تومان)', 'سررسید', 'در وجه', 'وضعیت'], bg=p['paid_head'], fg='#ffffff')}"
                    f"</thead><tbody>{paid_rows}</tbody></table>")

        # ---------- نیروی کار (اختیاری) ----------
        staff_block = ""
        if with_staff and staff_list:
            staff_rows = ""
            staff_sum = 0
            for i, m in enumerate(staff_list):
                w = int(m.get("wage") or 0)
                staff_sum += w
                staff_rows += InvoiceBuilder._row([
                    InvoiceBuilder._td(str(i + 1), size="7.5pt"),
                    InvoiceBuilder._td(m.get('name', '-'), align="right", size="7.5pt"),
                    InvoiceBuilder._td(m.get('role', '-'), size="7.5pt"),
                    InvoiceBuilder._td(f"{w:,}", size="7.5pt"),
                ])
            staff_rows += InvoiceBuilder._row([
                InvoiceBuilder._td("<b>جمع حقوق نیروی کار</b>", align="right",
                                   colspan=3, bg=p["alt_row"], size="8pt", bold=True),
                InvoiceBuilder._td(f"<b>{staff_sum:,}</b>", bg=p["alt_row"], size="8pt", bold=True),
            ])
            staff_block = (
                f"<div style='margin-top:5px; font-size:8.5pt; font-weight:bold; color:{p['accent']};'>"
                f"نیروی کار این قرارداد</div>"
                f"<table width='100%' style='border-collapse:collapse;'>"
                f"<thead>{InvoiceBuilder._head(['ردیف', 'نام', 'نقش', 'حقوق این قرارداد (تومان)'], bg=p['staff_head'], fg='#ffffff')}</thead>"
                f"<tbody>{staff_rows}</tbody></table>")

        html = f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {header}
          {info_table}
          {items_table}
          {calc_table}
          {deposits_block}
          {checks_block}
          {staff_block}
          <table width="100%" style="border-collapse:collapse; margin-top:14px;">
            <tr>
              {InvoiceBuilder._rev([
                InvoiceBuilder._td("مهر و امضای مشتری", align="center", size="8.5pt"),
                InvoiceBuilder._td("مهر و امضای استودیو", align="center", size="8.5pt"),
              ])}
            </tr>
          </table>
          <div style="text-align:center; margin-top:6px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}{(' | ' + studio_address) if studio_address else ''}
          </div>
        </div>
        """
        filename = f"فاکتور_{c[2]}_{c[3]}_{invoice_code}.pdf"
        return html, filename

    # ============================================================
    # ==============  فاکتور پروژه‌های تبلیغاتی  ==================
    # ============================================================
    @staticmethod
    def build_commercial_invoice(project_id, with_details=True, with_staff=False, mono=False):
        p = doc_palette(mono)
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT id, code, title, project_type, camera_count, total_amount,
                          paid_amount, project_date, description, client_name, client_phone
                          FROM commercial_projects WHERE id=?""", (project_id,))
        c = cursor.fetchone()
        conn.close()
        if not c:
            return "", ""

        studio_name, studio_phone, studio_address = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()
        invoice_code = f"CIV-{c[0]:05d}"
        total = c[5] or 0
        paid = c[6] or 0
        remain = max(0, total - paid)

        header = InvoiceBuilder._title_band(
            p, studio_name,
            f"<div style='font-size:9pt;'>فاکتور پروژه {c[3]}</div>",
            [f"تاریخ: {now.strftime('%Y/%m/%d')}   ساعت: {now.strftime('%H:%M')}",
             f"فاکتور: {invoice_code}", f"پروژه: {c[1] or '-'}"])

        info_rows = (
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>صورتحساب آقای / خانم:</b>", align="right",
                                   bg=p["box_bg"], size="8pt", width="26%"),
                InvoiceBuilder._td(f"{c[9] or '-'}", align="right", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>تلفن همراه:</b>", align="right",
                                   bg=p["box_bg"], size="8pt", width="26%"),
                InvoiceBuilder._td(f"{c[10] or '-'}", align="right", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>تاریخ پروژه:</b>", align="right",
                                   bg=p["box_bg"], size="8pt", width="26%"),
                InvoiceBuilder._td(f"{c[7] or '-'}", align="right", size="8pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>تعداد دوربین:</b>", align="right",
                                   bg=p["box_bg"], size="8pt", width="26%"),
                InvoiceBuilder._td(f"{c[4]}", align="right", size="8pt"),
            ])
        )
        items_rows = InvoiceBuilder._row([
            InvoiceBuilder._td("1", size="8pt"),
            InvoiceBuilder._td("-", size="8pt"),
            InvoiceBuilder._td(c[2], align="right", size="8pt"),
            InvoiceBuilder._td("1", size="8pt"),
            InvoiceBuilder._td(f"{total:,}", size="8pt"),
            InvoiceBuilder._td(f"{total:,}", size="8pt"),
        ])
        totals = (
            InvoiceBuilder._row([
                InvoiceBuilder._td("جمع کل", align="right", bg=p["box_bg"], size="8.5pt"),
                InvoiceBuilder._td(f"{total:,}", size="8.5pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("پرداختی", align="right", bg=p["box_bg"], size="8.5pt"),
                InvoiceBuilder._td(f"{paid:,}", size="8.5pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>مانده قابل پرداخت</b>", align="right",
                                   bg=p["total_bg"], size="9pt", bold=True),
                InvoiceBuilder._td(f"<b>{remain:,}</b>", size="9pt", bold=True),
            ])
        )
        words = (
            f"<div style='font-size:8pt;'><b>مبلغ کل به حروف:</b> "
            f"{number_to_persian_words(total)} تومان</div>"
            f"<div style='font-size:8pt; margin-top:2px;'><b>مانده به حروف:</b> "
            f"{number_to_persian_words(remain)} تومان</div>"
            f"<div style='font-size:7.5pt; margin-top:4px; color:{p['muted']};'>"
            f"<b>توضیحات:</b> {c[8] if c[8] else 'ندارد'}</div>"
        )

        html = f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {header}
          <table width="100%" style="border-collapse:collapse; margin-top:3px;">{info_rows}</table>
          <table width="100%" style="border-collapse:collapse; margin-top:4px;">
            <thead>{InvoiceBuilder._head(['ردیف', 'کد', 'عنوان خدمات', 'تعداد', 'بهای واحد (تومان)', 'مبلغ کل (تومان)'], bg=p['table_head'], fg=p['table_head_fg'])}</thead>
            <tbody>{items_rows}</tbody>
          </table>
          <table width="100%" style="border-collapse:collapse; margin-top:4px;">
            {InvoiceBuilder._row([InvoiceBuilder._td(words, align='right', size='8pt'),
                                  f"<table width='100%' style='border-collapse:collapse;'>{totals}</table>"])}
          </table>
          <table width="100%" style="border-collapse:collapse; margin-top:14px;">
            <tr>
              {InvoiceBuilder._rev([
                InvoiceBuilder._td("مهر و امضای مشتری", align="center", size="8.5pt"),
                InvoiceBuilder._td("مهر و امضای استودیو", align="center", size="8.5pt"),
              ])}
            </tr>
          </table>
          <div style="text-align:center; margin-top:6px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}{(' | ' + studio_address) if studio_address else ''}
          </div>
        </div>
        """
        filename = f"فاکتور_{c[2]}_{invoice_code}.pdf"
        return html, filename

    # ============================================================
    # =====================  رسید وجه  ==========================
    # ============================================================
    @staticmethod
    def build_receipt(name, amount, for_service, date_str, direction="دریافتی", mono=False):
        p = doc_palette(mono)
        studio_name, studio_phone, studio_address = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()
        amount_words = number_to_persian_words(amount)

        if direction == "دریافتی":
            title = "رسید دریافت وجه"
            party_label = "دریافت شد از آقا / خانم"
            action = "به حسابداری تحویل گردید."
        else:
            title = "رسید پرداخت وجه"
            party_label = "پرداخت شد به آقا / خانم"
            action = "از حساب استودیو پرداخت گردید."

        # برچسب‌ها سمت راست و کادرهای پرکردنی سمت چپ (سلول اول = مقدار)
        rows = (
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>" + party_label + ":</b>", align="right",
                                   bg=p["box_bg"], size="9pt", width="46%"),
                InvoiceBuilder._td(f"<b>{name}</b>", align="right", size="10pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>مبلغ (تومان):</b>", align="right",
                                   bg=p["box_bg"], size="9pt", width="46%"),
                InvoiceBuilder._td(f"<b>{amount:,}</b>", align="right", size="12pt", bold=True),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>مبلغ به حروف:</b>", align="right",
                                   bg=p["box_bg"], size="9pt", width="46%"),
                InvoiceBuilder._td(f"{amount_words} تومان", align="right", size="9pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>بابت:</b>", align="right",
                                   bg=p["box_bg"], size="9pt", width="46%"),
                InvoiceBuilder._td(for_service, align="right", size="9pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>شرح:</b>", align="right",
                                   bg=p["box_bg"], size="9pt", width="46%"),
                InvoiceBuilder._td(action, align="right", size="9pt"),
            ])
        )

        html = f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {InvoiceBuilder._title_band(
              p, studio_name,
              f"<div style='font-size:10pt;'>{title}</div>",
              [f"تاریخ: {date_str}", f"ساعت: {now.strftime('%H:%M')}",
               f"شماره رسید: RC-{now.strftime('%y%m%d%H%M')}"])}
          <table width="100%" style="border-collapse:collapse; margin-top:6px;">{rows}</table>
          <div style="margin-top:8px; border:1px solid {p['warn_border']}; background-color:{p['warn_bg']};
                      padding:5px; font-size:8.5pt;">
            <b>توجه:</b> بیعانه پرداختی پس داده نمی‌شود.
          </div>
          <table width="100%" style="border-collapse:collapse; margin-top:40px;">
            <tr>
              {InvoiceBuilder._rev([
                InvoiceBuilder._td("امضای حسابداری", align="center", size="9pt"),
                InvoiceBuilder._td("مهر استودیو", align="center", size="9pt"),
              ])}
            </tr>
          </table>
          <div style="text-align:center; margin-top:12px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}{(' | ' + studio_address) if studio_address else ''}
          </div>
        </div>
        """
        return html

    # ============================================================
    # ==================  گزارش ریز کارکرد فرد  =================
    # ============================================================
    @staticmethod
    def build_staff_report(person_id, period="ماهانه", year=None, month=None, mono=False):
        p = doc_palette(mono)
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT code, name, role, phone FROM persons WHERE id=?", (person_id,))
        pr = cursor.fetchone()
        if not pr:
            conn.close()
            return "", ""

        work_year = year or get_current_year()
        cursor.execute("""SELECT amount, year, month, day, description, category
                          FROM transactions WHERE person_id=? AND work_year=?
                          ORDER BY year DESC, month DESC, day DESC""",
                       (person_id, work_year))
        trans = cursor.fetchall()
        conn.close()

        total_salary = sum(t[0] for t in trans)
        trans_rows = ""
        for idx, t in enumerate(trans):
            trans_rows += InvoiceBuilder._row([
                InvoiceBuilder._td(str(idx + 1), size="7.5pt"),
                InvoiceBuilder._td(f"{t[1]}/{t[2]:02d}/{t[3]:02d}", size="7.5pt"),
                InvoiceBuilder._td(t[5], size="7.5pt"),
                InvoiceBuilder._td(t[4] or "-", align="right", size="7.5pt"),
                InvoiceBuilder._td(f"{t[0]:,}", size="7.5pt"),
            ])
        if not trans_rows:
            trans_rows = f"<tr>{InvoiceBuilder._td('تراکنشی یافت نشد', colspan=5, size='8pt')}</tr>"

        studio_name, studio_phone, _addr = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()

        info_rows = (
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>کد پرسنل:</b>", align="right", bg=p["box_bg"], size="8.5pt", width="22%"),
                InvoiceBuilder._td(pr[0] or "-", align="right", size="8.5pt"),
                InvoiceBuilder._td("<b>نام و نام خانوادگی:</b>", align="right", bg=p["box_bg"], size="8.5pt", width="22%"),
                InvoiceBuilder._td(pr[1], align="right", size="8.5pt"),
            ]) +
            InvoiceBuilder._row([
                InvoiceBuilder._td("<b>نقش / تخصص:</b>", align="right", bg=p["box_bg"], size="8.5pt", width="22%"),
                InvoiceBuilder._td(pr[2], align="right", size="8.5pt"),
                InvoiceBuilder._td("<b>تلفن:</b>", align="right", bg=p["box_bg"], size="8.5pt", width="22%"),
                InvoiceBuilder._td(pr[3] or "-", align="right", size="8.5pt"),
            ])
        )

        html = f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {InvoiceBuilder._title_band(
              p, studio_name,
              "<div style='font-size:10pt;'>گزارش ریز کارکرد پرسنل</div>",
              [f"تاریخ: {now.strftime('%Y/%m/%d')}", f"دوره: {period}",
               f"سال کاری: {work_year}"])}
          <table width="100%" style="border-collapse:collapse; margin-top:5px;">{info_rows}</table>
          <div style="margin-top:6px; font-size:9pt; font-weight:bold; color:{p['accent']};">
            ریز پرداختی‌ها ({period})
          </div>
          <table width="100%" style="border-collapse:collapse; margin-top:2px;">
            <thead>{InvoiceBuilder._head(['ردیف', 'تاریخ', 'دسته', 'شرح', 'مبلغ (تومان)'], bg=p['table_head'], fg=p['table_head_fg'])}</thead>
            <tbody>{trans_rows}</tbody>
            <tfoot>
              {InvoiceBuilder._row([
                  InvoiceBuilder._td("<b>جمع کل پرداختی به این فرد</b>", align="right",
                                     colspan=4, bg=p["total_bg"], size="9pt", bold=True),
                  InvoiceBuilder._td(f"<b>{total_salary:,}</b>", bg=p["total_bg"], size="9pt", bold=True),
              ])}
            </tfoot>
          </table>
          <div style="margin-top:6px; font-size:8.5pt;">
            <b>جمع کل به حروف:</b> {number_to_persian_words(total_salary)} تومان
          </div>
          <div style="text-align:center; margin-top:26px; font-size:9pt;">
            مهر و امضای مدیر / حسابداری
          </div>
          <div style="text-align:center; margin-top:10px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}
          </div>
        </div>
        """
        filename = f"ریز_کارکرد_{pr[1]}_{period}.pdf"
        return html, filename


class StudioAccountingApp(QMainWindow):
    ROLES = ["تدوینگر", "عکاس", "فیلمبردار", "هلی شات و FPV کار", "اوپراتور کرین"]
    PROJECT_TYPES = ["عروسی", "عقد", "تولد", "تبلیغاتی", "قبض و کرایه"]
    PERSIAN_MONTHS = ["فروردین", "اردیبهشت", "خرداد", "تیر", "مرداد", "شهریور",
                      "مهر", "آبان", "آذر", "دی", "بهمن", "اسفند"]

    def __init__(self):
        super().__init__()
        self.setWindowTitle(f"نرم‌افزار مدیریت مالی و حسابداری IMART STUDIO - نسخه {APP_VERSION}")
        self.setMinimumSize(1100, 700)
        self.resize(1400, 900)
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self._allow_close = False

        icon_path = resource_path("Accounting.png")
        if os.path.exists(icon_path):
            self.setWindowIcon(QIcon(icon_path))
        elif os.path.exists(resource_path("icon.png")):
            self.setWindowIcon(QIcon(resource_path("icon.png")))

        self.setFont(QFont(APP_FONT_FAMILY, 10))

        self.stack = QStackedWidget()
        self.setCentralWidget(self.stack)

        self.dashboard_screen = QWidget()
        self.stack.addWidget(self.dashboard_screen)

        self.main_app_screen = QWidget()
        self.stack.addWidget(self.main_app_screen)

        self.stack.setCurrentWidget(self.dashboard_screen)
        # اگر کاربر قبلاً ابعاد/مکان پنجره را تغییر داده باشد، همان حالت برمی‌گردد
        if not restore_geometry("main", self):
            self.center_on_screen()

    def center_on_screen(self):
        """قرار گرفتن پنجره برنامه در وسط صفحه نمایش"""
        try:
            screen = self.screen() or QApplication.primaryScreen()
            if not screen:
                return
            geo = screen.availableGeometry()
            w = min(self.width(), max(900, geo.width() - 60))
            h = min(self.height(), max(640, geo.height() - 80))
            self.resize(w, h)
            self.move(geo.x() + (geo.width() - w) // 2,
                      geo.y() + (geo.height() - h) // 2)
        except Exception as e:
            print("[WINDOW] center error:", e)

    def prompt_login(self):
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT value FROM settings WHERE key='app_password'")
        row = cursor.fetchone()
        saved_pass = row[0] if row else "123"
        conn.close()

        entered_pass, ok = QInputDialog.getText(
            None, "ورود به سیستم IMART STUDIO", "لطفاً رمز عبور را وارد کنید:",
            QLineEdit.EchoMode.Password
        )
        if ok and entered_pass == saved_pass:
            return True
        elif ok:
            QMessageBox.critical(None, "خطا", "رمز عبور اشتباه است!")
            return False
        return False

    def check_check_alerts(self):
        """اخطار ۳ روز مانده به سررسید چک‌ها"""
        try:
            today = jdatetime.date.today()
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("""SELECT check_number, bank_name, amount, due_date,
                              issuer_name, check_type
                              FROM checks WHERE is_passed=0""")
            rows = cursor.fetchall()
            conn.close()

            alerts = []
            for chk in rows:
                try:
                    due = parse_jalali(chk[3])
                    if due is None:
                        continue
                    days_left = (due - today).days
                    if 0 <= days_left <= 3:
                        alerts.append((chk, days_left))
                    elif days_left < 0:
                        alerts.append((chk, days_left))
                except Exception:
                    pass

            if alerts:
                msg = "<h3 style='color:#c0392b;'>⚠️ اخطار سررسید چک‌ها</h3><hr>"
                for chk, days in alerts:
                    if days < 0:
                        status = f"<span style='color:#c0392b;'><b>گذشته ({abs(days)} روز)</b></span>"
                    elif days == 0:
                        status = "<span style='color:#e67e22;'><b>امروز!</b></span>"
                    else:
                        status = f"<span style='color:#f39c12;'><b>{days} روز مانده</b></span>"

                    msg += f"""
                    <p style='font-size:11pt;'>
                        <b>شماره چک:</b> {chk[0]} | <b>بانک:</b> {chk[1] or '-'}<br>
                        <b>مبلغ:</b> {chk[2]:,} تومان | <b>سررسید:</b> {chk[3]} ({status})<br>
                        <b>صادرکننده:</b> {chk[4] or '-'} | <b>نوع:</b> {chk[5] or 'دریافتی'}
                    </p><hr>
                    """
                box = QMessageBox(self)
                box.setWindowTitle("اخطار چک")
                box.setTextFormat(Qt.TextFormat.RichText)
                box.setText(msg)
                box.setIcon(QMessageBox.Icon.Warning)
                box.exec()
        except Exception as e:
            print("Error in check alerts:", e)

    def make_dashboard_button(self, icon, title, color, index):
        """
        کارت گرافیکی منو با انیمیشن:
        با ورود موس، دکمه کمی بزرگ‌تر می‌شود و حلقه رنگی دور آیکون روشن می‌گردد.
        """
        base_w, base_h = 150, 128
        grow = 12
        m = grow // 2

        holder = QWidget()
        holder.setFixedSize(base_w + grow, base_h + grow)
        holder.setStyleSheet("background: transparent;")

        btn = HoverTileButton(holder)
        btn.setCursor(Qt.CursorShape.PointingHandCursor)
        btn.setGeometry(m, m, base_w, base_h)

        small_rect = QRect(m, m, base_w, base_h)
        big_rect = QRect(0, 0, base_w + grow, base_h + grow)

        ring_on = shade_color(color, 95)

        def tile_style(ring_color, ring_width):
            return f"""
            QPushButton {{
                background-color: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                    stop:0 {shade_color(color, 28)}, stop:0.55 {color},
                    stop:1 {shade_color(color, -40)});
                border: {ring_width}px solid {ring_color};
                border-radius: 17px;
            }}
            QLabel {{ background: transparent; color: #ffffff; border: none; }}
            """

        btn.setStyleSheet(tile_style("rgba(255,255,255,120)", 1))

        lay = QVBoxLayout(btn)
        lay.setContentsMargins(5, 10, 5, 9)
        lay.setSpacing(5)

        lbl_icon = QLabel(icon)
        lbl_icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_icon.setStyleSheet("font-size: 25pt; background: transparent;")
        lbl_icon.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)

        lbl_title = QLabel(title)
        lbl_title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_title.setWordWrap(True)
        lbl_title.setStyleSheet("font-size: 9pt; font-weight: bold; background: transparent;")
        lbl_title.setAttribute(Qt.WidgetAttribute.WA_TransparentForMouseEvents)

        lay.addWidget(lbl_icon, 3)
        lay.addWidget(lbl_title, 2)

        shadow = QGraphicsDropShadowEffect()
        shadow.setBlurRadius(18)
        shadow.setColor(QColor(31, 78, 120, 80))
        shadow.setOffset(0, 4)
        btn.setGraphicsEffect(shadow)

        anim = QPropertyAnimation(btn, b"geometry", btn)
        anim.setDuration(165)
        anim.setEasingCurve(QEasingCurve.Type.OutCubic)

        def animate_to(target):
            try:
                anim.stop()
                anim.setStartValue(btn.geometry())
                anim.setEndValue(target)
                anim.start()
            except Exception as e:
                print("[DASH] animation error:", e)

        def on_enter():
            c = QColor(ring_on)
            btn.setStyleSheet(tile_style(ring_on, 3))
            shadow.setBlurRadius(34)
            shadow.setColor(QColor(c.red(), c.green(), c.blue(), 200))
            lbl_icon.setStyleSheet("font-size: 29pt; background: transparent;")
            animate_to(big_rect)

        def on_leave():
            btn.setStyleSheet(tile_style("rgba(255,255,255,120)", 1))
            shadow.setBlurRadius(18)
            shadow.setColor(QColor(31, 78, 120, 80))
            lbl_icon.setStyleSheet("font-size: 25pt; background: transparent;")
            animate_to(small_rect)

        btn.on_enter = on_enter
        btn.on_leave = on_leave
        btn.clicked.connect(lambda _, i=index: self.open_tab_index(i))
        btn.setToolTip(f"ورود به بخش {title}")
        btn._anim = anim
        return holder

    def setup_dashboard_ui(self):
        root = QVBoxLayout()
        root.setContentsMargins(14, 12, 14, 12)
        root.setSpacing(8)
        root.addStretch(1)

        # ---------------- سربرگ گرافیکی ----------------
        header = QFrame()
        header.setStyleSheet("""
            QFrame {
                background-color: qlineargradient(x1:0, y1:0, x2:1, y2:0,
                    stop:0 #1F4E78, stop:0.5 #2c6699, stop:1 #1F4E78);
                border-radius: 16px;
                border: 1px solid #2980b9;
            }
            QLabel { background: transparent; color: #ffffff; }
        """)
        card_shadow(header, blur=26, dy=6, alpha=90)
        h_lay = QVBoxLayout(header)
        h_lay.setContentsMargins(18, 12, 18, 12)
        h_lay.setSpacing(3)

        lbl_title = QLabel(f"{ico('commercial')} سیستم جامع مدیریت مالی و حسابداری IMART STUDIO")
        lbl_title.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_title.setFont(QFont(APP_FONT_FAMILY, 16, QFont.Weight.Bold))
        h_lay.addWidget(lbl_title)

        lbl_sub = QLabel(f"نسخه {APP_VERSION}  |  طراح: {DEVELOPER_NAME}  |  "
                         f"پشتیبانی: {get_setting('studio_phone', '09173736618')}")
        lbl_sub.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_sub.setStyleSheet("color: #cfe4f7; font-size: 10pt; background: transparent;")
        h_lay.addWidget(lbl_sub)

        wrap_header = QHBoxLayout()
        wrap_header.addStretch(1)
        wrap_header.addWidget(header, 8)
        wrap_header.addStretch(1)
        root.addLayout(wrap_header)
        root.addSpacing(4)

        # ---------------- نوار سال کاری با سوییچ قبل/بعد ----------------
        year_card = QFrame()
        year_card.setStyleSheet("""
            QFrame { background-color: #ffffff; border: 1px solid #d6e0ec; border-radius: 14px; }
            QLabel { background: transparent; }
        """)
        year_lay = QHBoxLayout(year_card)
        year_lay.setContentsMargins(12, 8, 12, 8)
        year_lay.setSpacing(7)

        lbl_year = QLabel(f"{ico('calendar')} سال کاری:")
        lbl_year.setFont(QFont(APP_FONT_FAMILY, 12, QFont.Weight.Bold))
        lbl_year.setStyleSheet("color:#1F4E78;")
        year_lay.addWidget(lbl_year)

        self.btn_year_prev = QPushButton(f"{ico('prev')} سال قبل")
        self.btn_year_prev.setToolTip("نمایش اطلاعات سال کاری قبل")
        self.btn_year_prev.setStyleSheet("background-color:#7f8c8d; color:white; font-weight:bold; padding:7px 11px;")
        self.btn_year_prev.clicked.connect(lambda: self.step_work_year(-1))
        year_lay.addWidget(self.btn_year_prev)

        self.combo_work_year = QComboBox()
        self.combo_work_year.setFont(QFont(APP_FONT_FAMILY, 12, QFont.Weight.Bold))
        self.combo_work_year.setFixedWidth(120)
        self.combo_work_year.setStyleSheet("padding:6px; border:2px solid #2980b9; border-radius:8px;")
        self.load_work_years()
        current_year = get_current_year()
        idx = self.combo_work_year.findText(str(current_year))
        if idx >= 0:
            self.combo_work_year.setCurrentIndex(idx)
        self.combo_work_year.currentTextChanged.connect(self.switch_work_year)
        year_lay.addWidget(self.combo_work_year)

        self.btn_year_next = QPushButton(f"سال بعد {ico('next')}")
        self.btn_year_next.setToolTip("نمایش اطلاعات سال کاری بعد")
        self.btn_year_next.setStyleSheet("background-color:#7f8c8d; color:white; font-weight:bold; padding:7px 11px;")
        self.btn_year_next.clicked.connect(lambda: self.step_work_year(1))
        year_lay.addWidget(self.btn_year_next)

        btn_new_year = QPushButton(ico_text("add", "شروع سال کاری جدید"))
        btn_new_year.setStyleSheet("background-color:#27ae60; color:white; font-weight:bold; padding:7px 11px;")
        btn_new_year.clicked.connect(self.start_new_work_year)
        year_lay.addWidget(btn_new_year)

        year_lay.addStretch()

        wrap_year = QHBoxLayout()
        wrap_year.addStretch(1)
        wrap_year.addWidget(year_card, 9)
        wrap_year.addStretch(1)
        root.addLayout(wrap_year)

        # خلاصه وضعیت سال کاری در یک ردیف جداگانه تا کامل دیده شود
        self.lbl_dash_summary = QLabel("")
        self.lbl_dash_summary.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.lbl_dash_summary.setWordWrap(True)
        self.lbl_dash_summary.setStyleSheet(
            "color:#1F4E78; font-size:9.5pt; font-weight:bold; "
            "background-color:#ffffff; border:1px solid #d6e0ec; border-radius:10px; padding:6px;")
        root.addWidget(self.lbl_dash_summary)
        root.addSpacing(4)

        # ---------------- ۱۲ آیتم در دو ردیف ۶ تایی وسط‌چین ----------------
        buttons = [
            (ico('wedding'),    "پروژه‌های عروس و داماد\n(قراردادها)", 0,  "#8e44ad"),
            (ico('commercial'), "تبلیغاتی / بیوتی / تولدی\n(پروژه‌ها)", 1,  "#16a085"),
            (ico('staff'),      "کارکنان و حقوق\n(پرسنل)", 2,             "#2980b9"),
            (ico('expenses'),   "مدیریت هزینه‌ها\n(قبوض و...)", 3,         "#c0392b"),
            (ico('reports'),    "گزارش مالی جامع\n(نمودارها)", 4,         "#d35400"),
            (ico('search'),     "جستجوی پیشرفته\n(کلی)", 5,               "#00838f"),
            (ico('inventory'),  "انبار تجهیزات\n(موجودی)", 6,             "#2c3e50"),
            (ico('banks'),      "کارت‌های بانکی\n(بانک‌ها)", 7,            "#0e6655"),
            (ico('checks'),     "مدیریت چک‌ها\n(دریافتی/پرداختی)", 8,     "#e67e22"),
            (ico('receipt'),    "صدور رسید وجه\n(خدمات)", 9,              "#7f8c8d"),
            (ico('settings'),   "تنظیمات و سال کاری\n(سیستم)", 10,        "#5d6d7e"),
            (ico('about'),      "درباره برنامه\n(توسعه‌دهنده)", 11,       "#34495e"),
        ]

        grid = QGridLayout()
        grid.setSpacing(10)
        grid.setContentsMargins(0, 0, 0, 0)
        for i, (icon, text, index, color) in enumerate(buttons):
            r, c = divmod(i, 6)
            grid.addWidget(self.make_dashboard_button(icon, text, color, index), r, c)

        grid_wrap = QHBoxLayout()
        grid_wrap.addStretch(1)
        grid_wrap.addLayout(grid)
        grid_wrap.addStretch(1)
        root.addLayout(grid_wrap)

        root.addStretch(1)

        footer = QLabel(f"{ico('money')} IMART STUDIO  |  تلفن: 09173736618  |  نسخه {APP_VERSION}")
        footer.setAlignment(Qt.AlignmentFlag.AlignCenter)
        footer.setStyleSheet("color:#5d6d7e; font-size:9pt; padding-top:4px;")
        root.addWidget(footer)

        self.dashboard_screen.setLayout(root)
        self.refresh_dashboard_summary()

    def refresh_dashboard_summary(self):
        """خلاصه وضعیت سال کاری روی داشبورد"""
        try:
            if not hasattr(self, "lbl_dash_summary"):
                return
            wy = get_current_year()
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT COUNT(*), COALESCE(SUM(total_amount),0) FROM wedding_contracts WHERE work_year=?", (wy,))
            c_count, c_sum = cursor.fetchone()
            cursor.execute("SELECT COUNT(*) FROM commercial_projects WHERE work_year=?", (wy,))
            p_count = cursor.fetchone()[0]
            cursor.execute("SELECT COUNT(*) FROM checks WHERE work_year=? AND is_passed=0", (wy,))
            open_checks = cursor.fetchone()[0]
            conn.close()
            self.lbl_dash_summary.setText(
                f"قرارداد: {c_count}  |  پروژه: {p_count}  |  "
                f"جمع قراردادها: {c_sum:,} تومان  |  چک باز: {open_checks}"
            )
        except Exception as e:
            print("[DASH] summary error:", e)

    def step_work_year(self, direction):
        """سوییچ سریع به سال کاری قبل یا بعد"""
        if not hasattr(self, "combo_work_year"):
            return
        idx = self.combo_work_year.currentIndex() + direction
        if 0 <= idx < self.combo_work_year.count():
            self.combo_work_year.setCurrentIndex(idx)
        else:
            QMessageBox.information(self, "سال کاری",
                                    "سال کاری دیگری در لیست موجود نیست.\n"
                                    "برای افزودن سال جدید از دکمه «شروع سال کاری جدید» استفاده کنید.")

    def load_work_years(self):
        """بارگذاری لیست سال‌های کاری - با محافظ"""
        if not hasattr(self, 'combo_work_year'):
            return
        self.combo_work_year.blockSignals(True)
        current = self.combo_work_year.currentText()
        self.combo_work_year.clear()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        years = set()
        years.add(jdatetime.date.today().year)
        try:
            years.add(get_current_year())
        except Exception:
            pass

        for table in ["wedding_contracts", "commercial_projects", "expenses",
                      "transactions", "checks", "wedding_deposits", "persons"]:
            try:
                cursor.execute(f"SELECT DISTINCT work_year FROM {table} WHERE work_year IS NOT NULL")
                for (y,) in cursor.fetchall():
                    if y:
                        years.add(int(y))
            except Exception:
                pass

        conn.close()

        for y in sorted(years, reverse=True):
            self.combo_work_year.addItem(str(y))

        # سال جدید همیشه قابل انتخاب باشد (سوییچ به سال بعد از آخرین سال)
        try:
            newest = max(years) if years else jdatetime.date.today().year
            self.combo_work_year.insertItem(0, str(newest + 1))
        except Exception:
            pass

        target = current or str(get_current_year())
        idx = self.combo_work_year.findText(target)
        if idx >= 0:
            self.combo_work_year.setCurrentIndex(idx)
        self.combo_work_year.blockSignals(False)

    def switch_work_year(self, year_str):
        """تغییر سال با محافظ"""
        try:
            year = int(year_str)
        except Exception:
            return
        set_setting("current_work_year", year)
        recalc_inventory_usage()

        # ⚠️ فقط اگه پنجره اصلی ساخته شده، رفرش کن
        if hasattr(self, 'tabs') and hasattr(self, 'w_table'):
            try:
                self.load_wedding_contracts()
                self.load_commercial_projects()
                self.load_expenses_table()
                self.load_staff_table()
                self.load_persons_combo()
                self.load_person_events()
                self.load_inventory()
                self.load_checks_table()
                self.load_report_persons_combo()
                self.calculate_financial_report()
                self.refresh_settlement_cards()
            except Exception as e:
                print("Refresh error:", e)

        if hasattr(self, 'combo_settings_year'):
            try:
                self.load_settings_years()
            except Exception:
                pass

        self.refresh_dashboard_summary()

        if hasattr(self, 'tabs'):
            QMessageBox.information(self, "تغییر سال", f"سال کاری به {year} تغییر یافت.")

    def open_tab_index(self, index):
        if hasattr(self, 'tabs'):
            self.tabs.setCurrentIndex(index)
            self.stack.setCurrentWidget(self.main_app_screen)

    def _output_flow(self, builder, default_filename, title="سند",
                     allow_details=True, allow_staff=False):
        """
        مسیر مشترک خروجی همه اسناد:
        انتخاب چاپ (سیاه و سفید) یا PDF (رنگی)، با/بدون ریز پکیج، با/بدون نیروی کار.
        همه اسناد روی یک برگ A5 چاپ می‌شوند.
        """
        dlg = PrintChoiceDialog(self, title, allow_details, allow_staff)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            return
        is_print = (dlg.mode == "print")
        try:
            html = builder(dlg.with_details, dlg.with_staff, is_print)
        except TypeError:
            try:
                html = builder(dlg.with_details, dlg.with_staff)
            except TypeError:
                html = builder()
        if not html:
            QMessageBox.warning(self, "خطا", "سندی برای چاپ یافت نشد.")
            return

        if dlg.mode == "pdf":
            path, _ = QFileDialog.getSaveFileName(self, "ذخیره PDF", default_filename,
                                                  "PDF Files (*.pdf)")
            if not path:
                return
            ok, err = save_html_pdf(html, path, title)
            if ok:
                QMessageBox.information(self, "موفقیت", f"فایل PDF رنگی ذخیره شد:\n{path}")
            else:
                QMessageBox.critical(self, "خطا", f"ساخت PDF ناموفق بود:\n{err}")
        else:
            print_html_document(html, self, title)

    def _contract_print_flow(self, contract_id):
        _html, filename = InvoiceBuilder.build_wedding_invoice(contract_id)
        if not _html:
            QMessageBox.warning(self, "خطا", "قرارداد مورد نظر یافت نشد.")
            return
        self._output_flow(
            lambda wd, ws, mono: InvoiceBuilder.build_wedding_invoice(contract_id, wd, ws, mono)[0],
            filename, "فاکتور قرارداد عروس و داماد", allow_details=True, allow_staff=True)

    def _commercial_print_flow(self, project_id):
        html, filename = InvoiceBuilder.build_commercial_invoice(project_id)
        if not html:
            QMessageBox.warning(self, "خطا", "پروژه مورد نظر یافت نشد.")
            return
        self._output_flow(
            lambda wd, ws, mono: InvoiceBuilder.build_commercial_invoice(project_id, wd, ws, mono)[0],
            filename, "فاکتور پروژه تبلیغاتی", allow_details=False, allow_staff=False)

    def export_table_excel(self, table, title, default_name):
        """خروجی اکسل از یک جدول برنامه (openpyxl)"""
        path, _ = QFileDialog.getSaveFileName(self, "خروجی اکسل", default_name, "Excel Files (*.xlsx)")
        if not path:
            return
        headers = []
        for c in range(table.columnCount()):
            h = table.horizontalHeaderItem(c)
            headers.append(h.text() if h else f"ستون {c+1}")
        rows = []
        for r in range(table.rowCount()):
            row = []
            for c in range(table.columnCount()):
                if table.cellWidget(r, c) is not None:
                    row.append("")
                    continue
                it = table.item(r, c)
                txt = it.text() if it else ""
                raw = txt.replace(',', '').strip()
                if raw.lstrip('-').isdigit():
                    row.append(int(raw))
                else:
                    row.append(txt)
            rows.append(row)
        ok, err = export_rows_to_excel(headers, rows, path, title, title)
        if ok:
            QMessageBox.information(self, "موفقیت", f"خروجی اکسل ذخیره شد:\n{path}")
        else:
            QMessageBox.critical(self, "خطا", f"خروجی اکسل ناموفق بود:\n{err}")

    def exit_application(self):
        """
        دکمه خروج: دقیقاً مانند زدن ضربدر پنجره عمل می‌کند؛
        یعنی ابتدا منوی بک‌آپ باز می‌شود و بعد از ذخیره، برنامه بسته می‌شود.
        """
        self.close()

    def setup_main_app_ui(self):
        main_layout = QVBoxLayout()
        main_layout.setContentsMargins(10, 8, 10, 8)
        main_layout.setSpacing(8)

        # ---------------- نوار ابزار بالای برنامه ----------------
        top_card = QFrame()
        top_card.setStyleSheet("""
            QFrame {
                background-color: qlineargradient(x1:0, y1:0, x2:0, y2:1,
                    stop:0 #ffffff, stop:1 #eef4fb);
                border: 1px solid #d6e0ec;
                border-radius: 12px;
            }
        """)
        card_shadow(top_card, blur=16, dy=3, alpha=50)
        top_bar = QHBoxLayout(top_card)
        top_bar.setContentsMargins(12, 9, 12, 9)
        top_bar.setSpacing(8)

        def add_top(icon_key, text, slot, color, tip=""):
            b = QPushButton(ico_text(icon_key, text))
            b.setStyleSheet(
                f"background-color:{color}; color:white; font-weight:bold; padding:9px 13px;")
            b.clicked.connect(slot)
            if tip:
                b.setToolTip(tip)
            top_bar.addWidget(b)
            return b

        add_top("home", "منوی اصلی", lambda: self.stack.setCurrentWidget(self.dashboard_screen),
                "#34495e", "بازگشت به داشبورد اصلی")
        add_top("edit", "عنوان‌های پیش‌فرض", self.edit_default_titles, "#e67e22",
                "تغییر نام استودیو، تلفن و آدرس روی فاکتورها")
        add_top("backup", "پشتیبان‌گیری", self.backup_db, "#16a085",
                "ذخیره یک نسخه پشتیبان از اطلاعات")
        add_top("restore", "بازیابی بک‌آپ", self.restore_db, "#8e44ad",
                "برگرداندن اطلاعات از فایل پشتیبان")
        add_top("key", "تغییر رمز", self.change_password, "#c0392b",
                "تغییر رمز عبور برنامه")
        add_top("cancel", "خروج از برنامه", self.exit_application, "#7b241c",
                "ذخیره نسخه پشتیبان و بستن برنامه (مانند زدن ضربدر پنجره)")

        top_bar.addStretch()

        self.lbl_top_year = QLabel("")
        self.lbl_top_year.setStyleSheet("color:#1F4E78; font-weight:bold; font-size:11pt;")
        top_bar.addWidget(self.lbl_top_year)

        main_layout.addWidget(top_card)

        # ---------------- تب‌ها ----------------
        self.tabs = QTabWidget()
        self.tabs.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        self.tabs.setDocumentMode(True)

        self.tab_wedding = QWidget()
        self.tab_commercial = QWidget()
        self.tab_staff = QWidget()
        self.tab_expenses = QWidget()
        self.tab_reports = QWidget()
        self.tab_search = QWidget()
        self.tab_inventory = QWidget()
        self.tab_banks = QWidget()
        self.tab_checks = QWidget()
        self.tab_receipt = QWidget()
        self.tab_workyear = QWidget()
        self.tab_about = QWidget()

        self.tabs.addTab(self.tab_wedding, ico_text("wedding", "پروژه‌های عروس و داماد"))
        self.tabs.addTab(self.tab_commercial, ico_text("commercial", "تبلیغاتی / بیوتی"))
        self.tabs.addTab(self.tab_staff, ico_text("staff", "کارکنان و پرسنل"))
        self.tabs.addTab(self.tab_expenses, ico_text("expenses", "مدیریت هزینه‌ها"))
        self.tabs.addTab(self.tab_reports, ico_text("reports", "گزارش مالی جامع"))
        self.tabs.addTab(self.tab_search, ico_text("search", "جستجوی پیشرفته"))
        self.tabs.addTab(self.tab_inventory, ico_text("inventory", "انبار تجهیزات"))
        self.tabs.addTab(self.tab_banks, ico_text("banks", "کارت‌های بانکی"))
        self.tabs.addTab(self.tab_checks, ico_text("checks", "مدیریت چک‌ها"))
        self.tabs.addTab(self.tab_receipt, ico_text("receipt", "صدور رسید وجه"))
        self.tabs.addTab(self.tab_workyear, ico_text("settings", "تنظیمات و سال کاری"))
        self.tabs.addTab(self.tab_about, ico_text("about", "درباره برنامه"))

        main_layout.addWidget(self.tabs, 1)

        self.setup_wedding_tab()
        self.setup_commercial_tab()
        self.setup_staff_tab()
        self.setup_expenses_tab()
        self.setup_reports_tab()
        self.setup_search_tab()
        self.setup_inventory_tab()
        self.setup_banks_tab()
        self.setup_checks_tab()
        self.setup_receipt_tab()
        self.setup_workyear_tab()
        self.setup_about_tab()

        footer = QLabel(f"{ico('commercial')} IMART STUDIO  |  تلفن: 09173736618 "
                        f" |  نسخه {APP_VERSION}  |  طراح: {DEVELOPER_NAME}")
        footer.setAlignment(Qt.AlignmentFlag.AlignCenter)
        footer.setFont(QFont(APP_FONT_FAMILY, 9))
        footer.setStyleSheet("color: #5d6d7e; padding-top: 3px;")
        main_layout.addWidget(footer)

        self.main_app_screen.setLayout(main_layout)
        self.update_top_year_label()

    def update_top_year_label(self):
        try:
            if hasattr(self, "lbl_top_year"):
                self.lbl_top_year.setText(f"{ico('calendar')} سال کاری فعلی: {get_current_year()}")
            self.refresh_dashboard_summary()
        except Exception:
            pass

    def edit_default_titles(self):
        if not ask_security_password(self):
            return
        dlg = QDialog(self)
        dlg.setWindowTitle("تغییر عنوان‌های پیش‌فرض")
        dlg.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        dlg.resize(560, 460)
        dlg.setFont(QFont(APP_FONT_FAMILY, 10))

        v = QVBoxLayout()
        form = QFormLayout()

        txt_studio = QLineEdit(get_setting("studio_name", "IMART STUDIO"))
        txt_phone = QLineEdit(get_setting("studio_phone", "09173736618"))
        txt_address = QLineEdit(get_setting("studio_address", ""))

        form.addRow(f"{ico('commercial')} نام استودیو:", txt_studio)
        form.addRow(f"{ico('people')} تلفن:", txt_phone)
        form.addRow(f"{ico('home')} آدرس:", txt_address)

        v.addLayout(form)

        info = QLabel(f"{ico('detail')} <b>راهنما:</b> نام استودیو، تلفن و آدرس در فاکتورها، "
                      "رسیدها و گزارش‌های چاپی نمایش داده می‌شود.")
        info.setStyleSheet("background-color: #e8f8f5; padding: 10px; border-radius: 8px;")
        info.setWordWrap(True)
        v.addWidget(info)

        def save_titles():
            if not ask_security_password(self):
                return
            set_setting("studio_name", txt_studio.text())
            set_setting("studio_phone", txt_phone.text())
            set_setting("studio_address", txt_address.text())
            QMessageBox.information(dlg, "موفقیت", "عنوان‌ها به‌روزرسانی شد.")
            dlg.accept()

        btn = QPushButton(ico_text("save", "ذخیره"))
        btn.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 10px;")
        btn.clicked.connect(save_titles)
        v.addWidget(btn)

        dlg.setLayout(v)
        dlg.exec()

    # ============================================================
    # ==============  تب ۱: عروس و داماد  =======================
    # ============================================================
    def setup_wedding_tab(self):
        # ==================== فرم ثبت قرارداد ====================
        form_scroll = QScrollArea()
        form_scroll.setWidgetResizable(True)
        form_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        form_holder = QWidget()
        form_holder_layout = QVBoxLayout(form_holder)
        form_holder_layout.setContentsMargins(0, 0, 0, 0)

        form_box = QGroupBox("ثبت قرارداد جدید عروس و داماد")
        form_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        form_layout = QFormLayout()
        form_layout.setSpacing(5)

        self.w_groom = QLineEdit()
        self.w_bride = QLineEdit()
        self.w_groom_phone = QLineEdit()
        self.w_bride_phone = QLineEdit()
        self.w_contract_date = PersianDateEdit(date_str=jalali_today_str())
        self.w_ceremony_date = PersianDateEdit(date_str=jalali_today_str())
        self.w_venue = QLineEdit()
        self.w_venue.setPlaceholderText("نام تالار / باغ / مکان برگزاری مراسم")

        form_layout.addRow(f"{ico('wedding')} نام داماد <span style='color:red;'>*</span>:", self.w_groom)
        form_layout.addRow(f"{ico('wedding')} نام عروس <span style='color:red;'>*</span>:", self.w_bride)
        form_layout.addRow(f"{ico('people')} تلفن داماد:", self.w_groom_phone)
        form_layout.addRow(f"{ico('people')} تلفن عروس:", self.w_bride_phone)
        form_layout.addRow(f"{ico('calendar')} تاریخ قرارداد:", self.w_contract_date)
        form_layout.addRow(f"{ico('calendar')} تاریخ مراسم:", self.w_ceremony_date)
        form_layout.addRow(f"{ico('home')} مکان مراسم:", self.w_venue)

        btn_manage_items = QPushButton(ico_text("package", "مدیریت / افزودن پکیج، کالا و تجهیزات"))
        btn_manage_items.setStyleSheet("background-color: #e67e22; color: white; font-weight: bold; padding: 7px;")
        btn_manage_items.clicked.connect(self.open_manage_items)
        form_layout.addRow(btn_manage_items)

        # ---- لیست اقلام: بدون اسکرول افقی، ردیف‌ها ریز و متن‌ها کوتاه می‌شوند ----
        scroll_area = QScrollArea()
        scroll_area.setWidgetResizable(True)
        scroll_area.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        scroll_area.setFixedHeight(230)

        self.items_container = QWidget()
        self.items_vbox = QVBoxLayout()
        self.items_vbox.setContentsMargins(0, 0, 0, 0)
        self.items_vbox.setSpacing(1)
        self.items_container.setLayout(self.items_vbox)
        scroll_area.setWidget(self.items_container)

        items_box = QGroupBox("مفاد فاکتور (انتخاب موارد — روی پکیج دوبار کلیک کنید)")
        items_box_layout = QVBoxLayout()
        items_box_layout.addWidget(scroll_area)
        items_box.setLayout(items_box_layout)
        form_layout.addRow(items_box)

        self.w_lbl_total = QLabel("")
        self.w_lbl_total.setWordWrap(True)
        self.w_lbl_total.setStyleSheet("background-color:#e8f8f5; border:1px solid #16a085; "
                                       "border-radius:8px; padding:8px; font-size:10pt;")
        form_layout.addRow(self.w_lbl_total)

        self.w_discount = QLineEdit("0")
        self.w_discount.textChanged.connect(lambda t: self.on_discount_changed(t))
        self.w_first_deposit = QLineEdit("0")
        self.w_first_deposit.textChanged.connect(lambda t: self.on_deposit_changed(t))

        form_layout.addRow(f"{ico('money')} تخفیف (تومان):", self.w_discount)
        form_layout.addRow(f"{ico('deposit')} بیعانه اول (تومان):", self.w_first_deposit)

        self.w_bank_combo = QComboBox()
        self.load_bank_combo()
        form_layout.addRow(f"{ico('banks')} بانک واریزی بیعانه:", self.w_bank_combo)

        self.w_lbl_remain = QLabel("")
        self.w_lbl_remain.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        self.w_lbl_remain.setStyleSheet("color: #c0392b; font-size: 10pt; padding:3px;")
        form_layout.addRow(self.w_lbl_remain)

        # ---------- نیروی کار قرارداد ----------
        staff_box = QGroupBox("نیروی کار این قرارداد (تیک بزنید و حقوق همان قرارداد را وارد کنید)")
        staff_box_layout = QVBoxLayout()
        staff_scroll = QScrollArea()
        staff_scroll.setWidgetResizable(True)
        staff_scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)
        staff_scroll.setFixedHeight(140)
        self.staff_container = QWidget()
        self.staff_vbox = QVBoxLayout()
        self.staff_vbox.setContentsMargins(0, 0, 0, 0)
        self.staff_vbox.setSpacing(1)
        self.staff_container.setLayout(self.staff_vbox)
        staff_scroll.setWidget(self.staff_container)
        staff_box_layout.addWidget(staff_scroll)
        staff_box.setLayout(staff_box_layout)
        form_layout.addRow(staff_box)

        self.w_desc = QTextEdit()
        self.w_desc.setFixedHeight(54)
        self.w_desc.setPlaceholderText("توضیحات کامل قرارداد...")
        form_layout.addRow(f"{ico('edit')} توضیحات قرارداد:", self.w_desc)

        self.item_checkboxes = {}
        self.item_price_inputs = {}
        self.item_is_pkg = {}
        self.staff_checks = {}
        self.staff_wages = {}
        self.staff_meta = {}
        self.load_item_checkboxes()
        self.load_staff_checkboxes()

        self.btn_save_wedding = QPushButton(ico_text("save", "ثبت نهایی قرارداد و صدور کد یکتا"))
        self.btn_save_wedding.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        self.btn_save_wedding.setStyleSheet("background-color: #2980b9; color: white; padding: 10px;")
        self.btn_save_wedding.clicked.connect(self.save_wedding_contract)
        form_layout.addRow(self.btn_save_wedding)

        form_box.setLayout(form_layout)
        form_holder_layout.addWidget(form_box)
        form_scroll.setWidget(form_holder)

        # ==================== جدول قراردادها ====================
        table_box = QGroupBox("لیست قراردادها (برای جزئیات دوبار کلیک کنید)")
        table_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        table_layout = QVBoxLayout()

        tools = QHBoxLayout()
        btn_refresh = QPushButton(ico_text("refresh", "بروزرسانی"))
        btn_refresh.clicked.connect(self.load_wedding_contracts)
        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color:#1e8449; color:white; font-weight:bold; padding:6px;")
        btn_excel.clicked.connect(lambda: self.export_table_excel(
            self.w_table, "قراردادهای عروس و داماد",
            f"قراردادها_{get_current_year()}.xlsx"))
        tools.addWidget(btn_refresh)
        tools.addWidget(btn_excel)

        tools.addWidget(QLabel(f"{ico('filter')} دسته‌بندی:"))
        self.combo_sort = QComboBox()
        self.combo_sort.addItems([
            "ترتیب عادی (جدیدترین)",
            "نزدیک‌ترین تاریخ مراسم",
            "تسویه‌نشده‌ها",
            "مبلغ: کم به زیاد",
            "مبلغ: زیاد به کم",
        ])
        self.combo_sort.setToolTip("ترتیب و دسته‌بندی نمایش قراردادها")
        self.combo_sort.currentIndexChanged.connect(self.load_wedding_contracts)
        tools.addWidget(self.combo_sort)
        tools.addStretch()
        table_layout.addLayout(tools)

        self.w_table = QTableWidget()
        self.w_table.setColumnCount(17)
        self.w_table.setHorizontalHeaderLabels([
            "کد", "زوجین", "تلفن داماد", "تلفن عروس", "تاریخ مراسم", "مکان مراسم",
            "روزشمار مراسم", "جمع خام", "تخفیف", "جمع بعد از تخفیف", "دریافتی", "مانده",
            "وضعیت تسویه", "بیعانه‌ها", "چاپ / PDF", "ویرایش", "حذف"
        ])
        self.w_table.setAlternatingRowColors(True)
        self.w_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.w_table.verticalHeader().setDefaultSectionSize(32)
        self.w_table.doubleClicked.connect(self.on_wedding_double_click)
        persist_table(self.w_table, "wedding_contracts", QHeaderView.ResizeMode.ResizeToContents)
        table_layout.addWidget(self.w_table)
        table_box.setLayout(table_layout)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        form_holder_wrap = QWidget()
        fw = QVBoxLayout(form_holder_wrap)
        fw.setContentsMargins(0, 0, 0, 0)
        fw.addWidget(form_scroll)
        table_wrap = QWidget()
        tw = QVBoxLayout(table_wrap)
        tw.setContentsMargins(0, 0, 0, 0)
        tw.addWidget(table_box)
        splitter.addWidget(form_holder_wrap)
        splitter.addWidget(table_wrap)
        splitter.setSizes([560, 840])
        persist_splitter(splitter, "tab_wedding")

        outer = QHBoxLayout()
        outer.addWidget(splitter)
        self.tab_wedding.setLayout(outer)
        self.load_wedding_contracts()

    def open_manage_items(self):
        dlg = ManageItemsDialog(self)
        dlg.exec()
        self.load_item_checkboxes()
        self.load_inventory()

    def show_package_details_from_tab(self, name):
        dlg = PackageDetailsDialog(name, self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            self.load_item_checkboxes()

    def load_item_checkboxes(self):
        for i in reversed(range(self.items_vbox.count())):
            widget = self.items_vbox.itemAt(i).widget()
            if widget:
                widget.setParent(None)

        self.item_checkboxes = {}
        self.item_price_inputs = {}
        self.item_is_pkg = {}

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT item_name, code, price, is_package, parent_package FROM item_prices "
                       "ORDER BY is_package DESC, parent_package, item_name")
        rows = cursor.fetchall()
        conn.close()

        packages = [r for r in rows if r[3] == 1]
        singles = [r for r in rows if r[3] == 0]

        def add_row(item_name, code, price, is_pkg, indent=False):
            row_w = QWidget()
            row_w.setFixedHeight(26)
            row_l = QHBoxLayout(row_w)
            row_l.setContentsMargins(16 if indent else 0, 0, 0, 0)
            row_l.setSpacing(5)

            cb = ElidedCheckBox(f"{'• ' if indent else ''}[{code or '-'}] {item_name}")
            cb.setFont(QFont(APP_FONT_FAMILY, 9,
                             QFont.Weight.Normal if indent else QFont.Weight.Bold))
            cb.setToolTip(f"{item_name}  ({price:,} تومان)")
            cb.stateChanged.connect(self.calc_wedding_total)

            txt_p = QLineEdit(f"{price:,}")
            txt_p.setFixedWidth(90)
            txt_p.setFixedHeight(22)
            txt_p.setFont(QFont(APP_FONT_FAMILY, 9))
            txt_p.setToolTip("قیمت این مورد در این قرارداد")
            txt_p.textChanged.connect(lambda t, p=txt_p: p.setText(format_number(t)))
            txt_p.textChanged.connect(self.calc_wedding_total)

            row_l.addWidget(cb, 1)
            if is_pkg:
                btn_det = QPushButton(ico("detail"))
                btn_det.setFixedSize(24, 22)
                btn_det.setToolTip("مدیریت و جزئیات کامل این پکیج")
                btn_det.clicked.connect(lambda _, n=item_name: self.show_package_details_from_tab(n))
                row_l.addWidget(btn_det)
            row_l.addWidget(txt_p)
            row_w.setLayout(row_l)

            self.items_vbox.addWidget(row_w)
            self.item_checkboxes[item_name] = cb
            self.item_price_inputs[item_name] = txt_p
            self.item_is_pkg[item_name] = is_pkg

        for pkg in packages:
            add_row(pkg[0], pkg[1], pkg[2], 1, indent=False)
            for sub in [s for s in singles if s[4] == pkg[0]]:
                add_row(sub[0], sub[1], sub[2], 0, indent=True)

        for ind in [s for s in singles if not s[4]]:
            add_row(ind[0], ind[1], ind[2], 0, indent=False)

        self.items_vbox.addStretch()
        self.calc_wedding_total()

    def load_staff_checkboxes(self):
        """بارگذاری افراد بخش کارکنان برای انتخاب در قرارداد"""
        if not hasattr(self, "staff_vbox"):
            return
        for i in reversed(range(self.staff_vbox.count())):
            widget = self.staff_vbox.itemAt(i).widget()
            if widget:
                widget.setParent(None)

        self.staff_checks = {}
        self.staff_wages = {}
        self.staff_meta = {}

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT id, code, name, role FROM persons ORDER BY name")
        persons = cursor.fetchall()
        conn.close()

        if not persons:
            lbl = QLabel("هنوز نیرویی در بخش «کارکنان» تعریف نشده است. ابتدا از تب کارکنان نیرو اضافه کنید.")
            lbl.setStyleSheet("color:#7f8c8d; padding:6px;")
            lbl.setWordWrap(True)
            self.staff_vbox.addWidget(lbl)
            return

        for pid, pcode, pname, prole in persons:
            row_w = QWidget()
            row_w.setFixedHeight(26)
            row_l = QHBoxLayout(row_w)
            row_l.setContentsMargins(0, 0, 0, 0)
            row_l.setSpacing(5)
            cb = ElidedCheckBox(f"[{pcode or '-'}] {pname} ({prole})")
            cb.setFont(QFont(APP_FONT_FAMILY, 9))
            cb.stateChanged.connect(self.calc_wedding_total)
            wage = QLineEdit("0")
            wage.setFixedWidth(90)
            wage.setFixedHeight(22)
            wage.setFont(QFont(APP_FONT_FAMILY, 9))
            wage.setToolTip("حقوق این فرد در این قرارداد")
            wage.textChanged.connect(lambda t, p=wage: p.setText(format_number(t)))
            wage.textChanged.connect(self.calc_wedding_total)
            self.staff_checks[pid] = cb
            self.staff_wages[pid] = wage
            self.staff_meta[pid] = (pname, prole)
            row_l.addWidget(cb, 1)
            row_l.addWidget(wage)
            row_w.setLayout(row_l)
            self.staff_vbox.addWidget(row_w)

        self.staff_vbox.addStretch()

    def on_discount_changed(self, text):
        formatted = format_number(text)
        if formatted != text:
            self.w_discount.setText(formatted)
        self.calc_wedding_total()

    def on_deposit_changed(self, text):
        formatted = format_number(text)
        if formatted != text:
            self.w_first_deposit.setText(formatted)
        self.calc_wedding_total()

    def calc_wedding_total(self):
        if not hasattr(self, 'w_discount') or not hasattr(self, 'w_first_deposit'):
            return

        raw_sum = 0
        count_items = 0
        for item_name, cb in self.item_checkboxes.items():
            if cb.isChecked():
                count_items += 1
                raw_sum += parse_number(self.item_price_inputs[item_name].text())

        discount = parse_number(self.w_discount.text())
        deposit = parse_number(self.w_first_deposit.text())
        net_total = max(0, raw_sum - discount)
        remain = max(0, net_total - deposit)

        staff_total = 0
        staff_count = 0
        for pid, cb in getattr(self, "staff_checks", {}).items():
            if cb.isChecked():
                staff_count += 1
                staff_total += parse_number(self.staff_wages[pid].text())

        self.w_lbl_total.setText(
            f"<b>تعداد خدمات انتخاب‌شده:</b> {count_items} | "
            f"<b>جمع قیمت خام:</b> {raw_sum:,} تومان<br>"
            f"<b>تخفیف:</b> {discount:,} تومان | "
            f"<b>جمع کل بعد از تخفیف:</b> <span style='color:#2980b9;'>{net_total:,} تومان</span><br>"
            f"<b>مبلغ به حروف:</b> {number_to_persian_words(net_total)} تومان"
        )
        self.w_lbl_remain.setText(
            f"بیعانه: {deposit:,} تومان | باقی‌مانده حساب: {remain:,} تومان | "
            f"نیروی کار: {staff_count} نفر / {staff_total:,} تومان"
        )

    def load_bank_combo(self):
        self.w_bank_combo.clear()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT id, bank_name FROM bank_cards")
        rows = cursor.fetchall()
        conn.close()
        if not rows:
            self.w_bank_combo.addItem("--- بدون بانک ثبت‌شده ---", None)
        for b_id, b_name in rows:
            self.w_bank_combo.addItem(f"{ico('banks')} {b_name}", b_id)

    def save_wedding_contract(self):
        groom = self.w_groom.text().strip()
        bride = self.w_bride.text().strip()

        if not groom or not bride:
            QMessageBox.warning(self, "خطا", "لطفاً نام عروس و داماد را وارد کنید.")
            return

        selected_items = [item for item, cb in self.item_checkboxes.items() if cb.isChecked()]
        if not selected_items:
            QMessageBox.warning(self, "خطا", "حداقل یک مورد از فاکتور را انتخاب کنید.")
            return

        items_detail = {}
        raw_total = 0
        for it in selected_items:
            p = parse_number(self.item_price_inputs[it].text())
            items_detail[it] = p
            raw_total += p

        discount = parse_number(self.w_discount.text())
        net_total = max(0, raw_total - discount)
        first_deposit = parse_number(self.w_first_deposit.text())
        bank_id = self.w_bank_combo.currentData()
        bank_name = self.w_bank_combo.currentText().replace(ico('banks'), "").strip() if bank_id else ""
        desc = self.w_desc.toPlainText().strip()
        venue = self.w_venue.text().strip()
        work_year = get_current_year()

        staff_list = []
        for pid, cb in self.staff_checks.items():
            if cb.isChecked():
                pname, prole = self.staff_meta.get(pid, ("-", "-"))
                staff_list.append({
                    "id": pid, "name": pname, "role": prole,
                    "wage": parse_number(self.staff_wages[pid].text()),
                })

        remain = net_total - first_deposit
        is_settled = 1 if remain <= 0 else 0

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        # کد یکتا: هیچ دو قراردادی کد یکسان ندارند (حتی بعد از حذف رکوردها)
        code = generate_unique_code(f"W-{work_year}", "wedding_contracts", "code",
                                    start=1, width=4)

        cursor.execute('''
            INSERT INTO wedding_contracts
            (code, groom_name, bride_name, groom_phone, bride_phone, contract_date,
             ceremony_date, selected_items, total_amount, discount, paid_amount,
             bank_id, is_settled, description, work_year, items_detail, staff_ids, venue)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (code, groom, bride, self.w_groom_phone.text(), self.w_bride_phone.text(),
              self.w_contract_date.text(), self.w_ceremony_date.text(),
              ",".join(selected_items), raw_total, discount, first_deposit,
              bank_id, is_settled, desc, work_year,
              json.dumps(items_detail, ensure_ascii=False),
              json.dumps(staff_list, ensure_ascii=False), venue))

        contract_id = cursor.lastrowid
        if first_deposit > 0:
            cursor.execute('''
                INSERT INTO wedding_deposits (contract_id, amount, deposit_date, bank_name, work_year)
                VALUES (?, ?, ?, ?, ?)
            ''', (contract_id, first_deposit, self.w_contract_date.text(), bank_name, work_year))

        conn.commit()
        conn.close()

        # انبار، کارکنان و داشبورد فوراً هماهنگ می‌شوند
        recalc_inventory_usage()
        sync_contract_staff({
            "id": contract_id, "code": code, "groom_name": groom, "bride_name": bride,
            "ceremony_date": self.w_ceremony_date.text(), "staff": staff_list,
        }, work_year)

        QMessageBox.information(self, "موفقیت", f"قرارداد با کد یکتای {code} ثبت شد.")

        self.w_groom.clear()
        self.w_bride.clear()
        self.w_groom_phone.clear()
        self.w_bride_phone.clear()
        self.w_venue.clear()
        self.w_contract_date.set_date(jalali_today_str())
        self.w_ceremony_date.set_date(jalali_today_str())
        self.w_discount.setText("0")
        self.w_first_deposit.setText("0")
        self.w_desc.clear()
        for cb in self.item_checkboxes.values():
            cb.setChecked(False)
        for cb in self.staff_checks.values():
            cb.setChecked(False)
        for w in self.staff_wages.values():
            w.setText("0")

        self.load_wedding_contracts()
        self.load_inventory()
        self.load_work_years()
        self.load_staff_table()
        self.load_person_events()
        self.refresh_dashboard_summary()
        self.refresh_settlement_cards()

    def _sorted_contracts(self, rows):
        """اعمال دسته‌بندی انتخابی روی لیست قراردادها"""
        mode = self.combo_sort.currentText() if hasattr(self, "combo_sort") else "ترتیب عادی (جدیدترین)"
        data = []
        for row in rows:
            (c_id, code, groom, bride, g_phone, b_phone, cer_date, raw,
             discount, paid, settled, venue) = row
            raw = raw or 0
            discount = discount or 0
            paid = paid or 0
            net = max(0, raw - discount)
            state, remain = settlement_state(net, paid)
            data.append({
                "id": c_id, "code": code, "groom": groom, "bride": bride,
                "g_phone": g_phone, "b_phone": b_phone, "ceremony": cer_date,
                "venue": venue, "raw": raw, "discount": discount, "net": net,
                "paid": paid, "remain": remain, "state": state, "settled": settled,
            })

        today = jdatetime.date.today()

        def upcoming_key(d):
            jd = parse_jalali(d["ceremony"])
            if jd is None:
                return (2, 0)
            diff = (jd - today).days
            return (0, diff) if diff >= 0 else (1, -diff)

        if mode == "نزدیک‌ترین تاریخ مراسم":
            data.sort(key=upcoming_key)
        elif mode == "تسویه‌نشده‌ها":
            order = {"تسویه نشده": 0, "تسویه ناقص": 1, "تسویه کامل": 2}
            data.sort(key=lambda d: (order.get(d["state"], 3), upcoming_key(d)))
        elif mode == "مبلغ: کم به زیاد":
            data.sort(key=lambda d: d["net"])
        elif mode == "مبلغ: زیاد به کم":
            data.sort(key=lambda d: d["net"], reverse=True)
        else:
            data.sort(key=lambda d: d["id"], reverse=True)
        return data

    def load_wedding_contracts(self):
        if not hasattr(self, "w_table"):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        work_year = get_current_year()
        cursor.execute('''
            SELECT id, code, groom_name, bride_name, groom_phone, bride_phone,
                   ceremony_date, total_amount, discount, paid_amount, is_settled, venue
            FROM wedding_contracts WHERE work_year=? ORDER BY id DESC
        ''', (work_year,))
        rows = cursor.fetchall()
        conn.close()

        data = self._sorted_contracts(rows)
        self.w_table.setRowCount(0)
        for r_idx, d in enumerate(data):
            self.w_table.insertRow(r_idx)
            self.w_table.setItem(r_idx, 0, QTableWidgetItem(d["code"] or str(d["id"])))
            self.w_table.setItem(r_idx, 1, QTableWidgetItem(f"{d['groom']} و {d['bride']}"))
            self.w_table.setItem(r_idx, 2, QTableWidgetItem(d["g_phone"] or "-"))
            self.w_table.setItem(r_idx, 3, QTableWidgetItem(d["b_phone"] or "-"))
            self.w_table.setItem(r_idx, 4, QTableWidgetItem(d["ceremony"] or "-"))
            self.w_table.setItem(r_idx, 5, QTableWidgetItem(d["venue"] or "-"))

            days_item = QTableWidgetItem(days_left_short(d["ceremony"]))
            jd = parse_jalali(d["ceremony"])
            if jd is not None:
                diff = (jd - jdatetime.date.today()).days
                if diff < 0:
                    days_item.setForeground(QColor("#7f8c8d"))
                elif diff <= 7:
                    days_item.setForeground(QColor("#c0392b"))
                else:
                    days_item.setForeground(QColor("#1e8449"))
            self.w_table.setItem(r_idx, 6, days_item)

            self.w_table.setItem(r_idx, 7, QTableWidgetItem(f"{d['raw']:,}"))
            self.w_table.setItem(r_idx, 8, QTableWidgetItem(f"{d['discount']:,}"))

            net_item = QTableWidgetItem(f"{d['net']:,}")
            net_item.setForeground(QColor("#2980b9"))
            self.w_table.setItem(r_idx, 9, net_item)

            self.w_table.setItem(r_idx, 10, QTableWidgetItem(f"{d['paid']:,}"))

            remain_item = QTableWidgetItem(f"{d['remain']:,}")
            remain_item.setForeground(QColor("#c0392b" if d["remain"] > 0 else "#27ae60"))
            self.w_table.setItem(r_idx, 11, remain_item)

            state_item = QTableWidgetItem(d["state"])
            state_item.setForeground(QColor(
                "#27ae60" if d["state"] == "تسویه کامل"
                else ("#e67e22" if d["state"] == "تسویه ناقص" else "#c0392b")))
            f = state_item.font()
            f.setBold(True)
            state_item.setFont(f)
            self.w_table.setItem(r_idx, 12, state_item)

            btn_dep = QPushButton(ico_text("deposit", "بیعانه‌ها"))
            btn_dep.setStyleSheet("background-color:#16a085; color:white; font-weight:bold; padding:3px;")
            btn_dep.clicked.connect(lambda _, cid=d["id"]: self.open_deposits_dialog(cid))
            self.w_table.setCellWidget(r_idx, 13, btn_dep)

            btn_print = QPushButton(ico_text("print", "چاپ / PDF"))
            btn_print.setStyleSheet("background-color:#2980b9; color:white; font-weight:bold; padding:3px;")
            btn_print.setToolTip("چاپ A5 سیاه و سفید یا PDF رنگی، با/بدون ریز پکیج و نیروی کار")
            btn_print.clicked.connect(lambda _, cid=d["id"]: self._contract_print_flow(cid))
            self.w_table.setCellWidget(r_idx, 14, btn_print)

            btn_edit = QPushButton(ico("edit"))
            btn_edit.setToolTip("ویرایش قرارداد")
            btn_edit.setStyleSheet("background-color: #f39c12; color: white;")
            btn_edit.clicked.connect(lambda _, cid=d["id"]: self.edit_wedding_contract(cid))
            self.w_table.setCellWidget(r_idx, 15, btn_edit)

            btn_del = QPushButton(ico("delete"))
            btn_del.setToolTip("حذف قرارداد")
            btn_del.setStyleSheet("background-color: #e74c3c; color: white;")
            btn_del.clicked.connect(lambda _, cid=d["id"]: self.delete_wedding_contract(cid))
            self.w_table.setCellWidget(r_idx, 16, btn_del)

        self.refresh_dashboard_summary()

    def edit_wedding_contract(self, contract_id):
        dlg = EditContractDialog(contract_id, self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            self.load_wedding_contracts()
            self.load_inventory()
            self.load_staff_table()
            self.load_person_events()
            self.refresh_settlement_cards()

    def delete_wedding_contract(self, contract_id):
        if QMessageBox.question(
                self, "تایید حذف",
                "آیا از حذف این قرارداد مطمئن هستید؟\n"
                "بیعانه‌ها و حقوق ثبت‌شده این قرارداد هم حذف می‌شود و انبار به‌روز می‌گردد.",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        ) != QMessageBox.StandardButton.Yes:
            return
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM wedding_contracts WHERE id=?", (contract_id,))
        cursor.execute("DELETE FROM wedding_deposits WHERE contract_id=?", (contract_id,))
        cursor.execute("DELETE FROM transactions WHERE contract_id=?", (contract_id,))
        conn.commit()
        conn.close()
        recalc_inventory_usage()
        self.load_wedding_contracts()
        self.load_inventory()
        self.load_staff_table()
        self.load_person_events()
        self.refresh_settlement_cards()

    def on_wedding_double_click(self, index):
        row = index.row()
        item = self.w_table.item(row, 0)
        if not item:
            return
        code = item.text()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM wedding_contracts WHERE code=?", (code,))
        r = cursor.fetchone()
        conn.close()
        if r:
            dlg = ContractDetailsDialog(r[0], self)
            dlg.exec()

    def open_deposits_dialog(self, contract_id):
        dlg = DepositsDialog(contract_id, self)
        dlg.exec()
        self.load_wedding_contracts()

    def export_wedding_pdf(self, contract_id):
        """سازگاری با نسخه قبل: خروجی PDF رنگی با ریز پکیج"""
        html, filename = InvoiceBuilder.build_wedding_invoice(contract_id, True, True, False)
        if not html:
            return
        path, _ = QFileDialog.getSaveFileName(self, "ذخیره فاکتور PDF", filename, "PDF Files (*.pdf)")
        if not path:
            return
        ok, err = save_html_pdf(html, path, "فاکتور قرارداد")
        if ok:
            QMessageBox.information(self, "موفقیت", f"فایل PDF ذخیره شد:\n{path}")
        else:
            QMessageBox.critical(self, "خطا", f"خطا در ساخت PDF:\n{err}")

    def print_wedding(self, contract_id):
        """سازگاری با نسخه قبل: چاپ A5 سیاه و سفید"""
        html, _ = InvoiceBuilder.build_wedding_invoice(contract_id, True, True, True)
        if html:
            print_html_document(html, self, "فاکتور قرارداد عروس و داماد")

    # ============================================================
    # ==============  تب ۲: تبلیغاتی / بیوتی  ===================
    # ============================================================
    def setup_commercial_tab(self):
        layout = QHBoxLayout()
        layout.setSpacing(10)

        form_box = QGroupBox("ثبت پروژه جدید تبلیغاتی / بیوتی / تولدی")
        form_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        form_layout = QFormLayout()
        form_layout.setSpacing(7)

        self.c_title = QLineEdit()
        self.c_type = QComboBox()
        self.c_type.addItems(["تبلیغاتی", "بیوتی", "تولدی", "عروسی", "عقد"])
        self.c_client_name = QLineEdit()
        self.c_client_phone = QLineEdit()
        self.c_cameras = QSpinBox()
        self.c_cameras.setRange(1, 20)
        self.c_cameras.setValue(1)
        self.c_amount = QLineEdit()
        self.c_amount.textChanged.connect(lambda t: self.c_amount.setText(format_number(t)))
        self.c_amount.textChanged.connect(self.update_commercial_words)
        self.c_paid = QLineEdit()
        self.c_paid.textChanged.connect(lambda t: self.c_paid.setText(format_number(t)))
        self.c_date = PersianDateEdit(date_str=jalali_today_str())
        self.c_desc = QTextEdit()
        self.c_desc.setFixedHeight(60)

        form_layout.addRow(f"{ico('commercial')} عنوان پروژه <span style='color:red;'>*</span>:", self.c_title)
        form_layout.addRow("نوع پروژه:", self.c_type)
        form_layout.addRow(f"{ico('people')} نام مشتری:", self.c_client_name)
        form_layout.addRow(f"{ico('people')} تلفن مشتری:", self.c_client_phone)
        form_layout.addRow(f"{ico('camera')} تعداد دوربین:", self.c_cameras)
        form_layout.addRow(f"{ico('money')} مبلغ کل (تومان) <span style='color:red;'>*</span>:", self.c_amount)
        form_layout.addRow(f"{ico('deposit')} مبلغ پرداختی (تومان):", self.c_paid)
        form_layout.addRow(f"{ico('calendar')} تاریخ پروژه (روز/ماه/سال):", self.c_date)
        form_layout.addRow("توضیحات:", self.c_desc)

        self.c_lbl_words = QLabel("")
        self.c_lbl_words.setWordWrap(True)
        self.c_lbl_words.setStyleSheet("background-color:#e8f8f5; border:1px solid #16a085; "
                                       "border-radius:8px; padding:9px; font-size:10pt;")
        form_layout.addRow(self.c_lbl_words)

        btn_save = QPushButton(ico_text("save", "ثبت پروژه"))
        btn_save.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        btn_save.setStyleSheet("background-color: #27ae60; color: white; padding: 10px;")
        btn_save.clicked.connect(self.save_commercial_project)
        form_layout.addRow(btn_save)

        form_box.setLayout(form_layout)
        layout.addWidget(form_box, 2)

        table_box = QGroupBox("لیست پروژه‌ها")
        table_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        t_layout = QVBoxLayout()

        tools = QHBoxLayout()
        btn_refresh = QPushButton(ico_text("refresh", "بروزرسانی"))
        btn_refresh.clicked.connect(self.load_commercial_projects)
        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color:#1e8449; color:white; font-weight:bold; padding:7px;")
        btn_excel.clicked.connect(lambda: self.export_table_excel(
            self.c_table, "پروژه‌های تبلیغاتی", f"پروژه‌ها_{get_current_year()}.xlsx"))
        tools.addWidget(btn_refresh)
        tools.addWidget(btn_excel)
        tools.addStretch()
        t_layout.addLayout(tools)

        self.c_table = QTableWidget()
        self.c_table.setColumnCount(11)
        self.c_table.setHorizontalHeaderLabels([
            "کد", "عنوان", "نوع", "مشتری", "دوربین", "مبلغ کل",
            "پرداختی", "مانده", "تاریخ", "چاپ/PDF", "عملیات"
        ])
        self.c_table.setAlternatingRowColors(True)
        self.c_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        persist_table(self.c_table, "commercial", QHeaderView.ResizeMode.ResizeToContents)
        self.c_table.verticalHeader().setDefaultSectionSize(34)
        self.c_table.doubleClicked.connect(self.on_commercial_double_click)
        t_layout.addWidget(self.c_table)
        table_box.setLayout(t_layout)
        layout.addWidget(table_box, 5)

        self.tab_commercial.setLayout(layout)
        self.load_commercial_projects()

    def update_commercial_words(self):
        if not hasattr(self, "c_lbl_words"):
            return
        amount = parse_number(self.c_amount.text())
        paid = parse_number(self.c_paid.text())
        remain = max(0, amount - paid)
        self.c_lbl_words.setText(
            f"<b>مبلغ کل به حروف:</b> {number_to_persian_words(amount)} تومان<br>"
            f"<b>مانده:</b> {remain:,} تومان  |  "
            f"<b>مانده به حروف:</b> {number_to_persian_words(remain)} تومان"
        )

    def save_commercial_project(self):
        title = self.c_title.text().strip()
        amount = parse_number(self.c_amount.text())
        if not title or amount <= 0:
            QMessageBox.warning(self, "خطا", "لطفاً عنوان و مبلغ کل را وارد کنید.")
            return

        work_year = get_current_year()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        code = generate_unique_code(f"C-{work_year}", "commercial_projects", "code", start=1, width=4)

        cursor.execute('''
            INSERT INTO commercial_projects
            (code, title, project_type, camera_count, total_amount, paid_amount,
             project_date, description, client_name, client_phone, work_year)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ''', (code, title, self.c_type.currentText(), self.c_cameras.value(),
              amount, parse_number(self.c_paid.text()), self.c_date.text(),
              self.c_desc.toPlainText(), self.c_client_name.text(),
              self.c_client_phone.text(), work_year))

        conn.commit()
        conn.close()
        QMessageBox.information(self, "موفقیت", f"پروژه با کد یکتای {code} ثبت شد.")

        self.c_title.clear()
        self.c_client_name.clear()
        self.c_client_phone.clear()
        self.c_amount.clear()
        self.c_paid.clear()
        self.c_desc.clear()
        self.c_date.set_date(jalali_today_str())

        self.load_commercial_projects()
        self.load_work_years()
        self.refresh_dashboard_summary()
        self.refresh_settlement_cards()

    def load_commercial_projects(self):
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        work_year = get_current_year()
        cursor.execute("""SELECT id, code, title, project_type, camera_count,
                          total_amount, paid_amount, project_date, description,
                          client_name, client_phone
                          FROM commercial_projects WHERE work_year=? ORDER BY id DESC""",
                       (work_year,))
        rows = cursor.fetchall()
        conn.close()

        self.c_table.setRowCount(0)
        for r_idx, row in enumerate(rows):
            (p_id, code, title, ptype, cameras, total, paid, date, desc,
             client, phone) = row
            total = total or 0
            paid = paid or 0
            remain = max(0, total - paid)

            self.c_table.insertRow(r_idx)
            self.c_table.setItem(r_idx, 0, QTableWidgetItem(code or str(p_id)))
            self.c_table.setItem(r_idx, 1, QTableWidgetItem(title))
            self.c_table.setItem(r_idx, 2, QTableWidgetItem(ptype))
            self.c_table.setItem(r_idx, 3, QTableWidgetItem(f"{client or '-'}\n{phone or ''}"))
            self.c_table.setItem(r_idx, 4, QTableWidgetItem(str(cameras)))
            self.c_table.setItem(r_idx, 5, QTableWidgetItem(f"{total:,}"))
            self.c_table.setItem(r_idx, 6, QTableWidgetItem(f"{paid:,}"))
            remain_item = QTableWidgetItem(f"{remain:,}")
            remain_item.setForeground(QColor("#c0392b" if remain > 0 else "#27ae60"))
            self.c_table.setItem(r_idx, 7, remain_item)
            self.c_table.setItem(r_idx, 8, QTableWidgetItem(date or "-"))

            btn_print = QPushButton(ico_text("print", "چاپ / PDF"))
            btn_print.setStyleSheet("background-color:#2980b9; color:white; font-weight:bold; padding:4px;")
            btn_print.clicked.connect(lambda _, pid=p_id: self._commercial_print_flow(pid))
            self.c_table.setCellWidget(r_idx, 9, btn_print)

            aw = QWidget()
            al = QHBoxLayout()
            al.setContentsMargins(0, 0, 0, 0)
            al.setSpacing(2)
            btn_edit = QPushButton(ico("edit"))
            btn_edit.setToolTip("ویرایش پروژه")
            btn_edit.setStyleSheet("background-color: #f39c12; color: white;")
            btn_edit.clicked.connect(lambda _, pid=p_id: self.edit_commercial_project(pid))
            btn_del = QPushButton(ico("delete"))
            btn_del.setToolTip("حذف پروژه")
            btn_del.setStyleSheet("background-color: #e74c3c; color: white;")
            btn_del.clicked.connect(lambda _, pid=p_id: self.delete_commercial_project(pid))
            al.addWidget(btn_edit)
            al.addWidget(btn_del)
            aw.setLayout(al)
            self.c_table.setCellWidget(r_idx, 10, aw)

        self.refresh_settlement_cards()

    def edit_commercial_project(self, p_id):
        dlg = EditCommercialDialog(p_id, self)
        if dlg.exec() == QDialog.DialogCode.Accepted:
            self.load_commercial_projects()

    def on_commercial_double_click(self, index):
        row = index.row()
        item = self.c_table.item(row, 0)
        if not item:
            return
        code = item.text()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT id FROM commercial_projects WHERE code=?", (code,))
        r = cursor.fetchone()
        conn.close()
        if r:
            dlg = EditCommercialDialog(r[0], self)
            if dlg.exec() == QDialog.DialogCode.Accepted:
                self.load_commercial_projects()

    def export_commercial_pdf(self, p_id):
        html, filename = InvoiceBuilder.build_commercial_invoice(p_id)
        if not html:
            return
        path, _ = QFileDialog.getSaveFileName(self, "ذخیره فاکتور PDF", filename, "PDF Files (*.pdf)")
        if not path:
            return
        ok, err = save_html_pdf(html, path, "فاکتور پروژه")
        if ok:
            QMessageBox.information(self, "موفقیت", f"فایل PDF ذخیره شد:\n{path}")
        else:
            QMessageBox.critical(self, "خطا", f"خطا:\n{err}")

    def print_commercial(self, p_id):
        html, _ = InvoiceBuilder.build_commercial_invoice(p_id, mono=True)
        if html:
            print_html_document(html, self, "فاکتور پروژه تبلیغاتی")

    def delete_commercial_project(self, p_id):
        if QMessageBox.question(self, "تایید حذف", "آیا از حذف این پروژه مطمئن هستید؟",
                                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
                                ) != QMessageBox.StandardButton.Yes:
            return
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM commercial_projects WHERE id=?", (p_id,))
        conn.commit()
        conn.close()
        self.load_commercial_projects()
        self.refresh_settlement_cards()

    # ============================================================
    # ==============  تب ۳: پرسنل و کارکنان  ====================
    # ============================================================
    def setup_staff_tab(self):
        layout = QVBoxLayout()
        layout.setSpacing(8)

        # ==================== ردیف اول: تعریف نیرو و ثبت حقوق ====================
        forms_layout = QHBoxLayout()
        forms_layout.setSpacing(10)

        person_box = QGroupBox("تعریف نیروی جدید")
        person_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        person_form = QFormLayout()
        self.txt_person_code = QLineEdit()
        self.txt_person_code.setPlaceholderText("خالی بگذارید تا کد یکتا ساخته شود")
        self.txt_person_name = QLineEdit()
        self.txt_person_phone = QLineEdit()
        self.combo_person_role = QComboBox()
        self.combo_person_role.setEditable(True)
        self.combo_person_role.addItems(self.ROLES)
        btn_add_person = QPushButton(ico_text("add", "ثبت فرد جدید"))
        btn_add_person.setStyleSheet("background-color: #8e44ad; color: white; font-weight: bold; padding:8px;")
        btn_add_person.clicked.connect(self.add_person)

        person_form.addRow(f"{ico('people')} کد پرسنل (یکتا):", self.txt_person_code)
        person_form.addRow(f"{ico('people')} نام و نام خانوادگی <span style='color:red;'>*</span>:", self.txt_person_name)
        person_form.addRow(f"{ico('staff')} تخصص / نقش:", self.combo_person_role)
        person_form.addRow(f"{ico('people')} تلفن همراه:", self.txt_person_phone)
        person_form.addRow(btn_add_person)
        person_box.setLayout(person_form)

        trans_box = QGroupBox("ثبت حقوق و پرداختی مستقل (خارج از قرارداد)")
        trans_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        trans_form = QFormLayout()

        self.combo_staff_persons = QComboBox()
        self.txt_staff_amount = QLineEdit()
        self.txt_staff_amount.textChanged.connect(lambda t: self.txt_staff_amount.setText(format_number(t)))
        self.txt_staff_desc = QLineEdit()

        btn_add_trans = QPushButton(ico_text("money", "ثبت پرداخت حقوق"))
        btn_add_trans.setStyleSheet("background-color: #8e44ad; color: white; font-weight: bold; padding:8px;")
        btn_add_trans.clicked.connect(self.add_staff_payment)

        trans_form.addRow(f"{ico('people')} انتخاب فرد <span style='color:red;'>*</span>:", self.combo_staff_persons)
        trans_form.addRow(f"{ico('money')} مبلغ (تومان) <span style='color:red;'>*</span>:", self.txt_staff_amount)
        trans_form.addRow("بابت / توضیحات:", self.txt_staff_desc)
        trans_form.addRow(btn_add_trans)
        trans_box.setLayout(trans_form)

        forms_layout.addWidget(person_box, 2)
        forms_layout.addWidget(trans_box, 3)
        layout.addLayout(forms_layout)

        # ==================== گزارش دوره‌ای ====================
        report_box = QGroupBox("گزارش ریز کارکرد (هفتگی / ماهانه / سالانه)")
        report_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        rep_layout = QHBoxLayout()

        self.combo_report_person = QComboBox()
        self.combo_report_period = QComboBox()
        self.combo_report_period.addItems(["هفتگی", "ماهانه", "سالانه", "کل"])

        btn_show = QPushButton(ico_text("search", "نمایش گزارش"))
        btn_show.setStyleSheet("background-color: #2980b9; color: white; font-weight: bold; padding:7px;")
        btn_show.clicked.connect(self.show_staff_report)

        btn_rep_pdf = QPushButton(ico_text("pdf", "PDF"))
        btn_rep_pdf.setStyleSheet("background-color: #c0392b; color: white; font-weight: bold; padding:7px;")
        btn_rep_pdf.clicked.connect(self.export_staff_report_pdf)

        btn_rep_print = QPushButton(ico_text("print", "چاپ"))
        btn_rep_print.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding:7px;")
        btn_rep_print.clicked.connect(self.print_staff_report)

        rep_layout.addWidget(QLabel(f"{ico('people')} فرد:"))
        rep_layout.addWidget(self.combo_report_person, 2)
        rep_layout.addWidget(QLabel("دوره:"))
        rep_layout.addWidget(self.combo_report_period, 1)
        rep_layout.addWidget(btn_show)
        rep_layout.addWidget(btn_rep_pdf)
        rep_layout.addWidget(btn_rep_print)
        rep_layout.addStretch()
        report_box.setLayout(rep_layout)
        layout.addWidget(report_box)

        # ==================== ردیف دوم: مراسم‌ها و پرداختی‌ها ====================
        bottom_split = QHBoxLayout()
        bottom_split.setSpacing(10)

        events_box = QGroupBox("مراسم‌ها و قراردادهای هر فرد (حقوق دریافتی از هر قرارداد)")
        events_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        ev_layout = QVBoxLayout()

        ev_top = QHBoxLayout()
        ev_top.addWidget(QLabel(f"{ico('people')} انتخاب فرد:"))
        self.combo_person_events = QComboBox()
        self.combo_person_events.currentIndexChanged.connect(self.load_person_events)
        ev_top.addWidget(self.combo_person_events, 3)

        btn_del_person = QPushButton(ico_text("delete", "حذف این فرد"))
        btn_del_person.setStyleSheet("background-color:#c0392b; color:white; font-weight:bold; padding:7px;")
        btn_del_person.clicked.connect(self.delete_person)
        ev_top.addWidget(btn_del_person)
        ev_layout.addLayout(ev_top)

        self.events_table = QTableWidget()
        self.events_table.setColumnCount(6)
        self.events_table.setHorizontalHeaderLabels([
            "کد قرارداد", "زوجین", "تاریخ مراسم", "نقش در قرارداد", "حقوق این قرارداد", "مبلغ کل قرارداد"
        ])
        self.events_table.setAlternatingRowColors(True)
        self.events_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        persist_table(self.events_table, "staff_events", QHeaderView.ResizeMode.Stretch)
        ev_layout.addWidget(self.events_table)

        self.lbl_events_sum = QLabel("")
        self.lbl_events_sum.setWordWrap(True)
        self.lbl_events_sum.setStyleSheet("background-color:#f3ecfb; border:1px solid #8e44ad; "
                                          "border-radius:8px; padding:8px;")
        ev_layout.addWidget(self.lbl_events_sum)
        events_box.setLayout(ev_layout)

        pay_box = QGroupBox("لیست پرداختی‌ها و حقوق‌های ثبت‌شده")
        pay_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        pay_layout = QVBoxLayout()

        pay_tools = QHBoxLayout()
        btn_refresh = QPushButton(ico_text("refresh", "بروزرسانی"))
        btn_refresh.clicked.connect(self.load_staff_table)
        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color:#1e8449; color:white; font-weight:bold; padding:7px;")
        btn_excel.clicked.connect(lambda: self.export_table_excel(
            self.staff_table, "پرداختی پرسنل", f"پرداختی_پرسنل_{get_current_year()}.xlsx"))
        pay_tools.addWidget(btn_refresh)
        pay_tools.addWidget(btn_excel)
        pay_tools.addStretch()
        pay_layout.addLayout(pay_tools)

        self.staff_table = QTableWidget()
        self.staff_table.setColumnCount(7)
        self.staff_table.setHorizontalHeaderLabels([
            "کد", "نام فرد", "نقش/تخصص", "مبلغ پرداخت (تومان)",
            "توضیحات", "تاریخ", "حذف"
        ])
        self.staff_table.setAlternatingRowColors(True)
        self.staff_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        persist_table(self.staff_table, "staff_payments", QHeaderView.ResizeMode.Stretch)
        pay_layout.addWidget(self.staff_table)
        pay_box.setLayout(pay_layout)

        bottom_split.addWidget(events_box, 3)
        bottom_split.addWidget(pay_box, 3)
        layout.addLayout(bottom_split)

        self.tab_staff.setLayout(layout)
        self.load_persons_combo()
        self.load_staff_table()
        self.load_person_events()

    def add_person(self):
        name = self.txt_person_name.text().strip()
        if not name:
            QMessageBox.warning(self, "خطا", "لطفاً نام فرد را وارد کنید.")
            return

        code = self.txt_person_code.text().strip()
        work_year = get_current_year()
        role = self.combo_person_role.currentText().strip() or "سایر"

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        if not code:
            code = generate_unique_code(f"P-{work_year}", "persons", "code", start=1, width=4)
        else:
            cursor.execute("SELECT name FROM persons WHERE code=?", (code,))
            dup = cursor.fetchone()
            if dup:
                conn.close()
                QMessageBox.warning(self, "کد تکراری",
                                    f"کد «{code}» قبلاً برای «{dup[0]}» ثبت شده است.\n"
                                    "کد هر فرد باید یکتا باشد.")
                return

        cursor.execute("""INSERT INTO persons (code, name, role, phone, work_year)
                          VALUES (?, ?, ?, ?, ?)""",
                       (code, name, role, self.txt_person_phone.text(), work_year))
        conn.commit()
        conn.close()

        self.txt_person_code.clear()
        self.txt_person_name.clear()
        self.txt_person_phone.clear()
        self.load_persons_combo()
        self.load_person_events()
        # نیروی جدید بلافاصله در فرم قرارداد قابل انتخاب می‌شود
        self.load_staff_checkboxes()
        QMessageBox.information(self, "موفقیت", f"فرد با کد یکتای {code} ثبت شد.")

    def load_persons_combo(self):
        self.combo_staff_persons.clear()
        self.combo_report_person.clear()
        if hasattr(self, "combo_person_events"):
            current = self.combo_person_events.currentData()
            self.combo_person_events.blockSignals(True)
            self.combo_person_events.clear()

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT id, code, name, role FROM persons ORDER BY name")
        rows = cursor.fetchall()
        conn.close()

        for p_id, code, name, role in rows:
            label = f"[{code or '-'}] {name} ({role})"
            self.combo_staff_persons.addItem(label, p_id)
            self.combo_report_person.addItem(label, p_id)
            if hasattr(self, "combo_person_events"):
                self.combo_person_events.addItem(label, p_id)

        if hasattr(self, "combo_person_events"):
            if current is not None:
                idx = self.combo_person_events.findData(current)
                if idx >= 0:
                    self.combo_person_events.setCurrentIndex(idx)
            self.combo_person_events.blockSignals(False)

    def add_staff_payment(self):
        person_id = self.combo_staff_persons.currentData()
        amount = parse_number(self.txt_staff_amount.text())
        if not person_id or amount <= 0:
            QMessageBox.warning(self, "خطا", "لطفاً فرد و مبلغ معتبر وارد کنید.")
            return

        today = jdatetime.date.today()
        work_year = get_current_year()

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute('''
            INSERT INTO transactions
            (trans_type, category, person_id, amount, year, month, day, description, work_year)
            VALUES ('expense', 'حقوق کارکنان', ?, ?, ?, ?, ?, ?, ?)
        ''', (person_id, amount, today.year, today.month, today.day,
              self.txt_staff_desc.text(), work_year))
        conn.commit()
        conn.close()

        self.txt_staff_amount.clear()
        self.txt_staff_desc.clear()
        self.load_staff_table()
        self.load_person_events()

    def load_staff_table(self):
        if not hasattr(self, "staff_table"):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        work_year = get_current_year()
        cursor.execute('''
            SELECT t.id, p.code, p.name, p.role, t.amount, t.description,
                   t.year, t.month, t.day
            FROM transactions t
            JOIN persons p ON t.person_id = p.id
            WHERE t.work_year=?
            ORDER BY t.id DESC
        ''', (work_year,))
        rows = cursor.fetchall()
        conn.close()

        self.staff_table.setRowCount(0)
        for r_idx, row in enumerate(rows):
            (tid, code, name, role, amount, desc, yr, mo, dy) = row
            self.staff_table.insertRow(r_idx)
            self.staff_table.setItem(r_idx, 0, QTableWidgetItem(code or "-"))
            self.staff_table.setItem(r_idx, 1, QTableWidgetItem(name))
            self.staff_table.setItem(r_idx, 2, QTableWidgetItem(role))
            self.staff_table.setItem(r_idx, 3, QTableWidgetItem(f"{amount:,}"))
            self.staff_table.setItem(r_idx, 4, QTableWidgetItem(desc if desc else "-"))
            self.staff_table.setItem(r_idx, 5, QTableWidgetItem(f"{yr}/{mo:02d}/{dy:02d}"))

            btn_del = QPushButton(ico("delete"))
            btn_del.setToolTip("حذف این پرداختی")
            btn_del.setStyleSheet("background-color: #e74c3c; color: white;")
            btn_del.clicked.connect(lambda _, x=tid: self.delete_staff_trans(x))
            self.staff_table.setCellWidget(r_idx, 6, btn_del)

    def delete_staff_trans(self, tid):
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM transactions WHERE id=?", (tid,))
        conn.commit()
        conn.close()
        self.load_staff_table()
        self.load_person_events()

    def load_person_events(self):
        """نمایش مراسم‌ها و قراردادهایی که این فرد در آن‌ها حضور داشته"""
        if not hasattr(self, "events_table"):
            return
        pid = self.combo_person_events.currentData() if hasattr(self, "combo_person_events") else None
        self.events_table.setRowCount(0)
        if not pid:
            if hasattr(self, "lbl_events_sum"):
                self.lbl_events_sum.setText("فردی انتخاب نشده است.")
            return

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT code, groom_name, bride_name, ceremony_date, staff_ids, total_amount
                          FROM wedding_contracts WHERE staff_ids IS NOT NULL AND staff_ids<>''""")
        contracts = cursor.fetchall()
        conn.close()

        rows = []
        total_wage = 0
        for code, groom, bride, cer, staff_json, total in contracts:
            try:
                members = json.loads(staff_json) if staff_json else []
            except Exception:
                members = []
            for m in members:
                try:
                    if int(m.get("id")) == int(pid):
                        wage = int(m.get("wage") or 0)
                        total_wage += wage
                        rows.append((code or "-", f"{groom} و {bride}", cer or "-",
                                     m.get("role", "-"), wage, total or 0))
                        break
                except Exception:
                    continue

        rows.sort(key=lambda r: r[2], reverse=True)
        self.events_table.setRowCount(len(rows))
        for i, r in enumerate(rows):
            self.events_table.setItem(i, 0, QTableWidgetItem(str(r[0])))
            self.events_table.setItem(i, 1, QTableWidgetItem(str(r[1])))
            self.events_table.setItem(i, 2, QTableWidgetItem(str(r[2])))
            self.events_table.setItem(i, 3, QTableWidgetItem(str(r[3])))
            self.events_table.setItem(i, 4, QTableWidgetItem(f"{r[4]:,}"))
            self.events_table.setItem(i, 5, QTableWidgetItem(f"{r[5]:,}"))

        self.lbl_events_sum.setText(
            f"<b>تعداد مراسم‌های این فرد:</b> {len(rows)}  |  "
            f"<b>جمع حقوق قراردادی دریافتی:</b> {total_wage:,} تومان<br>"
            f"<b>به حروف:</b> {number_to_persian_words(total_wage)} تومان"
        )

    def delete_person(self):
        pid = self.combo_person_events.currentData() if hasattr(self, "combo_person_events") else None
        if not pid:
            QMessageBox.warning(self, "خطا", "لطفاً یک فرد را انتخاب کنید.")
            return
        if QMessageBox.question(
                self, "تایید حذف",
                "با حذف این فرد، همه پرداختی‌ها و حقوق‌های ثبت‌شده او هم حذف می‌شود.\n"
                "ادامه می‌دهید؟",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        ) != QMessageBox.StandardButton.Yes:
            return
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM transactions WHERE person_id=?", (pid,))
        cursor.execute("DELETE FROM persons WHERE id=?", (pid,))
        conn.commit()
        conn.close()
        self.load_persons_combo()
        self.load_staff_table()
        self.load_person_events()
        self.load_staff_checkboxes()

    def show_staff_report(self):
        person_id = self.combo_report_person.currentData()
        if not person_id:
            QMessageBox.warning(self, "خطا", "لطفاً یک فرد انتخاب کنید.")
            return

        period = self.combo_report_period.currentText()
        work_year = get_current_year()

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT amount, year, month, day, description, category
                          FROM transactions WHERE person_id=? AND work_year=?
                          ORDER BY year DESC, month DESC, day DESC""",
                       (person_id, work_year))
        all_trans = cursor.fetchall()
        conn.close()

        today = jdatetime.date.today()
        filtered = []
        for t in all_trans:
            try:
                tdate = jdatetime.date(t[1], t[2], t[3])
                diff = (today - tdate).days
                if period == "هفتگی" and 0 <= diff <= 7:
                    filtered.append(t)
                elif period == "ماهانه" and 0 <= diff <= 30:
                    filtered.append(t)
                elif period in ("سالانه", "کل"):
                    filtered.append(t)
            except Exception:
                filtered.append(t)

        if not filtered:
            QMessageBox.information(self, "گزارش", "تراکنشی در این بازه یافت نشد.")
            return

        dlg = QDialog(self)
        dlg.setWindowTitle(f"ریز کارکرد - {period}")
        dlg.resize(760, 540)
        dlg.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        dlg.setFont(QFont(APP_FONT_FAMILY, 10))
        v = QVBoxLayout()

        total = sum(t[0] for t in filtered)
        lbl = QLabel(
            f"<b>جمع کل پرداختی در {period}:</b> {total:,} تومان<br>"
            f"<b>به حروف:</b> {number_to_persian_words(total)} تومان"
        )
        lbl.setStyleSheet("background-color: #e8f8f5; padding: 10px; border-radius: 8px; font-size: 11pt;")
        v.addWidget(lbl)

        table = QTableWidget()
        table.setColumnCount(4)
        table.setHorizontalHeaderLabels(["تاریخ", "دسته", "شرح", "مبلغ (تومان)"])
        table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeMode.Stretch)
        table.setAlternatingRowColors(True)
        table.setRowCount(len(filtered))
        for idx, t in enumerate(filtered):
            table.setItem(idx, 0, QTableWidgetItem(f"{t[1]}/{t[2]:02d}/{t[3]:02d}"))
            table.setItem(idx, 1, QTableWidgetItem(t[5]))
            table.setItem(idx, 2, QTableWidgetItem(t[4] if t[4] else "-"))
            table.setItem(idx, 3, QTableWidgetItem(f"{t[0]:,}"))
        v.addWidget(table)
        dlg.setLayout(v)
        dlg.exec()

    def export_staff_report_pdf(self):
        person_id = self.combo_report_person.currentData()
        if not person_id:
            QMessageBox.warning(self, "خطا", "لطفاً یک فرد انتخاب کنید.")
            return
        period = self.combo_report_period.currentText()
        html, filename = InvoiceBuilder.build_staff_report(
            person_id, period=period, year=get_current_year())
        if not html:
            return
        path, _ = QFileDialog.getSaveFileName(self, "ذخیره PDF", filename, "PDF Files (*.pdf)")
        if not path:
            return
        ok, err = save_html_pdf(html, path, "ریز کارکرد پرسنل")
        if ok:
            QMessageBox.information(self, "موفقیت", f"فایل PDF ذخیره شد:\n{path}")
        else:
            QMessageBox.critical(self, "خطا", f"خطا:\n{err}")

    def print_staff_report(self):
        person_id = self.combo_report_person.currentData()
        if not person_id:
            QMessageBox.warning(self, "خطا", "لطفاً یک فرد انتخاب کنید.")
            return
        period = self.combo_report_period.currentText()
        html, _ = InvoiceBuilder.build_staff_report(
            person_id, period=period, year=get_current_year(), mono=True)
        if html:
            print_html_document(html, self, "ریز کارکرد پرسنل")

    # ============================================================
    # ==============  تب ۴: مدیریت هزینه‌ها  ====================
    # ============================================================
    def setup_expenses_tab(self):
        layout = QHBoxLayout()
        layout.setSpacing(8)

        left = QVBoxLayout()
        left.setSpacing(8)

        # ---------- کارت‌های وضعیت تسویه پروژه‌ها ----------
        settle_box = QGroupBox("وضعیت تسویه پروژه‌ها (برای دیدن لیست هر دسته، روی آن کلیک کنید)")
        settle_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        cards_row = QHBoxLayout()
        cards_row.setSpacing(6)
        self.settle_cards = {}
        for key, title, color in (
            ("all", "کل پروژه‌ها", "#2c3e50"),
            ("full", "تسویه کامل", "#1e8449"),
            ("partial", "تسویه ناقص", "#e67e22"),
            ("none", "تسویه نشده", "#c0392b"),
        ):
            b = QPushButton(title)
            b.setMinimumHeight(60)
            b.setCursor(Qt.CursorShape.PointingHandCursor)
            b.setStyleSheet(f"background-color:{color}; color:white; font-weight:bold; "
                            f"font-size:9.5pt; border-radius:10px; padding:5px;")
            b.clicked.connect(lambda _, k=key: self.show_settlement_list(k))
            self.settle_cards[key] = b
            b._title = title
            cards_row.addWidget(b)
        settle_box.setLayout(cards_row)
        left.addWidget(settle_box)

        form_box = QGroupBox("ثبت هزینه جدید")
        form_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        form_layout = QFormLayout()
        form_layout.setSpacing(6)

        self.exp_code = QLineEdit()
        self.exp_code.setPlaceholderText("خالی بگذارید تا کد یکتا ساخته شود")
        self.exp_title = QLineEdit()
        self.exp_cat = QComboBox()
        self.exp_cat.setEditable(True)
        self.exp_cat.addItems([
            "قبوض (آب، برق، گاز، تلفن)", "کرایه آتلیه / دفتر",
            "خرید تجهیزات و مصرفی", "تبلیغات و بازاریابی",
            "حقوق و دستمزد", "سایر هزینه‌ها"
        ])
        self.exp_amount = QLineEdit()
        self.exp_amount.textChanged.connect(lambda t: self.exp_amount.setText(format_number(t)))
        self.exp_amount.textChanged.connect(self.update_expense_words)
        self.exp_date = PersianDateEdit(date_str=jalali_today_str())
        self.exp_desc = QTextEdit()
        self.exp_desc.setFixedHeight(56)

        form_layout.addRow(f"{ico('list')} کد هزینه (یکتا):", self.exp_code)
        form_layout.addRow(f"{ico('expenses')} عنوان هزینه <span style='color:red;'>*</span>:", self.exp_title)
        form_layout.addRow("دسته‌بندی:", self.exp_cat)
        form_layout.addRow(f"{ico('money')} مبلغ (تومان) <span style='color:red;'>*</span>:", self.exp_amount)
        form_layout.addRow(f"{ico('calendar')} تاریخ هزینه (روز/ماه/سال):", self.exp_date)
        form_layout.addRow("توضیحات:", self.exp_desc)

        self.exp_lbl_words = QLabel("")
        self.exp_lbl_words.setWordWrap(True)
        self.exp_lbl_words.setStyleSheet("background-color:#fdedec; border:1px solid #c0392b; "
                                         "border-radius:8px; padding:8px; font-size:9.5pt;")
        form_layout.addRow(self.exp_lbl_words)

        btn_save = QPushButton(ico_text("save", "ثبت هزینه"))
        btn_save.setStyleSheet("background-color: #c0392b; color: white; font-weight: bold; padding: 9px;")
        btn_save.clicked.connect(self.save_expense)
        form_layout.addRow(btn_save)

        form_box.setLayout(form_layout)
        left.addWidget(form_box)
        left.addStretch()

        left_wrap = QWidget()
        left_wrap.setLayout(left)
        layout.addWidget(left_wrap, 3)

        table_box = QGroupBox("لیست هزینه‌ها")
        table_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        t_layout = QVBoxLayout()

        tools = QHBoxLayout()
        btn_refresh = QPushButton(ico_text("refresh", "بروزرسانی"))
        btn_refresh.clicked.connect(self.load_expenses_table)
        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color:#1e8449; color:white; font-weight:bold; padding:6px;")
        btn_excel.clicked.connect(lambda: self.export_table_excel(
            self.exp_table, "هزینه‌ها", f"هزینه‌ها_{get_current_year()}.xlsx"))
        tools.addWidget(btn_refresh)
        tools.addWidget(btn_excel)
        tools.addStretch()
        t_layout.addLayout(tools)

        self.exp_table = QTableWidget()
        self.exp_table.setColumnCount(8)
        self.exp_table.setHorizontalHeaderLabels([
            "کد", "عنوان", "دسته‌بندی", "مبلغ (تومان)",
            "تاریخ", "توضیحات", "چاپ", "حذف"
        ])
        self.exp_table.setAlternatingRowColors(True)
        self.exp_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.exp_table.verticalHeader().setDefaultSectionSize(30)
        persist_table(self.exp_table, "expenses", QHeaderView.ResizeMode.Stretch)
        t_layout.addWidget(self.exp_table)
        table_box.setLayout(t_layout)
        layout.addWidget(table_box, 6)

        self.tab_expenses.setLayout(layout)
        self.load_expenses_table()
        self.refresh_settlement_cards()

    # ---------------- کارت‌های وضعیت تسویه ----------------
    def _project_settlements(self):
        """فهرست وضعیت تسویه همه قراردادها و پروژه‌های سال کاری جاری"""
        work_year = get_current_year()
        out = []
        try:
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("""SELECT code, groom_name, bride_name, ceremony_date,
                                     total_amount, discount, paid_amount
                              FROM wedding_contracts WHERE work_year=?""", (work_year,))
            for r in cursor.fetchall():
                net = max(0, (r[4] or 0) - (r[5] or 0))
                state, remain = settlement_state(net, r[6])
                out.append({"kind": "قرارداد", "code": r[0],
                            "title": f"{r[1]} و {r[2]}", "date": r[3],
                            "net": net, "paid": r[6] or 0, "remain": remain, "state": state})
            cursor.execute("""SELECT code, title, client_name, project_date,
                                     total_amount, paid_amount
                              FROM commercial_projects WHERE work_year=?""", (work_year,))
            for r in cursor.fetchall():
                net = r[4] or 0
                state, remain = settlement_state(net, r[5])
                out.append({"kind": "پروژه", "code": r[0],
                            "title": f"{r[1]} - {r[2] or '-'}", "date": r[3],
                            "net": net, "paid": r[5] or 0, "remain": remain, "state": state})
            conn.close()
        except Exception as e:
            print("[SETTLE] error:", e)
        return out

    def refresh_settlement_cards(self):
        """به‌روزرسانی تعداد و مانده هر دسته روی کارت‌ها"""
        if not getattr(self, "settle_cards", None):
            return
        try:
            data = self._project_settlements()
            counts = {"all": len(data), "full": 0, "partial": 0, "none": 0}
            sums = {"all": 0, "full": 0, "partial": 0, "none": 0}
            for d in data:
                sums["all"] += d["remain"]
                key = {"تسویه کامل": "full", "تسویه ناقص": "partial",
                       "تسویه نشده": "none"}.get(d["state"], "none")
                counts[key] += 1
                sums[key] += d["remain"]
            for key, btn in self.settle_cards.items():
                title = getattr(btn, "_title", key)
                btn.setText(f"{title}\n{counts[key]} مورد\nمانده: {sums[key]:,} تومان")
        except Exception as e:
            print("[SETTLE] refresh error:", e)

    def show_settlement_list(self, kind):
        dlg = SettlementStatusDialog(kind, self)
        dlg.exec()
        self.refresh_settlement_cards()

    def update_expense_words(self):
        if not hasattr(self, "exp_lbl_words"):
            return
        amount = parse_number(self.exp_amount.text())
        self.exp_lbl_words.setText(f"<b>مبلغ به حروف:</b> {number_to_persian_words(amount)} تومان")

    def save_expense(self):
        title = self.exp_title.text().strip()
        amount = parse_number(self.exp_amount.text())
        if not title or amount <= 0:
            QMessageBox.warning(self, "خطا", "لطفاً عنوان و مبلغ را وارد کنید.")
            return

        work_year = get_current_year()
        code = self.exp_code.text().strip()

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        if not code:
            code = generate_unique_code(f"E-{work_year}", "expenses", "code", start=1, width=4)
        else:
            cursor.execute("SELECT title FROM expenses WHERE code=?", (code,))
            dup = cursor.fetchone()
            if dup:
                conn.close()
                QMessageBox.warning(self, "کد تکراری",
                                    f"کد «{code}» قبلاً برای «{dup[0]}» ثبت شده است.")
                return

        cursor.execute('''INSERT INTO expenses
            (code, title, category, amount, date_str, description, work_year)
            VALUES (?, ?, ?, ?, ?, ?, ?)''',
            (code, title, self.exp_cat.currentText(), amount,
             self.exp_date.text(), self.exp_desc.toPlainText(), work_year))
        conn.commit()
        conn.close()

        self.exp_code.clear()
        self.exp_title.clear()
        self.exp_amount.clear()
        self.exp_desc.clear()
        self.exp_date.set_date(jalali_today_str())
        self.load_expenses_table()
        self.calculate_financial_report()
        self.refresh_settlement_cards()

    def load_expenses_table(self):
        if not hasattr(self, "exp_table"):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        work_year = get_current_year()
        cursor.execute("""SELECT id, code, title, category, amount, date_str, description
                          FROM expenses WHERE work_year=? ORDER BY id DESC""",
                       (work_year,))
        rows = cursor.fetchall()
        conn.close()

        self.exp_table.setRowCount(0)
        for r_idx, row in enumerate(rows):
            (eid, code, title, cat, amount, date_str, desc) = row
            self.exp_table.insertRow(r_idx)
            self.exp_table.setItem(r_idx, 0, QTableWidgetItem(code or "-"))
            self.exp_table.setItem(r_idx, 1, QTableWidgetItem(title))
            self.exp_table.setItem(r_idx, 2, QTableWidgetItem(cat))
            self.exp_table.setItem(r_idx, 3, QTableWidgetItem(f"{amount:,}"))
            self.exp_table.setItem(r_idx, 4, QTableWidgetItem(date_str or "-"))
            self.exp_table.setItem(r_idx, 5, QTableWidgetItem(desc if desc else "-"))

            btn_print = QPushButton(ico("print"))
            btn_print.setToolTip("چاپ رسید هزینه")
            btn_print.setStyleSheet("background-color: #2980b9; color: white;")
            btn_print.clicked.connect(lambda _, x=eid: self.print_expense(x))
            self.exp_table.setCellWidget(r_idx, 6, btn_print)

            btn_del = QPushButton(ico("delete"))
            btn_del.setToolTip("حذف هزینه")
            btn_del.setStyleSheet("background-color: #e74c3c; color: white;")
            btn_del.clicked.connect(lambda _, x=eid: self.delete_expense(x))
            self.exp_table.setCellWidget(r_idx, 7, btn_del)

    def expense_html(self, eid, mono=False):
        p = doc_palette(mono)
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT code, title, category, amount, date_str, description "
                       "FROM expenses WHERE id=?", (eid,))
        r = cursor.fetchone()
        conn.close()
        if not r:
            return None
        studio_name, studio_phone, studio_address = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()

        def pair(label, value):
            return InvoiceBuilder._row([
                InvoiceBuilder._td(f"<b>{label}</b>", align="right",
                                   bg=p["box_bg"], size="8.5pt", width="26%"),
                InvoiceBuilder._td(value, align="right", size="8.5pt"),
            ])

        rows = (
            pair("کد:", r[0] or "-") +
            pair("عنوان:", r[1]) +
            pair("دسته‌بندی:", r[2]) +
            pair("مبلغ:", f"<b>{r[3]:,}</b> تومان") +
            pair("مبلغ به حروف:", f"{number_to_persian_words(r[3])} تومان") +
            pair("تاریخ:", r[4] or "-") +
            pair("توضیحات:", r[5] or "-")
        )

        return f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {InvoiceBuilder._title_band(
              p, studio_name, "<div style='font-size:10pt;'>رسید هزینه</div>",
              [f"تاریخ چاپ: {now.strftime('%Y/%m/%d')}", f"ساعت: {now.strftime('%H:%M')}"])}
          <table width="100%" style="border-collapse:collapse; margin-top:6px;">{rows}</table>
          <div style="text-align:center; margin-top:34px; font-size:9pt;">مهر و امضای حسابداری</div>
          <div style="text-align:center; margin-top:10px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}{(' | ' + studio_address) if studio_address else ''}
          </div>
        </div>
        """

    def print_expense(self, eid):
        html = self.expense_html(eid, mono=True)
        if html:
            print_html_document(html, self, "رسید هزینه")

    def delete_expense(self, eid):
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM expenses WHERE id=?", (eid,))
        conn.commit()
        conn.close()
        self.load_expenses_table()
        self.refresh_settlement_cards()

    # ============================================================
    # ==============  تب ۵: گزارش مالی جامع  ====================
    # ============================================================
    def setup_reports_tab(self):
        layout = QVBoxLayout()
        layout.setSpacing(8)

        top_layout = QHBoxLayout()
        top_layout.setSpacing(7)

        self.combo_report_type = QComboBox()
        self.combo_report_type.addItems(["گزارش کل", "گزارش ماهانه", "گزارش سالانه", "گزارش هر فرد"])
        top_layout.addWidget(QLabel("نوع گزارش:"))
        top_layout.addWidget(self.combo_report_type, 2)

        top_layout.addWidget(QLabel("فرد:"))
        self.combo_report_person_fin = QComboBox()
        self.load_report_persons_combo()
        top_layout.addWidget(self.combo_report_person_fin, 2)

        btn_calc = QPushButton(ico_text("calc", "محاسبه گزارش"))
        btn_calc.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 7px;")
        btn_calc.clicked.connect(self.calculate_financial_report)
        top_layout.addWidget(btn_calc)

        btn_print_rep = QPushButton(ico_text("print", "چاپ / PDF گزارش"))
        btn_print_rep.setStyleSheet("background-color: #8e44ad; color: white; font-weight: bold; padding: 7px;")
        btn_print_rep.clicked.connect(self.print_financial_report)
        top_layout.addWidget(btn_print_rep)

        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color: #1e8449; color: white; font-weight: bold; padding: 7px;")
        btn_excel.clicked.connect(lambda: self.export_table_excel(
            self.rep_table, "گزارش مالی جامع", f"گزارش_مالی_{get_current_year()}.xlsx"))
        top_layout.addWidget(btn_excel)

        layout.addLayout(top_layout)

        self.lbl_rep_summary = QLabel("ورودی کل: ۰ تومان | خروجی کل: ۰ تومان | سود خالص: ۰ تومان")
        self.lbl_rep_summary.setWordWrap(True)
        self.lbl_rep_summary.setStyleSheet(
            "font-size: 11pt; font-weight: bold; background-color: #e8f8f5; "
            "padding: 12px; border-radius: 8px; border: 1px solid #16a085;")
        layout.addWidget(self.lbl_rep_summary)

        # ---------------- بخش نمودارها ----------------
        chart_box = QGroupBox("نمودارهای گرافیکی (فارسی و راست‌چین)")
        chart_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        ch = QHBoxLayout()

        self.combo_chart_kind = QComboBox()
        self.chart_kinds = [
            ("نمودار درآمد و هزینه سال کاری", "income_expense"),
            ("نمودار درآمد ماهانه (به تفکیک ماه شمسی)", "monthly_income"),
            ("نمودار هزینه‌ها به تفکیک دسته‌بندی", "expense_categories"),
            ("نمودار حقوق و پرداختی پرسنل", "staff_salaries"),
            ("نمودار وضعیت موجودی انبار تجهیزات", "inventory_status"),
            ("نمودار وضعیت چک‌های دریافتی و پرداختی", "checks_status"),
            ("نمودار پلکانی درآمد ماهانه", "monthly_step"),
            ("نمودار نقطه‌ای مبلغ مراسم", "scatter"),
            ("نمودار خطی روند درآمد و هزینه", "line_trend"),
            ("نمودار دونات هزینه‌ها", "donut"),
            ("نمودار انباشته درآمد و هزینه", "stacked"),
            ("نمودار سطحی روند انباشته سود", "area"),
        ]
        for title, _key in self.chart_kinds:
            self.combo_chart_kind.addItem(f"{ico('chart')} {title}")
        ch.addWidget(QLabel("نوع نمودار:"))
        ch.addWidget(self.combo_chart_kind, 3)

        btn_chart = QPushButton(ico_text("chart", "نمایش نمودار"))
        btn_chart.setStyleSheet("background-color: #2980b9; color: white; font-weight: bold; padding: 7px;")
        btn_chart.clicked.connect(self.show_financial_chart)
        ch.addWidget(btn_chart)
        ch.addStretch()
        chart_box.setLayout(ch)
        layout.addWidget(chart_box)

        self.rep_table = QTableWidget()
        self.rep_table.setColumnCount(5)
        self.rep_table.setHorizontalHeaderLabels([
            "ردیف", "بخش", "نوع تراکنش", "مبلغ (تومان)", "تاریخ / شرح"
        ])
        self.rep_table.setAlternatingRowColors(True)
        self.rep_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        persist_table(self.rep_table, "financial_report", QHeaderView.ResizeMode.Stretch)
        layout.addWidget(self.rep_table)

        self.tab_reports.setLayout(layout)
        self.calculate_financial_report()

    def load_report_persons_combo(self):
        if not hasattr(self, "combo_report_person_fin"):
            return
        self.combo_report_person_fin.clear()
        self.combo_report_person_fin.addItem("--- همه ---", None)
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT id, code, name, role FROM persons ORDER BY name")
        for p_id, code, name, role in cursor.fetchall():
            self.combo_report_person_fin.addItem(f"[{code or '-'}] {name} ({role})", p_id)
        conn.close()

    def _report_totals(self):
        work_year = get_current_year()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT COALESCE(SUM(total_amount),0), COALESCE(SUM(discount),0), "
                       "COALESCE(SUM(paid_amount),0) FROM wedding_contracts WHERE work_year=?", (work_year,))
        w_raw, w_disc, w_paid = cursor.fetchone()
        cursor.execute("SELECT COALESCE(SUM(total_amount),0), COALESCE(SUM(paid_amount),0) "
                       "FROM commercial_projects WHERE work_year=?", (work_year,))
        c_in, c_paid = cursor.fetchone()
        cursor.execute("SELECT COALESCE(SUM(amount),0) FROM expenses WHERE work_year=?", (work_year,))
        exp_out = cursor.fetchone()[0]
        cursor.execute("SELECT COALESCE(SUM(amount),0) FROM transactions WHERE work_year=?", (work_year,))
        staff_out = cursor.fetchone()[0]
        cursor.execute("SELECT COUNT(*) FROM persons")
        persons_count = cursor.fetchone()[0]
        conn.close()
        w_net = max(0, w_raw - w_disc)
        return {
            "year": work_year, "w_raw": w_raw, "w_disc": w_disc, "w_net": w_net,
            "w_paid": w_paid, "c_in": c_in, "c_paid": c_paid,
            "exp_out": exp_out, "staff_out": staff_out,
            "total_in": w_net + c_in, "total_out": exp_out + staff_out,
            "persons": persons_count,
        }

    def calculate_financial_report(self):
        if not hasattr(self, "rep_table"):
            return
        work_year = get_current_year()
        report_type = self.combo_report_type.currentText()
        t = self._report_totals()

        total_in = t["total_in"]
        total_out = t["total_out"]
        profit = total_in - total_out

        self.lbl_rep_summary.setText(
            f"<b>سال کاری {work_year}</b>  |  "
            f"ورودی کل: <b>{total_in:,}</b> تومان  |  "
            f"خروجی کل: <b>{total_out:,}</b> تومان  |  "
            f"سود خالص: <span style='color:{'#27ae60' if profit >= 0 else '#c0392b'};'>"
            f"<b>{profit:,}</b> تومان</span><br>"
            f"<b>سود خالص به حروف:</b> {number_to_persian_words(abs(profit))} تومان "
            f"{'' if profit >= 0 else '(زیان)'}"
        )

        self.rep_table.setRowCount(0)
        records = [
            ("قراردادهای عروس و داماد (خام)", "درآمد (ورودی)", t["w_raw"], "جمع قیمت خام قراردادها"),
            ("تخفیف داده‌شده به مشتریان", "کاهش درآمد", t["w_disc"], "جمع تخفیف‌ها"),
            ("قراردادهای عروس و داماد (بعد از تخفیف)", "درآمد (ورودی)", t["w_net"], "مبلغ نهایی قراردادها"),
            ("پروژه‌های تبلیغاتی / بیوتی / تولدی", "درآمد (ورودی)", t["c_in"], "مجموع ورودی‌ها"),
            ("هزینه‌های جاری و قبوض", "هزینه (خروجی)", t["exp_out"], "مجموع هزینه‌ها"),
            ("حقوق و پرداختی پرسنل", "هزینه (خروجی)", t["staff_out"], "مجموع پرداختی‌ها"),
        ]

        person_id = self.combo_report_person_fin.currentData()
        if report_type == "گزارش هر فرد" and person_id:
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT name, role FROM persons WHERE id=?", (person_id,))
            p = cursor.fetchone()
            if p:
                cursor.execute("""SELECT amount, year, month, day, description
                                  FROM transactions WHERE person_id=? AND work_year=?
                                  ORDER BY year DESC, month DESC, day DESC""",
                               (person_id, work_year))
                trans = cursor.fetchall()
                cursor.execute("SELECT COALESCE(SUM(amount),0) FROM transactions "
                               "WHERE person_id=? AND work_year=?", (person_id, work_year))
                total_person = cursor.fetchone()[0]
                self.lbl_rep_summary.setText(
                    f"<b>گزارش فردی:</b> {p[0]} ({p[1]})  |  "
                    f"<b>جمع کل پرداختی:</b> {total_person:,} تومان  |  "
                    f"<b>تعداد تراکنش:</b> {len(trans)}<br>"
                    f"<b>به حروف:</b> {number_to_persian_words(total_person)} تومان"
                )
                records = []
                for tr in trans:
                    records.append((f"پرداختی {tr[1]}/{tr[2]:02d}/{tr[3]:02d}",
                                    "حقوق کارکنان", tr[0], tr[4] if tr[4] else "-"))
            conn.close()

        for idx, rec in enumerate(records):
            self.rep_table.insertRow(idx)
            self.rep_table.setItem(idx, 0, QTableWidgetItem(str(idx + 1)))
            self.rep_table.setItem(idx, 1, QTableWidgetItem(str(rec[0])))
            self.rep_table.setItem(idx, 2, QTableWidgetItem(str(rec[1])))
            amount_item = QTableWidgetItem(f"{rec[2]:,}")
            if "خروجی" in str(rec[1]) or "کاهش" in str(rec[1]):
                amount_item.setForeground(QColor("#c0392b"))
            else:
                amount_item.setForeground(QColor("#1e8449"))
            self.rep_table.setItem(idx, 3, amount_item)
            self.rep_table.setItem(idx, 4, QTableWidgetItem(str(rec[3])))

    # ============================================================
    # ====================  ساخت نمودارها  ======================
    # ============================================================
    def _new_figure(self):
        setup_matplotlib_font()
        fig = Figure(figsize=(9.0, 5.2), facecolor="#f8fbff")
        return fig

    def _style_axes(self, ax, money_axis=None, grid_axis="y"):
        """استایل مشترک نمودارها با محور پولی سه‌رقمی و راست‌چین"""
        ax.set_facecolor("#ffffff")
        for spine in ax.spines.values():
            spine.set_color("#c9d7e6")
        ax.grid(axis=grid_axis, alpha=0.25, linestyle="--", color="#7f8c8d")
        ax.grid(axis=("y" if grid_axis == "x" else "x"), visible=False)
        ax.tick_params(axis="both", labelsize=9, colors="#2c3e50")
        comma = plt.FuncFormatter(lambda x, p: f"{int(x):,}")
        if money_axis == "x":
            ax.xaxis.set_major_formatter(comma)
        elif money_axis == "y":
            ax.yaxis.set_major_formatter(comma)

    def _draw_barh(self, ax, labels, amounts, color, edge):
        """نوار افقی راست‌چین همراه برچسب مبلغ بدون همپوشانی"""
        bars = ax.barh(labels, amounts, color=color, edgecolor=edge, linewidth=1.1, height=0.6)
        ax.yaxis.tick_right()
        self._style_axes(ax, money_axis="x", grid_axis="x")
        maxv = max(amounts + [1])
        ax.set_xlim(0, maxv * 1.28)
        for bar, amt in zip(bars, amounts):
            ax.text(bar.get_width() + maxv * 0.02,
                    bar.get_y() + bar.get_height() / 2, f"{amt:,}",
                    va="center", ha="left", fontsize=9, fontweight="bold", color="#2c3e50")
        return bars

    def build_chart(self, kind):
        """ساخت شکل نمودار بر اساس نوع انتخاب‌شده — با متن فارسی درست"""
        t = self._report_totals()
        wy = t["year"]
        fig = self._new_figure()
        ax = fig.add_subplot(111)

        if kind == "income_expense":
            labels = [fa_shape("درآمد عروسی (بعد از تخفیف)"), fa_shape("درآمد تبلیغاتی"),
                      fa_shape("هزینه‌ها"), fa_shape("حقوق پرسنل")]
            amounts = [t["w_net"], t["c_in"], t["exp_out"], t["staff_out"]]
            labels = labels[::-1]
            amounts = amounts[::-1]
            self._draw_barh(ax, labels, amounts,
                            ["#8e44ad", "#c0392b", "#2980b9", "#27ae60"], "#2c3e50")
            ax.set_title(fa_shape(f"نمودار درآمد و هزینه - سال کاری {wy}"),
                         fontsize=13, fontweight="bold", pad=14, color="#1F4E78")
            ax.set_xlabel(fa_shape("مبلغ (تومان)"), fontsize=10, labelpad=6)

        elif kind == "monthly_income":
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT ceremony_date FROM wedding_contracts "
                           "WHERE work_year=? AND ceremony_date IS NOT NULL", (wy,))
            dates = [r[0] for r in cursor.fetchall()]
            cursor.execute("SELECT project_date FROM commercial_projects "
                           "WHERE work_year=? AND project_date IS NOT NULL", (wy,))
            dates += [r[0] for r in cursor.fetchall()]
            conn.close()
            buckets = {m: 0 for m in range(1, 13)}
            for d in dates:
                jd = parse_jalali(d)
                if jd:
                    buckets[jd.month] += 1
            labels = [fa_shape(jalali_month_name(m)) for m in range(1, 13)][::-1]
            counts = [buckets[m] for m in range(1, 13)][::-1]
            bars = ax.bar(labels, counts, color="#2980b9", edgecolor="#1F4E78", linewidth=1.0)
            for bar, cnt in zip(bars, counts):
                if cnt:
                    ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height() + 0.05,
                            str(cnt), ha="center", va="bottom", fontsize=9, fontweight="bold")
            ax.set_title(fa_shape(f"نمودار تعداد مراسم و پروژه‌ها به تفکیک ماه - سال {wy}"),
                         fontsize=13, fontweight="bold", pad=14, color="#1F4E78")
            ax.set_ylabel(fa_shape("تعداد"), fontsize=10)
            self._style_axes(ax)
            ax.yaxis.set_major_formatter(plt.FuncFormatter(lambda x, p: f"{int(x)}"))

        elif kind == "expense_categories":
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT category, COALESCE(SUM(amount),0) FROM expenses "
                           "WHERE work_year=? GROUP BY category ORDER BY 2 DESC", (wy,))
            rows = cursor.fetchall()
            conn.close()
            if not rows:
                labels = [fa_shape("هزینه‌ای ثبت نشده")]
                amounts = [1]
                colors = ["#dfe6e9"]
            else:
                labels = [fa_shape(r[0] or "-") for r in rows][::-1]
                amounts = [r[1] for r in rows][::-1]
                colors = plt.cm.Set3([i / max(1, len(rows) - 1) for i in range(len(rows))])[::-1]
            wedges, _txt, _auto = ax.pie(
                amounts, labels=None, autopct=lambda p: f"{p:.1f}%",
                colors=colors, startangle=90, textprops={"fontsize": 9})
            ax.legend(wedges, labels, loc="center right", bbox_to_anchor=(1.05, 0.5),
                      fontsize=9, frameon=False)
            ax.set_title(fa_shape(f"نمودار هزینه‌ها به تفکیک دسته‌بندی - سال {wy}"),
                         fontsize=13, fontweight="bold", pad=14, color="#1F4E78")
            ax.axis("equal")

        elif kind == "staff_salaries":
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("""SELECT p.name, COALESCE(SUM(t.amount),0) AS s
                              FROM transactions t JOIN persons p ON t.person_id=p.id
                              WHERE t.work_year=? GROUP BY p.id ORDER BY s DESC LIMIT 15""", (wy,))
            rows = cursor.fetchall()
            conn.close()
            if not rows:
                labels, amounts = [fa_shape("پرداختی ثبت نشده")], [0]
            else:
                labels = [fa_shape(r[0]) for r in rows][::-1]
                amounts = [r[1] for r in rows][::-1]
            self._draw_barh(ax, labels, amounts, "#8e44ad", "#5b2c8e")
            ax.set_title(fa_shape(f"نمودار حقوق و پرداختی پرسنل - سال {wy}"),
                         fontsize=13, fontweight="bold", pad=14, color="#1F4E78")
            ax.set_xlabel(fa_shape("مبلغ (تومان)"), fontsize=10, labelpad=6)

        elif kind == "inventory_status":
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT item_name, total_count, used_count FROM inventory ORDER BY item_name")
            rows = cursor.fetchall()
            conn.close()
            if not rows:
                rows = [("انبار خالی", 0, 0)]
            labels = [fa_shape(r[0]) for r in rows][::-1]
            totals = [max(0, (r[1] or 0) - (r[2] or 0)) for r in rows][::-1]
            used = [r[2] or 0 for r in rows][::-1]
            idx = range(len(labels))
            ax.barh(idx, totals, color="#27ae60", edgecolor="#1e8449", height=0.55, label=fa_shape("موجود"))
            ax.barh(idx, used, left=totals, color="#e67e22", edgecolor="#b9600f",
                    height=0.55, label=fa_shape("در استفاده"))
            ax.set_yticks(list(idx))
            ax.set_yticklabels(labels)
            ax.yaxis.tick_right()
            self._style_axes(ax, grid_axis="x")
            ax.legend(fontsize=9, frameon=False, loc="lower right", bbox_to_anchor=(1.0, -0.18), ncol=2)
            ax.set_title(fa_shape("نمودار وضعیت موجودی انبار تجهیزات"),
                         fontsize=13, fontweight="bold", pad=14, color="#1F4E78")
            ax.set_xlabel(fa_shape("تعداد"), fontsize=10)

        elif kind == "checks_status":
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT COALESCE(check_type,'دریافتی'), COALESCE(SUM(amount),0), COUNT(*) "
                           "FROM checks WHERE work_year=? GROUP BY check_type", (wy,))
            rows = cursor.fetchall()
            conn.close()
            if not rows:
                labels, amounts = [fa_shape("چکی ثبت نشده")], [1]
                colors = ["#dfe6e9"]
            else:
                labels = [fa_shape(r[0]) for r in rows]
                amounts = [r[1] for r in rows]
                colors = ["#27ae60" if "دریافتی" in r[0] else "#c0392b" for r in rows]
            ax.pie(amounts, labels=labels, autopct=lambda p: f"{p:.1f}%", colors=colors,
                   startangle=90, textprops={"fontsize": 10})
            ax.set_title(fa_shape(f"نمودار وضعیت چک‌های دریافتی و پرداختی - سال {wy}"),
                         fontsize=13, fontweight="bold", pad=14, color="#1F4E78")
            ax.axis("equal")

        elif kind == "monthly_step":
            wedding, comm, exp = self._monthly_series()
            months = [fa_shape(jalali_month_name(m)) for m in range(1, 13)][::-1]
            values = [(wedding[m] + comm[m]) for m in range(1, 13)][::-1]
            xs = list(range(len(values)))
            ax.step(xs, values, where="mid", color="#2980b9", linewidth=2.4)
            ax.fill_between(xs, values, step="mid", color="#2980b9", alpha=0.16)
            ax.set_xticks(xs)
            ax.set_xticklabels(months, fontsize=8)
            self._style_axes(ax, money_axis="y")
            ax.set_title(fa_shape(f"نمودار پلکانی درآمد ماهانه - سال {wy}"),
                         fontsize=12.5, fontweight="bold", pad=12, color="#1F4E78")
            ax.set_ylabel(fa_shape("مبلغ (تومان)"), fontsize=10)

        elif kind == "scatter":
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("""SELECT ceremony_date, total_amount, discount FROM wedding_contracts
                              WHERE work_year=?""", (wy,))
            pts = cursor.fetchall()
            conn.close()
            xs, ys = [], []
            for d, raw, disc in pts:
                jd = parse_jalali(d)
                if jd:
                    xs.append(jd.month + jd.day / 31.0)
                    ys.append(max(0, (raw or 0) - (disc or 0)))
            ax.scatter(xs, ys, s=70, color="#8e44ad", edgecolor="#5b2c8e", alpha=0.8)
            ax.set_xticks(range(1, 13))
            ax.set_xticklabels([fa_shape(jalali_month_name(m)) for m in range(1, 13)], fontsize=8)
            self._style_axes(ax, money_axis="y")
            ax.set_title(fa_shape(f"نمودار نقطه‌ای مبلغ مراسم به تفکیک ماه - سال {wy}"),
                         fontsize=12.5, fontweight="bold", pad=12, color="#1F4E78")
            ax.set_ylabel(fa_shape("مبلغ قرارداد (تومان)"), fontsize=10)

        elif kind == "line_trend":
            wedding, comm, exp = self._monthly_series()
            months = [fa_shape(jalali_month_name(m)) for m in range(1, 13)][::-1]
            income = [(wedding[m] + comm[m]) for m in range(1, 13)][::-1]
            costs = [exp[m] for m in range(1, 13)][::-1]
            xs = list(range(len(months)))
            ax.plot(xs, income, marker="o", linewidth=2.2, color="#1e8449",
                    label=fa_shape("درآمد"))
            ax.plot(xs, costs, marker="s", linewidth=2.2, color="#c0392b",
                    label=fa_shape("هزینه"))
            ax.set_xticks(xs)
            ax.set_xticklabels(months, fontsize=8)
            self._style_axes(ax, money_axis="y")
            ax.legend(fontsize=9, frameon=False)
            ax.set_title(fa_shape(f"نمودار خطی روند درآمد و هزینه - سال {wy}"),
                         fontsize=12.5, fontweight="bold", pad=12, color="#1F4E78")
            ax.set_ylabel(fa_shape("مبلغ (تومان)"), fontsize=10)

        elif kind == "donut":
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT category, COALESCE(SUM(amount),0) FROM expenses "
                           "WHERE work_year=? GROUP BY category ORDER BY 2 DESC", (wy,))
            rows = cursor.fetchall()
            conn.close()
            if not rows:
                labels, amounts = [fa_shape("هزینه‌ای ثبت نشده")], [1]
                colors = ["#dfe6e9"]
            else:
                labels = [fa_shape(r[0] or "-") for r in rows]
                amounts = [r[1] for r in rows]
                colors = list(plt.cm.Pastel1(range(len(rows))))
            wedges, _t, _a = ax.pie(amounts, labels=None, autopct=lambda p: f"{p:.1f}%",
                                    colors=colors, startangle=90,
                                    wedgeprops=dict(width=0.42, edgecolor="white"),
                                    textprops={"fontsize": 8.5})
            ax.legend(wedges, labels, loc="center right", bbox_to_anchor=(1.12, 0.5),
                      fontsize=8.5, frameon=False)
            ax.set_title(fa_shape(f"نمودار دونات هزینه‌ها به تفکیک دسته - سال {wy}"),
                         fontsize=12.5, fontweight="bold", pad=12, color="#1F4E78")
            ax.axis("equal")

        elif kind == "stacked":
            wedding, comm, exp = self._monthly_series()
            months = [fa_shape(jalali_month_name(m)) for m in range(1, 13)][::-1]
            xs = list(range(len(months)))
            w_vals = [wedding[m] for m in range(1, 13)][::-1]
            c_vals = [comm[m] for m in range(1, 13)][::-1]
            e_vals = [exp[m] for m in range(1, 13)][::-1]
            ax.bar(xs, w_vals, color="#27ae60", edgecolor="#1e8449", width=0.62,
                   label=fa_shape("درآمد عروسی"))
            ax.bar(xs, c_vals, bottom=w_vals, color="#2980b9", edgecolor="#1F4E78",
                   width=0.62, label=fa_shape("درآمد تبلیغاتی"))
            ax.bar(xs, e_vals, color="#c0392b", edgecolor="#a93226", width=0.32,
                   label=fa_shape("هزینه‌ها"))
            ax.set_xticks(xs)
            ax.set_xticklabels(months, fontsize=8)
            self._style_axes(ax, money_axis="y")
            ax.legend(fontsize=8.5, frameon=False, ncol=3)
            ax.set_title(fa_shape(f"نمودار انباشته درآمد و هزینه ماهانه - سال {wy}"),
                         fontsize=12.5, fontweight="bold", pad=12, color="#1F4E78")
            ax.set_ylabel(fa_shape("مبلغ (تومان)"), fontsize=10)

        elif kind == "area":
            wedding, comm, exp = self._monthly_series()
            months = [fa_shape(jalali_month_name(m)) for m in range(1, 13)][::-1]
            xs = list(range(len(months)))
            profit = [(wedding[m] + comm[m] - exp[m]) for m in range(1, 13)][::-1]
            cum = []
            run = 0
            for p_ in profit:
                run += p_
                cum.append(run)
            ax.fill_between(xs, cum, color="#16a085", alpha=0.35)
            ax.plot(xs, cum, color="#0e6655", linewidth=2.4, marker="o", markersize=4)
            ax.axhline(0, color="#7f8c8d", linewidth=1, linestyle="--")
            ax.set_xticks(xs)
            ax.set_xticklabels(months, fontsize=8)
            self._style_axes(ax, money_axis="y")
            ax.set_title(fa_shape(f"نمودار سطحی روند انباشته سود - سال {wy}"),
                         fontsize=12.5, fontweight="bold", pad=12, color="#1F4E78")
            ax.set_ylabel(fa_shape("سود انباشته (تومان)"), fontsize=10)

        fig.tight_layout()
        return fig

    def _monthly_series(self):
        """درآمد عروسی، درآمد تبلیغاتی و هزینه‌ها به تفکیک ماه شمسی"""
        wy = get_current_year()
        wedding = [0] * 13
        comm = [0] * 13
        exp = [0] * 13
        try:
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("SELECT ceremony_date, total_amount, discount FROM wedding_contracts "
                           "WHERE work_year=?", (wy,))
            for d, raw, disc in cursor.fetchall():
                jd = parse_jalali(d)
                if jd and 1 <= jd.month <= 12:
                    wedding[jd.month] += max(0, (raw or 0) - (disc or 0))
            cursor.execute("SELECT project_date, total_amount FROM commercial_projects "
                           "WHERE work_year=?", (wy,))
            for d, amt in cursor.fetchall():
                jd = parse_jalali(d)
                if jd and 1 <= jd.month <= 12:
                    comm[jd.month] += (amt or 0)
            cursor.execute("SELECT date_str, amount FROM expenses WHERE work_year=?", (wy,))
            for d, amt in cursor.fetchall():
                jd = parse_jalali(d)
                if jd and 1 <= jd.month <= 12:
                    exp[jd.month] += (amt or 0)
            conn.close()
        except Exception as e:
            print("[CHART] monthly series error:", e)
        return wedding, comm, exp


    def show_financial_chart(self):
        """نمایش نمودار انتخابی با عنوان فارسی درست (رفع برعکس و جدا بودن حروف)"""
        try:
            idx = self.combo_chart_kind.currentIndex()
            if idx < 0:
                idx = 0
            title, kind = self.chart_kinds[idx]
            title = f"{title} - سال {get_current_year()}"
            fig = self.build_chart(kind)
            dlg = ChartDialog(fig, title, self)
            dlg.exec()
        except Exception as e:
            QMessageBox.critical(self, "خطا در نمودار", f"ساخت نمودار ناموفق بود:\n{e}")

    def financial_report_html(self, mono=False):
        p = doc_palette(mono)
        t = self._report_totals()
        wy = t["year"]
        total_in, total_out = t["total_in"], t["total_out"]
        profit = total_in - total_out
        studio_name, studio_phone, studio_address = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()

        def row(i, title, kind, amount, bg=None):
            return InvoiceBuilder._row([
                InvoiceBuilder._td(str(i), size="8pt", bg=bg),
                InvoiceBuilder._td(title, align="right", size="8pt", bg=bg),
                InvoiceBuilder._td(kind, size="8pt", bg=bg),
                InvoiceBuilder._td(f"{amount:,}", size="8pt", bg=bg),
            ])

        rows = (row(1, "قراردادهای عروس و داماد (قیمت خام)", "درآمد", t["w_raw"]) +
                row(2, "تخفیف داده‌شده به مشتریان", "کاهش درآمد", t["w_disc"], p["alt_row"]) +
                row(3, "قراردادهای عروس و داماد (بعد از تخفیف)", "درآمد", t["w_net"]) +
                row(4, "پروژه‌های تبلیغاتی / بیوتی / تولدی", "درآمد", t["c_in"]) +
                row(5, "هزینه‌های جاری و قبوض", "هزینه", t["exp_out"], p["alt_row"]) +
                row(6, "حقوق و پرداختی پرسنل", "هزینه", t["staff_out"], p["alt_row"]))

        summary = InvoiceBuilder._row([
            InvoiceBuilder._td(f"<b>ورودی کل</b><br>{total_in:,} تومان", align="center",
                               bg=p["alt_row"], size="8.5pt"),
            InvoiceBuilder._td(f"<b>خروجی کل</b><br>{total_out:,} تومان", align="center",
                               bg=p["box_bg"], size="8.5pt"),
            InvoiceBuilder._td(f"<b>سود خالص</b><br>{profit:,} تومان", align="center",
                               bg=p["total_bg"], size="9pt", bold=True),
        ])

        return f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {InvoiceBuilder._title_band(
              p, studio_name,
              f"<div style='font-size:10pt;'>گزارش مالی جامع - سال کاری {wy}</div>",
              [f"تاریخ چاپ: {now.strftime('%Y/%m/%d')}", f"ساعت: {now.strftime('%H:%M')}",
               f"تعداد پرسنل: {t['persons']}"])}
          <table width="100%" style="border-collapse:collapse; margin-top:5px;">
            <thead>{InvoiceBuilder._head(['ردیف', 'بخش', 'نوع', 'مبلغ (تومان)'],
                                         bg=p['table_head'], fg=p['table_head_fg'])}</thead>
            <tbody>{rows}</tbody>
          </table>
          <table width="100%" style="border-collapse:collapse; margin-top:5px;">{summary}</table>
          <div style="margin-top:6px; font-size:8.5pt;">
            <b>سود خالص به حروف:</b> {number_to_persian_words(abs(profit))} تومان
            {'' if profit >= 0 else '(زیان)'}
          </div>
          <div style="text-align:center; margin-top:30px; font-size:9pt;">مهر و امضای مدیر / حسابداری</div>
          <div style="text-align:center; margin-top:10px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}{(' | ' + studio_address) if studio_address else ''}
          </div>
        </div>
        """

    def print_financial_report(self):
        self._output_flow(lambda wd, ws, mono: self.financial_report_html(mono),
                          f"گزارش_مالی_{get_current_year()}.pdf",
                          "گزارش مالی جامع", allow_details=False)

    # ============================================================
    # ==============  تب ۶: جستجوی پیشرفته  =====================
    # ============================================================
    def setup_search_tab(self):
        layout = QVBoxLayout()
        layout.setSpacing(8)

        search_bar = QHBoxLayout()
        search_bar.setSpacing(7)

        self.txt_search_keyword = QLineEdit()
        self.txt_search_keyword.setPlaceholderText("جستجو بر اساس نام، کد، تاریخ، عنوان...")
        self.txt_search_keyword.returnPressed.connect(self.perform_search)

        self.chk_search_by_date = QCheckBox("جستجو بر اساس تاریخ شمسی")
        self.chk_search_by_date.stateChanged.connect(self.toggle_date_search)

        self.date_search = PersianDateEdit(date_str=jalali_today_str())
        self.date_search.setEnabled(False)

        btn_search = QPushButton(ico_text("search", "جستجو"))
        btn_search.setStyleSheet("background-color: #16a085; color: white; font-weight: bold; padding: 7px;")
        btn_search.clicked.connect(self.perform_search)

        btn_clear = QPushButton(ico_text("filter", "پاک کردن"))
        btn_clear.setStyleSheet("background-color: #95a5a6; color: white; font-weight: bold; padding: 7px;")
        btn_clear.clicked.connect(self.clear_search)

        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color: #1e8449; color: white; font-weight: bold; padding: 7px;")
        btn_excel.clicked.connect(lambda: self.export_table_excel(
            self.search_table, "نتیجه جستجو", "جستجو.xlsx"))

        search_bar.addWidget(QLabel("عبارت:"))
        search_bar.addWidget(self.txt_search_keyword, 3)
        search_bar.addWidget(self.chk_search_by_date)
        search_bar.addWidget(self.date_search)
        search_bar.addWidget(btn_search)
        search_bar.addWidget(btn_clear)
        search_bar.addWidget(btn_excel)
        layout.addLayout(search_bar)

        lbl_hint = QLabel(f"{ico('calendar')} در کادر تاریخ، ابتدا <b>روز</b> را بنویسید؛ "
                          "پس از کامل شدن، نشانگر خودکار به <b>ماه</b> و سپس <b>سال</b> می‌رود.")
        lbl_hint.setStyleSheet("background-color:#eaf6ff; border:1px solid #a9d3ee; "
                               "border-radius:8px; padding:7px;")
        layout.addWidget(lbl_hint)

        self.search_table = QTableWidget()
        self.search_table.setColumnCount(7)
        self.search_table.setHorizontalHeaderLabels([
            "نوع", "کد", "عنوان / نام", "تاریخ", "مبلغ کل", "پرداختی", "مانده"
        ])
        self.search_table.setAlternatingRowColors(True)
        self.search_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        persist_table(self.search_table, "search_results", QHeaderView.ResizeMode.Stretch)
        self.search_table.doubleClicked.connect(self.on_search_double_click)
        layout.addWidget(self.search_table)

        detail_box = QGroupBox("جزئیات کامل (برای مشاهده، روی نتیجه دوبار کلیک کنید)")
        detail_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        dl = QVBoxLayout()
        self.lbl_search_detail = QLabel("نتیجه‌ای انتخاب نشده است.")
        self.lbl_search_detail.setWordWrap(True)
        self.lbl_search_detail.setStyleSheet(
            "background-color: #f8f9fa; padding: 12px; border-radius: 8px; "
            "border: 1px solid #bdc3c7; font-size: 11pt;")
        self.lbl_search_detail.setTextInteractionFlags(Qt.TextInteractionFlag.TextSelectableByMouse)
        dl.addWidget(self.lbl_search_detail)
        detail_box.setLayout(dl)
        layout.addWidget(detail_box, 1)

        self.tab_search.setLayout(layout)

    def toggle_date_search(self):
        checked = self.chk_search_by_date.isChecked()
        self.date_search.setEnabled(checked)
        self.txt_search_keyword.setEnabled(not checked)

    def clear_search(self):
        self.txt_search_keyword.clear()
        self.chk_search_by_date.setChecked(False)
        self.date_search.setEnabled(False)
        self.txt_search_keyword.setEnabled(True)
        self.search_table.setRowCount(0)
        self.lbl_search_detail.setText("نتیجه‌ای انتخاب نشده است.")

    def perform_search(self):
        self.search_table.setRowCount(0)
        self.lbl_search_detail.setText("در حال جستجو...")

        results = []

        if self.chk_search_by_date.isChecked():
            # تاریخ شمسی انتخاب‌شده (روز/ماه/سال)
            date_str = self.date_search.text()
            kw_pattern = f"%{date_str}%"
            use_kw = False
        else:
            kw = self.txt_search_keyword.text().strip()
            if not kw:
                QMessageBox.warning(self, "خطا", "عبارت جستجو را وارد کنید.")
                return
            kw_pattern = f"%{kw}%"
            use_kw = True

        work_year = get_current_year()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        # ---------- قراردادهای عروس و داماد ----------
        if use_kw:
            cursor.execute("""SELECT id, code, groom_name, bride_name, ceremony_date,
                              total_amount, discount, paid_amount
                              FROM wedding_contracts
                              WHERE work_year=? AND (
                                groom_name LIKE ? OR bride_name LIKE ? OR
                                code LIKE ? OR ceremony_date LIKE ? OR
                                contract_date LIKE ? OR
                                groom_phone LIKE ? OR bride_phone LIKE ?)""",
                           (work_year, kw_pattern, kw_pattern, kw_pattern, kw_pattern,
                            kw_pattern, kw_pattern, kw_pattern))
        else:
            cursor.execute("""SELECT id, code, groom_name, bride_name, ceremony_date,
                              total_amount, discount, paid_amount
                              FROM wedding_contracts
                              WHERE work_year=? AND (ceremony_date LIKE ? OR contract_date LIKE ?)""",
                           (work_year, kw_pattern, kw_pattern))
        for r in cursor.fetchall():
            net = max(0, (r[5] or 0) - (r[6] or 0))
            results.append({
                "type": "قرارداد عروسی",
                "id": r[0], "code": r[1], "title": f"{r[2]} و {r[3]}",
                "date": r[4], "total": net, "paid": r[7] or 0,
                "table": "wedding_contracts"})

        # ---------- پروژه‌های تبلیغاتی ----------
        if use_kw:
            cursor.execute("""SELECT id, code, title, client_name, project_date,
                              total_amount, paid_amount
                              FROM commercial_projects
                              WHERE work_year=? AND (
                                title LIKE ? OR client_name LIKE ? OR
                                code LIKE ? OR project_date LIKE ? OR client_phone LIKE ?)""",
                           (work_year, kw_pattern, kw_pattern, kw_pattern,
                            kw_pattern, kw_pattern))
        else:
            cursor.execute("""SELECT id, code, title, client_name, project_date,
                              total_amount, paid_amount
                              FROM commercial_projects
                              WHERE work_year=? AND project_date LIKE ?""",
                           (work_year, kw_pattern))
        for r in cursor.fetchall():
            results.append({
                "type": "پروژه تبلیغاتی",
                "id": r[0], "code": r[1], "title": r[2],
                "date": r[4], "total": r[5] or 0, "paid": r[6] or 0,
                "table": "commercial_projects"})

        # ---------- هزینه‌ها ----------
        if use_kw:
            cursor.execute("""SELECT id, code, title, category, date_str, amount
                              FROM expenses
                              WHERE work_year=? AND (
                                title LIKE ? OR category LIKE ? OR
                                code LIKE ? OR date_str LIKE ?)""",
                           (work_year, kw_pattern, kw_pattern, kw_pattern, kw_pattern))
        else:
            cursor.execute("""SELECT id, code, title, category, date_str, amount
                              FROM expenses WHERE work_year=? AND date_str LIKE ?""",
                           (work_year, kw_pattern))
        for r in cursor.fetchall():
            results.append({
                "type": "هزینه",
                "id": r[0], "code": r[1], "title": r[2],
                "date": r[4], "total": r[5] or 0, "paid": r[5] or 0,
                "table": "expenses"})

        # ---------- چک‌ها ----------
        if use_kw:
            cursor.execute("""SELECT id, check_number, bank_name, issuer_name, due_date, amount,
                              is_passed, COALESCE(check_type,'دریافتی')
                              FROM checks
                              WHERE work_year=? AND (
                                check_number LIKE ? OR bank_name LIKE ? OR
                                issuer_name LIKE ? OR due_date LIKE ? OR check_type LIKE ?)""",
                           (work_year, kw_pattern, kw_pattern, kw_pattern,
                            kw_pattern, kw_pattern))
        else:
            cursor.execute("""SELECT id, check_number, bank_name, issuer_name, due_date, amount,
                              is_passed, COALESCE(check_type,'دریافتی')
                              FROM checks WHERE work_year=? AND due_date LIKE ?""",
                           (work_year, kw_pattern))
        for r in cursor.fetchall():
            results.append({
                "type": f"چک {r[7]}",
                "id": r[0], "code": r[1], "title": f"{r[2]} - {r[3] or '-'}",
                "date": r[4], "total": r[5] or 0,
                "paid": (r[5] or 0) if r[6] else 0,
                "table": "checks"})

        # ---------- پرسنل ----------
        if use_kw:
            cursor.execute("""SELECT id, code, name, role, phone FROM persons
                              WHERE name LIKE ? OR code LIKE ? OR role LIKE ? OR phone LIKE ?""",
                           (kw_pattern, kw_pattern, kw_pattern, kw_pattern))
            for r in cursor.fetchall():
                results.append({
                    "type": "پرسنل",
                    "id": r[0], "code": r[1], "title": f"{r[2]} ({r[3]})",
                    "date": "-", "total": 0, "paid": 0,
                    "table": "persons"})

        # ---------- پرداختی پرسنل ----------
        if use_kw:
            cursor.execute("""SELECT t.id, p.code, p.name, t.description,
                                     t.year, t.month, t.day, t.amount
                              FROM transactions t
                              JOIN persons p ON t.person_id = p.id
                              WHERE t.work_year=? AND (
                                p.name LIKE ? OR p.code LIKE ? OR t.description LIKE ?)""",
                           (work_year, kw_pattern, kw_pattern, kw_pattern))
            for r in cursor.fetchall():
                results.append({
                    "type": "پرداختی پرسنل",
                    "id": r[0], "code": r[1], "title": f"پرداختی به {r[2]}",
                    "date": f"{r[4]}/{r[5]:02d}/{r[6]:02d}",
                    "total": r[7] or 0, "paid": r[7] or 0,
                    "table": "transactions"})

        conn.close()

        self.search_table.setRowCount(0)
        for r_idx, res in enumerate(results):
            self.search_table.insertRow(r_idx)
            self.search_table.setItem(r_idx, 0, QTableWidgetItem(res["type"]))
            self.search_table.setItem(r_idx, 1, QTableWidgetItem(str(res["code"] or "-")))
            self.search_table.setItem(r_idx, 2, QTableWidgetItem(res["title"]))
            self.search_table.setItem(r_idx, 3, QTableWidgetItem(str(res["date"])))
            self.search_table.setItem(r_idx, 4,
                                      QTableWidgetItem(f"{res['total']:,}" if res['total'] else "-"))
            remain = res['total'] - res['paid']
            self.search_table.setItem(r_idx, 5,
                                      QTableWidgetItem(f"{res['paid']:,}" if res['paid'] else "-"))
            self.search_table.setItem(r_idx, 6,
                                      QTableWidgetItem(f"{remain:,}" if res['total'] else "-"))
            self.search_table.item(r_idx, 0).setData(Qt.ItemDataRole.UserRole, res)

        self.txt_search_keyword.clear()
        self.lbl_search_detail.setText(
            f"<b>{len(results)}</b> نتیجه یافت شد. برای مشاهده جزئیات روی نتیجه دوبار کلیک کنید.")

    def on_search_double_click(self, index):
        row = index.row()
        item = self.search_table.item(row, 0)
        if not item:
            return
        res = item.data(Qt.ItemDataRole.UserRole)
        if not res:
            return

        detail_html = f"<h3 style='color:#2980b9;'>جزئیات: {res['type']}</h3><hr>"
        detail_html += f"<p><b>کد:</b> {res['code']}<br>"
        detail_html += f"<b>عنوان:</b> {res['title']}<br>"
        detail_html += f"<b>تاریخ:</b> {res['date']}</p>"

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        if res["table"] == "wedding_contracts":
            cursor.execute("""SELECT groom_phone, bride_phone, contract_date,
                                     ceremony_date, selected_items, description,
                                     discount, is_settled, total_amount, staff_ids
                                     FROM wedding_contracts WHERE id=?""", (res["id"],))
            r = cursor.fetchone()
            if r:
                detail_html += f"<p><b>تلفن داماد:</b> {r[0] or '-'} | <b>تلفن عروس:</b> {r[1] or '-'}<br>"
                detail_html += f"<b>تاریخ قرارداد:</b> {r[2]} | <b>تاریخ مراسم:</b> {r[3]}<br>"
                detail_html += f"<b>جمع خام:</b> {r[8]:,} تومان | <b>تخفیف:</b> {r[6]:,} تومان | "
                detail_html += f"<b>بعد از تخفیف:</b> {max(0,(r[8] or 0)-(r[6] or 0)):,} تومان<br>"
                detail_html += f"<b>وضعیت:</b> {'✅ تسویه' if r[7] else '⏳ در جریان'}</p>"
                detail_html += f"<p><b>توضیحات:</b> {r[5] or '-'}</p>"
                items = [x.strip() for x in (r[4].split(',') if r[4] else []) if x.strip()]
                detail_html += "<p><b>خدمات انتخاب شده:</b></p><ul>"
                for it in items:
                    detail_html += f"<li>{it}</li>"
                detail_html += "</ul>"
                try:
                    staff = json.loads(r[9]) if r[9] else []
                except Exception:
                    staff = []
                if staff:
                    detail_html += "<p><b>نیروی کار این قرارداد:</b></p><ul>"
                    for m in staff:
                        detail_html += f"<li>{m.get('name')} ({m.get('role')}) - " \
                                       f"{int(m.get('wage') or 0):,} تومان</li>"
                    detail_html += "</ul>"

                cursor.execute("SELECT amount, deposit_date, bank_name FROM wedding_deposits "
                               "WHERE contract_id=? ORDER BY id", (res["id"],))
                deps = cursor.fetchall()
                if deps:
                    detail_html += "<p><b>بیعانه‌ها:</b></p><ul>"
                    for d in deps:
                        detail_html += f"<li>{d[0]:,} تومان - {d[1]} - {d[2] or '-'}</li>"
                    detail_html += "</ul>"

        elif res["table"] == "commercial_projects":
            cursor.execute("""SELECT project_type, camera_count, description,
                                     client_name, client_phone, project_date
                                     FROM commercial_projects WHERE id=?""", (res["id"],))
            r = cursor.fetchone()
            if r:
                detail_html += f"<p><b>نوع:</b> {r[0]} | <b>دوربین:</b> {r[1]}<br>"
                detail_html += f"<b>مشتری:</b> {r[3] or '-'} | <b>تلفن:</b> {r[4] or '-'}<br>"
                detail_html += f"<b>تاریخ پروژه:</b> {r[5]}</p>"
                detail_html += f"<p><b>توضیحات:</b> {r[2] or '-'}</p>"

        elif res["table"] == "expenses":
            cursor.execute("SELECT category, description FROM expenses WHERE id=?", (res["id"],))
            r = cursor.fetchone()
            if r:
                detail_html += f"<p><b>دسته:</b> {r[0]}<br><b>توضیحات:</b> {r[1] or '-'}</p>"

        elif res["table"] == "checks":
            cursor.execute("SELECT check_type, description, contract_id FROM checks WHERE id=?",
                           (res["id"],))
            r = cursor.fetchone()
            if r:
                detail_html += f"<p><b>نوع چک:</b> {r[0] or 'دریافتی'}<br>"
                detail_html += f"<b>قرارداد مرتبط:</b> {r[2] if r[2] else 'ثبت نشده'}<br>"
                detail_html += f"<b>توضیحات:</b> {r[1] or '-'}</p>"

        elif res["table"] == "persons":
            cursor.execute("SELECT code, name, role, phone FROM persons WHERE id=?", (res["id"],))
            r = cursor.fetchone()
            if r:
                detail_html += f"<p><b>کد پرسنل:</b> {r[0]}<br><b>نام:</b> {r[1]}<br>"
                detail_html += f"<b>نقش:</b> {r[2]}<br><b>تلفن:</b> {r[3] or '-'}</p>"
                cursor.execute("""SELECT COALESCE(SUM(amount),0) FROM transactions
                                  WHERE person_id=? AND work_year=?""",
                               (res["id"], get_current_year()))
                total = cursor.fetchone()[0]
                detail_html += f"<p><b>جمع پرداختی در این سال:</b> {total:,} تومان</p>"

        conn.close()

        detail_html += f"<hr><p><b>جمع کل:</b> {res['total']:,} تومان | "
        detail_html += f"<b>پرداختی:</b> {res['paid']:,} تومان | "
        remain = res['total'] - res['paid']
        color = '#c0392b' if remain > 0 else 'green'
        detail_html += f"<b style='color:{color};'>مانده:</b> {remain:,} تومان</p>"

        self.lbl_search_detail.setText(detail_html)

    # ============================================================
    # ==============  تب ۷: انبار تجهیزات  =====================
    # ============================================================
    def setup_inventory_tab(self):
        layout = QVBoxLayout()
        layout.setSpacing(8)

        # ==================== افزودن تجهیز جدید ====================
        add_box = QGroupBox("افزودن تجهیز / کالای جدید (هماهنگ با فاکتور و قراردادها)")
        add_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        add_layout = QGridLayout()
        add_layout.setHorizontalSpacing(9)
        add_layout.setVerticalSpacing(7)

        self.inv_name = QLineEdit()
        self.inv_code = QLineEdit()
        self.inv_code.setPlaceholderText("خودکار (یکتا) در صورت خالی بودن")
        self.inv_category = QLineEdit()
        self.inv_category.setPlaceholderText("دوربین، لنز، نور، خدمات...")
        self.inv_total = QSpinBox()
        self.inv_total.setRange(0, 100000)
        self.inv_total.setValue(1)
        self.inv_price = QLineEdit()
        self.inv_price.textChanged.connect(lambda t: self.inv_price.setText(format_number(t)))
        self.inv_note = QLineEdit()
        self.inv_note.setPlaceholderText("سریال، مدل، مشخصات فنی...")

        add_layout.addWidget(QLabel(f"{ico('inventory')} نام تجهیز:"), 0, 0)
        add_layout.addWidget(self.inv_name, 0, 1)
        add_layout.addWidget(QLabel("کد (یکتا):"), 0, 2)
        add_layout.addWidget(self.inv_code, 0, 3)
        add_layout.addWidget(QLabel("دسته‌بندی:"), 0, 4)
        add_layout.addWidget(self.inv_category, 0, 5)

        add_layout.addWidget(QLabel("تعداد کل:"), 1, 0)
        add_layout.addWidget(self.inv_total, 1, 1)
        add_layout.addWidget(QLabel(f"{ico('money')} قیمت واحد (تومان):"), 1, 2)
        add_layout.addWidget(self.inv_price, 1, 3)
        add_layout.addWidget(QLabel("مشخصات:"), 1, 4)
        add_layout.addWidget(self.inv_note, 1, 5)

        btn_add = QPushButton(ico_text("add", "افزودن به انبار و فاکتور"))
        btn_add.setStyleSheet("background-color:#27ae60; color:white; font-weight:bold; padding:9px;")
        btn_add.clicked.connect(self.add_inventory_item)
        add_layout.addWidget(btn_add, 2, 0, 1, 6)

        add_box.setLayout(add_layout)
        layout.addWidget(add_box)

        # ==================== نوار ابزار انبار ====================
        top = QHBoxLayout()
        btn_refresh = QPushButton(ico_text("refresh", "بروزرسانی"))
        btn_refresh.setStyleSheet("background-color: #3498db; color: white; font-weight: bold; padding: 7px;")
        btn_refresh.clicked.connect(self.load_inventory)

        btn_save = QPushButton(ico_text("save", "ثبت تغییرات انبار"))
        btn_save.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 7px;")
        btn_save.setToolTip("تعدادهای جدید را در همه بخش‌ها (انبار، قرارداد و فاکتور) اعمال می‌کند")
        btn_save.clicked.connect(self.save_inventory_changes)

        btn_recalc = QPushButton(ico_text("calc", "بازمحاسبه مصرف از قراردادها"))
        btn_recalc.setStyleSheet("background-color: #e67e22; color: white; font-weight: bold; padding: 7px;")
        btn_recalc.setToolTip("«استفاده شده» هر قلم را از روی قراردادهای واقعی دوباره حساب می‌کند")
        btn_recalc.clicked.connect(self.recalc_inventory_click)

        btn_print = QPushButton(ico_text("print", "چاپ / PDF لیست انبار"))
        btn_print.setStyleSheet("background-color: #2980b9; color: white; font-weight: bold; padding: 7px;")
        btn_print.clicked.connect(self.print_inventory)

        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color: #1e8449; color: white; font-weight: bold; padding: 7px;")
        btn_excel.clicked.connect(lambda: self.export_table_excel(
            self.inv_table, "انبار تجهیزات", "انبار_تجهیزات.xlsx"))

        top.addWidget(btn_refresh)
        top.addWidget(btn_save)
        top.addWidget(btn_recalc)
        top.addWidget(btn_print)
        top.addWidget(btn_excel)
        top.addStretch()
        layout.addLayout(top)

        self.lbl_inv_summary = QLabel("")
        self.lbl_inv_summary.setWordWrap(True)
        self.lbl_inv_summary.setStyleSheet("background-color:#eaf6ff; border:1px solid #a9d3ee; "
                                           "border-radius:8px; padding:9px; font-size:10pt;")
        layout.addWidget(self.lbl_inv_summary)

        self.inv_table = QTableWidget()
        self.inv_table.setColumnCount(9)
        self.inv_table.setHorizontalHeaderLabels([
            "کد کالا", "نام آیتم / پکیج", "دسته‌بندی", "تعداد کل (قابل ویرایش)",
            "استفاده شده", "موجودی", "قیمت واحد (تومان)", "مشخصات", "حذف"
        ])
        self.inv_table.setAlternatingRowColors(True)
        self.inv_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        persist_table(self.inv_table, "inventory", QHeaderView.ResizeMode.ResizeToContents)
        self.inv_table.verticalHeader().setDefaultSectionSize(36)
        layout.addWidget(self.inv_table, 1)

        self.tab_inventory.setLayout(layout)
        self.inv_spins = {}
        self.load_inventory()

    def recalc_inventory_click(self):
        recalc_inventory_usage()
        self.load_inventory()
        QMessageBox.information(self, "انبار", "مصرف انبار از روی قراردادهای واقعی بازمحاسبه شد.")

    def add_inventory_item(self):
        """افزودن تجهیز جدید با جزئیات، هماهنگ با انبار، قراردادها و فاکتور"""
        name = self.inv_name.text().strip()
        if not name:
            QMessageBox.warning(self, "خطا", "لطفاً نام تجهیز را وارد کنید.")
            return

        code = self.inv_code.text().strip()
        if not code:
            code = generate_unique_code("EQP", "item_prices", "code", start=3001)

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT item_name FROM item_prices WHERE code=? AND item_name<>?", (code, name))
        dup = cursor.fetchone()
        if dup:
            conn.close()
            QMessageBox.warning(self, "کد تکراری",
                                f"کد «{code}» قبلاً برای «{dup[0]}» ثبت شده است.")
            return

        cursor.execute('''
            INSERT INTO item_prices (item_name, code, price, is_package, parent_package)
            VALUES (?, ?, ?, 0, '')
            ON CONFLICT(item_name) DO UPDATE SET
                code=excluded.code, price=excluded.price
        ''', (name, code, parse_number(self.inv_price.text())))

        cursor.execute('''
            INSERT INTO inventory (item_name, total_count, used_count, category, note)
            VALUES (?, ?, 0, ?, ?)
            ON CONFLICT(item_name) DO UPDATE SET
                total_count=excluded.total_count,
                category=excluded.category,
                note=excluded.note
        ''', (name, self.inv_total.value(),
              self.inv_category.text().strip(), self.inv_note.text().strip()))

        conn.commit()
        conn.close()

        self.inv_name.clear()
        self.inv_code.clear()
        self.inv_category.clear()
        self.inv_note.clear()
        self.inv_price.clear()
        self.inv_total.setValue(1)

        self.load_inventory()
        self.load_item_checkboxes()
        QMessageBox.information(self, "موفقیت",
                                f"تجهیز «{name}» با کد {code} به انبار و فاکتور اضافه شد.")

    def load_inventory(self):
        if not hasattr(self, "inv_table"):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT i.item_name, ip.code, i.total_count, i.used_count,
                                 ip.price, i.category, i.note
                          FROM inventory i
                          LEFT JOIN item_prices ip ON i.item_name = ip.item_name
                          ORDER BY i.item_name""")
        rows = cursor.fetchall()
        conn.close()

        self.inv_table.setRowCount(0)
        self.inv_spins = {}
        total_count_sum = 0
        total_remain_sum = 0
        total_value = 0
        low_stock = []

        for r_idx, (name, code, total, used, price, category, note) in enumerate(rows):
            total = total or 0
            used = used or 0
            remain = max(0, total - used)
            price = price or 0
            total_count_sum += total
            total_remain_sum += remain
            total_value += remain * price
            if remain <= 1:
                low_stock.append(name)

            self.inv_table.insertRow(r_idx)
            self.inv_table.setItem(r_idx, 0, QTableWidgetItem(code or "-"))
            self.inv_table.setItem(r_idx, 1, QTableWidgetItem(name))
            self.inv_table.setItem(r_idx, 2, QTableWidgetItem(category or "-"))

            spin = QSpinBox()
            spin.setRange(0, 100000)
            spin.setValue(total)
            spin.setAlignment(Qt.AlignmentFlag.AlignCenter)
            spin.setStyleSheet("background-color:#fffdf3; border:1px solid #f39c12; border-radius:6px;")
            spin.valueChanged.connect(lambda _v, n=name: self.on_inv_total_changed(n))
            self.inv_spins[name] = spin
            self.inv_table.setCellWidget(r_idx, 3, spin)

            used_item = QTableWidgetItem(str(used))
            used_item.setForeground(QColor("#e67e22"))
            self.inv_table.setItem(r_idx, 4, used_item)

            remain_item = QTableWidgetItem(str(remain))
            remain_item.setForeground(QColor("#c0392b" if remain <= 1 else "#1e8449"))
            f = remain_item.font()
            f.setBold(True)
            remain_item.setFont(f)
            self.inv_table.setItem(r_idx, 5, remain_item)

            self.inv_table.setItem(r_idx, 6, QTableWidgetItem(f"{price:,}"))
            self.inv_table.setItem(r_idx, 7, QTableWidgetItem(note or "-"))

            btn_del = QPushButton(ico("delete"))
            btn_del.setToolTip("حذف این قلم از انبار و فاکتور")
            btn_del.setStyleSheet("background-color:#e74c3c; color:white;")
            btn_del.clicked.connect(lambda _, n=name: self.delete_inventory_item(n))
            self.inv_table.setCellWidget(r_idx, 8, btn_del)

        self.lbl_inv_summary.setText(
            f"<b>تعداد اقلام:</b> {len(rows)}  |  "
            f"<b>جمع موجودی (تعداد):</b> {total_count_sum}  |  "
            f"<b>جمع موجودی قابل استفاده:</b> {total_remain_sum}  |  "
            f"<b>ارزش ریالی موجودی:</b> {total_value:,} تومان"
            + (f"<br><b style='color:#c0392b;'>⚠️ موجودی کم / تمام‌شده:</b> "
               f"{'، '.join(low_stock)}" if low_stock else "")
        )
        self.refresh_dashboard_summary()

    def on_inv_total_changed(self, item_name):
        """پیش‌نمایش زنده موجودی هنگام تغییر تعداد کل"""
        spin = self.inv_spins.get(item_name)
        if not spin:
            return
        for r in range(self.inv_table.rowCount()):
            if self.inv_table.item(r, 1) and self.inv_table.item(r, 1).text() == item_name:
                try:
                    used = int(self.inv_table.item(r, 4).text())
                except Exception:
                    used = 0
                remain = max(0, spin.value() - used)
                it = QTableWidgetItem(str(remain))
                it.setForeground(QColor("#c0392b" if remain <= 1 else "#1e8449"))
                f = it.font()
                f.setBold(True)
                it.setFont(f)
                self.inv_table.setItem(r, 5, it)
                break

    def save_inventory_changes(self):
        """ثبت تعدادهای جدید — همه بخش‌ها با هم هماهنگ می‌شوند"""
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        changed = 0
        for name, spin in self.inv_spins.items():
            cursor.execute("UPDATE inventory SET total_count=? WHERE item_name=?",
                           (spin.value(), name))
            if cursor.rowcount:
                changed += 1
        conn.commit()
        conn.close()
        recalc_inventory_usage()
        self.load_inventory()
        self.load_item_checkboxes()
        QMessageBox.information(self, "موفقیت",
                                f"تعداد کل {changed} قلم انبار به‌روزرسانی شد و "
                                "موجودی و فاکتورها هماهنگ شدند.")

    def delete_inventory_item(self, item_name):
        if QMessageBox.question(
                self, "تایید حذف",
                f"آیا «{item_name}» از انبار و لیست فاکتورها حذف شود؟",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        ) != QMessageBox.StandardButton.Yes:
            return
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM inventory WHERE item_name=?", (item_name,))
        cursor.execute("DELETE FROM item_prices WHERE item_name=?", (item_name,))
        conn.commit()
        conn.close()
        recalc_inventory_usage()
        self.load_inventory()
        self.load_item_checkboxes()

    def inventory_html(self, mono=False):
        p = doc_palette(mono)
        studio_name, studio_phone, _addr = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT i.item_name, ip.code, i.total_count, i.used_count,
                                 ip.price, i.category, i.note
                          FROM inventory i
                          LEFT JOIN item_prices ip ON i.item_name = ip.item_name
                          ORDER BY i.item_name""")
        rows = cursor.fetchall()
        conn.close()

        body = ""
        total_value = 0
        total_count = total_used = total_remain = 0
        for idx, (name, code, total, used, price, category, note) in enumerate(rows, start=1):
            total = total or 0
            used = used or 0
            remain = max(0, total - used)
            price = price or 0
            total_value += remain * price
            total_count += total
            total_used += used
            total_remain += remain
            body += InvoiceBuilder._row([
                InvoiceBuilder._td(str(idx), size="7.5pt"),
                InvoiceBuilder._td(code or "-", size="7pt"),
                InvoiceBuilder._td(name, align="right", size="7.5pt"),
                InvoiceBuilder._td(category or "-", size="7pt"),
                InvoiceBuilder._td(str(total), size="7.5pt"),
                InvoiceBuilder._td(str(used), size="7.5pt"),
                InvoiceBuilder._td(f"<b>{remain}</b>", size="7.5pt"),
                InvoiceBuilder._td(f"{price:,}", size="7.5pt"),
                InvoiceBuilder._td(note or "-", align="right", size="7pt"),
            ])
        if not body:
            body = f"<tr>{InvoiceBuilder._td('انبار خالی است', colspan=9, size='9pt')}</tr>"

        return f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {InvoiceBuilder._title_band(
              p, studio_name,
              "<div style='font-size:10pt;'>لیست انبار تجهیزات</div>",
              [f"تاریخ چاپ: {now.strftime('%Y/%m/%d')}", f"ساعت: {now.strftime('%H:%M')}",
               f"تعداد اقلام: {len(rows)}"])}
          <table width="100%" style="border-collapse:collapse; margin-top:5px;">
            <thead>{InvoiceBuilder._head(['ردیف', 'کد', 'نام', 'دسته', 'تعداد کل', 'استفاده شده',
                                          'موجودی', 'قیمت واحد', 'مشخصات'],
                                         bg=p['table_head'], fg=p['table_head_fg'])}</thead>
            <tbody>{body}</tbody>
            <tfoot>
              {InvoiceBuilder._row([
                  InvoiceBuilder._td('<b>جمع کل</b>', align='right', colspan=4,
                                     bg=p['total_bg'], size='8pt', bold=True),
                  InvoiceBuilder._td(f"<b>{total_count}</b>", bg=p['total_bg'], size='8pt', bold=True),
                  InvoiceBuilder._td(f"<b>{total_used}</b>", bg=p['total_bg'], size='8pt', bold=True),
                  InvoiceBuilder._td(f"<b>{total_remain}</b>", bg=p['total_bg'], size='8pt', bold=True),
                  InvoiceBuilder._td("", bg=p['total_bg'], size='8pt'),
                  InvoiceBuilder._td("", bg=p['total_bg'], size='8pt'),
              ])}
            </tfoot>
          </table>
          <div style="margin-top:5px; font-size:8.5pt;">
            <b>ارزش ریالی موجودی انبار:</b> {total_value:,} تومان |
            <b>به حروف:</b> {number_to_persian_words(total_value)} تومان
          </div>
          <div style="text-align:center; margin-top:24px; font-size:9pt;">مهر و امضای انباردار / مدیر</div>
          <div style="text-align:center; margin-top:8px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}
          </div>
        </div>
        """

    def print_inventory(self):
        self._output_flow(lambda wd, ws, mono: self.inventory_html(mono),
                          "لیست_انبار.pdf", "لیست انبار تجهیزات", allow_details=False)

    # ============================================================
    # ==============  تب ۸: کارت‌های بانکی  ====================
    # ============================================================
    def setup_banks_tab(self):
        layout = QHBoxLayout()
        layout.setSpacing(10)

        form_box = QGroupBox("افزودن حساب بانکی")
        form_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        form_layout = QFormLayout()
        form_layout.setSpacing(7)

        self.b_name = QLineEdit()
        self.b_card = QLineEdit()
        self.b_card.setPlaceholderText("مثلاً 6037-XXXX-XXXX-XXXX")
        self.b_sheba = QLineEdit()
        self.b_sheba.setPlaceholderText("IR...")

        form_layout.addRow(f"{ico('banks')} نام بانک <span style='color:red;'>*</span>:", self.b_name)
        form_layout.addRow(f"{ico('banks')} شماره کارت:", self.b_card)
        form_layout.addRow(f"{ico('banks')} شماره شبا:", self.b_sheba)

        btn_save = QPushButton(ico_text("save", "ذخیره کارت"))
        btn_save.setStyleSheet("background-color: #16a085; color: white; font-weight: bold; padding: 9px;")
        btn_save.clicked.connect(self.save_bank_card)
        form_layout.addRow(btn_save)

        info = QLabel(f"{ico('detail')} بانک‌های ثبت‌شده در بخش بیعانه‌های قرارداد و "
                      "فاکتورها قابل انتخاب هستند.")
        info.setWordWrap(True)
        info.setStyleSheet("background-color:#eaf6ff; border:1px solid #a9d3ee; "
                           "border-radius:8px; padding:9px;")
        form_layout.addRow(info)

        form_box.setLayout(form_layout)
        layout.addWidget(form_box, 2)

        table_box = QGroupBox("حساب‌های بانکی")
        table_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        tl = QVBoxLayout()

        self.b_table = QTableWidget()
        self.b_table.setColumnCount(4)
        self.b_table.setHorizontalHeaderLabels(["نام بانک", "شماره کارت", "شماره شبا", "حذف"])
        self.b_table.setAlternatingRowColors(True)
        self.b_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        persist_table(self.b_table, "banks", QHeaderView.ResizeMode.Stretch)
        self.b_table.verticalHeader().setDefaultSectionSize(32)
        tl.addWidget(self.b_table)
        table_box.setLayout(tl)
        layout.addWidget(table_box, 4)

        self.tab_banks.setLayout(layout)
        self.load_bank_cards()

    def save_bank_card(self):
        if not self.b_name.text().strip():
            QMessageBox.warning(self, "خطا", "نام بانک را وارد کنید.")
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("INSERT INTO bank_cards (bank_name, card_number, sheba_number) VALUES (?, ?, ?)",
                       (self.b_name.text().strip(), self.b_card.text().strip(), self.b_sheba.text().strip()))
        conn.commit()
        conn.close()
        self.b_name.clear()
        self.b_card.clear()
        self.b_sheba.clear()
        self.load_bank_cards()
        self.load_bank_combo()
        QMessageBox.information(self, "موفقیت", "حساب بانکی ذخیره شد.")

    def load_bank_cards(self):
        if not hasattr(self, "b_table"):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("SELECT id, bank_name, card_number, sheba_number FROM bank_cards")
        rows = cursor.fetchall()
        conn.close()

        self.b_table.setRowCount(0)
        for r_idx, (bid, name, card, sheba) in enumerate(rows):
            self.b_table.insertRow(r_idx)
            self.b_table.setItem(r_idx, 0, QTableWidgetItem(f"{ico('banks')} {name}"))
            self.b_table.setItem(r_idx, 1, QTableWidgetItem(str(card or "-")))
            self.b_table.setItem(r_idx, 2, QTableWidgetItem(str(sheba or "-")))
            btn_del = QPushButton(ico("delete"))
            btn_del.setToolTip("حذف این حساب")
            btn_del.setStyleSheet("background-color: #e74c3c; color: white;")
            btn_del.clicked.connect(lambda _, x=bid: self.delete_bank_card(x))
            self.b_table.setCellWidget(r_idx, 3, btn_del)

    def delete_bank_card(self, bid):
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM bank_cards WHERE id=?", (bid,))
        conn.commit()
        conn.close()
        self.load_bank_cards()
        self.load_bank_combo()

    # ============================================================
    # ==============  تب ۹: مدیریت چک‌ها  =====================
    # ============================================================
    def setup_checks_tab(self):
        layout = QHBoxLayout()
        layout.setSpacing(10)

        form_box = QGroupBox("ثبت چک جدید")
        form_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        form_layout = QFormLayout()
        form_layout.setSpacing(7)

        self.chk_num = QLineEdit()
        self.chk_bank = QLineEdit()
        self.chk_issuer = QLineEdit()
        self.chk_type = QComboBox()
        # واژه «صادراتی» طبق درخواست به «پرداختی» تغییر کرد
        self.chk_type.addItems(["دریافتی", "پرداختی"])
        self.chk_type.currentTextChanged.connect(self.update_check_labels)
        self.chk_amount = QLineEdit()
        self.chk_amount.textChanged.connect(lambda t: self.chk_amount.setText(format_number(t)))
        self.chk_amount.textChanged.connect(self.update_check_words)
        self.chk_date = PersianDateEdit(date_str=jalali_today_str())
        self.chk_desc = QLineEdit()

        self.chk_contract_combo = QComboBox()
        self.load_check_contract_combo()

        form_layout.addRow(f"{ico('checks')} شماره چک <span style='color:red;'>*</span>:", self.chk_num)
        form_layout.addRow(f"{ico('banks')} بانک:", self.chk_bank)
        form_layout.addRow("نام طرف حساب:", self.chk_issuer)
        self.lbl_issuer = form_layout.labelForField(self.chk_issuer)
        form_layout.addRow("نوع چک:", self.chk_type)
        form_layout.addRow(f"{ico('money')} مبلغ (تومان) <span style='color:red;'>*</span>:", self.chk_amount)
        form_layout.addRow(f"{ico('calendar')} تاریخ سررسید (روز/ماه/سال):", self.chk_date)
        form_layout.addRow(f"{ico('contract')} قرارداد مرتبط (اختیاری):", self.chk_contract_combo)
        form_layout.addRow("توضیحات:", self.chk_desc)

        self.chk_lbl_words = QLabel("")
        self.chk_lbl_words.setWordWrap(True)
        self.chk_lbl_words.setStyleSheet("background-color:#fff9e6; border:1px solid #f39c12; "
                                         "border-radius:8px; padding:9px; font-size:10pt;")
        form_layout.addRow(self.chk_lbl_words)

        self.btn_save_check = QPushButton(ico_text("save", "ثبت چک دریافتی"))
        self.btn_save_check.setStyleSheet("background-color: #2980b9; color: white; font-weight: bold; padding: 9px;")
        self.btn_save_check.clicked.connect(self.save_check)
        form_layout.addRow(self.btn_save_check)

        form_box.setLayout(form_layout)
        layout.addWidget(form_box, 2)

        table_box = QGroupBox("لیست چک‌ها (با روزشمار سررسید)")
        table_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        tl = QVBoxLayout()

        tools = QHBoxLayout()
        btn_refresh = QPushButton(ico_text("refresh", "بروزرسانی"))
        btn_refresh.clicked.connect(self.load_checks_table)
        btn_excel = QPushButton(ico_text("excel", "خروجی اکسل"))
        btn_excel.setStyleSheet("background-color:#1e8449; color:white; font-weight:bold; padding:7px;")
        btn_excel.clicked.connect(lambda: self.export_table_excel(
            self.chk_table, "چک‌ها", f"چک‌ها_{get_current_year()}.xlsx"))
        self.combo_check_filter = QComboBox()
        self.combo_check_filter.addItems(["همه چک‌ها", "فقط دریافتی", "فقط پرداختی"])
        self.combo_check_filter.currentIndexChanged.connect(self.load_checks_table)
        tools.addWidget(btn_refresh)
        tools.addWidget(btn_excel)
        tools.addWidget(QLabel("نمایش:"))
        tools.addWidget(self.combo_check_filter)
        tools.addStretch()
        tl.addLayout(tools)

        self.lbl_checks_summary = QLabel("")
        self.lbl_checks_summary.setWordWrap(True)
        self.lbl_checks_summary.setStyleSheet("background-color:#eaf6ff; border:1px solid #a9d3ee; "
                                              "border-radius:8px; padding:9px; font-size:10pt;")
        tl.addWidget(self.lbl_checks_summary)

        self.chk_table = QTableWidget()
        self.chk_table.setColumnCount(11)
        self.chk_table.setHorizontalHeaderLabels([
            "شماره", "بانک", "طرف حساب", "نوع", "مبلغ",
            "سررسید", "روزشمار", "قرارداد", "وضعیت", "چاپ", "حذف"
        ])
        self.chk_table.setAlternatingRowColors(True)
        self.chk_table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        persist_table(self.chk_table, "checks", QHeaderView.ResizeMode.ResizeToContents)
        self.chk_table.verticalHeader().setDefaultSectionSize(36)
        tl.addWidget(self.chk_table)
        table_box.setLayout(tl)
        layout.addWidget(table_box, 5)

        self.tab_checks.setLayout(layout)
        self.update_check_labels()
        self.load_checks_table()

    def update_check_labels(self, text=None):
        """منوی هر نوع چک مخصوص خودش می‌شود"""
        t = text or (self.chk_type.currentText() if hasattr(self, "chk_type") else "دریافتی")
        if getattr(self, "lbl_issuer", None):
            self.lbl_issuer.setText("در وجه (پرداختی به):" if t == "پرداختی" else "دریافتی از:")
        if hasattr(self, "chk_issuer"):
            if t == "پرداختی":
                self.chk_issuer.setPlaceholderText("نام شخص یا شرکتی که چک در وجه او صادر شده است")
            else:
                self.chk_issuer.setPlaceholderText("نام شخص یا شرکتی که چک را به ما تحویل داده است")
        if hasattr(self, "btn_save_check"):
            self.btn_save_check.setText(
                ico_text("save", "ثبت چک پرداختی" if t == "پرداختی" else "ثبت چک دریافتی"))

    def update_check_words(self):
        if not hasattr(self, "chk_lbl_words"):
            return
        amount = parse_number(self.chk_amount.text())
        self.chk_lbl_words.setText(f"<b>مبلغ به حروف:</b> {number_to_persian_words(amount)} تومان")

    def load_check_contract_combo(self):
        if not hasattr(self, "chk_contract_combo"):
            return
        self.chk_contract_combo.clear()
        self.chk_contract_combo.addItem("--- بدون قرارداد ---", None)
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT id, code, groom_name, bride_name FROM wedding_contracts
                          WHERE work_year=? ORDER BY id DESC""", (get_current_year(),))
        rows = cursor.fetchall()
        conn.close()
        for cid, code, groom, bride in rows:
            self.chk_contract_combo.addItem(f"{code or cid} - {groom} و {bride}", cid)

    def save_check(self):
        num = self.chk_num.text().strip()
        amount = parse_number(self.chk_amount.text())
        if not num or amount <= 0:
            QMessageBox.warning(self, "خطا", "اطلاعات چک را کامل کنید.")
            return

        work_year = get_current_year()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""INSERT INTO checks
            (check_number, bank_name, amount, due_date, description,
             issuer_name, check_type, work_year, contract_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (num, self.chk_bank.text(), amount, self.chk_date.text(),
             self.chk_desc.text(), self.chk_issuer.text(),
             self.chk_type.currentText(), work_year,
             self.chk_contract_combo.currentData()))
        conn.commit()
        conn.close()

        self.chk_num.clear()
        self.chk_bank.clear()
        self.chk_issuer.clear()
        self.chk_amount.clear()
        self.chk_desc.clear()
        self.chk_date.set_date(jalali_today_str())
        self.load_checks_table()
        QMessageBox.information(self, "موفقیت", "چک با موفقیت ثبت شد.")

    def load_checks_table(self):
        if not hasattr(self, "chk_table"):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        work_year = get_current_year()
        cursor.execute("""SELECT id, check_number, bank_name, amount, due_date,
                          is_passed, issuer_name, check_type, description, contract_id
                          FROM checks WHERE work_year=? ORDER BY due_date""", (work_year,))
        rows = cursor.fetchall()
        contract_codes = {}
        cursor.execute("SELECT id, code FROM wedding_contracts")
        for cid, ccode in cursor.fetchall():
            contract_codes[cid] = ccode
        conn.close()

        selected_filter = self.combo_check_filter.currentText() if hasattr(self, "combo_check_filter") else "همه چک‌ها"
        if selected_filter == "فقط دریافتی":
            rows = [r for r in rows if (r[7] or "دریافتی") == "دریافتی"]
        elif selected_filter == "فقط پرداختی":
            rows = [r for r in rows if (r[7] or "دریافتی") == "پرداختی"]

        today = jdatetime.date.today()
        self.chk_table.setRowCount(0)
        sum_recv = sum_paid = 0
        passed = 0

        for r_idx, row in enumerate(rows):
            (cid, num, bank, amount, due, is_passed,
             issuer, ctype, desc, contract_id) = row
            ctype = ctype or "دریافتی"
            if ctype == "دریافتی":
                sum_recv += amount or 0
            else:
                sum_paid += amount or 0
            if is_passed:
                passed += 1

            try:
                due_date = parse_jalali(due)
                days_left = (due_date - today).days if due_date else None
                if is_passed:
                    countdown = "✅ پاس شده"
                elif days_left is None:
                    countdown = "-"
                elif days_left < 0:
                    countdown = f"🔴 {abs(days_left)} روز گذشته"
                elif days_left == 0:
                    countdown = "🟠 امروز!"
                elif days_left <= 3:
                    countdown = f"🟡 {days_left} روز مانده"
                else:
                    countdown = f"🟢 {days_left} روز مانده"
            except Exception:
                countdown = "-"

            self.chk_table.insertRow(r_idx)
            self.chk_table.setItem(r_idx, 0, QTableWidgetItem(num))
            self.chk_table.setItem(r_idx, 1, QTableWidgetItem(bank or "-"))
            self.chk_table.setItem(r_idx, 2, QTableWidgetItem(issuer or "-"))

            type_item = QTableWidgetItem(ctype)
            type_item.setForeground(QColor("#1e8449" if ctype == "دریافتی" else "#c0392b"))
            f = type_item.font()
            f.setBold(True)
            type_item.setFont(f)
            self.chk_table.setItem(r_idx, 3, type_item)

            amt_item = QTableWidgetItem(f"{amount:,}")
            amt_item.setForeground(QColor("#1e8449" if ctype == "دریافتی" else "#c0392b"))
            self.chk_table.setItem(r_idx, 4, amt_item)
            self.chk_table.setItem(r_idx, 5, QTableWidgetItem(due or "-"))
            self.chk_table.setItem(r_idx, 6, QTableWidgetItem(countdown))
            self.chk_table.setItem(r_idx, 7, QTableWidgetItem(contract_codes.get(contract_id, "-")))

            btn_pass = QPushButton("✅ پاس شده" if is_passed else "⏳ پاس‌کردن")
            btn_pass.setStyleSheet(
                "background-color: #1e8449; color: white; font-weight:bold;"
                if is_passed else "background-color: #e67e22; color: white; font-weight:bold;")
            btn_pass.clicked.connect(lambda _, x=cid, s=is_passed: self.toggle_check_pass(x, s))
            self.chk_table.setCellWidget(r_idx, 8, btn_pass)

            btn_print = QPushButton(ico("print"))
            btn_print.setToolTip("چاپ / PDF رسید چک")
            btn_print.setStyleSheet("background-color: #2980b9; color: white;")
            btn_print.clicked.connect(lambda _, x=cid: self.print_check(x))
            self.chk_table.setCellWidget(r_idx, 9, btn_print)

            btn_del = QPushButton(ico("delete"))
            btn_del.setToolTip("حذف چک")
            btn_del.setStyleSheet("background-color: #e74c3c; color: white;")
            btn_del.clicked.connect(lambda _, x=cid: self.delete_check(x))
            self.chk_table.setCellWidget(r_idx, 10, btn_del)

        self.lbl_checks_summary.setText(
            f"<b>تعداد چک‌ها:</b> {len(rows)}  |  "
            f"<b>پاس‌شده:</b> {passed}  |  "
            f"<b>جمع چک‌های دریافتی:</b> <span style='color:#1e8449;'>{sum_recv:,}</span> تومان  |  "
            f"<b>جمع چک‌های پرداختی:</b> <span style='color:#c0392b;'>{sum_paid:,}</span> تومان"
        )

    def toggle_check_pass(self, cid, current_state):
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("UPDATE checks SET is_passed=? WHERE id=?", (0 if current_state else 1, cid))
        conn.commit()
        conn.close()
        self.load_checks_table()

    def check_html(self, cid, mono=False):
        p = doc_palette(mono)
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("""SELECT check_number, bank_name, amount, due_date,
                          is_passed, issuer_name, check_type, description, contract_id
                          FROM checks WHERE id=?""", (cid,))
        r = cursor.fetchone()
        conn.close()
        if not r:
            return None

        studio_name, studio_phone, studio_address = InvoiceBuilder._studio_info()
        now = jdatetime.datetime.now()
        ctype = r[6] or "دریافتی"
        status = "پاس شده" if r[4] else "پاس نشده"
        party_label = "در وجه (پرداختی به)" if ctype == "پرداختی" else "دریافتی از"

        def pair(label, value):
            return InvoiceBuilder._row([
                InvoiceBuilder._td(f"<b>{label}</b>", align="right",
                                   bg=p["box_bg"], size="9pt", width="28%"),
                InvoiceBuilder._td(value, align="right", size="9pt"),
            ])

        rows = (pair("شماره چک:", r[0]) +
                pair("بانک:", r[1] or "-") +
                pair(party_label + ":", r[5] or "-") +
                pair("نوع چک:", ctype) +
                pair("مبلغ:", f"<b>{r[2]:,}</b> تومان") +
                pair("مبلغ به حروف:", f"{number_to_persian_words(r[2])} تومان") +
                pair("تاریخ سررسید:", r[3] or "-") +
                pair("وضعیت:", status) +
                pair("توضیحات:", r[7] or "-"))

        return f"""
        <div dir="rtl" style="font-family:'{INVOICE_FONT_FAMILY}', Tahoma; font-size:9pt;">
          {InvoiceBuilder._title_band(
              p, studio_name,
              f"<div style='font-size:10pt;'>رسید چک {ctype}</div>",
              [f"تاریخ چاپ: {now.strftime('%Y/%m/%d')}", f"ساعت: {now.strftime('%H:%M')}",
               f"وضعیت: {status}"])}
          <table width="100%" style="border-collapse:collapse; margin-top:6px;">{rows}</table>
          <div style="text-align:center; margin-top:36px; font-size:9pt;">مهر و امضای حسابداری</div>
          <div style="text-align:center; margin-top:10px; font-size:7pt; color:{p['muted']};">
            {studio_name} | تلفن: {studio_phone}{(' | ' + studio_address) if studio_address else ''}
          </div>
        </div>
        """

    def print_check(self, cid):
        self._output_flow(lambda wd, ws, mono: self.check_html(cid, mono),
                          "رسید_چک.pdf", "رسید چک", allow_details=False)

    def delete_check(self, cid):
        if not ask_security_password(self):
            return
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM checks WHERE id=?", (cid,))
        conn.commit()
        conn.close()
        self.load_checks_table()

    # ============================================================
    # ==============  تب ۱۰: صدور رسید وجه  ====================
    # ============================================================
    def setup_receipt_tab(self):
        layout = QVBoxLayout()
        layout.setSpacing(8)

        form_box = QGroupBox("صدور رسید وجه / بیعانه")
        form_box.setFont(QFont(APP_FONT_FAMILY, 11, QFont.Weight.Bold))
        form_layout = QFormLayout()
        form_layout.setSpacing(7)

        # انتخاب سریع مشتری از قراردادها (هماهنگی بین بخش‌ها)
        self.rc_customer_combo = QComboBox()
        self.load_receipt_customers()
        self.rc_customer_combo.currentIndexChanged.connect(self.on_receipt_customer_selected)
        form_layout.addRow(f"{ico('people')} انتخاب سریع از مشتریان:", self.rc_customer_combo)

        self.rc_name = QLineEdit()
        self.rc_amount = QLineEdit()
        self.rc_amount.textChanged.connect(lambda t: self.rc_amount.setText(format_number(t)))
        self.rc_direction = QComboBox()
        self.rc_direction.addItems(["دریافتی (از مشتری)", "پرداختی (به شخص)"])
        self.rc_for = QLineEdit("خدمات فیلم و عکس")
        self.rc_date = PersianDateEdit(date_str=jalali_today_str())

        for w in (self.rc_name, self.rc_amount, self.rc_for):
            w.textChanged.connect(self.update_receipt_preview)
        self.rc_date.dateChanged.connect(self.update_receipt_preview)
        self.rc_direction.currentIndexChanged.connect(self.update_receipt_preview)

        form_layout.addRow(f"{ico('people')} نام طرف حساب <span style='color:red;'>*</span>:", self.rc_name)
        form_layout.addRow(f"{ico('money')} مبلغ (تومان) <span style='color:red;'>*</span>:", self.rc_amount)
        form_layout.addRow("نوع رسید:", self.rc_direction)
        form_layout.addRow("بابت خدمات:", self.rc_for)
        form_layout.addRow(f"{ico('calendar')} تاریخ (روز/ماه/سال):", self.rc_date)

        self.rc_lbl_words = QLabel("")
        self.rc_lbl_words.setWordWrap(True)
        self.rc_lbl_words.setStyleSheet("background-color:#e8f8f5; border:1px solid #16a085; "
                                        "border-radius:8px; padding:9px; font-size:11pt;")
        form_layout.addRow(self.rc_lbl_words)

        btn_row = QHBoxLayout()
        btn_print = QPushButton(ico_text("print", "چاپ رسید"))
        btn_print.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 11px; font-size: 11pt;")
        btn_print.clicked.connect(self.print_receipt)

        btn_pdf = QPushButton(ico_text("pdf", "ذخیره PDF"))
        btn_pdf.setStyleSheet("background-color: #c0392b; color: white; font-weight: bold; padding: 11px; font-size: 11pt;")
        btn_pdf.clicked.connect(self.export_receipt_pdf)

        btn_row.addWidget(btn_print)
        btn_row.addWidget(btn_pdf)
        form_layout.addRow(btn_row)

        form_box.setLayout(form_layout)
        layout.addWidget(form_box)

        preview_box = QGroupBox("پیش‌نمایش زنده رسید")
        preview_box.setFont(QFont(APP_FONT_FAMILY, 10, QFont.Weight.Bold))
        pv = QVBoxLayout()
        self.rc_preview = QTextEdit()
        self.rc_preview.setReadOnly(True)
        self.rc_preview.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.rc_preview.setStyleSheet("background-color: #ffffff; font-size: 11pt; "
                                      "border:1px solid #d6e0ec; border-radius:8px;")
        self.rc_preview.setFixedHeight(320)
        pv.addWidget(self.rc_preview)
        preview_box.setLayout(pv)
        layout.addWidget(preview_box)

        layout.addStretch()
        self.tab_receipt.setLayout(layout)
        self.update_receipt_preview()

    def load_receipt_customers(self):
        """لیست مشتریان قراردادها برای انتخاب سریع در رسید"""
        if not hasattr(self, "rc_customer_combo"):
            return
        self.rc_customer_combo.clear()
        self.rc_customer_combo.addItem("--- انتخاب دستی ---", None)
        try:
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            cursor.execute("""SELECT id, code, groom_name, bride_name FROM wedding_contracts
                              WHERE work_year=? ORDER BY id DESC""", (get_current_year(),))
            contracts = cursor.fetchall()
            cursor.execute("SELECT client_name FROM commercial_projects "
                           "WHERE work_year=? AND client_name IS NOT NULL AND client_name<>''",
                           (get_current_year(),))
            clients = [r[0] for r in cursor.fetchall()]
            conn.close()
            for cid, code, groom, bride in contracts:
                self.rc_customer_combo.addItem(f"{code or cid} - {groom} و {bride}", f"{groom} و {bride}")
            for name in clients:
                self.rc_customer_combo.addItem(f"مشتری پروژه - {name}", name)
        except Exception as e:
            print("[RECEIPT] load customers error:", e)

    def on_receipt_customer_selected(self, _idx=None):
        data = self.rc_customer_combo.currentData() if hasattr(self, "rc_customer_combo") else None
        if data:
            self.rc_name.setText(str(data))
            self.update_receipt_preview()

    def update_receipt_preview(self):
        """
        رفع اشکال «پیش‌نمایش رسید وجه نمایش داده نمی‌شود»:
        پیش‌نمایش به همه تغییرات فرم به‌صورت زنده متصل شد.
        """
        if not hasattr(self, "rc_preview"):
            return
        amount = parse_number(self.rc_amount.text())
        if hasattr(self, "rc_lbl_words"):
            self.rc_lbl_words.setText(f"<b>مبلغ به حروف:</b> {number_to_persian_words(amount)} تومان")

        name = self.rc_name.text().strip()
        if not name and amount <= 0:
            self.rc_preview.setHtml(
                f"<div dir='rtl' style=\"font-family:'{INVOICE_FONT_FAMILY}', Tahoma; "
                "padding:20px; color:#7f8c8d;\">برای مشاهده پیش‌نمایش، نام طرف حساب و مبلغ را وارد کنید.</div>")
            return

        direction = "دریافتی" if self.rc_direction.currentIndex() == 0 else "پرداختی"
        html = InvoiceBuilder.build_receipt(
            name or "-", amount, self.rc_for.text(), self.rc_date.text(), direction)
        self.rc_preview.setHtml(html)

    def get_receipt_html(self, mono=False):
        name = self.rc_name.text().strip()
        amount = parse_number(self.rc_amount.text())
        if not name or amount <= 0:
            return None
        direction = "دریافتی" if self.rc_direction.currentIndex() == 0 else "پرداختی"
        return InvoiceBuilder.build_receipt(
            name, amount, self.rc_for.text(), self.rc_date.text(), direction, mono)

    def print_receipt(self):
        html = self.get_receipt_html(mono=True)
        if not html:
            QMessageBox.warning(self, "خطا", "لطفاً نام طرف حساب و مبلغ را وارد کنید.")
            return
        print_html_document(html, self, "رسید وجه")

    def export_receipt_pdf(self):
        html = self.get_receipt_html()
        if not html:
            QMessageBox.warning(self, "خطا", "لطفاً نام طرف حساب و مبلغ را وارد کنید.")
            return
        name = self.rc_name.text().strip()
        default_name = f"رسید_{name}_{jdatetime.date.today().strftime('%Y%m%d')}.pdf"
        path, _ = QFileDialog.getSaveFileName(self, "ذخیره رسید PDF", default_name, "PDF Files (*.pdf)")
        if not path:
            return
        ok, err = save_html_pdf(html, path, "رسید وجه")
        if ok:
            QMessageBox.information(self, "موفقیت", f"رسید PDF ذخیره شد:\n{path}")
        else:
            QMessageBox.critical(self, "خطا", f"خطا:\n{err}")

    # ============================================================
    # ==============  تب ۱۱: تنظیمات و سال کاری  ===============
    # ============================================================
    def setup_workyear_tab(self):
        layout = QVBoxLayout()
        layout.setSpacing(9)

        year_box = QGroupBox("مدیریت سال‌های کاری (سوییچ به سال‌های قبل و بعد)")
        year_box.setFont(QFont(APP_FONT_FAMILY, 11, QFont.Weight.Bold))
        yb = QVBoxLayout()

        year_row = QHBoxLayout()
        year_row.addWidget(QLabel(f"{ico('calendar')} سال کاری فعلی:"))

        btn_prev = QPushButton(f"{ico('prev')} سال قبل")
        btn_prev.setStyleSheet("background-color: #7f8c8d; color: white; font-weight: bold; padding: 8px 14px;")
        btn_prev.clicked.connect(lambda: self.step_work_year(-1))
        year_row.addWidget(btn_prev)

        self.combo_settings_year = QComboBox()
        self.combo_settings_year.setFont(QFont(APP_FONT_FAMILY, 11, QFont.Weight.Bold))
        self.combo_settings_year.setFixedWidth(130)
        self.combo_settings_year.setStyleSheet("padding:6px; border:2px solid #2980b9; border-radius:8px;")
        year_row.addWidget(self.combo_settings_year)

        btn_next = QPushButton(f"سال بعد {ico('next')}")
        btn_next.setStyleSheet("background-color: #7f8c8d; color: white; font-weight: bold; padding: 8px 14px;")
        btn_next.clicked.connect(lambda: self.step_work_year(1))
        year_row.addWidget(btn_next)

        btn_switch = QPushButton(ico_text("refresh", "تغییر به این سال"))
        btn_switch.setStyleSheet("background-color: #2980b9; color: white; font-weight: bold; padding: 8px;")
        btn_switch.clicked.connect(self.switch_year_from_settings)
        year_row.addWidget(btn_switch)

        btn_new_year = QPushButton(ico_text("add", "شروع سال کاری جدید"))
        btn_new_year.setStyleSheet("background-color: #27ae60; color: white; font-weight: bold; padding: 8px;")
        btn_new_year.clicked.connect(self.start_new_work_year)
        year_row.addWidget(btn_new_year)

        year_row.addStretch()
        yb.addLayout(year_row)

        self.lbl_year_stats = QLabel("")
        self.lbl_year_stats.setWordWrap(True)
        self.lbl_year_stats.setStyleSheet("background-color:#eaf6ff; border:1px solid #a9d3ee; "
                                          "border-radius:8px; padding:10px; font-size:10pt;")
        yb.addWidget(self.lbl_year_stats)

        info = QLabel(
            f"{ico('detail')} با تغییر سال کاری، تمام تب‌ها به‌طور خودکار فقط اطلاعات همان سال را "
            "نمایش می‌دهند. شروع سال جدید به معنی پاک شدن اطلاعات سال قبل نیست — همه سال‌ها محفوظ می‌مانند."
        )
        info.setWordWrap(True)
        info.setStyleSheet("background-color: #e8f8f5; padding: 10px; border-radius: 8px;")
        yb.addWidget(info)

        year_box.setLayout(yb)
        layout.addWidget(year_box)

        backup_box = QGroupBox("پشتیبان‌گیری و بازیابی")
        backup_box.setFont(QFont(APP_FONT_FAMILY, 11, QFont.Weight.Bold))
        bb = QVBoxLayout()

        row_b = QHBoxLayout()
        btn_backup = QPushButton(ico_text("backup", "پشتیبان‌گیری دستی"))
        btn_backup.setStyleSheet("background-color: #16a085; color: white; font-weight: bold; padding: 10px;")
        btn_backup.clicked.connect(self.backup_db)
        btn_restore = QPushButton(ico_text("restore", "بازیابی از فایل پشتیبان"))
        btn_restore.setStyleSheet("background-color: #8e44ad; color: white; font-weight: bold; padding: 10px;")
        btn_restore.clicked.connect(self.restore_db)
        btn_folder = QPushButton(ico_text("list", "تنظیم پوشه پیش‌فرض بک‌آپ"))
        btn_folder.setStyleSheet("background-color: #2980b9; color: white; font-weight: bold; padding: 10px;")
        btn_folder.clicked.connect(self.choose_default_backup_dir)
        row_b.addWidget(btn_backup)
        row_b.addWidget(btn_restore)
        row_b.addWidget(btn_folder)
        bb.addLayout(row_b)

        self.lbl_backup_dir = QLabel("")
        self.lbl_backup_dir.setWordWrap(True)
        self.lbl_backup_dir.setStyleSheet("background-color:#f4f8fd; border:1px solid #d6e0ec; "
                                          "border-radius:8px; padding:10px;")
        bb.addWidget(self.lbl_backup_dir)

        backup_box.setLayout(bb)
        layout.addWidget(backup_box)

        close_box = QGroupBox("بستن سال کاری (آرشیو)")
        close_box.setFont(QFont(APP_FONT_FAMILY, 11, QFont.Weight.Bold))
        cb = QVBoxLayout()

        warn = QLabel(
            f"{ico('warn')} <b>اخطار:</b> بستن سال کاری باعث خالی شدن جداول قراردادها، پروژه‌ها، "
            "هزینه‌ها و پرداختی‌های همان سال می‌شود. پیش از این کار به‌صورت خودکار نسخه پشتیبان "
            "گرفته می‌شود ولی این عمل بازگشت‌پذیر نیست."
        )
        warn.setWordWrap(True)
        warn.setStyleSheet("background-color: #fdedec; padding: 10px; border-radius: 8px; "
                           "border: 1px solid #e74c3c;")
        cb.addWidget(warn)

        btn_close_year = QPushButton(ico_text("closed", "بستن سال کاری (آرشیو + خالی کردن)"))
        btn_close_year.setStyleSheet("background-color: #c0392b; color: white; font-weight: bold; "
                                     "padding: 12px; font-size: 11pt;")
        btn_close_year.clicked.connect(self.close_work_year)
        cb.addWidget(btn_close_year)

        close_box.setLayout(cb)
        layout.addWidget(close_box)

        settings_box = QGroupBox("تنظیمات عمومی")
        settings_box.setFont(QFont(APP_FONT_FAMILY, 11, QFont.Weight.Bold))
        sb = QVBoxLayout()

        row_s = QHBoxLayout()
        btn_titles = QPushButton(ico_text("edit", "تغییر عنوان‌های پیش‌فرض استودیو"))
        btn_titles.setStyleSheet("background-color: #e67e22; color: white; font-weight: bold; padding: 9px;")
        btn_titles.clicked.connect(self.edit_default_titles)
        btn_pass = QPushButton(ico_text("key", "تغییر رمز عبور برنامه"))
        btn_pass.setStyleSheet("background-color: #c0392b; color: white; font-weight: bold; padding: 9px;")
        btn_pass.clicked.connect(self.change_password)
        row_s.addWidget(btn_titles)
        row_s.addWidget(btn_pass)
        row_s.addStretch()
        sb.addLayout(row_s)

        settings_box.setLayout(sb)
        layout.addWidget(settings_box)

        layout.addStretch()
        self.tab_workyear.setLayout(layout)
        self.load_settings_years()

    def load_settings_years(self):
        if not hasattr(self, 'combo_settings_year'):
            return
        self.combo_settings_year.blockSignals(True)
        self.combo_settings_year.clear()
        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()

        years = set()
        years.add(jdatetime.date.today().year)
        try:
            years.add(get_current_year())
        except Exception:
            pass

        for table in ["wedding_contracts", "commercial_projects", "expenses",
                      "transactions", "checks", "persons"]:
            try:
                cursor.execute(f"SELECT DISTINCT work_year FROM {table} WHERE work_year IS NOT NULL")
                for (y,) in cursor.fetchall():
                    if y:
                        years.add(int(y))
            except Exception:
                pass

        conn.close()

        for y in sorted(years, reverse=True):
            self.combo_settings_year.addItem(str(y))
        try:
            self.combo_settings_year.insertItem(0, str(max(years) + 1))
        except Exception:
            pass

        current = get_current_year()
        idx = self.combo_settings_year.findText(str(current))
        if idx >= 0:
            self.combo_settings_year.setCurrentIndex(idx)
        self.combo_settings_year.blockSignals(False)

        self.update_year_stats()
        self.update_backup_dir_label()

    def update_year_stats(self):
        if not hasattr(self, "lbl_year_stats"):
            return
        try:
            wy = get_current_year()
            conn = sqlite3.connect(DB_NAME)
            cursor = conn.cursor()
            stats = {}
            for key, table in (("قرارداد", "wedding_contracts"), ("پروژه", "commercial_projects"),
                               ("هزینه", "expenses"), ("پرداختی پرسنل", "transactions"),
                               ("چک", "checks")):
                cursor.execute(f"SELECT COUNT(*) FROM {table} WHERE work_year=?", (wy,))
                stats[key] = cursor.fetchone()[0]
            cursor.execute("SELECT COALESCE(SUM(total_amount),0) FROM wedding_contracts WHERE work_year=?", (wy,))
            raw = cursor.fetchone()[0]
            conn.close()
            self.lbl_year_stats.setText(
                f"<b>سال کاری {wy}:</b>   " +
                "  |  ".join(f"{k}: {v}" for k, v in stats.items()) +
                f"  |  جمع خام قراردادها: {raw:,} تومان"
            )
        except Exception as e:
            print("[YEAR] stats error:", e)

    def update_backup_dir_label(self):
        if hasattr(self, "lbl_backup_dir"):
            self.lbl_backup_dir.setText(
                f"{ico('backup')} <b>پوشه پیش‌فرض بک‌آپ:</b><br>{get_setting('backup_dir', os.path.join(APP_DIR, 'backups'))}"
            )

    def choose_default_backup_dir(self):
        current = get_setting("backup_dir", os.path.join(APP_DIR, "backups"))
        path = QFileDialog.getExistingDirectory(self, "انتخاب پوشه پیش‌فرض بک‌آپ", current)
        if path:
            set_setting("backup_dir", path)
            self.update_backup_dir_label()
            QMessageBox.information(self, "موفقیت", f"پوشه پیش‌فرض بک‌آپ ثبت شد:\n{path}")

    def switch_year_from_settings(self):
        try:
            new_year = int(self.combo_settings_year.currentText())
        except Exception:
            return
        set_setting("current_work_year", new_year)

        if hasattr(self, 'w_table'):
            self.load_wedding_contracts()
            self.load_commercial_projects()
            self.load_expenses_table()
            self.load_staff_table()
            self.load_person_events()
            self.load_inventory()
            self.load_checks_table()
            self.load_check_contract_combo()
            self.load_receipt_customers()
            self.calculate_financial_report()

        if hasattr(self, 'combo_work_year'):
            self.combo_work_year.blockSignals(True)
            idx = self.combo_work_year.findText(str(new_year))
            if idx >= 0:
                self.combo_work_year.setCurrentIndex(idx)
            self.combo_work_year.blockSignals(False)

        recalc_inventory_usage()
        self.load_settings_years()
        self.update_top_year_label()
        QMessageBox.information(self, "تغییر سال", f"سال کاری به {new_year} تغییر یافت.")

    def start_new_work_year(self):
        if not ask_security_password(self):
            return
        new_year, ok = QInputDialog.getInt(
            self, "سال کاری جدید", "سال جدید را وارد کنید:",
            get_current_year() + 1, 1300, 1500
        )
        if not ok:
            return
        set_setting("current_work_year", new_year)
        self.load_work_years()
        self.load_settings_years()
        if hasattr(self, 'w_table'):
            self.load_wedding_contracts()
            self.load_commercial_projects()
            self.load_expenses_table()
            self.load_staff_table()
            self.load_person_events()
            self.load_inventory()
            self.load_checks_table()
            self.load_check_contract_combo()
            self.load_receipt_customers()
        self.update_top_year_label()
        QMessageBox.information(self, "موفقیت",
                                f"سال کاری جدید {new_year} شروع شد. همه داده‌های سال‌های قبل محفوظ هستند.")

    def close_work_year(self):
        if not ask_security_password(self):
            return

        reply = QMessageBox.question(
            self, "تایید نهایی",
            "آیا مطمئن هستید که می‌خواهید سال کاری فعلی را ببندید؟\n\n"
            f"{ico('warn')} اطلاعات جاری پاک می‌شود (قبل از آن بک‌آپ گرفته می‌شود).",
            QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
        )
        if reply != QMessageBox.StandardButton.Yes:
            return

        today_str = jdatetime.date.today().strftime("%Y_%m_%d")
        work_year = get_current_year()
        backup_dir = get_setting("backup_dir", os.path.join(APP_DIR, "backups"))
        try:
            if not os.path.exists(backup_dir):
                os.makedirs(backup_dir)
        except Exception:
            backup_dir = APP_DIR
        backup_name = os.path.join(backup_dir, f"studio_archive_{work_year}_{today_str}.db")
        try:
            shutil.copy(DB_NAME, backup_name)
        except Exception as e:
            QMessageBox.critical(self, "خطا", f"گرفتن نسخه پشتیبان ناموفق بود:\n{e}")
            return

        conn = sqlite3.connect(DB_NAME)
        cursor = conn.cursor()
        cursor.execute("DELETE FROM wedding_contracts WHERE work_year=?", (work_year,))
        cursor.execute("DELETE FROM wedding_deposits WHERE work_year=?", (work_year,))
        cursor.execute("DELETE FROM commercial_projects WHERE work_year=?", (work_year,))
        cursor.execute("DELETE FROM expenses WHERE work_year=?", (work_year,))
        cursor.execute("DELETE FROM transactions WHERE work_year=?", (work_year,))
        cursor.execute("DELETE FROM checks WHERE work_year=?", (work_year,))
        cursor.execute("UPDATE inventory SET used_count=0")
        conn.commit()
        conn.close()

        recalc_inventory_usage()
        QMessageBox.information(self, "موفقیت",
                                f"سال کاری {work_year} بسته شد.\n\nفایل آرشیو:\n{backup_name}")
        if hasattr(self, 'w_table'):
            self.load_wedding_contracts()
            self.load_commercial_projects()
            self.load_expenses_table()
            self.load_staff_table()
            self.load_person_events()
            self.load_inventory()
            self.load_checks_table()
            self.load_check_contract_combo()
            self.calculate_financial_report()
        self.load_settings_years()

    # ============================================================
    # ==============  تب ۱۲: درباره برنامه  ====================
    # ============================================================
    def setup_about_tab(self):
        layout = QVBoxLayout()
        layout.setAlignment(Qt.AlignmentFlag.AlignCenter)

        card = QGroupBox("درباره نرم‌افزار IMART STUDIO")
        card.setFont(QFont(APP_FONT_FAMILY, 12, QFont.Weight.Bold))
        card_layout = QVBoxLayout()
        card_layout.setAlignment(Qt.AlignmentFlag.AlignCenter)

        lbl_title = QLabel(f"{ico('commercial')} IMART STUDIO")
        lbl_title.setFont(QFont(APP_FONT_FAMILY, 24, QFont.Weight.Bold))
        lbl_title.setStyleSheet("color: #1F4E78;")
        lbl_title.setAlignment(Qt.AlignmentFlag.AlignCenter)

        lbl_version = QLabel(f"نسخه {APP_VERSION}")
        lbl_version.setFont(QFont(APP_FONT_FAMILY, 13, QFont.Weight.Bold))
        lbl_version.setStyleSheet("color: #7f8c8d;")
        lbl_version.setAlignment(Qt.AlignmentFlag.AlignCenter)

        lbl_desc = QLabel(
            "نرم‌افزار جامع مدیریت مالی، حسابداری، پرسنلی و انبار\n"
            "مخصوص آتلیه‌ها، استودیوهای فیلمبرداری و پروژه‌های تولید محتوا\n\n"
            f"{ico('people')} طراح و توسعه‌دهنده: {DEVELOPER_NAME}\n"
            f"{ico('money')} شماره تماس پشتیبانی: 09173736618\n\n"
            "قابلیت‌های کلیدی این نسخه:\n"
            "• فاکتور، رسید و گزارش‌های چاپی و PDF فارسی با فونت یکان / نازنین\n"
            "• تاریخ شمسی در همه بخش‌ها با پرش خودکار روز به ماه و سال\n"
            "• کد یکتا برای هر فرد، خدمت، قرارداد، پروژه و هزینه\n"
            "• انبار، فاکتور، کارکنان و چک‌ها به‌صورت کامل هماهنگ\n"
            "• نمودارهای فارسی راست‌چین با عنوان متناسب هر نمودار\n\n"
            "تمامی حقوق این نرم‌افزار محفوظ می‌باشد."
        )
        lbl_desc.setFont(QFont(APP_FONT_FAMILY, 11))
        lbl_desc.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lbl_desc.setWordWrap(True)

        card_layout.addWidget(lbl_title)
        card_layout.addSpacing(8)
        card_layout.addWidget(lbl_version)
        card_layout.addSpacing(16)
        card_layout.addWidget(lbl_desc)

        card.setLayout(card_layout)
        card.setFixedWidth(680)

        layout.addWidget(card, alignment=Qt.AlignmentFlag.AlignCenter)
        self.tab_about.setLayout(layout)

    # ============================================================
    # ==============  متدهای عمومی  ============================
    # ============================================================
    def backup_db(self):
        default_dir = get_setting("backup_dir", os.path.join(APP_DIR, "backups"))
        try:
            if not os.path.exists(default_dir):
                os.makedirs(default_dir)
        except Exception:
            default_dir = APP_DIR
        default_name = os.path.join(
            default_dir, f"backup_{jdatetime.date.today().strftime('%Y_%m_%d_%H_%M')}.db")
        file_path, _ = QFileDialog.getSaveFileName(
            self, "ذخیره فایل پشتیبان", default_name, "Database Files (*.db)")
        if file_path:
            try:
                shutil.copy(DB_NAME, file_path)
                set_setting("backup_dir", os.path.dirname(file_path))
                self.update_backup_dir_label()
                QMessageBox.information(self, "موفقیت", f"پشتیبان‌گیری با موفقیت انجام شد.\n{file_path}")
            except Exception as e:
                QMessageBox.critical(self, "خطا", f"پشتیبان‌گیری ناموفق بود:\n{e}")

    def restore_db(self):
        if not ask_security_password(self):
            return
        default_dir = get_setting("backup_dir", APP_DIR)
        file_path, _ = QFileDialog.getOpenFileName(
            self, "انتخاب فایل پشتیبان", default_dir, "Database Files (*.db)")
        if not file_path:
            return
        try:
            safety = os.path.join(APP_DIR, f"before_restore_{jdatetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.db")
            shutil.copy(DB_NAME, safety)
            shutil.copy(file_path, DB_NAME)
            QMessageBox.information(
                self, "موفقیت",
                "پایگاه داده با موفقیت بازیابی شد.\n"
                "یک نسخه از اطلاعات قبلی هم ذخیره شد:\n"
                f"{safety}\n\nبرای اعمال کامل، برنامه را دوباره اجرا کنید.")
        except Exception as e:
            QMessageBox.critical(self, "خطا", f"بازیابی ناموفق بود:\n{e}")

    def change_password(self):
        if not ask_security_password(self):
            return
        new_pass, ok = QInputDialog.getText(
            self, "تغییر رمز عبور", "رمز عبور جدید را وارد کنید:",
            QLineEdit.EchoMode.Password)
        if ok and new_pass.strip():
            set_setting("app_password", new_pass.strip())
            QMessageBox.information(self, "موفقیت", "رمز عبور جدید با موفقیت ثبت گردید.")

    def closeEvent(self, event):
        """
        هنگام بستن برنامه، منوی بک‌آپ خودکار باز می‌شود و
        آدرس پیش‌فرض ذخیره بک‌آپ هر بار از کاربر پرسیده می‌شود.
        """
        if getattr(self, "_allow_close", False):
            self._run_auto_backup_rotation()
            event.accept()
            return

        dlg = ExitBackupDialog(self)
        if dlg.exec() != QDialog.DialogCode.Accepted:
            event.ignore()
            return

        if dlg.do_backup:
            target_dir = dlg.backup_dir
            try:
                if not os.path.exists(target_dir):
                    os.makedirs(target_dir)
                set_setting("backup_dir", target_dir)
                ts = jdatetime.datetime.now().strftime("%Y_%m_%d_%H_%M_%S")
                backup_path = os.path.join(target_dir, f"backup_exit_{ts}.db")
                shutil.copy(DB_NAME, backup_path)
                QMessageBox.information(self, "پشتیبان‌گیری", f"نسخه پشتیبان ذخیره شد:\n{backup_path}")
            except Exception as e:
                if QMessageBox.question(
                        self, "خطا در پشتیبان‌گیری",
                        f"پشتیبان‌گیری ناموفق بود:\n{e}\n\nباز هم از برنامه خارج می‌شوید؟",
                        QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No
                ) != QMessageBox.StandardButton.Yes:
                    event.ignore()
                    return

        self._run_auto_backup_rotation()
        event.accept()

    def _run_auto_backup_rotation(self):
        """نگهداری نسخه‌های داخلی و حذف قدیمی‌ترها"""
        try:
            save_geometry("main", self)
        except Exception:
            pass
        try:
            if not os.path.exists(DB_NAME):
                return
            backup_dir = os.path.join(APP_DIR, "auto_backups")
            if not os.path.exists(backup_dir):
                os.makedirs(backup_dir)
            ts = jdatetime.datetime.now().strftime("%Y_%m_%d_%H_%M_%S")
            shutil.copy(DB_NAME, os.path.join(backup_dir, f"auto_backup_{ts}.db"))
            backups = sorted(f for f in os.listdir(backup_dir) if f.startswith("auto_backup_"))
            if len(backups) > 10:
                for old in backups[:-10]:
                    try:
                        os.remove(os.path.join(backup_dir, old))
                    except Exception:
                        pass
        except Exception as e:
            print("Auto backup error:", e)


# ============================================================
# ===============  ExitBackupDialog (بک‌آپ خروج)  ============
# ============================================================
class ExitBackupDialog(QDialog):
    """
    منوی بک‌آپ خودکار هنگام بستن برنامه.
    آدرس پیش‌فرض ذخیره بک‌آپ هر بار به کاربر نشان داده و پرسیده می‌شود.
    """

    def __init__(self, parent=None):
        super().__init__(parent)
        self.setWindowTitle("پشتیبان‌گیری پیش از خروج")
        self.setLayoutDirection(Qt.LayoutDirection.RightToLeft)
        self.setFont(QFont(APP_FONT_FAMILY, 10))
        self.setMinimumWidth(560)
        self.do_backup = True

        default_dir = get_setting("backup_dir", os.path.join(APP_DIR, "backups"))
        self.backup_dir = default_dir

        lay = QVBoxLayout(self)
        lay.setSpacing(10)

        head = QLabel(f"{ico('backup')}  پشتیبان‌گیری خودکار قبل از بستن برنامه")
        head.setProperty("heading", True)
        head.setAlignment(Qt.AlignmentFlag.AlignCenter)
        lay.addWidget(head)

        info = QLabel(
            "برای جلوگیری از هرگونه از دست رفتن اطلاعات، پیش از خروج یک نسخه پشتیبان "
            "از پایگاه داده گرفته می‌شود.\n"
            "لطفاً <b>آدرس پیش‌فرض ذخیره بک‌آپ</b> را بررسی و در صورت نیاز تغییر دهید."
        )
        info.setWordWrap(True)
        info.setStyleSheet("background-color:#eaf6ff; border:1px solid #a9d3ee; "
                           "border-radius:8px; padding:10px;")
        lay.addWidget(info)

        path_row = QHBoxLayout()
        self.txt_dir = QLineEdit(default_dir)
        self.txt_dir.setReadOnly(True)
        self.txt_dir.setStyleSheet("background-color:#ffffff; border:1px solid #cbd7e6; "
                                   "border-radius:8px; padding:7px;")
        btn_pick = QPushButton(ico_text("list", "تغییر پوشه..."))
        btn_pick.setStyleSheet("background-color:#2980b9; color:white; font-weight:bold; padding:8px;")
        btn_pick.clicked.connect(self.pick_dir)
        path_row.addWidget(QLabel("آدرس ذخیره:"))
        path_row.addWidget(self.txt_dir, 3)
        path_row.addWidget(btn_pick)
        lay.addLayout(path_row)

        self.chk_backup = QCheckBox("پشتیبان‌گیری کن و سپس خارج شو (پیشنهاد می‌شود)")
        self.chk_backup.setChecked(True)
        lay.addWidget(self.chk_backup)

        self.chk_skip = QCheckBox("بدون پشتیبان‌گیری خارج شو")
        lay.addWidget(self.chk_skip)

        self.chk_backup.toggled.connect(self._sync_checks)
        self.chk_skip.toggled.connect(self._sync_checks)

        btn_row = QHBoxLayout()
        btn_ok = QPushButton(ico_text("ok", "تایید و خروج"))
        btn_ok.setStyleSheet("background-color:#27ae60; color:white; font-weight:bold; padding:10px;")
        btn_ok.clicked.connect(self._accept)
        btn_cancel = QPushButton(ico_text("cancel", "انصراف و بازگشت به برنامه"))
        btn_cancel.setStyleSheet("background-color:#7f8c8d; color:white; font-weight:bold; padding:10px;")
        btn_cancel.clicked.connect(self.reject)
        btn_row.addWidget(btn_ok, 2)
        btn_row.addWidget(btn_cancel, 2)
        lay.addLayout(btn_row)

    def _sync_checks(self):
        sender = self.sender()
        if sender is self.chk_backup and self.chk_backup.isChecked():
            self.chk_skip.setChecked(False)
        elif sender is self.chk_skip and self.chk_skip.isChecked():
            self.chk_backup.setChecked(False)
        if not self.chk_backup.isChecked() and not self.chk_skip.isChecked():
            self.chk_backup.setChecked(True)

    def pick_dir(self):
        path = QFileDialog.getExistingDirectory(self, "انتخاب پوشه ذخیره بک‌آپ", self.backup_dir)
        if path:
            self.backup_dir = path
            self.txt_dir.setText(path)

    def _accept(self):
        self.do_backup = self.chk_backup.isChecked()
        self.backup_dir = self.txt_dir.text().strip() or self.backup_dir
        self.accept()

# ============================================================
# ====================  نقطه ورود  ==========================
# ============================================================
def setup_logging():
    """
    در حالت اجرای ویندوزی (PyInstaller --windowed) مقدار sys.stdout برابر None است
    و هر print() باعث خطای پنهان می‌شود. اینجا خروجی به فایل لاگ هدایت می‌شود.
    """
    try:
        if sys.stdout is None or sys.stderr is None:
            log_path = os.path.join(APP_DIR, "imart_studio.log")
            f = open(log_path, "a", encoding="utf-8", buffering=1)
            if sys.stdout is None:
                sys.stdout = f
            if sys.stderr is None:
                sys.stderr = f
    except Exception:
        pass


def install_exception_hook(app_ref=None):
    """نمایش خطاهای پیش‌بینی‌نشده به‌جای بسته شدن بی‌صدا"""
    def _hook(exctype, value, tb):
        import traceback
        text = "".join(traceback.format_exception(exctype, value, tb))
        try:
            print("[UNHANDLED]\n" + text)
        except Exception:
            pass
        try:
            QMessageBox.critical(None, "خطای غیرمنتظره",
                                 f"یک خطای پیش‌بینی‌نشده رخ داد:\n\n{value}\n\n"
                                 "جزئیات در فایل imart_studio.log ذخیره شد.")
        except Exception:
            pass
    sys.excepthook = _hook


# ============================================================
# ====================  نقطه ورود  ==========================
# ============================================================
if __name__ == '__main__':
    setup_logging()

    # رفع «ذخیره نشدن تغییرات بعد از بستن برنامه»:
    # دیتابیس نسخه قبلی به مسیر ثابت برنامه منتقل می‌شود.
    migrate_legacy_database()

    app = QApplication(sys.argv)
    app.setApplicationName("IMART STUDIO")
    app.setApplicationVersion(APP_VERSION)
    install_exception_hook(app)

    APP_FONT_FAMILY = setup_fonts()
    app.setFont(QFont(APP_FONT_FAMILY, 10))
    apply_global_theme(app)

    icon_path = resource_path("Accounting.png")
    if os.path.exists(icon_path):
        app.setWindowIcon(QIcon(icon_path))

    # آماده‌سازی فونت نمودار
    setup_matplotlib_font()

    loading = LoadingScreen()
    loading.show()

    steps = [
        ("راه‌اندازی موتور برنامه...", 5),
        ("بارگذاری فونت‌های فارسی...", 12),
        ("اتصال به پایگاه داده SQLite...", 22),
        ("ایجاد جداول و کدهای یکتا...", 35),
        ("بارگذاری تنظیمات امنیتی...", 45),
        ("آماده‌سازی داشبورد...", 55),
        ("بارگذاری تب‌های برنامه...", 70),
        ("بارگذاری قراردادها و پروژه‌ها...", 82),
        ("بارگذاری انبار، بانک و چک‌ها...", 92),
        ("آماده‌سازی نهایی...", 98),
        ("ورود به سیستم...", 100),
    ]

    main_win = StudioAccountingApp()
    current_step = {"index": 0}

    def finish_loading():
        loading.close()
        if main_win.prompt_login():
            main_win.stack.setCurrentWidget(main_win.dashboard_screen)
            main_win.center_on_screen()
            main_win.show()
            main_win.raise_()
            main_win.activateWindow()
            QTimer.singleShot(2500, main_win.check_check_alerts)
        else:
            main_win._allow_close = True
            app.quit()

    def run_step():
        if current_step["index"] < len(steps):
            text, percent = steps[current_step["index"]]
            loading.set_progress(percent, text)

            if current_step["index"] == 2:
                migrate_legacy_database()
            elif current_step["index"] == 3:
                init_db()
            elif current_step["index"] == 5:
                main_win.setup_dashboard_ui()
            elif current_step["index"] == 6:
                main_win.setup_main_app_ui()

            current_step["index"] += 1
            QTimer.singleShot(250, run_step)
        else:
            QTimer.singleShot(400, finish_loading)

    QTimer.singleShot(400, run_step)

    sys.exit(app.exec())